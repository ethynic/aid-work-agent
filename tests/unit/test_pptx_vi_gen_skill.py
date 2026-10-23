"""pptx-vi-gen 技能测试：注册发现 + CLI 校验与端到端生成 + 编辑往返。

端到端用例实际生成 pptx（依赖容器内 Noto CJK 字体），
小内容下耗时秒级。产出为 python-pptx 原生可编辑 PPT。
"""

import json
import subprocess
import sys
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SKILL_DIR = PROJECT_ROOT / "src" / "skills" / "pptx-vi-gen-1.0.0"
GENERATE = SKILL_DIR / "scripts" / "generate.py"
EDIT = SKILL_DIR / "scripts" / "edit.py"
TEMPLATE_FILL = SKILL_DIR / "scripts" / "template_fill.py"

CONTENT = {
    "slides": [
        {"role": "cover", "title": "测试演示", "subtitle": "副标题",
         "date": "2026 年 10 月", "org": "测试部"},
        {"role": "bullets", "title": "要点页", "lead": "导语。",
         "items": [{"h": "要点一", "d": "具体描述一"},
                   {"h": "要点二", "d": "具体描述二"}]},
        {"role": "end", "title": "谢谢观看"},
    ]
}


def _run_cli(args, cwd, content=None, stdin_data=None):
    env = {
        "PATH": "/usr/local/bin:/usr/bin:/bin",
        "PROJECT_ROOT": str(PROJECT_ROOT),
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "HOME": "/tmp",
    }
    if content is not None:
        content_path = cwd / "content.json"
        content_path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
        args = args + ["--content", "content.json"]
    cmd = [sys.executable, str(GENERATE)] + args
    proc = subprocess.run(
        cmd, cwd=str(cwd), env=env, capture_output=True, text=True,
        input=stdin_data, timeout=120,
    )
    return proc


class TestPptxViGenSkill(unittest.TestCase):
    def test_skill_discovered_by_registry(self):
        from src.core.skill_registry import SkillRegistry
        registry = SkillRegistry()
        registry.load_from_directory(str(PROJECT_ROOT / "src" / "skills"))
        self.assertIn("pptx-vi-gen", registry.list_skills())
        skill = registry.get("pptx-vi-gen")
        self.assertIsNotNone(skill)

    def test_cli_end_to_end_default_vi(self):
        workdir = PROJECT_ROOT / "storage" / "tmp" / "test_pptx_vi_gen"
        workdir.mkdir(parents=True, exist_ok=True)
        proc = _run_cli(["--output", "out_test.pptx", "--preview"], workdir, content=CONTENT)
        result = json.loads(proc.stdout)
        self.assertTrue(result["success"], msg=proc.stdout + proc.stderr)
        pptx_path = Path(result["pptx"])
        self.assertTrue(pptx_path.exists())
        self.assertGreater(pptx_path.stat().st_size, 10_000)
        self.assertEqual(result["slides"], 3)
        self.assertTrue(result["vi"]["has_logo"] is False)
        preview = Path(result["preview_dir"])
        self.assertEqual(len(list(preview.glob("slide_*.png"))), 3)
        # 原生可编辑：文本框内存在 content 中的标题（非整页贴图）
        from pptx import Presentation
        prs = Presentation(str(pptx_path))
        all_text = "\n".join(
            sh.text_frame.text
            for slide in prs.slides for sh in slide.shapes
            if sh.has_text_frame
        )
        self.assertIn("测试演示", all_text)
        self.assertIn("要点一", all_text)
        pptx_path.unlink(missing_ok=True)
        import shutil
        shutil.rmtree(preview, ignore_errors=True)

    def test_cli_stdin_content(self):
        workdir = PROJECT_ROOT / "storage" / "tmp" / "test_pptx_vi_gen"
        workdir.mkdir(parents=True, exist_ok=True)
        proc = _run_cli(["--output", "out_stdin.pptx", "--content", "-"], workdir,
                        stdin_data=json.dumps(CONTENT, ensure_ascii=False))
        result = json.loads(proc.stdout)
        self.assertTrue(result["success"], msg=proc.stdout + proc.stderr)
        Path(result["pptx"]).unlink(missing_ok=True)

    def test_cli_rejects_unknown_role(self):
        workdir = PROJECT_ROOT / "storage" / "tmp" / "test_pptx_vi_gen"
        workdir.mkdir(parents=True, exist_ok=True)
        bad = {"slides": [{"role": "chart", "title": "x"}]}
        proc = _run_cli([], workdir, content=bad)
        result = json.loads(proc.stdout)
        self.assertFalse(result["success"])
        self.assertIn("role 非法", result["error"])

    def test_cli_rejects_missing_title(self):
        workdir = PROJECT_ROOT / "storage" / "tmp" / "test_pptx_vi_gen"
        workdir.mkdir(parents=True, exist_ok=True)
        bad = {"slides": [{"role": "cover", "subtitle": "no title"}]}
        proc = _run_cli([], workdir, content=bad)
        result = json.loads(proc.stdout)
        self.assertFalse(result["success"])
        self.assertIn("缺少 title", result["error"])

    def test_cli_rejects_missing_required_fields(self):
        workdir = PROJECT_ROOT / "storage" / "tmp" / "test_pptx_vi_gen"
        workdir.mkdir(parents=True, exist_ok=True)
        cases = [
            ({"role": "stats", "title": "x"}, "缺少必需字段 stats"),
            ({"role": "twocol", "title": "x",
              "left": {"h": "a", "items": ["1"]}}, "缺少必需字段 right"),
            ({"role": "bullets", "title": "x", "items": [{"d": "no h"}]},
             "缺少 h（小标题）"),
            ({"role": "stats", "title": "x",
              "stats": [{"num": str(k), "label": "l"} for k in range(5)]},
             "超过 4 张卡"),
        ]
        for content, expect in cases:
            proc = _run_cli([], workdir, content={"slides": [content]})
            result = json.loads(proc.stdout)
            self.assertFalse(result["success"], msg=proc.stdout)
            self.assertIn(expect, result["error"], msg=proc.stdout)


    def test_edit_cli_round_trip(self):
        """生成 -> edit.py 改字/删页，断言修改生效、.bak 生成。"""
        workdir = PROJECT_ROOT / "storage" / "tmp" / "test_pptx_vi_gen"
        workdir.mkdir(parents=True, exist_ok=True)
        proc = _run_cli(["--output", "out_edit.pptx"], workdir, content=CONTENT)
        result = json.loads(proc.stdout)
        self.assertTrue(result["success"], msg=proc.stdout + proc.stderr)
        deck = Path(result["pptx"])

        spec = {"ops": [
            {"op": "replace_text", "find": "测试演示", "replace": "修改后的标题", "scope": "deck"},
            {"op": "delete_slide", "slide": 2},
        ]}
        (workdir / "edit_spec.json").write_text(
            json.dumps(spec, ensure_ascii=False), encoding="utf-8")
        cmd = [sys.executable, str(EDIT), "--pptx", "out_edit.pptx", "--spec", "edit_spec.json"]
        proc = subprocess.run(
            cmd, cwd=str(workdir),
            env={"PATH": "/usr/local/bin:/usr/bin:/bin", "PROJECT_ROOT": str(PROJECT_ROOT),
                 "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1", "HOME": "/tmp"},
            capture_output=True, text=True, timeout=120)
        result = json.loads(proc.stdout)
        self.assertTrue(result["success"], msg=proc.stdout + proc.stderr)
        self.assertEqual(result["slides"], 2)
        self.assertTrue((workdir / "out_edit.pptx.bak").exists())

        from pptx import Presentation
        prs = Presentation(str(deck))
        all_text = "\n".join(
            sh.text_frame.text
            for slide in prs.slides for sh in slide.shapes
            if sh.has_text_frame
        )
        self.assertIn("修改后的标题", all_text)
        self.assertNotIn("测试演示", all_text)
        self.assertNotIn("要点页", all_text)   # 第 2 页已删

        deck.unlink(missing_ok=True)
        Path(str(deck) + ".bak").unlink(missing_ok=True)
        (workdir / "edit_spec.json").unlink(missing_ok=True)

    def test_cli_new_roles_end_to_end(self):
        """上游 2026-10-22 新增角色端到端：toc/table/timeline/process/quote/stats。

        同时验证溢出保护：stats 大数字 fit=True + max_lines=1 -> word_wrap=False。
        """
        workdir = PROJECT_ROOT / "storage" / "tmp" / "test_pptx_vi_gen"
        workdir.mkdir(parents=True, exist_ok=True)
        content = {"slides": [
            {"role": "toc", "items": ["第一章", "第二章", "第三章"]},
            {"role": "stats", "title": "关键指标",
             "stats": [{"num": "44.63万", "label": "成交额", "desc": "同比 +18%"},
                       {"num": "1,234", "label": "订单数", "desc": "环比 +6%"},
                       {"num": "31%", "label": "毛利率", "desc": "稳中有升"}]},
            {"role": "table", "title": "季度数据", "head": ["指标", "Q1", "Q2"],
             "rows": [["营收", "120", "150"], ["毛利率", "31%", "33%"]],
             "caption": "数据来源：测试"},
            {"role": "timeline", "title": "里程碑",
             "items": [{"t": "1 月", "h": "立项", "d": "完成调研"},
                       {"t": "2 月", "h": "开发", "d": "Alpha 版本"},
                       {"t": "3 月", "h": "上线", "d": "全量发布"}]},
            {"role": "process", "title": "实施流程",
             "steps": [{"h": "需求", "d": "澄清范围"},
                       {"h": "设计", "d": "原型评审"},
                       {"h": "交付", "d": "验收上线"}]},
            {"role": "quote", "quote": "把简单留给用户。", "author": "测试团队"},
            {"role": "end", "title": "谢谢观看"},
        ]}
        proc = _run_cli(["--output", "out_roles.pptx"], workdir, content=content)
        result = json.loads(proc.stdout)
        self.assertTrue(result["success"], msg=proc.stdout + proc.stderr)
        deck = Path(result["pptx"])
        try:
            from pptx import Presentation
            prs = Presentation(str(deck))
            all_text = "\n".join(
                sh.text_frame.text
                for slide in prs.slides for sh in slide.shapes
                if sh.has_text_frame
            )
            self.assertIn("关键指标", all_text)
            self.assertIn("44.63万", all_text)
            self.assertIn("里程碑", all_text)
            self.assertIn("把简单留给用户。", all_text)
            # 新版式产出原生表格（GraphicFrame，非贴图）
            self.assertTrue(
                any(sh.has_table for slide in prs.slides for sh in slide.shapes))
            # 溢出保护：fit=True + max_lines=1 的文本框必须禁换行
            num_frames = [
                sh.text_frame for slide in prs.slides for sh in slide.shapes
                if sh.has_text_frame and "44.63万" in sh.text_frame.text
            ]
            self.assertTrue(num_frames)
            self.assertFalse(num_frames[0].word_wrap)
        finally:
            deck.unlink(missing_ok=True)

    def test_native_measure_uses_real_font(self):
        """measure() 必须走真实字体文件，而不是静默退回 em 估算。"""
        import sys as _sys
        _sys.path.insert(0, str(SKILL_DIR / "scripts"))
        from pptx_gen import native
        from pptx_gen.native import measure
        s = "44.63万"
        m = measure(s, 94, bold=True)
        em = sum(native._em(ch) for ch in s) * 94
        self.assertGreater(m, 0)
        # 字体缺失时 measure 精确等于 em 估算 -> 此断言失败即为字体未生效
        self.assertNotAlmostEqual(m, em, delta=max(1.0, em * 0.005))

    def test_template_fill_analyze_smoke(self):
        """路线一 analyze 冒烟：临时构造最小模板，断言解析出版式与占位符。"""
        from pptx import Presentation
        workdir = PROJECT_ROOT / "storage" / "tmp" / "test_pptx_vi_gen"
        workdir.mkdir(parents=True, exist_ok=True)
        tpl = workdir / "tpl_smoke.pptx"
        prs = Presentation()
        s1 = prs.slides.add_slide(prs.slide_layouts[0])
        s1.shapes.title.text = "模板标题"
        prs.save(str(tpl))

        proc = subprocess.run(
            [sys.executable, str(TEMPLATE_FILL), "analyze", "--template", "tpl_smoke.pptx"],
            cwd=str(workdir),
            env={"PATH": "/usr/local/bin:/usr/bin:/bin", "PROJECT_ROOT": str(PROJECT_ROOT),
                 "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1", "HOME": "/tmp"},
            capture_output=True, text=True, timeout=120)
        result = json.loads(proc.stdout)
        self.assertTrue(result["success"], msg=proc.stdout + proc.stderr)
        layouts = result["analysis"]["layouts"]
        self.assertTrue(layouts)
        idx0 = layouts[0]["placeholders"]
        self.assertIn(0, [p["idx"] for p in idx0])
        tpl.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
