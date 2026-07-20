#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel 智能模板填充 — ExcelProcessTool 接线集成测试（Phase C）

覆盖 execute() 入口的 data + 样例 → 智能填充全链路、租户注入、旧 variables 兼容。
"""

import json
import sys
from pathlib import Path

import openpyxl
import pytest
from unittest.mock import patch

PROJECT_ROOT = Path(__file__).resolve().parents[3]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.tools.excel.excel_process_tool import ExcelProcessTool


def _build_sample(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报价"
    ws["A1"] = "表头1"
    ws["A2"] = "表头2"
    for i, h in enumerate(["类别", "项目", "金额"], 1):
        ws.cell(row=3, column=i, value=h)
    example = [["住宿", "示例A", 100], ["门票", "示例B", 50]]
    for r, row in enumerate(example, start=4):
        for i, v in enumerate(row, 1):
            ws.cell(row=r, column=i, value=v)
    ws["A6"] = "合计"
    ws["C6"] = 150
    wb.save(str(path))
    return path


STRUCTURE = {
    "columns": [
        {"col": 1, "header": "类别", "bind": "category"},
        {"col": 2, "header": "项目", "bind": "name"},
        {"col": 3, "header": "金额", "bind": "amount"},
    ],
    "detail_first_row": 4,
    "detail_last_row": 5,
    "detail_template_row": 4,
    "totals": [{"row": 6, "col": 3, "bind": "grand_total"}],
}


def _mock_llm():
    payload = json.dumps(STRUCTURE, ensure_ascii=False)
    return lambda prompt: payload


@pytest.fixture
def sample_path(tmp_path):
    return _build_sample(tmp_path / "sample.xlsx")


@pytest.mark.asyncio
async def test_execute_fill_with_data(sample_path, tmp_path):
    """execute(data + file_paths) → 确定性路由到 fill_template 智能填充"""
    data = {
        "rows": [
            {"category": "住宿", "name": "酒店X", "amount": 200},
            {"category": "门票", "name": "景点Y", "amount": 80},
        ],
        "totals": {"grand_total": 280},
    }
    tool = ExcelProcessTool()
    with patch("src.tools.excel.excel_template_ai._default_llm", _mock_llm()):
        res = await tool.execute(data=data, file_paths=[str(sample_path)], output_name="out.xlsx")

    assert res["success"], res
    assert res["file_path"]
    assert res["rows_rendered"] == 2

    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    assert ws["A4"].value == "住宿" and ws["B4"].value == "酒店X" and ws["C4"].value == 200
    assert ws["C6"].value == 280  # 合计


@pytest.mark.asyncio
async def test_execute_fill_m_greater_than_k(sample_path, tmp_path):
    """data 行多于样例 → 插入行，合计下移"""
    data = {
        "rows": [
            {"category": "A", "name": "n1", "amount": 10},
            {"category": "B", "name": "n2", "amount": 20},
            {"category": "C", "name": "n3", "amount": 30},
            {"category": "D", "name": "n4", "amount": 40},
        ],
        "totals": {"grand_total": 100},
    }
    tool = ExcelProcessTool()
    with patch("src.tools.excel.excel_template_ai._default_llm", _mock_llm()):
        res = await tool.execute(data=data, file_paths=[str(sample_path)])
    assert res["success"] and res["rows_rendered"] == 4
    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # 明细 4-7，合计下移到 8
    assert ws["A7"].value == "D"
    assert ws["A8"].value == "合计" and ws["C8"].value == 100


@pytest.mark.asyncio
async def test_execute_fill_tenant_injection(sample_path):
    """set_tenant_id/set_user_id 注入后，输出落到租户目录"""
    tool = ExcelProcessTool()
    tool.set_tenant_id("tenant_test_tt")
    tool.set_user_id("user_test_uu")
    data = {"rows": [{"category": "A", "name": "n", "amount": 1}], "totals": {}}
    with patch("src.tools.excel.excel_template_ai._default_llm", _mock_llm()):
        res = await tool.execute(data=data, file_paths=[str(sample_path)])
    assert res["success"]
    fp = Path(res["file_path"])
    assert "tenant_test_tt" in fp.parts
    assert "user_test_uu" in fp.parts


@pytest.mark.asyncio
async def test_execute_data_without_file_paths_errors(sample_path):
    """data 但无样例附件 → 友好报错（不静默）"""
    tool = ExcelProcessTool()
    data = {"rows": [{"category": "A", "name": "n", "amount": 1}]}
    with patch("src.tools.excel.excel_template_ai._default_llm", _mock_llm()):
        res = await tool.execute(data=data, file_paths=None)
    assert not res["success"]


@pytest.mark.asyncio
async def test_input_schema_has_data_field():
    """ExcelProcessInput 暴露 data 字段供 agent 调用"""
    schema = ExcelProcessTool().to_tool_definition()["input_schema"]
    assert "data" in schema.get("properties", {})


@pytest.mark.asyncio
async def test_legacy_variables_path_still_works(sample_path, tmp_path):
    """旧 variables 占位符路径不被新 data 分支破坏（直接调 handler）"""
    # 构造一个带 {{var}} 的模板
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "标题：{{title}}"
    ws["A2"] = "{{name}}"
    legacy = tmp_path / "legacy.xlsx"
    wb.save(str(legacy))

    tool = ExcelProcessTool()
    from src.tools.excel.excel_process_tool import PipelineContext
    ctx = PipelineContext(file_paths=[str(legacy)])
    res = await tool._handle_fill_template(
        ctx, {"variables": {"title": "测试标题", "name": "张三"}}
    )
    assert res["success"]
    wb2 = openpyxl.load_workbook(res["file_path"])
    assert wb2.active["A1"].value == "标题：测试标题"
    assert wb2.active["A2"].value == "张三"
