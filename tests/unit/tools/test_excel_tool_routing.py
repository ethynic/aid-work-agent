import sys
from pathlib import Path

import openpyxl
import pytest
from unittest.mock import AsyncMock, patch


PROJECT_ROOT = Path(__file__).resolve().parents[3]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


def test_excel_process_schema_has_structured_input_fields():
    from src.tools.excel.excel_process_tool import ExcelProcessTool

    schema = ExcelProcessTool().to_tool_definition()["input_schema"]
    properties = schema.get("properties", {})

    assert "instruction" in properties
    assert "content" in properties
    assert "content_type" in properties
    assert "output_name" in properties
    assert "context" in properties
    assert "file_paths" in properties


@pytest.mark.asyncio
async def test_excel_router_exports_markdown_table_without_llm():
    from src.tools.excel.excel_router import ExcelRouter

    router = ExcelRouter()
    context = """请导出Excel文件

| 类别 | 项目 | 金额 |
|------|------|------|
| 门票 | 成人票 | 120 |
"""

    result = await router.route(context)

    assert result["task"] == "export"
    assert result["params"]["data_type"] == "markdown"


@pytest.mark.asyncio
async def test_excel_process_export_structured_markdown_content():
    from src.tools.excel.excel_process_tool import ExcelProcessTool

    tool = ExcelProcessTool()
    mock_router = AsyncMock()
    mock_router.route.side_effect = AssertionError("确定性路由不应调用内部 LLM")
    tool._router = mock_router

    table = "| 项目 | 金额 |\n|------|------|\n| 门票 | 120 |"
    mock_result = {
        "success": True,
        "file_path": "/tmp/quote.xlsx",
        "file_name": "quote.xlsx",
        "file_size": 1024,
        "row_count": 1,
    }

    with patch("src.tools.excel.excel_writer.create_excel", return_value=mock_result) as mock_create:
        result = await tool.execute(
            instruction="导出Excel文件",
            content=table,
            content_type="markdown",
            output_name="报价单.xlsx",
        )

    assert result["success"] is True
    kwargs = mock_create.call_args.kwargs
    assert kwargs["data"] == table
    assert kwargs["data_type"] == "markdown"
    assert kwargs["file_name"] == "报价单.xlsx"


@pytest.mark.asyncio
async def test_excel_process_export_legacy_context_strips_instruction_prefix():
    from src.tools.excel.excel_process_tool import ExcelProcessTool

    tool = ExcelProcessTool()
    mock_router = AsyncMock()
    mock_router.route.side_effect = AssertionError("确定性路由不应调用内部 LLM")
    tool._router = mock_router

    context = "请导出Excel文件，下面是Markdown表格：\n\n| 项目 | 金额 |\n|------|------|\n| 住宿 | 300 |"
    mock_result = {
        "success": True,
        "file_path": "/tmp/export.xlsx",
        "file_name": "export.xlsx",
        "file_size": 1024,
        "row_count": 1,
    }

    with patch("src.tools.excel.excel_writer.create_excel", return_value=mock_result) as mock_create:
        result = await tool.execute(context=context)

    assert result["success"] is True
    data = mock_create.call_args.kwargs["data"]
    assert data.startswith("| 项目 | 金额 |")
    assert "请导出" not in data


@pytest.mark.asyncio
async def test_excel_process_export_structured_csv_content():
    from src.tools.excel.excel_process_tool import ExcelProcessTool

    tool = ExcelProcessTool()
    mock_router = AsyncMock()
    mock_router.route.side_effect = AssertionError("确定性路由不应调用内部 LLM")
    tool._router = mock_router

    csv_text = "项目,金额\n门票,120\n住宿,300"
    mock_result = {
        "success": True,
        "file_path": "/tmp/csv.xlsx",
        "file_name": "csv.xlsx",
        "file_size": 1024,
        "row_count": 2,
    }

    with patch("src.tools.excel.excel_writer.create_excel", return_value=mock_result) as mock_create:
        result = await tool.execute(
            instruction="导出Excel文件",
            content=csv_text,
            content_type="csv",
        )

    assert result["success"] is True
    kwargs = mock_create.call_args.kwargs
    assert kwargs["data"] == csv_text
    assert kwargs["data_type"] == "csv"


@pytest.mark.asyncio
async def test_excel_process_export_structured_json_content():
    from src.tools.excel.excel_process_tool import ExcelProcessTool

    tool = ExcelProcessTool()
    mock_router = AsyncMock()
    mock_router.route.side_effect = AssertionError("确定性路由不应调用内部 LLM")
    tool._router = mock_router

    json_text = '[{"项目": "门票", "金额": 120}, {"项目": "住宿", "金额": 300}]'
    mock_result = {
        "success": True,
        "file_path": "/tmp/json.xlsx",
        "file_name": "json.xlsx",
        "file_size": 1024,
        "row_count": 2,
    }

    with patch("src.tools.excel.excel_writer.create_excel", return_value=mock_result) as mock_create:
        result = await tool.execute(
            instruction="导出Excel文件",
            content=json_text,
            content_type="json",
        )

    assert result["success"] is True
    kwargs = mock_create.call_args.kwargs
    assert kwargs["data"] == [{"项目": "门票", "金额": 120}, {"项目": "住宿", "金额": 300}]
    assert kwargs["data_type"] == "dict_list"


@pytest.mark.asyncio
async def test_excel_process_export_needs_data_when_only_intent():
    from src.tools.excel.excel_process_tool import ExcelProcessTool

    tool = ExcelProcessTool()
    mock_router = AsyncMock()
    mock_router.route.side_effect = AssertionError("确定性路由不应调用内部 LLM")
    tool._router = mock_router

    result = await tool.execute(instruction="帮我导出Excel文件")

    assert result["success"] is False
    assert result["needs_data"] is True


@pytest.mark.asyncio
async def test_excel_process_csv_attachment_converts_without_needs_data(tmp_path):
    from src.tools.excel.excel_process_tool import ExcelProcessTool

    tool = ExcelProcessTool()
    mock_router = AsyncMock()
    mock_router.route.side_effect = AssertionError("确定性路由不应调用内部 LLM")
    tool._router = mock_router

    csv_path = tmp_path / "quote.csv"
    csv_path.write_text("项目,金额\n门票,120\n", encoding="utf-8")
    mock_result = {
        "success": True,
        "file_path": "/tmp/quote.xlsx",
        "file_name": "quote.xlsx",
        "file_size": 1024,
    }

    with patch("src.tools.excel.excel_writer.convert_format", return_value=mock_result) as mock_convert:
        result = await tool.execute(
            instruction="把CSV转成Excel",
            file_paths=[str(csv_path)],
            output_name="报价单.xlsx",
        )

    assert result["success"] is True
    kwargs = mock_convert.call_args.kwargs
    assert kwargs["target_format"] == "excel"
    assert kwargs["output_name"] == "报价单.xlsx"


@pytest.mark.asyncio
async def test_excel_router_converts_csv_attachment_without_llm(tmp_path):
    from src.tools.excel.excel_router import ExcelRouter

    csv_path = tmp_path / "quote.csv"
    csv_path.write_text("项目,金额\n门票,120\n", encoding="utf-8")

    result = await ExcelRouter().route("有EXCEL格式吗", [str(csv_path)])

    assert result["task"] == "convert"
    assert result["params"]["source_format"] == "csv"
    assert result["params"]["target_format"] == "excel"
    assert result["params"]["output_name"] == "quote.xlsx"


def test_convert_csv_to_excel_creates_readable_workbook(tmp_path, monkeypatch):
    from src.tools.excel.excel_lib import ExcelFileHandler
    from src.tools.excel.excel_writer import convert_format

    csv_path = tmp_path / "quote.csv"
    csv_path.write_text("项目,金额\n门票,120\n住宿,300\n", encoding="utf-8")
    monkeypatch.setattr(ExcelFileHandler, "get_session_dir", staticmethod(lambda: tmp_path))

    result = convert_format(str(csv_path), target_format="excel", output_name="quote.xlsx")

    assert result["success"] is True
    assert result["file_name"] == "quote.xlsx"

    wb = openpyxl.load_workbook(result["file_path"])
    ws = wb.active
    assert ws["A1"].value == "项目"
    assert ws["B1"].value == "金额"
    assert ws["A2"].value == "门票"
    assert ws["B2"].value == 120
    wb.close()


def test_excel_save_temp_does_not_overwrite_existing_file(tmp_path, monkeypatch):
    from src.tools.excel.excel_lib import ExcelFileHandler
    from src.tools.excel.excel_writer import create_excel

    monkeypatch.setattr(ExcelFileHandler, "get_session_dir", staticmethod(lambda: tmp_path))

    first = create_excel(
        data="| 项目 | 金额 |\n|------|------|\n| 门票 | 120 |",
        data_type="markdown",
        file_name="quote.xlsx",
    )
    second = create_excel(
        data="| 项目 | 金额 |\n|------|------|\n| 住宿 | 300 |",
        data_type="markdown",
        file_name="quote.xlsx",
    )

    assert first["success"] is True
    assert second["success"] is True
    assert first["file_path"] != second["file_path"]
    assert Path(first["file_path"]).exists()
    assert Path(second["file_path"]).exists()


def test_create_excel_writes_currency_text_as_number(tmp_path, monkeypatch):
    from src.tools.excel.excel_lib import ExcelFileHandler
    from src.tools.excel.excel_writer import create_excel

    monkeypatch.setattr(ExcelFileHandler, "get_session_dir", staticmethod(lambda: tmp_path))

    result = create_excel(
        data="| 项目 | 金额 |\n|------|------|\n| 合计 | ¥1,446.07 |",
        data_type="markdown",
        file_name="amount.xlsx",
    )

    assert result["success"] is True

    wb = openpyxl.load_workbook(result["file_path"])
    ws = wb.active
    assert ws["B2"].value == 1446.07
    wb.close()


# ============================================================
# 规则路由：修改类指令不得被 to_md 劫持（2026-09-15 研学行程事故回归）
# ============================================================


def test_rule_route_modify_intent_falls_through_to_llm():
    """修改类指令（含"内容"/"文本"字样）不命中 to_md 规则，交给 LLM 路由生成 modify"""
    from src.tools.excel.excel_router import ExcelRouter

    router = ExcelRouter()
    modify_contexts = [
        "修改该表格：1）把第一行标题单元格A1的内容从A改为B；2）删除重复行",
        "对表格做文本查找替换：把全文中的X替换为Y",
        "只做一件事：把 A1 单元格的文字改为 X",
        "把表格里的价格更新一下",
    ]
    for ctx in modify_contexts:
        res = router._rule_based_route(ctx, ["a.xlsx"])
        assert res.get("task") == "", f"修改指令被规则路由劫持: {ctx} -> {res}"


def test_rule_route_to_md_still_works_for_read_intent():
    """非修改指令的 to_md 短路保持不变"""
    from src.tools.excel.excel_router import ExcelRouter

    router = ExcelRouter()
    res = router._rule_based_route("读取这个附件的内容", ["a.xlsx"])
    assert res["task"] == "to_md"
