#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel 智能模板填充工具（Phase A）单测

覆盖：
- 行数不匹配三种（M==K / M>K / M<K）：输出无样例残留 + 数据正确
- 样式保留：非数据区字体/填充/数字格式/列宽/合并 位级一致；插入行复制行高+样式
- 合并单元格在行增删后正确重锚定
- analyze_structure（mock LLM）解析结构
- _verify_render 残留违例 fail-loud
- fill_with_sample 端到端 + 异常分支
"""

import json
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill, Alignment

PROJECT_ROOT = Path(__file__).resolve().parents[3]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.tools.excel.excel_template_ai import (
    ColumnBinding,
    FillData,
    SheetStructure,
    analyze_structure,
    fill_with_sample,
    _verify_render,
)


# ============================================================
# 测试样例 + 结构
# ============================================================

def _build_sample(path: Path) -> Path:
    """构造带样式的样例：标题合并 + meta + 表头(粗体填充) + 3 行明细(数字格式/行高) + 合计行"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报价"

    ws.merge_cells("A1:E1")
    ws["A1"] = "报价单标题"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A2"] = "客户："
    ws["B2"] = "（示例客户）"
    ws["D2"] = "日期："
    ws["E2"] = "2026-01-01"

    for i, h in enumerate(["类别", "项目", "单价", "数量", "金额"], 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(fill_type="solid", start_color="4472C4", end_color="4472C4")

    example = [
        ["住宿", "酒店A", 100, 2, 200],
        ["门票", "景点B", 50, 3, 150],
        ["餐饮", "午餐C", 30, 3, 90],
    ]
    for r, row in enumerate(example, start=4):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            if i in (3, 5):
                c.number_format = "#,##0.00"
    ws.row_dimensions[4].height = 30

    ws["A7"] = "合计"
    ws["E7"] = 440
    ws["E7"].number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 15

    wb.save(str(path))
    return path


# 与 _build_sample 对应的结构（mock LLM 返回这个）
SAMPLE_STRUCTURE = {
    "title": {"row": 1, "col": 1, "merge": "A1:E1"},
    "meta_fields": [
        {"row": 2, "col": 1, "bind": "customer_name"},
        {"row": 2, "col": 4, "bind": "date"},
    ],
    "columns": [
        {"col": 1, "header": "类别", "bind": "category"},
        {"col": 2, "header": "项目", "bind": "name"},
        {"col": 3, "header": "单价", "bind": "unit_price"},
        {"col": 4, "header": "数量", "bind": "quantity"},
        {"col": 5, "header": "金额", "bind": "amount"},
    ],
    "detail_first_row": 4,
    "detail_last_row": 6,
    "detail_template_row": 4,
    "totals": [{"row": 7, "col": 5, "bind": "grand_total"}],
}


def _mock_llm(structure_dict=None):
    payload = json.dumps(structure_dict or SAMPLE_STRUCTURE, ensure_ascii=False)
    return lambda prompt: payload


def _data(rows, meta=None, totals=None):
    return FillData(
        rows=rows,
        meta=meta or {"customer_name": "张三", "date": "2026-08-01"},
        totals=totals or {"grand_total": 9999},
    )


# ============================================================
# 行数不匹配 + 残留 + 样式
# ============================================================


@pytest.fixture
def sample_path(tmp_path):
    return _build_sample(tmp_path / "sample.xlsx")


def test_fill_m_equals_k_no_residue(sample_path, tmp_path):
    """M==K：原位替换，无残留，样式/合并/列宽保留"""
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "amount": 400},
        {"category": "门票", "name": "景点Y", "unit_price": 80, "quantity": 3, "amount": 240},
        {"category": "餐饮", "name": "午餐Z", "unit_price": 40, "quantity": 3, "amount": 120},
    ]
    res = fill_with_sample(
        str(sample_path), _data(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"], res
    assert res["rows_rendered"] == 3

    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # 数据正确（无残留）
    assert ws["A4"].value == "住宿" and ws["B4"].value == "酒店X" and ws["E4"].value == 400
    assert ws["A6"].value == "餐饮" and ws["E6"].value == 120
    # 合计填入
    assert ws["E7"].value == 9999
    # meta 填入
    assert ws["B2"].value == "张三" and ws["E2"].value == "2026-08-01"
    # 样式保留：表头粗体 + 填充
    assert ws["A3"].font.bold is True
    assert ws["A3"].fill.start_color.rgb in ("004472C4", "4472C4")
    # 列宽保留
    assert ws.column_dimensions["A"].width == 15
    # 标题合并保留
    assert any(mr.min_row == 1 and mr.max_row == 1 and mr.min_col == 1 and mr.max_col == 5
               for mr in ws.merged_cells.ranges)
    # 数字格式保留
    assert ws["E4"].number_format == "#,##0.00"


def test_meta_left_right_label_not_overwritten(sample_path, tmp_path):
    """左右结构：标题格保留，值写到标题右侧相邻格（修 tr_643d42f978664bb9：值覆盖了标题）"""
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "amount": 400},
        {"category": "门票", "name": "景点Y", "unit_price": 80, "quantity": 3, "amount": 240},
        {"category": "餐饮", "name": "午餐Z", "unit_price": 40, "quantity": 3, "amount": 120},
    ]
    res = fill_with_sample(
        str(sample_path), _data(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"]
    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # 标题格 A2/D2 必须保留，绝不能被值覆盖
    assert ws["A2"].value == "客户：", f"标题被覆盖: A2={ws['A2'].value!r}"
    assert ws["D2"].value == "日期：", f"标题被覆盖: D2={ws['D2'].value!r}"
    # 值在标题右侧相邻格 B2/E2
    assert ws["B2"].value == "张三"
    assert ws["E2"].value == "2026-08-01"


def test_meta_value_col_override(sample_path, tmp_path):
    """value_col 显式指定值格（值不在标题紧邻右侧时）"""
    structure = dict(SAMPLE_STRUCTURE)
    structure = json.loads(json.dumps(SAMPLE_STRUCTURE))  # 深拷贝
    structure["meta_fields"] = [{"row": 2, "col": 1, "bind": "customer_name", "value_col": 3}]
    rows = [{"category": "A", "name": "n", "unit_price": 1, "quantity": 1, "amount": 1},
            {"category": "B", "name": "n", "unit_price": 1, "quantity": 1, "amount": 1},
            {"category": "C", "name": "n", "unit_price": 1, "quantity": 1, "amount": 1}]
    res = fill_with_sample(
        str(sample_path), _data(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(structure),
    )
    assert res["success"]
    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    assert ws["A2"].value == "客户："   # 标题保留
    assert ws["C2"].value == "张三"      # 值在 value_col=3（而非默认 col+1=B2）


def test_meta_merged_label_value_after_merge(tmp_path):
    """标题本身是合并单元格时：值写到合并区右侧，标题保留，不报 MergedCell 只读（用户实测场景）"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A2:B2")
    ws["A2"] = "合同编号："  # 标题合并 A2:B2
    for i, h in enumerate(["类别", "金额"], 1):
        ws.cell(row=3, column=i, value=h)
    ws.cell(row=4, column=1, value="示例X"); ws.cell(row=4, column=2, value=100)
    ws.cell(row=5, column=1, value="示例Y"); ws.cell(row=5, column=2, value=100)
    p = tmp_path / "merged_label.xlsx"
    wb.save(str(p))

    structure = {
        "meta_fields": [{"row": 2, "col": 1, "bind": "contract"}],
        "columns": [{"col": 1, "bind": "category"}, {"col": 2, "bind": "amount"}],
        "detail_first_row": 4, "detail_last_row": 5, "detail_template_row": 4,
    }
    data = FillData(rows=[{"category": "A", "amount": 1}, {"category": "B", "amount": 2}],
                    meta={"contract": "2025SG001"})
    res = fill_with_sample(str(p), data, output_dir=str(tmp_path), llm_callable=_mock_llm(structure))
    assert res["success"], res

    wb2 = openpyxl.load_workbook(res["file_path"])
    ws2 = wb2.active
    assert ws2["A2"].value == "合同编号：", f"合并标题被覆盖: A2={ws2['A2'].value!r}"
    assert ws2["C2"].value == "2025SG001", f"值应在合并区右侧 C2: C2={ws2['C2'].value!r}"
    # 明细正常
    assert ws2["A4"].value == "A" and ws2["B5"].value == 2


def test_no_fill_horizontal_merge_preserved(tmp_path):
    """无内容填充的横向合并不被解除（用户实测：没填内容的合并单元格被解除合并了）"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:E1"); ws["A1"] = "标题"           # 标题合并（不填）
    ws.merge_cells("A2:E2"); ws["A2"] = "副标题：说明文字"  # meta 区合并（不填，无 bind）
    for i, h in enumerate(["类别", "金额"], 1):
        ws.cell(row=3, column=i, value=h)
    ws.cell(row=4, column=1, value="示例"); ws.cell(row=4, column=2, value=100)
    ws.cell(row=5, column=1, value="示例"); ws.cell(row=5, column=2, value=100)
    p = tmp_path / "nofill_merge.xlsx"
    wb.save(str(p))

    structure = {
        "columns": [{"col": 1, "bind": "category"}, {"col": 2, "bind": "amount"}],
        "detail_first_row": 4, "detail_last_row": 5, "detail_template_row": 4,
        # 无 meta_fields / title / totals —— A1:E1 和 A2:E2 都不填
    }
    data = FillData(rows=[{"category": "A", "amount": 1}, {"category": "B", "amount": 2}])
    res = fill_with_sample(str(p), data, output_dir=str(tmp_path), llm_callable=_mock_llm(structure))
    assert res["success"], res

    wb2 = openpyxl.load_workbook(res["file_path"])
    ws2 = wb2.active
    merges = {(mr.min_row, mr.max_row, mr.min_col, mr.max_col) for mr in ws2.merged_cells.ranges}
    assert (1, 1, 1, 5) in merges, f"标题合并被解除: {merges}"
    assert (2, 2, 1, 5) in merges, f"副标题合并被解除: {merges}"


def test_fill_m_greater_than_k_inserts_rows_with_style(sample_path, tmp_path):
    """M>K：插入行，续填，行高+样式复制，合计下移，合并保留"""
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "amount": 400},
        {"category": "门票", "name": "景点Y", "unit_price": 80, "quantity": 3, "amount": 240},
        {"category": "餐饮", "name": "午餐Z", "unit_price": 40, "quantity": 3, "amount": 120},
        {"category": "用车", "name": "大巴W", "unit_price": 500, "quantity": 1, "amount": 500},
        {"category": "导游", "name": "导游V", "unit_price": 300, "quantity": 1, "amount": 300},
    ]
    res = fill_with_sample(
        str(sample_path), _data(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"] and res["rows_rendered"] == 5

    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # 明细现在占 4-8 行，全部为 data
    assert ws["A4"].value == "住宿" and ws["A8"].value == "导游"
    assert ws["E8"].value == 300
    # 插入行（7,8）继承模板行 4 的数字格式 + 行高
    assert ws["E7"].number_format == "#,##0.00"
    assert ws["E8"].number_format == "#,##0.00"
    assert ws.row_dimensions[7].height == 30
    assert ws.row_dimensions[8].height == 30
    # 合计行原 7 → 9（delta=+2）
    assert ws["A9"].value == "合计"
    assert ws["E9"].value == 9999
    # 标题合并保留
    assert any(mr.min_row == 1 and mr.max_col == 5 for mr in ws.merged_cells.ranges)
    # 表头样式保留
    assert ws["A3"].font.bold is True


def test_fill_m_less_than_k_deletes_rows_no_residue(sample_path, tmp_path):
    """M<K：删除多余样例行，合计上移，无样例残留"""
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "amount": 400},
        {"category": "门票", "name": "景点Y", "unit_price": 80, "quantity": 3, "amount": 240},
    ]
    res = fill_with_sample(
        str(sample_path), _data(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"] and res["rows_rendered"] == 2

    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # 明细只剩 4-5
    assert ws["A4"].value == "住宿" and ws["A5"].value == "门票"
    # 原第 3 条样例"餐饮/午餐C/90"不得残留：合计行上移到 6
    assert ws["A6"].value == "合计"
    assert ws["E6"].value == 9999
    # 全表扫描：不应出现样例示例值"午餐C"或 90（作为餐饮行残留）
    for row in ws.iter_rows(values_only=True):
        assert "午餐C" not in [v for v in row if v is not None]
    # 90 已被替换/删除：E 列不再有 90
    e_values = [ws.cell(row=r, column=5).value for r in range(1, ws.max_row + 1)]
    assert 90 not in e_values


def test_inserted_row_inherits_border_fill(sample_path, tmp_path):
    """插入行继承模板行的填充与边框（样式完整复制）"""
    # 给模板行 4 加边框 + 填充，再验证插入行继承
    wb = openpyxl.load_workbook(str(sample_path))
    ws = wb.active
    from openpyxl.styles import Border, Side
    bd = Border(left=Side(style="thin"), right=Side(style="thin"))
    for col in range(1, 6):
        c = ws.cell(row=4, column=col)
        c.border = bd
        c.fill = PatternFill(fill_type="solid", start_color="FFF2CC", end_color="FFF2CC")
    wb.save(str(sample_path))

    rows = [
        {"category": "A", "name": "n1", "unit_price": 1, "quantity": 1, "amount": 1},
        {"category": "B", "name": "n2", "unit_price": 1, "quantity": 1, "amount": 1},
        {"category": "C", "name": "n3", "unit_price": 1, "quantity": 1, "amount": 1},
        {"category": "D", "name": "n4", "unit_price": 1, "quantity": 1, "amount": 1},
    ]
    res = fill_with_sample(
        str(sample_path), _data(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"]
    wb2 = openpyxl.load_workbook(res["file_path"])
    ws2 = wb2.active
    # 插入行 7 应继承行 4 的填充
    assert ws2.cell(row=7, column=1).fill.start_color.rgb in ("00FFF2CC", "FFF2CC")
    assert ws2.cell(row=7, column=1).border.left.style == "thin"


# ============================================================
# analyze_structure（mock LLM）
# ============================================================


def test_analyze_structure_parses_mock_llm(sample_path):
    data = _data([{"category": "x"}])
    s = analyze_structure(str(sample_path), data, llm_callable=_mock_llm())
    assert isinstance(s, SheetStructure)
    assert s.detail_first_row == 4 and s.detail_last_row == 6
    assert len(s.columns) == 5
    assert s.columns[0].bind == "category"
    assert len(s.totals) == 1 and s.totals[0].bind == "grand_total"
    assert s.title == {"row": 1, "col": 1, "merge": "A1:E1"}


def test_analyze_structure_rejects_missing_detail(sample_path):
    """缺明细行区时 fail-loud"""
    bad = {"columns": [{"col": 1, "bind": "x"}]}  # 无 detail_first_row
    data = _data([{"x": 1}])
    with pytest.raises(ValueError):
        analyze_structure(str(sample_path), data, llm_callable=_mock_llm(bad))


def test_analyze_structure_rejects_bad_json(sample_path):
    data = _data([{"x": 1}])
    with pytest.raises(ValueError):
        analyze_structure(str(sample_path), data, llm_callable=lambda p: "不是JSON")


def test_analyze_structure_rejects_bad_group_rows(sample_path):
    """分组 detail_first_row > detail_last_row 时 fail-loud（避免 K<=0 插入超额行）"""
    bad = {
        "columns": [{"col": 1, "bind": "x"}],
        "groups": [{"name": "g", "detail_first_row": 8, "detail_last_row": 5}],
    }
    data = _data([{"x": 1}])
    with pytest.raises(ValueError):
        analyze_structure(str(sample_path), data, llm_callable=_mock_llm(bad))


def test_analyze_tolerates_null_fields_from_llm(sample_path):
    """LLM 对无表头列返回 header:null / 未绑定列 bind:null 不应崩溃（生产 trace tr_f93f6a0e 根因）"""
    structure_with_nulls = {
        "columns": [
            {"col": 1, "header": None, "bind": "category"},
            {"col": 2, "header": "项目", "bind": "name"},
            {"col": 5, "header": "金额", "bind": None},
        ],
        "detail_first_row": 4,
        "detail_last_row": 6,
        "totals": [{"row": 7, "col": 5, "bind": None}],
        "meta_fields": [{"row": 2, "col": 2, "label": None, "bind": "customer_name"}],
    }
    data = _data([{"category": "x"}])
    s = analyze_structure(str(sample_path), data, llm_callable=_mock_llm(structure_with_nulls))
    assert len(s.columns) == 3  # 三列都保留，None 被容忍
    assert s.columns[2].bind is None
    assert len(s.totals) == 1


def test_analyze_skips_malformed_records(sample_path):
    """缺必填/类型错的记录被跳过，不让单条坏数据搞垮整个分析"""
    bad = {
        "columns": [
            {"col": 1, "bind": "x"},
            {"header": "无col"},
            "not_a_dict",
            {"col": 2, "bind": "y"},
        ],
        "detail_first_row": 4,
        "detail_last_row": 6,
    }
    data = _data([{"x": 1}])
    s = analyze_structure(str(sample_path), data, llm_callable=_mock_llm(bad))
    assert len(s.columns) == 2  # 只剩 col=1 和 col=2 两条有效


# ============================================================
# 不变式校验 fail-loud
# ============================================================


def test_verify_render_raises_on_residue(tmp_path):
    """明细单元格残留样例值时 _verify_render 报错"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A4"] = "残留值"  # 与 data 不符
    structure = SheetStructure(
        columns=[ColumnBinding(col=1, bind="x")],
        detail_first_row=4, detail_last_row=4, detail_template_row=4,
    )
    data = FillData(rows=[{"x": "正确值"}])
    with pytest.raises(AssertionError):
        _verify_render(ws, structure, data)


# ============================================================
# 异常分支
# ============================================================


def test_fill_missing_sample_returns_error(tmp_path):
    res = fill_with_sample(str(tmp_path / "nope.xlsx"), _data([{"x": 1}]),
                           llm_callable=_mock_llm())
    assert not res["success"]
    assert "不存在" in res["error"]


def test_fill_empty_rows_returns_error(sample_path, tmp_path):
    res = fill_with_sample(str(sample_path), _data([]),
                           output_dir=str(tmp_path), llm_callable=_mock_llm())
    assert not res["success"]
    assert "rows" in res["error"]


# ============================================================
# 非标量值前置拦截（线上事故回归：嵌套 dict 写单元格 → "Cannot convert ... to Excel"）
# ============================================================


def _llm_must_not_be_called(prompt):
    raise AssertionError("校验失败不应触发 LLM 结构分析调用")


def test_nested_totals_dict_rejected_before_llm(sample_path):
    """交叉表合计按列分档传嵌套 dict → 前置拦截并给出拆平指引，不触发 LLM、不崩 openpyxl"""
    data = {
        "rows": [{"费用项目": "研学课程服务费", "40人（元）": 7120}],
        "totals": {"合计总价": {"40人（元）": 9520, "45人（元）": 10560, "49人（元）": 11392}},
    }
    res = fill_with_sample(str(sample_path), data, llm_callable=_llm_must_not_be_called)
    assert not res["success"]
    assert "data.totals.合计总价" in res["error"]
    assert "标量" in res["error"]
    assert "拆平" in res["error"]
    assert "Cannot convert" not in res["error"]


def test_nested_row_value_rejected(sample_path):
    data = {
        "rows": [{"费用项目": {"课程": 1}, "40人（元）": 7120}],
        "totals": {},
    }
    res = fill_with_sample(str(sample_path), data, llm_callable=_llm_must_not_be_called)
    assert not res["success"]
    assert "data.rows[1].费用项目" in res["error"]
    assert "标量" in res["error"]


def test_nested_meta_and_per_capita_rejected(sample_path):
    data = {
        "rows": [{"a": 1}],
        "meta": {"报价日期": ["2026-08-18"]},
        "totals": {"per_capita": {"人均费用": {"40人（元）": 238}}},
    }
    res = fill_with_sample(str(sample_path), data, llm_callable=_llm_must_not_be_called)
    assert not res["success"]
    # meta 先于 totals 被检出
    assert "data.meta.报价日期" in res["error"]

    data2 = {
        "rows": [{"a": 1}],
        "meta": {},
        "totals": {"per_capita": {"人均费用": {"40人（元）": 238}}},
    }
    res2 = fill_with_sample(str(sample_path), data2, llm_callable=_llm_must_not_be_called)
    assert not res2["success"]
    assert "data.totals.per_capita.人均费用" in res2["error"]


def test_scalar_data_still_passes_validation(sample_path, tmp_path):
    """标量数据（含 None/bool/str/int/float）不受新校验影响"""
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "amount": 400},
        {"category": None, "name": "含None", "unit_price": 0.5, "quantity": True, "amount": 0},
    ]
    data = FillData(rows=rows, meta={"customer_name": "张三", "date": "2026-08-18"},
                    totals={"grand_total": 9999, "per_capita": {"成人人均": 238}})
    res = fill_with_sample(str(sample_path), data,
                           output_dir=str(tmp_path), llm_callable=_mock_llm())
    assert res["success"], res


# ============================================================
# 显式 columns 优先（脱离 LLM 列推断）
# ============================================================


def test_explicit_columns_override_llm(sample_path, tmp_path):
    """data.columns 显式提供时，覆盖 LLM 推断的列绑定"""
    rows = [{"cat": "住宿", "amt": 400}]
    data = FillData(
        rows=rows, meta={}, totals={},
        columns=[{"col": 1, "bind": "cat"}, {"col": 5, "bind": "amt"}],
    )
    res = fill_with_sample(
        str(sample_path), data,
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"]
    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    assert ws["A4"].value == "住宿"
    assert ws["E4"].value == 400
    # 未显式绑定的列（2,3,4）应被清空，不残留样例"酒店A/100/2"
    assert ws["B4"].value in (None, "")
    assert ws["C4"].value in (None, "")


# ============================================================
# Phase B：多分组 + 小计 + 合计 + 人均
# ============================================================


def _build_grouped_sample(path: Path) -> Path:
    """中科悦行风格：2 个分组（房餐车 / 门票），各带小计行 + 合计行 + 人均行"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报价"

    ws.merge_cells("A1:F1")
    ws["A1"] = "文旅报价表"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A2"] = "客户："
    ws["B2"] = "（示例客户）"
    ws["D2"] = "日期："
    ws["E2"] = "2026-01-01"

    for i, h in enumerate(["类别", "项目", "单价", "数量", "单位", "金额"], 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(fill_type="solid", start_color="4472C4", end_color="4472C4")

    # 分组1：房餐车（示例 2 行）行4-5，小计行6
    g1 = [["住宿", "示例酒店A", 500, 2, "间", 1000],
          ["餐饮", "示例午餐B", 50, 3, "人", 150]]
    for r, row in enumerate(g1, start=4):
        for i, v in enumerate(row, 1):
            ws.cell(row=r, column=i, value=v)
    ws["A6"] = "房餐车小计"
    ws["F6"] = 1150
    ws["F6"].number_format = "#,##0.00"
    ws.row_dimensions[4].height = 28

    # 分组2：门票（示例 2 行）行7-8，小计行9
    g2 = [["门票/项目", "示例景点C", 160, 2, "人", 320],
          ["门票/项目", "示例项目D", 80, 2, "人", 160]]
    for r, row in enumerate(g2, start=7):
        for i, v in enumerate(row, 1):
            ws.cell(row=r, column=i, value=v)
    ws["A9"] = "门票小计"
    ws["F9"] = 480
    ws["F9"].number_format = "#,##0.00"

    # 合计行10 + 人均行11
    ws["A10"] = "合计"
    ws["F10"] = 1630
    ws["F10"].number_format = "#,##0.00"
    ws["A11"] = "成人人均"
    ws["B11"] = 815
    ws["C11"] = "儿童人均"
    ws["D11"] = 407.5

    ws.column_dimensions["A"].width = 14
    wb.save(str(path))
    return path


GROUPED_STRUCTURE = {
    "title": {"row": 1, "col": 1, "merge": "A1:F1"},
    "meta_fields": [
        {"row": 2, "col": 1, "bind": "customer_name"},
        {"row": 2, "col": 4, "bind": "date"},
    ],
    "columns": [
        {"col": 1, "header": "类别", "bind": "category"},
        {"col": 2, "header": "项目", "bind": "name"},
        {"col": 3, "header": "单价", "bind": "unit_price"},
        {"col": 4, "header": "数量", "bind": "quantity"},
        {"col": 5, "header": "单位", "bind": "unit"},
        {"col": 6, "header": "金额", "bind": "amount"},
    ],
    "groups": [
        {"name": "房餐车小计", "detail_first_row": 4, "detail_last_row": 5,
         "detail_template_row": 4, "subtotal_row": 6, "subtotal_col": 6,
         "match": {"category": ["住宿", "餐饮", "用车"]}},
        {"name": "门票小计", "detail_first_row": 7, "detail_last_row": 8,
         "detail_template_row": 7, "subtotal_row": 9, "subtotal_col": 6,
         "match": {"category": ["门票/项目"]}},
    ],
    "totals": [
        {"row": 10, "col": 6, "bind": "grand_total"},
        {"row": 11, "col": 2, "bind": "成人人均"},
        {"row": 11, "col": 4, "bind": "儿童人均"},
    ],
}


@pytest.fixture
def grouped_path(tmp_path):
    return _build_grouped_sample(tmp_path / "grouped.xlsx")


def _gdata(rows, *, group_subtotals=None, totals=None, meta=None):
    if group_subtotals is None:
        group_subtotals = {"房餐车小计": 1000, "门票小计": 200}
    if totals is None:
        totals = {"grand_total": 1200, "per_capita": {"成人人均": 600, "儿童人均": 300}}
    return FillData(
        rows=rows,
        meta=meta or {"customer_name": "张三", "date": "2026-08-01"},
        group_subtotals=group_subtotals,
        totals=totals,
    )


def test_multi_group_mixed_deltas(grouped_path, tmp_path):
    """分组1 M=3(>K=2 插入)、分组2 M=1(<K=2 删除)：小计/合计/人均正确重定位，无残留"""
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "unit": "间", "amount": 400},
        {"category": "餐饮", "name": "午餐Y", "unit_price": 40, "quantity": 3, "unit": "人", "amount": 120},
        {"category": "用车", "name": "大巴Z", "unit_price": 480, "quantity": 1, "unit": "辆", "amount": 480},
        {"category": "门票/项目", "name": "景点W", "unit_price": 200, "quantity": 1, "unit": "团", "amount": 200},
    ]
    res = fill_with_sample(
        str(grouped_path), _gdata(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(GROUPED_STRUCTURE),
    )
    assert res["success"] and res["rows_rendered"] == 4

    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active

    # 分组1：3 行（4-6），末行是用车
    assert ws["A4"].value == "住宿" and ws["A5"].value == "餐饮" and ws["A6"].value == "用车"
    assert ws["F6"].value == 480
    # 分组1 小计在行 7（原 6 + 分组1 delta +1）
    assert ws["A7"].value == "房餐车小计" and ws["F7"].value == 1000
    # 分组2：1 行在行 8（原 7 + 分组1 delta +1）
    assert ws["A8"].value == "门票/项目" and ws["F8"].value == 200
    # 分组2 小计在行 9（原 9，+1-1=0 偏移）
    assert ws["A9"].value == "门票小计" and ws["F9"].value == 200
    # 合计 + 人均
    assert ws["A10"].value == "合计" and ws["F10"].value == 1200
    assert ws["B11"].value == 600 and ws["D11"].value == 300
    # meta
    assert ws["B2"].value == "张三" and ws["E2"].value == "2026-08-01"

    # 无样例残留：示例项目名/示例小计值不得出现
    all_vals = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert not any("示例" in v for v in all_vals), "样例示例数据残留"
    # 旧样例小计值（房餐车 1150）不得残留（门票小计 480 已被 F9==200 覆盖，见上行断言）
    f6_values = [ws.cell(row=r, column=6).value for r in range(1, ws.max_row + 1)]
    assert 1150 not in f6_values

    # 样式保留：标题合并、表头粗体、小计数字格式、插入行继承行高
    assert any(mr.min_row == 1 and mr.max_col == 6 for mr in ws.merged_cells.ranges)
    assert ws["A3"].font.bold is True
    assert ws["F7"].number_format == "#,##0.00"  # 小计行数字格式（行6样式随内容移到行7）
    assert ws.row_dimensions[6].height == 28  # 插入的用车行继承模板行4的行高


def test_group_subtotals_numeric_fallback(grouped_path, tmp_path):
    """不传 group_subtotals 时，按 subtotal_col(amount) 数值求和"""
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "unit": "间", "amount": 400},
        {"category": "餐饮", "name": "午餐Y", "unit_price": 40, "quantity": 3, "unit": "人", "amount": 120},
        {"category": "门票/项目", "name": "景点W", "unit_price": 200, "quantity": 1, "unit": "团", "amount": 200},
    ]
    data = _gdata(rows, group_subtotals={}, totals={"grand_total": 740})
    res = fill_with_sample(
        str(grouped_path), data,
        output_dir=str(tmp_path), llm_callable=_mock_llm(GROUPED_STRUCTURE),
    )
    assert res["success"]
    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # 分组1（住宿+餐饮）小计 = 400+120 = 520；M=2==K=2 无偏移，仍在行6
    assert ws["A6"].value == "房餐车小计" and ws["F6"].value == 520
    # 分组2（门票）M=1<K=2 删1行：明细上移到行7，小计行从9上移到8
    assert ws["A7"].value == "门票/项目"
    assert ws["A8"].value == "门票小计" and ws["F8"].value == 200


def test_column_style_subtotal_misjudged_skipped(grouped_path, tmp_path):
    """列式小计误判为 subtotal_row（小计行=明细行8）→ 防御跳过写入：
    明细数据不被覆盖、渲染不硬报错；合法小计（行6）照常写入"""
    structure = {
        **GROUPED_STRUCTURE,
        "groups": [
            GROUPED_STRUCTURE["groups"][0],
            {**GROUPED_STRUCTURE["groups"][1], "subtotal_row": 8},  # 误判：把明细行8当小计行
        ],
    }
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "unit": "间", "amount": 400},
        {"category": "餐饮", "name": "午餐Y", "unit_price": 40, "quantity": 3, "unit": "人", "amount": 120},
        {"category": "门票/项目", "name": "景点V", "unit_price": 160, "quantity": 2, "unit": "人", "amount": 320},
        {"category": "门票/项目", "name": "景点W", "unit_price": 80, "quantity": 2, "unit": "人", "amount": 160},
    ]
    res = fill_with_sample(
        str(grouped_path), _gdata(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(structure),
    )
    assert res["success"], res
    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # 合法小计（行6，不在任何明细区）正常写入
    assert ws["F6"].value == 1000
    # 误判小计的行8仍是明细数据，未被小计值 480 覆盖
    assert ws["B8"].value == "景点W" and ws["F8"].value == 160


def test_unmatched_rows_go_to_first_group(grouped_path, tmp_path):
    """category 不匹配任何分组的行并入首组（不静默丢数据）"""
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "unit": "间", "amount": 400},
        {"category": "未知类", "name": "神秘项", "unit_price": 10, "quantity": 1, "unit": "个", "amount": 10},
    ]
    res = fill_with_sample(
        str(grouped_path), _gdata(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(GROUPED_STRUCTURE),
    )
    assert res["success"]
    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # 两行都进了分组1（行4-5）
    assert ws["A4"].value == "住宿" and ws["A5"].value == "未知类"


def test_groups_empty_falls_back_to_single_region(sample_path, tmp_path):
    """groups 为空时退回单明细区（Phase A 兼容路径）"""
    rows = [
        {"category": "A", "name": "n1", "unit_price": 1, "quantity": 1, "amount": 1},
        {"category": "B", "name": "n2", "unit_price": 1, "quantity": 1, "amount": 1},
    ]
    res = fill_with_sample(
        str(sample_path), _data(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"] and res["rows_rendered"] == 2


# ============================================================
# 回归：行增删后下方已存在行的行高保留 + meta/totals 无值清空
# ============================================================


def test_shifted_rows_retain_height(sample_path, tmp_path):
    """openpyxl insert_rows 不平移 row_dimensions：下方已存在行（合计）的行高必须显式恢复。

    样例合计行（行7）有自定义行高 25；M>K 插入 2 行后合计移到行9，行高应保留。
    （修复前：row_dimensions[9].height 为 None，丢失样例行高）
    """
    wb = openpyxl.load_workbook(str(sample_path))
    ws = wb.active
    ws.row_dimensions[7].height = 25  # 合计行自定义高度
    wb.save(str(sample_path))

    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "amount": 400},
        {"category": "门票", "name": "景点Y", "unit_price": 80, "quantity": 3, "amount": 240},
        {"category": "餐饮", "name": "午餐Z", "unit_price": 40, "quantity": 3, "amount": 120},
        {"category": "用车", "name": "大巴W", "unit_price": 500, "quantity": 1, "amount": 500},
        {"category": "导游", "name": "导游V", "unit_price": 300, "quantity": 1, "amount": 300},
    ]
    res = fill_with_sample(
        str(sample_path), _data(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"]
    wb2 = openpyxl.load_workbook(res["file_path"])
    ws2 = wb2.active
    # 合计行原 7 → 9（delta=+2），行高必须保留
    assert ws2["A9"].value == "合计"
    assert ws2.row_dimensions[9].height == 25, "下方已存在行的行高在 insert_rows 后丢失"


def test_meta_totals_cleared_when_not_provided(sample_path, tmp_path):
    """caller 未提供 meta/totals 值时，样例值必须清空（无残留契约）。

    样例 B2=（示例客户）、E7=440；data.meta={}、data.totals={}。
    修复前：B2 残留"（示例客户）"、E7 残留 440。
    """
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "amount": 400},
        {"category": "门票", "name": "景点Y", "unit_price": 80, "quantity": 3, "amount": 240},
        {"category": "餐饮", "name": "午餐Z", "unit_price": 40, "quantity": 3, "amount": 120},
    ]
    data = FillData(rows=rows, meta={}, totals={})
    res = fill_with_sample(
        str(sample_path), data,
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"]
    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # B2 (customer_name meta) 清空，不残留"（示例客户）"
    assert ws["B2"].value in (None, ""), f"meta 残留: B2={ws['B2'].value!r}"
    # E7 (grand_total total) 清空，不残留 440
    assert ws["E7"].value in (None, ""), f"totals 残留: E7={ws['E7'].value!r}"


# ============================================================
# 竖向合并：明细区内"同组值列"按相邻相同值合并居中
# ============================================================


def test_vertical_merge_same_adjacent_values(tmp_path):
    """样例明细区某列有竖向合并（表示该列同组值相同）；填充后相邻相同值应重新合并居中。

    修复前：明细区内竖向合并被 _compute_final_merges 当作"无法平移"直接丢弃，
    填充后每行都写了相同值却没有合并居中（用户实测场景）。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, h in enumerate(["组别", "项目", "金额"], 1):
        ws.cell(row=1, column=i, value=h)
    # 样例明细 2-4：A2:A3 合并（"甲组"跨2行 = 同组指示），A4 单独（"乙组"）
    ws.cell(row=2, column=1, value="甲组")
    ws.cell(row=2, column=2, value="示例1"); ws.cell(row=2, column=3, value=100)
    ws.cell(row=3, column=2, value="示例2"); ws.cell(row=3, column=3, value=100)
    ws.cell(row=4, column=1, value="乙组")
    ws.cell(row=4, column=2, value="示例3"); ws.cell(row=4, column=3, value=200)
    ws.merge_cells("A2:A3")  # 竖向合并——指示该列按相邻相同值合并
    p = tmp_path / "vmerge.xlsx"
    wb.save(str(p))

    structure = {
        "columns": [
            {"col": 1, "bind": "group"},
            {"col": 2, "bind": "name"},
            {"col": 3, "bind": "amount"},
        ],
        "detail_first_row": 2, "detail_last_row": 4, "detail_template_row": 2,
    }
    # 填充 4 行：甲组x2 + 乙组x2（M=4 > K=3，插入1行）
    rows = [
        {"group": "甲组", "name": "项1", "amount": 10},
        {"group": "甲组", "name": "项2", "amount": 20},
        {"group": "乙组", "name": "项3", "amount": 30},
        {"group": "乙组", "name": "项4", "amount": 40},
    ]
    data = FillData(rows=rows, meta={}, totals={})
    res = fill_with_sample(str(p), data, output_dir=str(tmp_path), llm_callable=_mock_llm(structure))
    assert res["success"], res

    wb2 = openpyxl.load_workbook(res["file_path"])
    ws2 = wb2.active
    # 明细现在 2-5（插入1行）；A2:A3 合并(甲组)、A4:A5 合并(乙组)
    # 合并后非锚点格(A3/A5) value=None，值只在锚点(A2/A4)
    assert ws2["A2"].value == "甲组"
    assert ws2["A4"].value == "乙组"
    merges = {(mr.min_row, mr.max_row, mr.min_col, mr.max_col) for mr in ws2.merged_cells.ranges}
    # 甲组相邻相同 → A2:A3 合并；乙组相邻相同 → A4:A5 合并
    assert (2, 3, 1, 1) in merges, f"甲组未竖向合并: {merges}"
    assert (4, 5, 1, 1) in merges, f"乙组未竖向合并: {merges}"
    # 甲/乙值不同 → A3:A4 不应合并
    assert (3, 4, 1, 1) not in merges, f"不同值被误合并: {merges}"
    # 锚点居中
    assert ws2["A2"].alignment.vertical == "center"
    assert ws2["A2"].alignment.horizontal == "center"
    assert ws2["A4"].alignment.vertical == "center"


def test_vertical_merge_not_triggered_without_sample_merge(sample_path, tmp_path):
    """样例明细区该列没有竖向合并时，即使填充后相邻值相同也不合并（避免误合并不该合并的列）。"""
    # sample_path 的明细区 A 列（类别）无竖向合并
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "amount": 400},
        {"category": "住宿", "name": "酒店Y", "unit_price": 150, "quantity": 1, "amount": 150},
        {"category": "住宿", "name": "酒店Z", "unit_price": 100, "quantity": 1, "amount": 100},
    ]
    res = fill_with_sample(
        str(sample_path), _data(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"]
    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # A4:A6 都是"住宿"（相邻相同），但样例 A 列无竖向合并指示 → 不应被合并
    merges = {(mr.min_row, mr.max_row, mr.min_col, mr.max_col) for mr in ws.merged_cells.ranges}
    assert not any(m[2] == 1 and m[3] == 1 and m[0] >= 4 and m[1] <= 6 for m in merges), \
        f"无竖向合并指示的列被误合并: {merges}"


def test_vertical_merge_within_each_group_not_cross(tmp_path):
    """多分组：每组明细内相邻相同值竖向合并；小计行（明细区外）不被并入合并。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, h in enumerate(["组别", "项目", "金额"], 1):
        ws.cell(row=1, column=i, value=h)
    # 分组1 明细 2-3（A2:A3 合并"甲"），小计行4
    ws.cell(row=2, column=1, value="甲")
    ws.cell(row=2, column=2, value="s1"); ws.cell(row=2, column=3, value=10)
    ws.cell(row=3, column=2, value="s2"); ws.cell(row=3, column=3, value=20)
    ws.merge_cells("A2:A3")
    ws.cell(row=4, column=1, value="甲小计"); ws.cell(row=4, column=3, value=30)
    # 分组2 明细 5-6（A5:A6 合并"乙"），小计行7
    ws.cell(row=5, column=1, value="乙")
    ws.cell(row=5, column=2, value="s3"); ws.cell(row=5, column=3, value=40)
    ws.cell(row=6, column=2, value="s4"); ws.cell(row=6, column=3, value=50)
    ws.merge_cells("A5:A6")
    ws.cell(row=7, column=1, value="乙小计"); ws.cell(row=7, column=3, value=90)
    p = tmp_path / "vmerge_grouped.xlsx"
    wb.save(str(p))

    structure = {
        "columns": [
            {"col": 1, "bind": "group"},
            {"col": 2, "bind": "name"},
            {"col": 3, "bind": "amount"},
        ],
        "groups": [
            {"name": "甲小计", "detail_first_row": 2, "detail_last_row": 3,
             "detail_template_row": 2, "subtotal_row": 4, "subtotal_col": 3,
             "match": {"group": ["甲"]}},
            {"name": "乙小计", "detail_first_row": 5, "detail_last_row": 6,
             "detail_template_row": 5, "subtotal_row": 7, "subtotal_col": 3,
             "match": {"group": ["乙"]}},
        ],
    }
    rows = [
        {"group": "甲", "name": "n1", "amount": 1},
        {"group": "甲", "name": "n2", "amount": 2},
        {"group": "乙", "name": "n3", "amount": 3},
        {"group": "乙", "name": "n4", "amount": 4},
    ]
    data = FillData(rows=rows, meta={}, totals={}, group_subtotals={"甲小计": 3, "乙小计": 7})
    res = fill_with_sample(str(p), data, output_dir=str(tmp_path), llm_callable=_mock_llm(structure))
    assert res["success"], res

    wb2 = openpyxl.load_workbook(res["file_path"])
    ws2 = wb2.active
    # M==K 各2行无偏移：分组1明细2-3(A2:A3合并)，小计4；分组2明细5-6(A5:A6合并)，小计7
    # 合并后非锚点格(A3/A6) value=None，值只在锚点(A2/A5)
    assert ws2["A2"].value == "甲"
    assert ws2["A5"].value == "乙"
    merges = {(mr.min_row, mr.max_row, mr.min_col, mr.max_col) for mr in ws2.merged_cells.ranges}
    # 每组内合并
    assert (2, 3, 1, 1) in merges, f"甲组内未合并: {merges}"
    assert (5, 6, 1, 1) in merges, f"乙组内未合并: {merges}"
    # 小计行（4、7）的 A 列不被并入任何竖向合并（合并不跨明细区/小计行）
    assert not any(m[2] == 1 and m[3] == 1 and m[0] <= 4 <= m[1] for m in merges), \
        f"小计行被并入合并: {merges}"
    assert not any(m[2] == 1 and m[3] == 1 and m[0] <= 7 <= m[1] for m in merges), \
        f"小计行被并入合并: {merges}"


# ============================================================
# 真实样例回归：大团分项报价（用户实测，钉死版式合并不变式）
# ============================================================

# 真实样例文件（含竖向类别合并 A4:A7/A10:A17/A18:A21 + 每行备注横向合并 G:J + 标题/meta/小计合并）
_REAL_TEMPLATE = PROJECT_ROOT / "tests" / "fixtures" / "excel" / "大团分项报价_调整后格式.xlsx"


def _real_quote_rows():
    """贵州天眼 4 天 3 晚行程成本明细（用户实测数据）：1 用车 + 12 门票/项目 + 1 住宿 = 14 行"""
    rows = [
        {"category": "用车", "name": "7座商务车", "unit_price": 1039.38, "quantity": 1, "unit": "辆", "amount": 519.69, "remark": "7座商务车、按公里计费"},
        {"category": "门票/项目", "name": "中国天眼科普基地(成人票)", "unit_price": 140.00, "quantity": 2, "unit": "人", "amount": 140.00, "remark": "挂牌价"},
        {"category": "门票/项目", "name": "南仁东先进事迹馆", "unit_price": 120.00, "quantity": 1, "unit": "团", "amount": 60.00, "remark": ""},
        {"category": "门票/项目", "name": "天文体验馆参观", "unit_price": 120.00, "quantity": 1, "unit": "团", "amount": 60.00, "remark": ""},
        {"category": "门票/项目", "name": "FAST观测体验", "unit_price": 30.00, "quantity": 2, "unit": "人", "amount": 30.00, "remark": ""},
        {"category": "门票/项目", "name": "天眼瞭望台直通车", "unit_price": 40.00, "quantity": 2, "unit": "人", "amount": 40.00, "remark": ""},
        {"category": "门票/项目", "name": "天象影院", "unit_price": 40.00, "quantity": 2, "unit": "人", "amount": 40.00, "remark": ""},
        {"category": "门票/项目", "name": "夜游望远镜观星", "unit_price": 50.00, "quantity": 2, "unit": "人", "amount": 50.00, "remark": ""},
        {"category": "门票/项目", "name": "桥梁科普馆", "unit_price": 300.00, "quantity": 1, "unit": "团", "amount": 150.00, "remark": ""},
        {"category": "门票/项目", "name": "基础探洞体验", "unit_price": 588.00, "quantity": 2, "unit": "人", "amount": 588.00, "remark": ""},
        {"category": "门票/项目", "name": "以上产品打包价", "unit_price": 168.00, "quantity": 2, "unit": "人", "amount": 168.00, "remark": ""},
        {"category": "门票/项目", "name": "制陶技艺拉坯", "unit_price": 55.00, "quantity": 2, "unit": "人", "amount": 55.00, "remark": ""},
        {"category": "门票/项目", "name": "制陶技艺捏塑", "unit_price": 20.00, "quantity": 2, "unit": "人", "amount": 20.00, "remark": ""},
        {"category": "住宿", "name": "平塘丰业大酒店（商务房型）", "unit_price": 338.00, "quantity": 1, "unit": "间", "amount": 507.00, "remark": "平塘3晚"},
    ]
    return rows


def test_real_quote_template_preserves_full_layout(tmp_path):
    """真实样例「大团分项报价」+ 14 行实测数据：输出版式必须与样例一致。

    用户实测回归（ce3d6b8 后格式全崩）：备注列 G:J 横向合并全丢（每行备注只占 G 列、
    长文本被截断）、成本类别 A 列竖向合并错乱。根因是明细区内的横向合并被当作"无法平移"
    一并丢弃。本测试钉死：每个明细行的横向合并保留 + 竖向类别合并按数据重建 + 标题/meta/
    小计合并保留并随删行上移。
    """
    if not _REAL_TEMPLATE.exists():
        pytest.skip(f"真实样例缺失: {_REAL_TEMPLATE}")

    # 单明细区 rows 4-21（样例成本明细），合计行 25；14 行 < 18 行 → 删 4 行，合计上移到 21
    structure = {
        "title": {"row": 1, "col": 1, "merge": "A1:J1"},
        "meta_fields": [
            {"row": 2, "col": 1, "bind": "project"},
            {"row": 2, "col": 4, "bind": "date"},
            {"row": 2, "col": 7, "bind": "people"},
        ],
        "columns": [
            {"col": 1, "bind": "category"},
            {"col": 2, "bind": "name"},
            {"col": 3, "bind": "unit_price"},
            {"col": 4, "bind": "quantity"},
            {"col": 5, "bind": "unit"},
            {"col": 6, "bind": "amount"},
            {"col": 7, "bind": "remark"},
        ],
        "detail_first_row": 4, "detail_last_row": 21, "detail_template_row": 4,
        "totals": [{"row": 25, "col": 6, "bind": "grand_total"}],
    }
    data = FillData(
        rows=_real_quote_rows(),
        meta={"project": "贵州天眼4天3晚行程", "date": "2026-07-18", "people": "2"},
        totals={"grand_total": 2427.69},
    )
    res = fill_with_sample(
        str(_REAL_TEMPLATE), data,
        output_dir=str(tmp_path), llm_callable=_mock_llm(structure),
    )
    assert res["success"], res

    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    merges = {(mr.min_row, mr.min_col, mr.max_row, mr.max_col) for mr in ws.merged_cells.ranges}

    # 1. 标题 / meta / 表头合并保留（均在明细区以上，不应被动）
    assert (1, 1, 1, 10) in merges, f"标题 A1:J1 丢失: {merges}"
    assert (2, 2, 2, 3) in merges, f"meta B2:C2 丢失: {merges}"
    assert (2, 5, 2, 6) in merges, f"meta E2:F2 丢失: {merges}"
    assert (2, 8, 2, 10) in merges, f"meta H2:J2 丢失: {merges}"
    assert (3, 7, 3, 10) in merges, f"表头 G3:J3 丢失: {merges}"

    # 2. 【核心回归点】备注列 G:J 横向合并：每个最终明细行（4-17）都保留
    #    修复前这些全被丢弃，导致备注只占 G 列、版式崩掉
    missing_gj = [r for r in range(4, 18) if (r, 7, r, 10) not in merges]
    assert not missing_gj, f"备注横向合并 G:J 丢失的行: {missing_gj}"

    # 3. 成本类别 A 列竖向合并：12 行"门票/项目"按相邻相同值合并居中（A5:A16）
    assert (5, 1, 16, 1) in merges, f"类别竖向合并 A5:A16 丢失: {merges}"
    # 用车(row4)、住宿(row17)各只 1 行，不被并入合并
    assert not any(m[1] == 1 and m[3] == 1 and m[0] <= 4 <= m[2] and (m[0], m[2]) != (4, 4)
                   for m in merges), f"用车行被误并入竖向合并: {merges}"
    assert not any(m[1] == 1 and m[3] == 1 and m[0] <= 17 <= m[2] and (m[0], m[2]) != (17, 17)
                   for m in merges), f"住宿行被误并入竖向合并: {merges}"

    # 4. 小计/合计行 A:E 合并保留并随删行上移（原 22-25 → 18-21）
    for r in (18, 19, 20, 21):
        assert (r, 1, r, 5) in merges, f"小计/合计 A{r}:E{r} 丢失: {merges}"
    assert (18, 9, 18, 10) in merges, f"I18:J18 丢失: {merges}"
    assert (20, 9, 20, 10) in merges, f"I20:J20 丢失: {merges}"

    # 5. 值正确 + 无样例残留
    assert ws["A4"].value == "用车"
    assert ws["B4"].value == "7座商务车"
    assert ws["A5"].value == "门票/项目"   # 竖向合并锚点
    assert ws["A17"].value == "住宿"
    assert ws["F21"].value == 2427.69      # 合计上移到 21
    # meta：标题格保留，值落右侧
    assert ws["A2"].value == "项目名称"
    assert ws["B2"].value == "贵州天眼4天3晚行程"
    assert ws["E2"].value == "2026-07-18"
    assert ws["H2"].value == "2"

    # 6. 成本类别 A 列整列样式归一：用车(A4)/门票·项目(A5)/住宿(A17) 都应继承类别锚点样式
    #    （bold），修复前只有第一个值（落在样例锚点行上）样式对，其余落在续行格上不粗体
    assert ws["A4"].font.bold is True, "A4 类别样式丢失"
    assert ws["A5"].font.bold is True, f"A5(门票/项目) 应归一到类别锚点粗体: bold={ws['A5'].font.bold}"
    assert ws["A17"].font.bold is True, f"A17(住宿) 应归一到类别锚点粗体: bold={ws['A17'].font.bold}"


# ============================================================
# 数据 ↔ 模板 覆盖：不可映射不填 / 未提供清空 / 提示用户
# ============================================================


def test_vmerge_column_style_normalized(tmp_path):
    """竖向合并列（类别列）整列明细格归一到锚点样式。

    样例里竖向合并的锚点格（每组首行）是粗体/居中，续行格是默认样式（合并后不可见）。
    M<K 删行后，数据值可能落到样例续行格上 → 继承默认样式，同列类别值有的粗体有的不粗体。
    归一后整列统一成锚点样式。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="类别")
    ws.cell(row=1, column=2, value="金额")
    # 样例明细 2-5：A2:A3="甲"（A2 锚点 bold/居中，A3 续行默认），A4:A5="乙"（同理）
    a2 = ws.cell(row=2, column=1, value="甲")
    a2.font = Font(bold=True, size=14)
    a2.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=2, column=2, value=10)
    ws.cell(row=3, column=2, value=20)  # A3 续行：默认样式
    a4 = ws.cell(row=4, column=1, value="乙")
    a4.font = Font(bold=True, size=14)
    a4.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=4, column=2, value=30)
    ws.cell(row=5, column=2, value=40)  # A5 续行：默认样式
    ws.merge_cells("A2:A3")
    ws.merge_cells("A4:A5")
    p = tmp_path / "vmerge_style.xlsx"
    wb.save(str(p))

    structure = {
        "columns": [{"col": 1, "bind": "category"}, {"col": 2, "bind": "amount"}],
        "detail_first_row": 2, "detail_last_row": 5, "detail_template_row": 2,
    }
    # 填 2 行（M=2 < K=4，删 2 行）：甲、乙 → A2=甲、A3=乙
    rows = [{"category": "甲", "amount": 1}, {"category": "乙", "amount": 2}]
    res = fill_with_sample(str(p), FillData(rows=rows), output_dir=str(tmp_path), llm_callable=_mock_llm(structure))
    assert res["success"], res

    wb2 = openpyxl.load_workbook(res["file_path"])
    ws2 = wb2.active
    # A3(乙) 落在样例续行格 A3 上；修复前继承默认样式（非粗体），归一后应为粗体
    assert ws2["A2"].font.bold is True
    assert ws2["A3"].font.bold is True, f"A3(乙) 应归一到类别锚点粗体样式: bold={ws2['A3'].font.bold}"


def test_unused_data_keys_ignored_and_reported(sample_path, tmp_path):
    """用户给的数据键若模板没对应列：不硬塞到别的列，并在 unused_data_keys 里报告。"""
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "amount": 400,
         "次数": 1, "随队老师": "张老师"},  # 次数/随队老师 模板无对应列
        {"category": "门票", "name": "景点Y", "unit_price": 80, "quantity": 3, "amount": 240},
        {"category": "餐饮", "name": "午餐Z", "unit_price": 40, "quantity": 3, "amount": 120},
    ]
    res = fill_with_sample(
        str(sample_path), _data(rows),
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"], res

    # 不可映射的键被报告
    assert "次数" in res["unused_data_keys"], res["unused_data_keys"]
    assert "随队老师" in res["unused_data_keys"], res["unused_data_keys"]
    # 没有硬塞：B4（项目列 = name）填的是"酒店X"，不是"次数"或"随队老师"的值
    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    assert ws["B4"].value == "酒店X"


def test_missing_fields_cleared_and_reported(sample_path, tmp_path):
    """模板字段绑定了但用户没给值：清空（不残留样例示例值），并在 missing_fields 里报告。"""
    rows = [
        {"category": "住宿", "name": "酒店X", "unit_price": 200, "quantity": 2, "amount": 400},
        {"category": "门票", "name": "景点Y", "unit_price": 80, "quantity": 3, "amount": 240},
        {"category": "餐饮", "name": "午餐Z", "unit_price": 40, "quantity": 3, "amount": 120},
    ]
    # 只给 customer_name，缺 date；不给 totals（缺 grand_total）
    data = FillData(rows=rows, meta={"customer_name": "张三"}, totals={})
    res = fill_with_sample(
        str(sample_path), data,
        output_dir=str(tmp_path), llm_callable=_mock_llm(),
    )
    assert res["success"], res

    # 缺失字段被报告
    binds = {m["bind"] for m in res["missing_fields"]}
    assert "date" in binds, res["missing_fields"]
    assert "grand_total" in binds, res["missing_fields"]

    wb = openpyxl.load_workbook(res["file_path"])
    ws = wb.active
    # 未提供的字段被清空，不残留样例的示例值（E2 原为 2026-01-01，E7 原为 440）
    assert ws["E2"].value in (None, ""), f"未提供的 date 应清空: E2={ws['E2'].value!r}"
    assert ws["E7"].value in (None, ""), f"未提供的 grand_total 应清空: E7={ws['E7'].value!r}"
    # message 里应提示缺失
    assert "未提供" in res["message"] or "清空" in res["message"], res["message"]
