"""
render_llm_view（Excel ETL M1 渲染层）单测

覆盖 D5 渲染契约与 gap-analysis §2 的逐脏点（D20）：
- 值渲染：Excel 序列日期、datetime（有无时间）、百分比格式、字符串照抄、超长截断
- 脏结构原样保留：两行表头、分段重复表头、标题/说明/T0 前置行、备注页脚
- 结构处理：全空行跳过、空列保留（列对齐）、合并单元格非锚点空串、重复列名、换行表头
- 黄金快照：5 个夹具全部 sheet 与 expected_render.json 逐 sheet 全等
- 异常路径：文件不存在 / 非 xlsx
"""

import datetime
import json
from pathlib import Path

import openpyxl
import pytest

from src.tools.excel.excel_reader import render_llm_view

pytestmark = [pytest.mark.tools]

FIXTURE_DIR = Path(__file__).resolve().parents[2] / "fixtures" / "excel_etl"

# 夹具文件名（与 gap-analysis §2 表格一致）
F_DELIQIDIAN = "202608德勤派单增减人员.xlsx"
F_WANBAO = "万宝盛华人员变动通知.xlsx"
F_HUOYUE = "活悦安徽分增减表.xlsx"
F_WAIGUAN = "外管离职导出（上海中企）.xlsx"
F_SHANGHAI = "上海信息数据（社保&公积金）.xlsx"


def _build_xlsx(path, sheet_name, rows, formats=None):
    """在 tmp_path 下生成小 xlsx：rows 为二维数组，formats 为 {(行, 列): number_format}（1 起始）"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            if val is not None:
                ws.cell(row=r_idx, column=c_idx, value=val)
    for (r_idx, c_idx), fmt in (formats or {}).items():
        ws.cell(row=r_idx, column=c_idx).number_format = fmt
    wb.save(path)
    return str(path)


def _render_fixture(file_name):
    """渲染夹具文件并返回 {sheet名: sheet dict}"""
    result = render_llm_view(str(FIXTURE_DIR / file_name))
    assert result["success"] is True, result
    return {s["name"]: s for s in result["sheets"]}


class TestValueRendering:
    """D5 值渲染契约：按 number_format/is_date 渲染，LLM 看见文件真实语义"""

    def test_serial_date_rendered_as_date(self, tmp_path):
        """Excel 序列日期（int + 日期格式，如 46239）渲染为日期字符串而非数字"""
        path = _build_xlsx(
            tmp_path / "serial.xlsx", "Sheet1",
            [["姓名", "派单日期"], ["张三", 46239]],
            formats={(2, 2): "yyyy-mm-dd"},
        )
        text = render_llm_view(path)["sheets"][0]["text"]
        assert "2026-08-05" in text
        assert "46239" not in text

    def test_datetime_without_time_renders_date_only(self, tmp_path):
        """datetime 无时间部分 -> YYYY-MM-DD（不含 00:00:00）"""
        path = _build_xlsx(
            tmp_path / "dt.xlsx", "Sheet1",
            [["入职日期"], [datetime.datetime(2026, 9, 1)]],
        )
        text = render_llm_view(path)["sheets"][0]["text"]
        assert "2026-09-01" in text
        assert "00:00:00" not in text

    def test_datetime_with_time_renders_full(self, tmp_path):
        """datetime 带时间 -> YYYY-MM-DD HH:MM:SS"""
        path = _build_xlsx(
            tmp_path / "dt2.xlsx", "Sheet1",
            [["打卡时间"], [datetime.datetime(2026, 9, 1, 8, 30, 5)]],
        )
        text = render_llm_view(path)["sheets"][0]["text"]
        assert "2026-09-01 08:30:05" in text

    def test_percent_formats(self, tmp_path):
        """百分比格式按格式小数位数渲染明文；字符串 5%:5% 原样照抄"""
        path = _build_xlsx(
            tmp_path / "pct.xlsx", "Sheet1",
            [["a", "b", "c"], [0.05, 0.055, "5%:5%"]],
            formats={(2, 1): "0%", (2, 2): "0.0%"},
        )
        text = render_llm_view(path)["sheets"][0]["text"]
        assert "| 5% | 5.5% | 5%:5% |" in text

    def test_long_cell_truncated(self, tmp_path):
        """超 300 字符单元格截断为 300 字 + …(截断)"""
        long_text = "长" * 301
        path = _build_xlsx(tmp_path / "long.xlsx", "Sheet1", [["备注"], [long_text]])
        text = render_llm_view(path)["sheets"][0]["text"]
        assert "…(截断)" in text
        # 输出结构：标题 / 空行 / 表头 / 分隔行 / 数据行
        data_line = text.splitlines()[4]
        assert data_line == "| " + "长" * 300 + "…(截断) |"

    def test_none_empty_and_strip(self, tmp_path):
        """None -> 空串；字符串去首尾空白"""
        path = _build_xlsx(tmp_path / "nil.xlsx", "Sheet1", [["a", "b"], [None, "  x  "]])
        text = render_llm_view(path)["sheets"][0]["text"]
        assert "|  | x |" in text


class TestDirtyStructurePreserved:
    """脏结构原样保留在输出中（不清洗、不删行）——gap-analysis §2 逐脏点"""

    def test_two_row_header_preserved(self):
        """活悦：R1 说明行 + R2 真表头，两行都原样出现在输出"""
        sheet = _render_fixture(F_HUOYUE)["减员"]
        lines = sheet["text"].splitlines()
        # markdown 表第一行 = sheet 第一行（说明行），真表头是第二行表体
        assert "2026年8月安徽分公司社保减员名单" in lines[2]
        assert "序号" in lines[4] and "社保最后缴纳月" in lines[4]
        assert sheet["row_count"] == 3

    def test_segment_and_repeated_header_preserved(self):
        """万宝减员：0806补充 分段行与重复表头原样保留"""
        sheet = _render_fixture(F_WANBAO)["减员"]
        text = sheet["text"]
        assert "| 0806补充 |" in text
        # 表头行出现两次（首行 + 分段后重复表头）
        assert text.count("雇员姓名") == 2
        assert sheet["row_count"] == 4  # 1 数据 + 分段行 + 重复表头 + 1 数据

    def test_title_rows_and_t0_preserved(self):
        """外管：标题/说明/T0 三行前置原样保留，真表头在数据区第 4 行"""
        sheet = _render_fixture(F_WAIGUAN)["减员"]
        table_lines = [ln for ln in sheet["text"].splitlines() if ln.startswith("|")]
        # 表行顺序：标题（作 markdown 表头）/ 分隔行 / 说明 / T0 / 真表头 / 2 行数据
        assert "上海中企人力资源服务有限公司离职人员名单" in table_lines[0]
        assert "说明：本表数据截至2026年8月" in table_lines[2]
        assert table_lines[3].startswith("| T0 |")
        assert "下岗原因" in table_lines[4]
        assert sheet["row_count"] == 5

    def test_footer_rows_preserved(self):
        """上海增员表：数据区后 2 行备注页脚原样保留并计入 row_count"""
        sheet = _render_fixture(F_SHANGHAI)["增员表"]
        text = sheet["text"]
        assert "备注：以上人员自2026年9月起缴纳社保" in text
        assert "数据来源：上海信息（社保&公积金）" in text
        assert sheet["row_count"] == 4  # 2 数据 + 2 页脚

    def test_empty_rows_skipped(self):
        """德勤新增：表头后 30 行全空行被跳过，row_count 只数有效行"""
        sheet = _render_fixture(F_DELIQIDIAN)["新增"]
        # 源文件 max_row=33（表头 + 30 空行 + 2 数据行）
        wb = openpyxl.load_workbook(FIXTURE_DIR / F_DELIQIDIAN)
        assert wb["新增"].max_row == 33
        wb.close()
        # 渲染后仅 4 行 markdown 表（表头 + 分隔 + 2 数据）
        table_lines = [ln for ln in sheet["text"].splitlines() if ln.startswith("|")]
        assert len(table_lines) == 4
        assert sheet["row_count"] == 2

    def test_merged_non_anchor_cells_empty(self):
        """外管标题行 A1:E1 合并：非锚点格渲染为空串，锚点保留标题"""
        sheet = _render_fixture(F_WAIGUAN)["减员"]
        assert sheet["merged_cells"] == ["A1:E1"]
        title_line = sheet["text"].splitlines()[2]
        cells = title_line.split("|")[1:-1]
        assert cells[0].strip() == "上海中企人力资源服务有限公司离职人员名单"
        assert all(c.strip() == "" for c in cells[1:5])  # B1:E1 合并区非锚点

    def test_duplicate_column_names_kept(self):
        """德勤：重复列名 派单日期 两列都保留（不去重列名）"""
        sheet = _render_fixture(F_DELIQIDIAN)["新增"]
        assert sheet["text"].count("派单日期") == 2

    def test_newline_header_preserved(self):
        """德勤：换行复合表头单元格原样（含 \\n）"""
        sheet = _render_fixture(F_DELIQIDIAN)["新增"]
        assert "社保缴至月份\n（如6月继续缴纳则填）" in sheet["text"]

    def test_empty_column_kept_and_aligned(self):
        """德勤新增：仅设格式的全空列保留，分隔行与数据行都是 11 列对齐"""
        sheet = _render_fixture(F_DELIQIDIAN)["新增"]
        text = sheet["text"]
        # 分隔行 11 列（10 个有值列 + 1 个仅设格式的全空列）
        sep_line = next(ln for ln in text.splitlines() if ln.startswith("| ---"))
        assert sep_line == "| " + " | ".join(["---"] * 11) + " |"
        # 数据行末尾：备注为空 + 全空列保留（尾部两个空单元格）
        assert "| 2026-08-12 |  |  |" in text


class TestGoldenSnapshot:
    """黄金快照：5 个夹具全部 sheet 与 expected_render.json 逐 sheet 全等"""

    @pytest.mark.parametrize("file_name", [
        F_DELIQIDIAN, F_WANBAO, F_HUOYUE, F_WAIGUAN, F_SHANGHAI,
    ])
    def test_fixture_matches_expected_render(self, file_name):
        expected_all = json.loads((FIXTURE_DIR / "expected_render.json").read_text(encoding="utf-8"))
        expected_sheets = expected_all[file_name]

        result = render_llm_view(str(FIXTURE_DIR / file_name))
        assert result["success"] is True, result
        assert result["file_name"] == file_name

        got_sheets = {s["name"]: s for s in result["sheets"]}
        # sheet 顺序与集合一致
        assert [s["name"] for s in result["sheets"]] == list(expected_sheets)
        for name, exp in expected_sheets.items():
            got = got_sheets[name]
            assert got["text"] == exp["text"], f"{file_name}#{name} text 不一致"
            assert got["row_count"] == exp["row_count"], f"{file_name}#{name} row_count 不一致"
            assert got["merged_cells"] == exp["merged_cells"], f"{file_name}#{name} merged_cells 不一致"


class TestErrorPaths:
    """异常路径：文件不存在 / 非 xlsx 返回 success False"""

    def test_missing_file(self):
        result = render_llm_view("/nonexistent/path/不存在.xlsx")
        assert result["success"] is False
        assert "文件不存在" in result["error"]

    def test_csv_rejected(self, tmp_path):
        csv_file = tmp_path / "data.csv"
        csv_file.write_text("a,b\n1,2\n", encoding="utf-8")
        result = render_llm_view(str(csv_file))
        assert result["success"] is False
        assert "xlsx" in result["error"]

    def test_xls_rejected(self, tmp_path):
        xls_file = tmp_path / "old.xls"
        xls_file.write_bytes(b"\xd0\xcf\x11\xe0")
        result = render_llm_view(str(xls_file))
        assert result["success"] is False
