"""
excel-to-template pipeline 单测（Phase 3 / M5 编排，模块导入方式不起子进程）

覆盖：
- 文件入口 e2e（--llm-fake）：zip 夹具 → 模板指纹识别 → 抽取 → 填充 → 报告，
  输出 xlsx 读回与黄金 14 条字段级全等；D24 marker 不涉及
- 模板识别（D3）：指纹零命中/多命中报错、显式指认优先
- 邮件入口（mock email_lib + fake LLM）：D23 两级判定（关键词直下 + LLM 批量判定）、
  D24 UID 水位二次运行跳过
- zip 路径穿越防护（恶意成员名拒绝）
- DB 池初始化失败降级不崩（计量不落库管线照跑）
- schema 指纹不匹配会话内重抽（D14 只读降级提示进报告）
- .xls 老格式来源失败记录进结果（D4）
"""

import importlib.util
import json
import os
import zipfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

pytestmark = [pytest.mark.tools]

REPO_ROOT = Path(__file__).resolve().parents[3]
FIXTURE_DIR = Path(__file__).resolve().parents[2] / "fixtures" / "excel_etl"
PIPELINE_PATH = REPO_ROOT / "src" / "skills" / "excel-to-template-1.0.0" / "scripts" / "pipeline.py"

# 以模块方式加载 pipeline（文件名含 '-'，无法常规 import）
_spec = importlib.util.spec_from_file_location("excel_etl_pipeline", PIPELINE_PATH)
pipeline = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pipeline)

GOLDEN_FILES = [
    "202608德勤派单增减人员.xlsx",
    "万宝盛华人员变动通知.xlsx",
    "活悦安徽分增减表.xlsx",
    "外管离职导出（上海中企）.xlsx",
    "上海信息数据（社保&公积金）.xlsx",
]
ZIP_FIXTURE = FIXTURE_DIR / "ExcelAI测试样本.zip"


@pytest.fixture(scope="module")
def golden():
    return json.loads((FIXTURE_DIR / "golden_records.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def schema():
    from src.tools.excel.excel_extract import load_default_schema
    return load_default_schema()


@pytest.fixture(scope="module")
def zip_fixture():
    """zip 夹具（*.zip 全局 gitignore，缺失时确定性重建）"""
    if not ZIP_FIXTURE.exists():
        spec = importlib.util.spec_from_file_location(
            "build_fixtures", FIXTURE_DIR / "build_fixtures.py"
        )
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        module.build_zip()
    return ZIP_FIXTURE


def _sequential_fake_responses(golden):
    """按 zip 解包后的 sheet 渲染顺序构造 FakeLLM 顺序响应（extract 调用顺序确定）"""
    from src.tools.excel.excel_reader import render_llm_view

    by_source = {}
    for rec in golden:
        by_source.setdefault(rec["_source"], []).append(rec)
    responses = []
    for name in GOLDEN_FILES:
        rendered = render_llm_view(str(FIXTURE_DIR / name))
        assert rendered["success"] is True
        for sheet in rendered["sheets"]:
            label = f"{name}#{sheet['name']}"
            responses.append(json.dumps(
                {"records": by_source[label]}, ensure_ascii=False
            ))
    return responses


def _llm_dispatch_by_source(golden):
    """按 prompt 中的 _source 标签分发黄金记录（与 test_excel_extract 同款 mock）"""
    by_source = {}
    for rec in golden:
        by_source.setdefault(rec["_source"], []).append(rec)

    def llm(prompt):
        for label, records in by_source.items():
            if label in prompt:
                return (json.dumps({"records": records}, ensure_ascii=False),
                        {"prompt_tokens": 100, "completion_tokens": 50,
                         "total_tokens": 150, "model": "deepseek-flash"})
        raise AssertionError(f"mock LLM 无法识别 prompt: {prompt[:120]}")

    return llm


def _read_back_records(xlsx_path, schema):
    """openpyxl 读回输出 xlsx → 18 字段记录列表"""
    field_names = [f["name"] for f in schema["fields"]]
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.worksheets[0]
    header_map = {str(c.value).strip(): c.column for c in ws[1]}
    return [
        {name: ws.cell(row=r, column=header_map[name]).value
         for name in field_names if name in header_map}
        for r in range(2, ws.max_row + 1)
    ]


# ============================================================
# 文件入口 e2e（mock LLM）
# ============================================================


class TestFileEntryEndToEnd:
    def test_zip_entry_golden_e2e(self, zip_fixture, golden, schema, tmp_path, capsys):
        """zip 全链路：识别模板 → 抽取 14 条 → 填充读回字段级全等 → 报告完整"""
        fake = tmp_path / "fake_llm.json"
        fake.write_text(json.dumps(
            _sequential_fake_responses(golden), ensure_ascii=False
        ), encoding="utf-8")

        rc = pipeline.main([
            "--zip", str(zip_fixture),
            "--output-dir", str(tmp_path),
            "--llm-fake", str(fake),
        ])
        assert rc == 0
        result = json.loads(capsys.readouterr().out)
        assert result["success"] is True
        assert result["stats"]["extracted"] == 14
        assert result["stats"]["written"] == 14
        assert result["stats"]["sources_ok"] == 5
        assert all(result["fill_checks"].values())

        # 读回字段级全等（空值归一后与黄金集合比对）
        rows = _read_back_records(Path(result["xlsx"]), schema)
        assert len(rows) == 14
        field_names = [f["name"] for f in schema["fields"]]

        def _key(rec):
            return (rec.get("姓名"), rec.get("身份证"))

        golden_by_key = {_key(g): g for g in golden}
        for row in rows:
            g = golden_by_key[_key(row)]
            for f in field_names:
                assert (row.get(f) or None) == (g.get(f) or None), \
                    f"{row.get('姓名')}.{f}: {row.get(f)!r} != {g.get(f)!r}"

        # 报告：摘要 / 明细 / 同人多条 / 计量小节齐备；无人工清单
        report = Path(result["report"]).read_text(encoding="utf-8")
        for section in ("一、摘要", "二、明细", "三、同人多条", "计量小计"):
            assert section in report
        # 黄金记录存在合法 warning（如减员参保年月为 null 照写）但无人工清单条目
        assert "warning：" in report
        assert "error：" not in report
        assert "→ 人工处理：" not in report
        # 文件入口不涉及邮件水位
        assert not (tmp_path / pipeline.WATERMARK_FILENAME).exists()

    def test_files_entry_with_explicit_template(self, golden, schema, tmp_path):
        """--files 多文件 + 显式 --template：run_etl 注入 mock LLM（与 main 同一编排）"""
        sources = [FIXTURE_DIR / n for n in GOLDEN_FILES]
        template = FIXTURE_DIR / "生成标准模板格式.xlsx"
        if not template.exists():  # gitignore 排除，确定性重建
            spec = importlib.util.spec_from_file_location(
                "build_fixtures", FIXTURE_DIR / "build_fixtures.py")
            m = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(m)
            m.build_standard_template()

        metering = pipeline.Metering(record_usage=False)
        result = pipeline.run_etl(
            sources, template, tmp_path,
            llm_callable=_llm_dispatch_by_source(golden), metering=metering,
        )
        assert result["success"] is True
        assert result["stats"]["written"] == 14
        assert metering.llm_calls == 8  # 8 个 sheet 各 1 次 extract
        assert metering.summary()["total_tokens"] == 8 * 150


# ============================================================
# 模板识别（D3）
# ============================================================


class TestIdentifyTemplate:
    def test_fingerprint_auto_detect(self, schema, zip_fixture, tmp_path):
        files = pipeline.safe_unzip(zip_fixture, tmp_path / "unpacked")
        template, sources = pipeline.identify_template(files, schema, None)
        assert template.name == "生成标准模板格式.xlsx"
        assert sorted(f.name for f in sources) == sorted(GOLDEN_FILES)

    def test_no_match_raises(self, schema, tmp_path):
        files = [FIXTURE_DIR / n for n in GOLDEN_FILES]
        with pytest.raises(ValueError, match="--template"):
            pipeline.identify_template(files, schema, None)

    def test_multiple_match_raises(self, schema, tmp_path):
        src = FIXTURE_DIR / "生成标准模板格式.xlsx"
        if not src.exists():
            spec = importlib.util.spec_from_file_location(
                "build_fixtures", FIXTURE_DIR / "build_fixtures.py")
            m = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(m)
            m.build_standard_template()
        dup = tmp_path / "模板副本.xlsx"
        dup.write_bytes(src.read_bytes())
        with pytest.raises(ValueError, match="多个文件"):
            pipeline.identify_template([src, dup], schema, None)

    def test_explicit_template_wins(self, schema, tmp_path):
        src = FIXTURE_DIR / "生成标准模板格式.xlsx"
        explicit = tmp_path / "指认模板.xlsx"
        explicit.write_bytes(src.read_bytes())
        template, sources = pipeline.identify_template([src, explicit], schema, explicit)
        assert template == explicit
        assert sources == [src]


# ============================================================
# 邮件入口（D21~D24）
# ============================================================


class TestEmailEntry:
    def _emails(self):
        return [
            {  # 第一级：附件名命中关键词"增减" → 直接下载，不走 LLM
                "uid": "10", "subject": "8月报表",
                "body_preview": "见附件", "attachments": [
                    {"filename": "202608德勤派单增减人员.xlsx",
                     "content_type": "application/octet-stream", "size": 10240},
                ],
            },
            {  # 附件名无关键词（august_roster.xlsx）→ 第二级 LLM 判定
                "uid": "20", "subject": "8月社保增减员名单",
                "body_preview": "本月增减员人员名单见附件，请查收",
                "attachments": [
                    {"filename": "august_roster.xlsx",
                     "content_type": "application/octet-stream", "size": 20480},
                ],
            },
        ]

    def _patch_email_lib(self, emails, downloads):
        read = patch.object(
            pipeline.email_lib, "read_emails",
            return_value={"success": True, "emails": emails, "folders": [], "count": len(emails)},
        )
        download = patch.object(
            pipeline.email_lib, "download_attachments",
            side_effect=lambda ue, uid, d, filenames=None, folder="INBOX": {
                "success": True, "uid": str(uid),
                "files": [{"filename": fn, "path": str(Path(d) / fn), "size": 1}
                          for fn in (filenames or [])],
                "skipped": [],
            },
        )
        return read, download

    def test_two_stage_judgment_and_watermark(self, tmp_path):
        """D23：关键词直下 + LLM 批量判定；D24：二次运行跳过已处理 UID"""
        emails = self._emails()
        read_p, download_p = self._patch_email_lib(emails, [])
        fake = pipeline.FakeLLM([
            json.dumps({"judgments": [
                {"uid": "20", "downloads": ["august_roster.xlsx"], "reason": "8月增减员名单"},
            ]}, ensure_ascii=False),
        ])

        with read_p as mock_read, download_p as mock_download:
            result = pipeline.collect_email_files(
                user_email=object(), email_since="2026-08-01",
                output_dir=tmp_path, llm_callable=fake,
            )
            assert result["success"] is True
            assert result["emails_processed"] == 2
            assert result["judged_emails"] == 1
            # read 走服务器端 SINCE
            assert mock_read.call_args.kwargs["since"] == "2026-08-01"
            # 两封各一次下载调用
            downloaded_uids = [str(c.args[1]) for c in mock_download.call_args_list]
            assert downloaded_uids == ["10", "20"]
            assert mock_download.call_args_list[0].kwargs["filenames"] == \
                ["202608德勤派单增减人员.xlsx"]
            assert mock_download.call_args_list[1].kwargs["filenames"] == ["august_roster.xlsx"]

        # 水位已写入两条 UID
        marker = tmp_path / pipeline.WATERMARK_FILENAME
        assert marker.exists()
        assert set(json.loads(marker.read_text(encoding="utf-8"))["processed_uids"]) == {"10", "20"}

        # 二次运行：全部命中水位 → 不判定不下载（fake 无剩余响应也不会被调用）
        fake2 = pipeline.FakeLLM([])  # 若被调用会抛"响应已耗尽"，反证未调用
        read_p2, download_p2 = self._patch_email_lib(emails, [])
        with read_p2, download_p2 as mock_download2:
            result2 = pipeline.collect_email_files(
                user_email=object(), email_since="2026-08-01",
                output_dir=tmp_path, llm_callable=fake2,
            )
            assert result2["success"] is True
            assert result2["emails_processed"] == 0
            assert result2["files"] == []
            mock_download2.assert_not_called()

    def test_keyword_stage_skips_irrelevant_extension(self, tmp_path):
        """第一级预筛：关键词命中但扩展名不在 {.xlsx,.xls,.zip}（如 .pdf）不下载"""
        emails = [{
            "uid": "30", "subject": "通知", "body_preview": "",
            "attachments": [{"filename": "增减说明.pdf",
                             "content_type": "application/pdf", "size": 10}],
        }]
        read_p, download_p = self._patch_email_lib(emails, [])
        with read_p, download_p as mock_download:
            result = pipeline.collect_email_files(
                user_email=object(), email_since="2026-08-01",
                output_dir=tmp_path, llm_callable=pipeline.FakeLLM([
                    json.dumps({"judgments": []}, ensure_ascii=False),
                ]),
            )
            assert result["success"] is True
            mock_download.assert_not_called()  # PDF 未命中第一级，LLM 判定也无下载

    def test_judge_batch_output_contract(self):
        emails = [{
            "uid": "40", "subject": "s", "body_preview": "b",
            "attachments": [{"filename": "x.xlsx", "content_type": "a", "size": 1}],
        }]
        fake = pipeline.FakeLLM([
            json.dumps({"judgments": [
                {"uid": "40", "downloads": ["x.xlsx"], "reason": "r"},
                {"uid": "99", "downloads": [], "reason": "不相关"},
                "垃圾项",
            ]}, ensure_ascii=False),
        ])
        verdict = pipeline.judge_attachments_batch(emails, fake)
        assert verdict == {"40": ["x.xlsx"]}

    def test_main_email_entry_missing_credentials(self, monkeypatch):
        """AID_USER_ID 缺失 → resolve 阶段明确报错（不静默空跑）"""
        monkeypatch.delenv("AID_USER_ID", raising=False)
        with pytest.raises(ValueError, match="AID_USER_ID"):
            pipeline.resolve_email_credentials()

    def test_main_email_entry_no_new_attachments_message(self, tmp_path, capsys):
        """邮件入口跑了但零新附件（D24 水位重跑）：报错不误导为"没有输入文件"，
        并带 email 摘要与重试指引"""
        emails = [{
            "uid": "10", "subject": "8月报表", "body_preview": "", "attachments": [],
        }]
        with patch.object(pipeline, "resolve_email_credentials", return_value=object()), \
             patch.object(pipeline.email_lib, "read_emails", return_value={
                 "success": True, "emails": emails, "folders": [], "count": 1,
             }), patch.object(pipeline.email_lib, "download_attachments") as mock_dl:
            rc = pipeline.main([
                "--email-since", "2026-08-01", "--output-dir", str(tmp_path),
            ])
        assert rc == 1
        result = json.loads(capsys.readouterr().out)
        assert result["success"] is False
        assert "没有新附件" in result["error"] and "水位" in result["error"]
        assert "没有输入文件（--files" not in result["error"]
        assert result["email"]["emails_seen"] == 1
        mock_dl.assert_not_called()  # 无附件邮件不触发下载


# ============================================================
# zip 路径穿越防护
# ============================================================


class TestInputResolution:
    def test_bare_filename_glob_stays_in_cwd_tenant_isolated(self, tmp_path, monkeypatch):
        """纯文件名递归兜底只在 cwd（本租户技能工作区）内：项目根 storage/tenants
        下他租户的同名文件不可被解析到（租户隔离）"""
        ws = tmp_path / "ws"
        ws.mkdir()
        other_tenant = tmp_path / "storage" / "tenants" / "other" / "conversation"
        other_tenant.mkdir(parents=True)
        (other_tenant / "人员变动通知.xlsx").write_bytes(b"other tenant data")

        monkeypatch.setattr(pipeline, "PROJECT_ROOT", str(tmp_path))
        monkeypatch.chdir(ws)
        # cwd 内无此文件、项目根精确路径也无 → 不落在他租户文件上，明确报不存在
        with pytest.raises(FileNotFoundError, match="输入文件不存在"):
            pipeline._resolve_input_path("人员变动通知.xlsx")

        # cwd 子目录内有同名文件 → 正常解析（兜底仍服务本工作区）
        sub = ws / "sub"
        sub.mkdir()
        (sub / "人员变动通知.xlsx").write_bytes(b"mine")
        assert pipeline._resolve_input_path("人员变动通知.xlsx") == (sub / "人员变动通知.xlsx").resolve()


class TestZipSafety:
    def test_safe_unzip_rejects_traversal_member(self, tmp_path):
        evil = tmp_path / "evil.zip"
        with zipfile.ZipFile(evil, "w") as zf:
            zf.writestr("正常.xlsx", b"ok")
            zf.writestr("../evil.xlsx", b"bad")
        with pytest.raises(ValueError, match="路径穿越"):
            pipeline.safe_unzip(evil, tmp_path / "out")

    def test_safe_unzip_rejects_absolute_member(self, tmp_path):
        evil = tmp_path / "evil2.zip"
        with zipfile.ZipFile(evil, "w") as zf:
            zf.writestr("/tmp/evil.xlsx", b"bad")
        with pytest.raises(ValueError, match="非法"):
            pipeline.safe_unzip(evil, tmp_path / "out")

    def test_main_traversal_zip_fails_gracefully(self, tmp_path, capsys):
        evil = tmp_path / "evil.zip"
        with zipfile.ZipFile(evil, "w") as zf:
            zf.writestr("../evil.xlsx", b"bad")
        rc = pipeline.main(["--zip", str(evil), "--output-dir", str(tmp_path)])
        assert rc == 1
        result = json.loads(capsys.readouterr().out)
        assert result["success"] is False
        assert "路径穿越" in result["error"]
        assert not (tmp_path.parent / "evil.xlsx").exists()

    def test_safe_unzip_only_takes_excel(self, tmp_path):
        pack = tmp_path / "mix.zip"
        with zipfile.ZipFile(pack, "w") as zf:
            zf.writestr("a.xlsx", b"1")
            zf.writestr("说明.txt", b"2")
            zf.writestr("b.xls", b"3")
        extracted = pipeline.safe_unzip(pack, tmp_path / "out")
        assert sorted(f.name for f in extracted) == ["a.xlsx", "b.xls"]


# ============================================================
# DB 池降级与计量
# ============================================================


class TestDbDegrade:
    def test_init_db_pool_swallows_failure(self, monkeypatch):
        """DB 初始化异常 → 返回 False 不抛（管线降级继续）"""
        import src.db.database as db_module
        monkeypatch.setattr(db_module, "get_postgres_pool",
                            lambda: (_ for _ in ()).throw(RuntimeError("no db")))
        assert pipeline._init_db_pool() is False

    def test_truth_recheck_demotion_updates_stats(self, golden, tmp_path):
        """unmask 真值复检（第 5 步）：校验位错误行进人工清单，
        stats.manual_review_count 以最终清单计（含 run_extraction 之后的降级）"""
        template = FIXTURE_DIR / "生成标准模板格式.xlsx"
        if not template.exists():
            spec = importlib.util.spec_from_file_location(
                "build_fixtures", FIXTURE_DIR / "build_fixtures.py")
            m = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(m)
            m.build_standard_template()

        deedu = [dict(g) for g in golden if g["_source"].startswith(GOLDEN_FILES[0])]
        bad = dict(deedu[0])
        bad["身份证"] = bad["身份证"][:-1] + ("0" if bad["身份证"][-1] != "0" else "1")
        deedu[0] = bad  # 一条校验位错误 → 真值复检降级人工清单

        with patch.object(pipeline, "run_extraction", return_value={
            "success": True, "records": deedu,
            "record_warnings": [[] for _ in deedu],
            "manual_review": [],  # run_extraction 阶段零人工（脱敏域占位符放行）
            "duplicate_groups": [],
            "stats": {"extracted": len(deedu), "repaired": 0,
                      "manual_review_count": 0, "dropped_items": 0, "llm_calls": 1},
        }):
            result = pipeline.run_etl(
                [FIXTURE_DIR / GOLDEN_FILES[0]], template, tmp_path, llm_callable=None,
            )
        assert result["success"] is True
        assert result["stats"]["written"] == len(deedu) - 1
        # 关键：stats 计入真值复检降级（run_extraction 原始计数为 0，会偏少）
        assert result["stats"]["manual_review_count"] == 1
        report = Path(result["report"]).read_text(encoding="utf-8")
        assert "人工处理清单：1 条" in report
        assert "身份证校验位错误" in report

    def test_run_etl_without_db(self, golden, schema, tmp_path):
        """record_usage=False（DB 不可用）：计量只累计不落库，管线正常完成"""
        template = FIXTURE_DIR / "生成标准模板格式.xlsx"
        if not template.exists():
            spec = importlib.util.spec_from_file_location(
                "build_fixtures", FIXTURE_DIR / "build_fixtures.py")
            m = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(m)
            m.build_standard_template()
        metering = pipeline.Metering(record_usage=False)
        result = pipeline.run_etl(
            [FIXTURE_DIR / GOLDEN_FILES[0]], template, tmp_path,
            llm_callable=_llm_dispatch_by_source(
                [g for g in golden if g["_source"].startswith(GOLDEN_FILES[0])]
            ),
            metering=metering,
        )
        assert result["success"] is True
        assert result["stats"]["written"] == 3  # 德勤 3 条
        assert metering.record_usage is False
        assert metering.llm_calls == 2  # 新增 + 减少 两个 sheet


# ============================================================
# schema 指纹不匹配重抽（D14 只读降级）
# ============================================================


class TestSchemaReextract:
    def test_mismatch_template_reextracts_in_session(self, golden, schema, tmp_path):
        """模板表头改别名 → 指纹不匹配 → 一次 LLM 重抽 schema（会话内），报告带提示"""
        template = FIXTURE_DIR / "生成标准模板格式.xlsx"
        if not template.exists():
            spec = importlib.util.spec_from_file_location(
                "build_fixtures", FIXTURE_DIR / "build_fixtures.py")
            m = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(m)
            m.build_standard_template()
        wb = openpyxl.load_workbook(template)
        wb.worksheets[0].cell(row=1, column=1, value="操作类型")  # 别名（指纹变化）
        variant = tmp_path / "变体模板.xlsx"
        wb.save(variant)
        wb.close()

        from src.tools.excel.excel_extract import schema_matches
        assert schema_matches(variant, schema) is False

        stages = []

        def llm(prompt):
            if "字段定义专家" in prompt:
                stages.append("schema")
                return (json.dumps({"fields": schema["fields"]}, ensure_ascii=False),
                        {"prompt_tokens": 500, "completion_tokens": 100,
                         "total_tokens": 600, "model": "deepseek-flash"})
            stages.append("extract")
            by_source = {}
            for g in golden:
                by_source.setdefault(g["_source"], []).append(g)
            for label, records in by_source.items():
                if label in prompt:
                    return (json.dumps({"records": records}, ensure_ascii=False),
                            {"prompt_tokens": 100, "completion_tokens": 50,
                             "total_tokens": 150, "model": "deepseek-flash"})
            raise AssertionError(f"无法识别 prompt: {prompt[:120]}")

        result = pipeline.run_etl(
            [FIXTURE_DIR / GOLDEN_FILES[0]], variant, tmp_path, llm_callable=llm,
        )
        assert result["success"] is True
        assert stages[0] == "schema"  # 先重抽再抽取
        report = Path(result["report"]).read_text(encoding="utf-8")
        assert "指纹不匹配" in report and "重抽" in report
        # 变体表头（别名"操作类型"）仍能绑定填充
        rows = _read_back_records(Path(result["xlsx"]), schema)
        assert len(rows) == 3


# ============================================================
# .xls 老格式（D4）
# ============================================================


class TestXlsRejected:
    def test_xls_source_recorded_as_failed(self, golden, tmp_path):
        """.xls 渲染失败 → failed_sources 记录，其余来源照常"""
        template = FIXTURE_DIR / "生成标准模板格式.xlsx"
        if not template.exists():
            spec = importlib.util.spec_from_file_location(
                "build_fixtures", FIXTURE_DIR / "build_fixtures.py")
            m = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(m)
            m.build_standard_template()
        fake_xls = tmp_path / "老格式.xls"
        fake_xls.write_bytes(b"not really xls")

        result = pipeline.run_etl(
            [fake_xls, FIXTURE_DIR / GOLDEN_FILES[0]], template, tmp_path,
            llm_callable=_llm_dispatch_by_source(
                [g for g in golden if g["_source"].startswith(GOLDEN_FILES[0])]
            ),
        )
        assert result["success"] is True
        assert result["stats"]["written"] == 3
        assert len(result["failed_sources"]) == 1
        assert result["failed_sources"][0]["file"] == "老格式.xls"
        assert "xlsx" in result["failed_sources"][0]["error"]

    def test_all_sources_failed(self, tmp_path):
        template = FIXTURE_DIR / "生成标准模板格式.xlsx"
        if not template.exists():
            spec = importlib.util.spec_from_file_location(
                "build_fixtures", FIXTURE_DIR / "build_fixtures.py")
            m = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(m)
            m.build_standard_template()
        fake_xls = tmp_path / "a.xls"
        fake_xls.write_bytes(b"x")
        result = pipeline.run_etl([fake_xls], template, tmp_path, llm_callable=None)
        assert result["success"] is False
        assert "没有可抽取" in result["error"]
