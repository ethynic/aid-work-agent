"""Phase 5 editable HTML PPTX export tests with real Playwright."""

from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from src.tools.ppt.html_exporter import HtmlExporter
from src.tools.ppt.ppt_capabilities import get_capabilities
from src.tools.ppt.spec import RasterLayerNode, SlideDeckSpec, SlideSpec

pytestmark = pytest.mark.tools


def _require_playwright():
    capability = get_capabilities()["html_export"]
    if not capability["available"]:
        pytest.skip(f"真实 Playwright 测试跳过: {capability['reason']}")


def _simple_html() -> str:
    return """<!doctype html><html><head><style>
html,body { margin:0; width:100%; height:100%; overflow:hidden; }
section.slide { position:relative; width:100vw; height:100vh; background:#f7f8fa; }
h1 { position:absolute; left:40px; top:20px; margin:0; font:700 40px Arial; color:#123456; z-index:3; }
p { position:absolute; left:40px; top:90px; width:420px; margin:0; font:20px Arial; z-index:3; }
.card { position:absolute; left:35px; top:140px; width:300px; height:120px;
  background:#ffffff; border:2px solid #335577; border-radius:12px; z-index:1; }
.divider { position:absolute; left:-20px; top:290px; width:900px; height:2px; background:#445566; }
.tag { position:absolute; left:370px; top:150px; padding:8px; background:#22aa77; z-index:5; }
table { position:absolute; left:370px; top:210px; width:360px; height:100px; border:1px solid #999; }
svg { position:absolute; left:700px; top:30px; width:80px; height:80px; }
</style></head><body><section class="slide">
<h1 id="title">可编辑标题</h1><p>正文内容</p><div class="card"></div>
<div class="divider"></div><div class="tag">标签</div>
<table><tr><th>项目</th><th>值</th></tr><tr><td>A</td><td>10</td></tr></table>
<svg viewBox="0 0 10 10"><circle cx="5" cy="5" r="4" fill="red"/></svg>
</section></body></html>"""


def test_real_playwright_maps_simple_dom_to_editable_nodes(tmp_path):
    _require_playwright()
    result = HtmlExporter(viewport_width=800, viewport_height=450).export(
        _simple_html(), tmp_path, include_editable=True
    )

    spec = result.editable_spec
    assert spec is not None
    node_types = [node.type for node in spec.slides[0].nodes]
    assert node_types.count("text") >= 2
    assert node_types.count("shape") >= 3
    assert node_types.count("image") == 1
    assert node_types.count("table") == 1
    assert result.editability["raster_layer_count"] == 0
    assert result.editability["editable_ratio"] == 1.0
    assert any(item["selector"] == "#title" for item in result.dom_manifest[0])
    z_indexes = [item["z_index"] for item in result.dom_manifest[0]]
    assert z_indexes == sorted(z_indexes)
    for node in spec.slides[0].nodes:
        assert node.x + node.w <= spec.width + 1e-6
        assert node.y + node.h <= spec.height + 1e-6


def test_real_playwright_rasterizes_canvas_and_complex_filters(tmp_path):
    _require_playwright()
    html = """<html><head><style>
body { margin:0; background:#eef; } h1 { position:absolute; left:20px; top:10px; }
canvas { position:absolute; left:40px; top:100px; width:300px; height:180px; }
.glass { position:absolute; left:400px; top:80px; width:300px; height:200px;
background:rgba(255,255,255,.5); backdrop-filter:blur(8px); mix-blend-mode:multiply; }
</style></head><body><h1>混合页面</h1><canvas id="c" width="300" height="180"></canvas>
<div class="glass"><p>应随复杂层截图</p></div><script>
const c=document.getElementById('c').getContext('2d'); c.fillStyle='#1677ff'; c.fillRect(0,0,300,180);
</script></body></html>"""
    result = HtmlExporter(viewport_width=800, viewport_height=450).export(
        html, tmp_path, include_editable=True
    )

    assert result.editability["native_text_count"] == 1
    assert result.editability["raster_layer_count"] == 2
    assert 0 < result.editability["raster_area_ratio"] < 1
    assert result.editability["editable_ratio"] < 1
    raster_nodes = [
        node for node in result.editable_spec.slides[0].nodes if node.type == "raster"
    ]
    assert all(Path(node.path).is_file() for node in raster_nodes)


def test_real_playwright_avoids_nested_text_duplicates_and_expands_spans(tmp_path):
    _require_playwright()
    html = """<html><body style="margin:0"><section class="slide">
<h1 style="position:absolute;left:20px;top:10px"><p>只保留一次</p></h1>
<table style="position:absolute;left:20px;top:100px;width:400px;height:180px">
<tr><th colspan="2">标题</th></tr>
<tr><td rowspan="2">A</td><td>1</td></tr><tr><td>2</td></tr>
</table></section></body></html>"""
    result = HtmlExporter(viewport_width=800, viewport_height=450).export(
        html, tmp_path, include_editable=True
    )

    nodes = result.editable_spec.slides[0].nodes
    text_nodes = [node for node in nodes if node.type == "text"]
    table = next(node for node in nodes if node.type == "table")
    assert [node.text for node in text_nodes] == ["只保留一次"]
    assert table.rows == [["标题", ""], ["A", "1"], ["", "2"]]


def test_real_playwright_image_layer_does_not_capture_background(tmp_path):
    _require_playwright()
    html = """<html><body style="margin:0;background:#ff0000">
<svg style="position:absolute;left:20px;top:20px;width:80px;height:80px"
viewBox="0 0 10 10"><circle cx="5" cy="5" r="3" fill="#0000ff"/></svg>
</body></html>"""
    result = HtmlExporter(viewport_width=800, viewport_height=450).export(
        html, tmp_path, include_editable=True
    )
    image_node = next(
        node for node in result.editable_spec.slides[0].nodes if node.type == "image"
    )
    with Image.open(image_node.path) as image:
        assert image.mode == "RGBA"
        assert image.getpixel((0, 0))[3] == 0


def test_editability_ratio_uses_union_of_overlapping_raster_layers():
    spec = SlideDeckSpec(
        title="overlap",
        width=10,
        height=10,
        slides=[
            SlideSpec(
                id="slide-1",
                nodes=[
                    RasterLayerNode(x=0, y=0, w=5, h=5, path="a.png"),
                    RasterLayerNode(x=0, y=0, w=5, h=5, path="b.png"),
                ],
            )
        ],
    )

    summary = HtmlExporter._editability_summary(spec)

    assert summary["raster_area_ratio"] == 0.25
    assert summary["editable_ratio"] == 0.75


@pytest.mark.asyncio
async def test_both_mode_creates_reopenable_distinct_pptx_files(tmp_path, monkeypatch):
    from src.tools.ppt.ppt_process_tool import PptProcessTool

    if not get_capabilities()["node_renderer"]["available"]:
        pytest.skip("真实 Node renderer 测试跳过")
    monkeypatch.setenv("PPT_ENABLE_HTML_EXPORT", "true")
    monkeypatch.setenv("PPT_HTML_VIEWPORT_WIDTH", "800")
    monkeypatch.setenv("PPT_HTML_VIEWPORT_HEIGHT", "450")
    monkeypatch.setattr(PptProcessTool, "_get_output_dir", lambda self: tmp_path)

    result = await PptProcessTool().execute(
        content=_simple_html(),
        content_type="html",
        output_name="Phase5稳定输出",
        export_mode="both",
    )

    assert result["success"] is True
    assert result["export_mode"] == "both"
    primary = Path(result["file_path"])
    alternate = Path(result["alternate_file_path"])
    assert primary.name == "Phase5稳定输出_editable.pptx"
    assert alternate.name == "Phase5稳定输出_high_fidelity.pptx"
    assert primary != alternate
    assert len(Presentation(primary).slides) == 1
    assert len(Presentation(alternate).slides) == 1
    assert result["qa_summary"]["editable_ratio"] == 1.0
    assert result["qa_summary"]["native_text_count"] >= 2
    assert not list(tmp_path.glob(".*_html_assets-*"))
