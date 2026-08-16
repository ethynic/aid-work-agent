"""Word 客户模板格式参考生成单元测试

覆盖「场景二：用户上传 Word 模板 docx，按模板整体格式生成新文档」链路：
- _extract_reference_default_cjk_font: 提取 Normal 样式 eastAsia / 异常路径回退 None
- _ensure_cjk_fonts 样式感知：样式链定义 eastAsia 时不覆盖 run，未定义时写默认值
- _convert_sync 集成：客户模板 Normal eastAsia=KaiTi 时输出正文中文字体为 KaiTi（真实 Pandoc）
- _handle_md_to_word 的 template_file：解析成功传 convert / 失败 fail loud / 不传行为不变
- md 来源守卫：file_paths 只有 .docx 时不把二进制当 md 读
- 路由 prompt 与工具描述更新
"""

from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

pytestmark = [pytest.mark.tools]


# =============================================================================
# 测试辅助
# =============================================================================

def _set_style_east_asia(style, font):
    """给样式设置 rPr/rFonts/eastAsia（XML 层操作）。"""
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), font)


def _make_reference_docx(path, east_asia=None):
    """构造参考 docx，可选在 Normal 样式上设置 eastAsia 中文字体。"""
    doc = Document()
    if east_asia:
        _set_style_east_asia(doc.styles["Normal"], east_asia)
    doc.save(str(path))
    return path


def _run_east_asia(run):
    """读 run 级 eastAsia（未设置为 None）。"""
    rPr = run._element.rPr
    if rPr is None:
        return None
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        return None
    return rFonts.get(qn('w:eastAsia'))


def _effective_east_asia(paragraph):
    """段落首个 run 生效的 eastAsia：run 级优先，否则沿样式链（至多 3 层）追溯。"""
    if paragraph.runs:
        run_ea = _run_east_asia(paragraph.runs[0])
        if run_ea:
            return run_ea
    style = paragraph.style
    for _ in range(3):
        if style is None:
            return None
        element = getattr(style, "element", None)
        if element is not None:
            rPr = element.find(qn('w:rPr'))
            if rPr is not None:
                rFonts = rPr.find(qn('w:rFonts'))
                if rFonts is not None:
                    ea = rFonts.get(qn('w:eastAsia'))
                    if ea:
                        return ea
        style = style.base_style
    return None


def _find_paragraph(doc, text):
    """按精确文本找段落。"""
    for para in doc.paragraphs:
        if para.text.strip() == text:
            return para
    return None


# =============================================================================
# _extract_reference_default_cjk_font 测试
# =============================================================================

class TestExtractReferenceDefaultCjkFont:
    """从 reference doc Normal 样式提取默认中文字体"""

    def test_returns_normal_style_east_asia(self, tmp_path):
        """Normal 样式 eastAsia=KaiTi 的 docx → 返回 KaiTi"""
        from src.tools.word.md_to_word import _extract_reference_default_cjk_font

        ref = _make_reference_docx(tmp_path / "ref_kaiti.docx", east_asia="KaiTi")
        assert _extract_reference_default_cjk_font(str(ref)) == "KaiTi"

    def test_missing_file_returns_none(self, tmp_path):
        """路径不存在 → 返回 None（不抛异常）"""
        from src.tools.word.md_to_word import _extract_reference_default_cjk_font

        assert _extract_reference_default_cjk_font(str(tmp_path / "nope.docx")) is None

    def test_plain_docx_returns_none(self, tmp_path):
        """未设置 eastAsia 的普通 docx → 返回 None"""
        from src.tools.word.md_to_word import _extract_reference_default_cjk_font

        ref = _make_reference_docx(tmp_path / "ref_plain.docx")
        assert _extract_reference_default_cjk_font(str(ref)) is None

    def test_none_reference_returns_none(self):
        """reference 为 None（无参考文档）→ 返回 None"""
        from src.tools.word.md_to_word import _extract_reference_default_cjk_font

        assert _extract_reference_default_cjk_font(None) is None


# =============================================================================
# _ensure_cjk_fonts 样式感知测试
# =============================================================================

class TestEnsureCjkFontsStyleAware:
    """_ensure_cjk_fonts 仅在样式链未定义 eastAsia 时写 run 级默认值"""

    def test_style_with_east_asia_not_overridden(self):
        """段落用带 eastAsia 的自定义样式 → run 不被写 SimSun"""
        from src.tools.word.md_to_word import _ensure_cjk_fonts

        doc = Document()
        custom = doc.styles.add_style("CustomCJK", WD_STYLE_TYPE.PARAGRAPH)
        _set_style_east_asia(custom, "KaiTi")
        doc.add_paragraph("测试正文", style=custom)

        _ensure_cjk_fonts(doc)

        assert _run_east_asia(doc.paragraphs[0].runs[0]) is None

    def test_base_style_east_asia_not_overridden(self):
        """样式自身未定义但 base style 定义了 eastAsia → run 不被写默认值"""
        from src.tools.word.md_to_word import _ensure_cjk_fonts

        doc = Document()
        parent = doc.styles.add_style("ParentCJK", WD_STYLE_TYPE.PARAGRAPH)
        _set_style_east_asia(parent, "FangSong")
        child = doc.styles.add_style("ChildCJK", WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = parent
        doc.add_paragraph("继承字体", style=child)

        _ensure_cjk_fonts(doc)

        assert _run_east_asia(doc.paragraphs[0].runs[0]) is None

    def test_no_style_definition_gets_default(self):
        """无样式定义 eastAsia → run 被写默认 SimSun（向后兼容）"""
        from src.tools.word.md_to_word import _ensure_cjk_fonts

        doc = Document()
        doc.add_paragraph("普通正文")

        _ensure_cjk_fonts(doc)

        assert _run_east_asia(doc.paragraphs[0].runs[0]) == "SimSun"

    def test_custom_default_font_param(self):
        """传入自定义默认字体参数 → run 写该字体而非 SimSun"""
        from src.tools.word.md_to_word import _ensure_cjk_fonts

        doc = Document()
        doc.add_paragraph("普通正文")

        _ensure_cjk_fonts(doc, cjk_font="KaiTi")

        assert _run_east_asia(doc.paragraphs[0].runs[0]) == "KaiTi"

    def test_table_cell_paragraphs_style_aware(self):
        """表格 cell 段落同样做样式感知处理"""
        from src.tools.word.md_to_word import _ensure_cjk_fonts

        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        cell_plain = table.rows[0].cells[0]
        cell_plain.paragraphs[0].add_run("普通单元格")

        custom = doc.styles.add_style("CellCJK", WD_STYLE_TYPE.PARAGRAPH)
        _set_style_east_asia(custom, "KaiTi")
        cell_styled = table.rows[0].cells[1]
        para = cell_styled.paragraphs[0]
        para.style = custom
        para.add_run("样式单元格")

        _ensure_cjk_fonts(doc)

        assert _run_east_asia(cell_plain.paragraphs[0].runs[0]) == "SimSun"
        assert _run_east_asia(cell_styled.paragraphs[0].runs[0]) is None


# =============================================================================
# _convert_sync 集成测试（真实 Pandoc）
# =============================================================================

class TestConvertSyncWithReferenceTemplate:
    """客户模板 Normal eastAsia 传导到输出文档（真实 Pandoc 转换）"""

    def test_body_run_uses_reference_cjk_font(self, tmp_path):
        """reference.docx Normal eastAsia=KaiTi → 输出正文生效字体为 KaiTi 而非 SimSun"""
        from src.tools.word.md_to_word import _convert_sync

        ref = _make_reference_docx(tmp_path / "reference.docx", east_asia="KaiTi")
        doc = _convert_sync("# 标题\n\n正文段落", template=str(ref))

        body = _find_paragraph(doc, "正文段落")
        assert body is not None
        # run 级未被写 SimSun，生效 eastAsia 来自样式继承（或 run 级 KaiTi）
        assert _run_east_asia(body.runs[0]) != "SimSun"
        assert _effective_east_asia(body) == "KaiTi"

    def test_default_template_still_writes_sim_sun(self):
        """无客户模板（内置默认）→ run 级仍写 SimSun（既有行为不变）"""
        from src.tools.word.md_to_word import _convert_sync

        doc = _convert_sync("# 标题\n\n正文段落", template="default")

        body = _find_paragraph(doc, "正文段落")
        assert body is not None
        assert _run_east_asia(body.runs[0]) == "SimSun"


# =============================================================================
# _handle_md_to_word 的 template_file 测试
# =============================================================================

@patch("src.tools.word.md_to_word.save_as", return_value={"file_path": "/tmp/x.docx"})
@patch("src.tools.word.md_to_word.convert_async", new_callable=AsyncMock)
class TestHandleMdToWordTemplateFile:
    """_handle_md_to_word 的 template_file 解析与传递"""

    @pytest.mark.asyncio
    async def test_template_file_resolved_and_passed(self, mock_convert, mock_save, tmp_path):
        """template_file 指向存在的 docx → 解析后作为 template 传给 convert_async"""
        from src.tools.word.word_process_tool import PipelineContext, WordProcessTool

        ref = _make_reference_docx(tmp_path / "客户模板.docx", east_asia="KaiTi")
        tool = WordProcessTool()
        mock_convert.return_value = MagicMock(name="Document")

        result = await tool._handle_md_to_word(
            PipelineContext(content="# 标题\n\n正文"),
            {"template_file": str(ref)},
        )

        assert result["success"] is True
        template_arg = mock_convert.call_args.kwargs["template"]
        assert Path(template_arg).resolve() == ref.resolve()

    @pytest.mark.asyncio
    async def test_template_file_takes_precedence_over_named_template(
        self, mock_convert, mock_save, tmp_path
    ):
        """template_file 与 template 同时出现 → template_file 优先"""
        from src.tools.word.word_process_tool import PipelineContext, WordProcessTool

        ref = _make_reference_docx(tmp_path / "客户模板.docx", east_asia="KaiTi")
        tool = WordProcessTool()
        mock_convert.return_value = MagicMock(name="Document")

        await tool._handle_md_to_word(
            PipelineContext(content="# 标题\n\n正文"),
            {"template": "formal", "template_file": str(ref)},
        )

        template_arg = mock_convert.call_args.kwargs["template"]
        assert Path(template_arg).resolve() == ref.resolve()

    @pytest.mark.asyncio
    async def test_template_file_missing_returns_error(self, mock_convert, mock_save, tmp_path):
        """template_file 指向不存在文件 → fail loud 返回格式参考模板错误"""
        from src.tools.word.word_process_tool import PipelineContext, WordProcessTool

        tool = WordProcessTool()
        ghost = tmp_path / "ghost.docx"

        result = await tool._handle_md_to_word(
            PipelineContext(content="# 标题\n\n正文"),
            {"template_file": str(ghost)},
        )

        assert result["success"] is False
        assert "格式参考模板" in result["error"]
        mock_convert.assert_not_called()

    @pytest.mark.asyncio
    async def test_no_template_params_behavior_unchanged(self, mock_convert, mock_save):
        """不传 template_file/template → template=None（行为不变）"""
        from src.tools.word.word_process_tool import PipelineContext, WordProcessTool

        tool = WordProcessTool()
        mock_convert.return_value = MagicMock(name="Document")

        result = await tool._handle_md_to_word(
            PipelineContext(content="# 标题\n\n正文"), {}
        )

        assert result["success"] is True
        assert mock_convert.call_args.kwargs["template"] is None

    @pytest.mark.asyncio
    async def test_named_template_still_passed_through(self, mock_convert, mock_save):
        """仅传命名模板名 → 原样传给 convert（静默回退逻辑保留在 _convert_sync）"""
        from src.tools.word.word_process_tool import PipelineContext, WordProcessTool

        tool = WordProcessTool()
        mock_convert.return_value = MagicMock(name="Document")

        await tool._handle_md_to_word(
            PipelineContext(content="# 标题\n\n正文"), {"template": "nonexistent"}
        )

        assert mock_convert.call_args.kwargs["template"] == "nonexistent"


# =============================================================================
# md 来源守卫测试
# =============================================================================

class TestMdSourceGuard:
    """file_paths 中的 .docx 是格式参考，不作为 Markdown 源读取"""

    @pytest.mark.asyncio
    async def test_docx_only_file_paths_requires_context(self, tmp_path):
        """ctx 无 content、file_paths=[一个 .docx] → 返回需要提供 context 的错误"""
        from src.tools.word.word_process_tool import PipelineContext, WordProcessTool

        ref = _make_reference_docx(tmp_path / "格式参考.docx")
        tool = WordProcessTool()
        ctx = PipelineContext(file_paths=[str(ref)])

        with patch("src.tools.word.md_to_word.convert_async", new_callable=AsyncMock) as mock_conv:
            result = await tool._handle_md_to_word(ctx, {"template_file": str(ref)})

        assert result["success"] is False
        assert "context" in result["error"]
        mock_conv.assert_not_called()


# =============================================================================
# 路由 prompt 与工具描述测试
# =============================================================================

class TestRoutingPromptAndDescription:
    """路由 prompt 新规则与工具描述更新"""

    def test_routing_prompt_contains_template_file_rule(self):
        """ROUTING_PROMPT_PREFIX 含 template_file 规则与格式参考关键词"""
        from src.tools.word.word_router import ROUTING_PROMPT_PREFIX

        assert "template_file" in ROUTING_PROMPT_PREFIX
        assert "格式" in ROUTING_PROMPT_PREFIX
        assert "docx" in ROUTING_PROMPT_PREFIX

    def test_routing_prompt_contains_template_file_example(self):
        """示例区有 template_file 用法示例 JSON"""
        from src.tools.word.word_router import ROUTING_PROMPT_PREFIX

        assert '"template_file"' in ROUTING_PROMPT_PREFIX

    def test_tool_description_no_forbid_statement(self):
        """工具描述不再包含「不要向用户索要模板路径」"""
        from src.tools.word.word_process_tool import TOOL_DESCRIPTION

        assert "不要向用户索要模板路径" not in TOOL_DESCRIPTION
        assert "格式参考" in TOOL_DESCRIPTION
