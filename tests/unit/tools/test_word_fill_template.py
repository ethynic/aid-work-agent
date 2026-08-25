"""
Word 模板占位符填充增强（场景一）单元测试

覆盖：
- 多占位符语法（{{x}} / {{ x }} / {x} / 【x】 / [x] / %x%）替换与误报过滤
- 扫描盲区：嵌套表格、页眉中的表格、文本框（w:txbxContent）、内容控件（w:sdt）、超链接内 run
- scan_placeholders 扫描结果聚合与去重
- fill_template 反馈：total / per_variable / unmatched_variables / remaining_placeholders
- 表格行循环展开：{{#列表名}} ... {{/列表名}}
- word_process 接线：TaskType.ALL / 路由 prompt / handler
"""

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

import pytest

pytestmark = pytest.mark.tools


# =============================================================================
# 测试辅助
# =============================================================================

def _para_text(p_el) -> str:
    """拼接 XML 元素下所有 w:t 文本（避免 xpath 命名空间前缀问题）"""
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def _all_text(doc) -> str:
    """全文档（含页眉页脚）文本，用于断言占位符已彻底替换"""
    from src.tools.word.word_lib import iter_all_paragraphs, paragraph_text_runs
    return "".join(
        "".join(r.text for r in paragraph_text_runs(para))
        for para in iter_all_paragraphs(doc)
    )


def _append_textbox(doc, text):
    """在 body 下构造简化的 w:txbxContent 文本框（内含 w:p），验证段落迭代覆盖能力"""
    txbx = parse_xml(
        '<w:txbxContent %s><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:txbxContent>'
        % (nsdecls("w"), text)
    )
    doc.element.body.append(txbx)
    return txbx


def _append_sdt(doc, text):
    """在 body 下构造块级内容控件 w:sdt（内含 w:p）"""
    sdt = parse_xml(
        '<w:sdt %s><w:sdtContent><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:sdtContent></w:sdt>'
        % (nsdecls("w"), text)
    )
    doc.element.body.append(sdt)
    return sdt


def _append_hyperlink_paragraph(doc, text):
    """构造含 w:hyperlink（内含 w:r）的段落，para.runs 不含其中 run"""
    p = doc.add_paragraph()
    hyperlink = parse_xml(
        '<w:hyperlink %s w:anchor="bookmark1"><w:r><w:t>%s</w:t></w:r></w:hyperlink>'
        % (nsdecls("w"), text)
    )
    p._p.append(hyperlink)
    return p


# =============================================================================
# 占位符语法替换
# =============================================================================

class TestPlaceholderSyntaxFill:
    """fill_template 多语法替换测试"""

    def test_double_brace(self):
        """{{变量名}} 基本替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("甲方：{{甲方}}")
        result = fill_template(doc, {"甲方": "XX公司"})

        assert result["total"] == 1
        assert doc.paragraphs[0].text == "甲方：XX公司"

    def test_double_brace_with_inner_spaces(self):
        """{{ 变量名 }} 带内侧空格同样替换（按原文实际出现形式逐串替换）"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("乙方：{{ 乙方 }}")
        result = fill_template(doc, {"乙方": "YY公司"})

        assert result["total"] == 1
        assert doc.paragraphs[0].text == "乙方：YY公司"

    def test_single_brace(self):
        """{变量名} 语法替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("编号：{编号}")
        result = fill_template(doc, {"编号": "SO-001"})

        assert result["total"] == 1
        assert doc.paragraphs[0].text == "编号：SO-001"

    def test_fullwidth_bracket(self):
        """【变量名】语法替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("日期：【日期】")
        result = fill_template(doc, {"日期": "2026年8月"})

        assert result["total"] == 1
        assert doc.paragraphs[0].text == "日期：2026年8月"

    def test_square_bracket(self):
        """[变量名] 语法替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("备注：[备注]")
        result = fill_template(doc, {"备注": "无"})

        assert result["total"] == 1
        assert doc.paragraphs[0].text == "备注：无"

    def test_percent(self):
        """%变量名% 语法替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("税率：%税率%")
        result = fill_template(doc, {"税率": "6%"})

        assert result["total"] == 1
        assert doc.paragraphs[0].text == "税率：6%"

    def test_plain_bracket_text_not_touched(self):
        """正文普通方括号文本在替换时不被误伤（只替换显式提供的变量名）"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("参考文献 [1] 和 [参考文献] 见文末")
        result = fill_template(doc, {"备注": "无"})

        assert result["total"] == 0
        assert doc.paragraphs[0].text == "参考文献 [1] 和 [参考文献] 见文末"

    def test_multiple_occurrences_counted(self):
        """同一变量多次出现全部替换并正确计数"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("{{客户}}与{{客户}}签约")
        doc.add_paragraph("客户仍是{{客户}}")
        result = fill_template(doc, {"客户": "A公司"})

        assert result["total"] == 3
        assert result["per_variable"]["客户"] == 3
        assert "{{客户}}" not in _all_text(doc)

    def test_non_string_scalar_value(self):
        """数值型变量值自动转字符串"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("金额：{{金额}}")
        result = fill_template(doc, {"金额": 5000})

        assert result["total"] == 1
        assert doc.paragraphs[0].text == "金额：5000"

    def test_cross_run_placeholder(self):
        """占位符被 Word 拆成多个 run（常见：编辑导致 {{姓名}} 断开）仍可替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        p = doc.add_paragraph()
        p.add_run("甲方：{{")
        p.add_run("甲方")
        p.add_run("}}")

        result = fill_template(doc, {"甲方": "XX公司"})

        assert result["total"] == 1
        assert p.text == "甲方：XX公司"
        assert "{{" not in _all_text(doc)

    def test_cross_run_hyperlink_placeholder(self):
        """超链接内 run 参与跨 run 拼接（占位符一半在超链接内、一半在外）"""
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        from src.tools.word.template_manager import fill_template
        from src.tools.word.word_lib import paragraph_text_runs

        doc = Document()
        p = doc.add_paragraph()
        p.add_run("详情见%")
        hyperlink = parse_xml(
            '<w:hyperlink %s w:anchor="bookmark1"><w:r><w:t>链接变量%%</w:t></w:r></w:hyperlink>'
            % nsdecls("w")
        )
        p._p.append(hyperlink)

        result = fill_template(doc, {"链接变量": "官网"})

        assert result["total"] == 1
        assert "".join(r.text for r in paragraph_text_runs(p)) == "详情见官网"

    def test_value_with_xml_special_chars(self):
        """替换值含 XML 特殊字符（<>&）时安全转义，保存重开不损坏"""
        import tempfile
        import os

        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("备注：{{备注}}")
        result = fill_template(doc, {"备注": "a<b>&c</b>含%与{}"})

        assert result["total"] == 1
        assert doc.paragraphs[0].text == "备注：a<b>&c</b>含%与{}"
        # 保存重开验证 XML 转义正确（lxml 序列化时转义 <>&）
        fd, path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)
        try:
            doc.save(path)
            reopened = Document(path)
            assert reopened.paragraphs[0].text == "备注：a<b>&c</b>含%与{}"
        finally:
            os.remove(path)


# =============================================================================
# 占位符扫描
# =============================================================================

class TestScanPlaceholders:
    """scan_placeholders 扫描测试"""

    def test_mixed_syntax_scan(self):
        """多语法混合文档扫描：name/syntax/count 正确，按 count 降序"""
        from src.tools.word.template_manager import scan_placeholders

        doc = Document()
        doc.add_paragraph("{{甲方}}与{{甲方}}签约")
        doc.add_paragraph("{编号}、【日期】、[备注]、%税率%")

        result = scan_placeholders(doc)
        by_key = {(p["name"], p["syntax"]): p["count"] for p in result}

        assert by_key[("甲方", "double_brace")] == 2
        assert by_key[("编号", "single_brace")] == 1
        assert by_key[("日期", "fullwidth_bracket")] == 1
        assert by_key[("备注", "square_bracket")] == 1
        assert by_key[("税率", "percent")] == 1
        # 按 count 降序
        counts = [p["count"] for p in result]
        assert counts == sorted(counts, reverse=True)

    def test_false_positives_filtered(self):
        """[1] 方括号引用与 100% 百分数不误报"""
        from src.tools.word.template_manager import scan_placeholders

        doc = Document()
        doc.add_paragraph("完成率 100%，参考文献 [1] 见文末")

        names = [p["name"] for p in scan_placeholders(doc)]
        assert names == []

    def test_percent_range_not_matched(self):
        """50%-60% 百分数区间不误报出 -60 变量"""
        from src.tools.word.template_manager import scan_placeholders

        doc = Document()
        doc.add_paragraph("区间 50%-60% 之间")

        assert scan_placeholders(doc) == []

    def test_bare_percent_before_real_percent_placeholder(self):
        """裸百分数出现在 %变量% 之前时，真占位符不被跨匹配吞掉"""
        from src.tools.word.template_manager import fill_template, scan_placeholders

        doc = Document()
        doc.add_paragraph("立省10%，折扣：%折扣%off")

        # 扫描：识别出真变量 折扣，不产生 "，折扣：" 之类的垃圾变量名
        scanned = scan_placeholders(doc)
        assert [p["name"] for p in scanned] == ["折扣"]
        assert scanned[0]["syntax"] == "percent"

        # 填充：真占位符被替换
        result = fill_template(doc, {"折扣": "8折"})
        assert doc.paragraphs[0].text == "立省10%，折扣：8折off"
        assert result["per_variable"] == {"折扣": 1}
        assert result["unmatched_variables"] == []
        assert result["remaining_placeholders"] == []

    def test_empty_document(self):
        """空文档返回空列表"""
        from src.tools.word.template_manager import scan_placeholders

        assert scan_placeholders(Document()) == []

    def test_same_name_merged_in_variable_names(self):
        """不同语法同名变量在变量名列表中合并去重"""
        from src.tools.word.template_manager import scan_placeholders, unique_variable_names

        doc = Document()
        doc.add_paragraph("{{客户}}与{客户}签约")

        names = unique_variable_names(scan_placeholders(doc))
        assert names == ["客户"]

    def test_double_brace_not_double_counted(self):
        """{{name}} 内部 {name} 不被单花括号语法重复识别"""
        from src.tools.word.template_manager import scan_placeholders

        doc = Document()
        doc.add_paragraph("{{客户}}")

        result = scan_placeholders(doc)
        assert len(result) == 1
        assert result[0]["syntax"] == "double_brace"


# =============================================================================
# 扫描盲区（嵌套表格 / 页眉表格 / 文本框 / 内容控件 / 超链接）
# =============================================================================

class TestBlindSpotFill:
    """fill_template 盲区覆盖测试"""

    def test_nested_table_cell(self):
        """嵌套表格（cell 里的表格）占位符替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        outer = doc.add_table(rows=2, cols=2)
        outer.cell(0, 0).text = "{{表头}}"
        inner = outer.cell(1, 0).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "{{嵌套值}}"

        result = fill_template(doc, {"表头": "标题", "嵌套值": "内容"})

        assert result["total"] == 2
        assert outer.cell(0, 0).text == "标题"
        assert inner.cell(0, 0).text == "内容"

    def test_header_table(self):
        """页眉中的表格占位符替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        header = doc.sections[0].header
        header_table = header.add_table(rows=1, cols=2, width=100)
        header_table.cell(0, 0).text = "{{页眉变量}}"

        result = fill_template(doc, {"页眉变量": "机密"})

        assert result["total"] == 1
        assert header_table.cell(0, 0).text == "机密"

    def test_footer_paragraph(self):
        """页脚段落占位符替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        footer = doc.sections[0].footer
        footer.paragraphs[0].text = "{{公司名}}"

        result = fill_template(doc, {"公司名": "A公司"})

        assert result["total"] == 1
        assert footer.paragraphs[0].text == "A公司"

    def test_textbox_paragraph(self):
        """文本框（w:txbxContent 内 w:p）占位符替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        txbx = _append_textbox(doc, "项目：{{项目名}}")

        result = fill_template(doc, {"项目名": "智慧园区"})

        assert result["total"] == 1
        assert _para_text(txbx) == "项目：智慧园区"

    def test_sdt_content_control(self):
        """内容控件（w:sdt 内 w:p）占位符替换"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        sdt = _append_sdt(doc, "部门：{{部门}}")

        result = fill_template(doc, {"部门": "研发部"})

        assert result["total"] == 1
        assert _para_text(sdt) == "部门：研发部"

    def test_hyperlink_run(self):
        """超链接内 run 的占位符替换（paragraph.runs 不含 w:hyperlink 里的 run）"""
        from src.tools.word.template_manager import fill_template
        from src.tools.word.word_lib import paragraph_text_runs

        doc = Document()
        p = _append_hyperlink_paragraph(doc, "详情见{{链接变量}}")

        # 前置确认：python-docx 原生 runs 确实不含超链接内 run
        assert "{{链接变量}}" not in "".join(r.text for r in p.runs)

        result = fill_template(doc, {"链接变量": "官网"})

        assert result["total"] == 1
        assert "".join(r.text for r in paragraph_text_runs(p)) == "详情见官网"


# =============================================================================
# 填充反馈
# =============================================================================

class TestFillFeedback:
    """fill_template 返回反馈测试"""

    def test_unmatched_variables(self):
        """提供了但全文没匹配到的变量进入 unmatched_variables"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("{{甲方}}")
        result = fill_template(doc, {"甲方": "A公司", "幽灵变量": "x"})

        assert result["unmatched_variables"] == ["幽灵变量"]

    def test_remaining_placeholders(self):
        """模板有但没传值的占位符留在文档并进入 remaining_placeholders"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("{{甲方}}与{{乙方}}")
        result = fill_template(doc, {"甲方": "A公司"})

        assert doc.paragraphs[0].text == "A公司与{{乙方}}"
        assert result["remaining_placeholders"] == ["乙方"]

    def test_total_and_per_variable(self):
        """total 与 per_variable 计数正确"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        doc.add_paragraph("{{甲方}}、{{甲方}}、【金额】")
        result = fill_template(doc, {"甲方": "A公司", "金额": "100万"})

        assert result["per_variable"] == {"甲方": 2, "金额": 1}
        assert result["total"] == 3
        assert result["unmatched_variables"] == []
        assert result["remaining_placeholders"] == []


# =============================================================================
# 表格行循环展开
# =============================================================================

def _build_order_doc():
    """构造订单模板：表头行 + 循环模板行（{{#items}} 与 {{/items}} 同行）"""
    doc = Document()
    doc.add_paragraph("订单号：{{订单号}}")
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "名称"
    table.rows[0].cells[1].text = "数量"
    table.rows[0].cells[2].text = "单价"
    table.rows[1].cells[0].text = "{{#items}}{{名称}}"
    table.rows[1].cells[1].text = "{{数量}}"
    table.rows[1].cells[2].text = "{{单价}}{{/items}}"
    return doc


class TestRowLoop:
    """表格行循环展开测试"""

    ITEMS = [
        {"名称": "A商品", "数量": "2", "单价": "10元"},
        {"名称": "B商品", "数量": "5", "单价": "20元"},
        {"名称": "C商品", "数量": "8", "单价": "30元"},
    ]

    def test_three_entries_expanded(self):
        """3 条数据展开为 3 行：标记清除、字段替换、行数正确"""
        from src.tools.word.template_manager import fill_template

        doc = _build_order_doc()
        result = fill_template(doc, {"items": self.ITEMS, "订单号": "SO-001"})

        table = doc.tables[0]
        assert len(table.rows) == 4  # 表头 + 3 条明细
        data_rows = [[c.text for c in row.cells] for row in table.rows[1:]]
        assert data_rows == [
            ["A商品", "2", "10元"],
            ["B商品", "5", "20元"],
            ["C商品", "8", "30元"],
        ]
        # 标记清除干净，正文变量替换正常
        assert "{{" not in _all_text(doc)
        assert doc.paragraphs[0].text == "订单号：SO-001"
        assert result["unmatched_variables"] == []
        assert result["remaining_placeholders"] == []
        # 明细字段计入 per_variable
        assert result["per_variable"]["名称"] == 3
        assert result["per_variable"]["订单号"] == 1

    def test_empty_list_deletes_rows(self):
        """空列表 → 整个区间行删除"""
        from src.tools.word.template_manager import fill_template

        doc = _build_order_doc()
        result = fill_template(doc, {"items": [], "订单号": "SO-002"})

        table = doc.tables[0]
        assert len(table.rows) == 1  # 仅剩表头
        assert "{{" not in _all_text(doc)
        # items 已消费（区间行被删除），不应报 unmatched
        assert result["unmatched_variables"] == []

    def test_cross_row_interval(self):
        """开始/结束标记在不同行：区间 = [开始行, 结束行]，整块复制"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        table.rows[0].cells[0].text = "名称"
        table.rows[0].cells[1].text = "数量"
        table.rows[1].cells[0].text = "{{#items}}{{名称}}"
        table.rows[1].cells[1].text = "说明行"
        table.rows[2].cells[0].text = "{{数量}}"
        table.rows[2].cells[1].text = "{{/items}}"

        fill_template(doc, {"items": self.ITEMS[:2]})

        # 2 条数据 × 2 行模板 = 4 行明细 + 1 表头
        assert len(table.rows) == 5
        texts = [[c.text for c in row.cells] for row in table.rows[1:]]
        assert texts == [
            ["A商品", "说明行"],
            ["2", ""],
            ["B商品", "说明行"],
            ["5", ""],
        ]
        assert "{{" not in _all_text(doc)

    def test_global_variable_in_row(self):
        """非列表全局变量在行内可用（取全局 variables 兜底）"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "名称"
        table.rows[0].cells[1].text = "公司"
        table.rows[1].cells[0].text = "{{#items}}{{名称}}{{/items}}"
        table.rows[1].cells[1].text = "{{公司}}"

        fill_template(doc, {"items": self.ITEMS[:2], "公司": "XX公司"})

        assert len(table.rows) == 3
        assert table.rows[1].cells[1].text == "XX公司"
        assert table.rows[2].cells[1].text == "XX公司"

    def test_same_name_second_loop_skipped(self):
        """同名循环只处理第一个，第二个保留并进入 remaining_placeholders"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.rows[0].cells[0].text = "{{#items}}{{名称}}{{/items}}"
        table.rows[1].cells[0].text = "{{#items}}{{名称}}{{/items}}"

        result = fill_template(doc, {"items": self.ITEMS[:2]})

        # 第一个循环展开 2 行，第二个保持原样
        texts = [row.cells[0].text for row in table.rows]
        assert texts[:2] == ["A商品", "B商品"]
        assert "{{#items}}{{名称}}{{/items}}" in texts
        assert "#items" in result["remaining_placeholders"]

    def test_list_not_provided_no_expansion(self):
        """未提供列表数据时不展开，标记与字段占位符留在文档"""
        from src.tools.word.template_manager import fill_template

        doc = _build_order_doc()
        result = fill_template(doc, {"订单号": "SO-003"})

        assert len(doc.tables[0].rows) == 2  # 模板行保持原样
        # items 未出现在 variables 中：标记留在文档进入 remaining 反馈
        assert result["unmatched_variables"] == []
        assert "#items" in result["remaining_placeholders"]
        assert "名称" in result["remaining_placeholders"]

    def test_non_list_variable_no_expansion(self):
        """列表变量传了非列表值（如字符串）时不展开，计入 unmatched_variables"""
        from src.tools.word.template_manager import fill_template

        doc = _build_order_doc()
        result = fill_template(doc, {"items": "不是列表", "订单号": "SO-004"})

        assert len(doc.tables[0].rows) == 2
        assert result["unmatched_variables"] == ["items"]

    def test_two_loops_in_one_table(self):
        """同表两个不同列表循环：前一个展开增删行不影响后一个的行区间"""
        from src.tools.word.template_manager import fill_template

        doc = Document()
        table = doc.add_table(rows=6, cols=1)
        table.rows[0].cells[0].text = "表头"
        # 循环 a（行 1）：3 条数据展开为 3 行，行数净增使后续行号偏移
        table.rows[1].cells[0].text = "{{#a}}{{a名}}{{/a}}"
        table.rows[2].cells[0].text = "a附言行"
        table.rows[3].cells[0].text = "分隔行"
        # 循环 b（行 4）：2 条数据
        table.rows[4].cells[0].text = "{{#b}}{{b名}}{{/b}}"
        table.rows[5].cells[0].text = "b附言行"

        result = fill_template(doc, {
            "a": [{"a名": n} for n in ("a1", "a2", "a3")],
            "b": [{"b名": n} for n in ("b1", "b2")],
        })

        texts = [row.cells[0].text for row in table.rows]
        # 表头 + a×3 + a附言 + 分隔 + b×2 + b附言
        assert texts == ["表头", "a1", "a2", "a3", "a附言行", "分隔行", "b1", "b2", "b附言行"]
        assert "{{" not in _all_text(doc)
        assert result["remaining_placeholders"] == []
        assert result["unmatched_variables"] == []
        assert result["per_variable"] == {"a名": 3, "b名": 2}


# =============================================================================
# word_process 接线
# =============================================================================

class TestWordProcessWiring:
    """WordProcessTool 接线测试"""

    def test_task_type_all_contains_scan_placeholders(self):
        """TaskType.ALL 包含 scan_placeholders"""
        from src.tools.word.word_process_tool import TaskType

        assert TaskType.SCAN_PLACEHOLDERS == "scan_placeholders"
        assert "scan_placeholders" in TaskType.ALL

    def test_router_prompt_mentions_scan_placeholders(self):
        """路由 prompt 含 scan_placeholders 及多语法说明"""
        from src.tools.word.word_router import ROUTING_PROMPT_PREFIX

        assert "scan_placeholders" in ROUTING_PROMPT_PREFIX
        assert "%变量名%" in ROUTING_PROMPT_PREFIX
        assert "表格行循环" in ROUTING_PROMPT_PREFIX

    def test_router_parse_response_accepts_scan_placeholders(self):
        """路由器解析接受 scan_placeholders 任务"""
        from src.tools.word.word_router import WordRouter

        result = WordRouter()._parse_response(
            '{"task": "scan_placeholders,fill_template", "params": {}, "reason": "先扫描再填充"}'
        )
        assert result["task"] == "scan_placeholders,fill_template"

    def test_tool_description_mentions_placeholder_capabilities(self):
        """工具描述补充多语法 / 行循环 / 扫描说明"""
        from src.tools.word.word_process_tool import TOOL_DESCRIPTION

        assert "【变量】" in TOOL_DESCRIPTION
        assert "表格行循环" in TOOL_DESCRIPTION
        assert "scan_placeholders" in TOOL_DESCRIPTION

    @pytest.mark.asyncio
    async def test_handle_fill_template_returns_feedback(self):
        """_handle_fill_template 适配新 Dict 返回：total/per_variable/unmatched/remaining"""
        from src.tools.word.word_lib import WordFileHandler
        from src.tools.word.word_process_tool import PipelineContext, WordProcessTool

        doc = _build_order_doc()
        save = WordFileHandler.save_temp(doc, file_name="wft_feedback.docx")

        tool = WordProcessTool()
        ctx = PipelineContext(file_paths=[save["file_path"]])
        result = await tool._handle_fill_template(
            ctx,
            {"variables": {"items": TestRowLoop.ITEMS, "订单号": "SO-001", "幽灵": "x"}},
        )

        assert result["success"] is True
        assert result["variables_replaced"] == result["total"] == 10
        assert result["per_variable"]["名称"] == 3
        assert result["unmatched_variables"] == ["幽灵"]
        assert result["remaining_placeholders"] == []
        assert "占位符" in result["message"]
        assert "10" in result["message"]
        # 生成了新文件
        assert result["file_path"].endswith(".docx")

    @pytest.mark.asyncio
    async def test_handle_scan_placeholders(self):
        """_handle_scan_placeholders 返回占位符与变量名列表（只读，无新文件）"""
        from src.tools.word.word_lib import WordFileHandler
        from src.tools.word.word_process_tool import PipelineContext, WordProcessTool

        doc = Document()
        doc.add_paragraph("{{甲方}}与{编号}签约")
        save = WordFileHandler.save_temp(doc, file_name="wft_scan.docx")

        tool = WordProcessTool()
        ctx = PipelineContext(file_paths=[save["file_path"]])
        result = await tool._handle_scan_placeholders(ctx, {})

        assert result["success"] is True
        assert {p["name"] for p in result["placeholders"]} == {"甲方", "编号"}
        assert sorted(result["variable_names"]) == ["甲方", "编号"]
        assert "file_path" not in result  # 只读操作不落盘新文件

    @pytest.mark.asyncio
    async def test_handle_scan_placeholders_missing_file(self):
        """scan_placeholders 文件不存在时返回错误"""
        from unittest.mock import patch

        from src.tools.word.word_process_tool import PipelineContext, WordProcessTool

        tool = WordProcessTool()
        ctx = PipelineContext(file_paths=["file_ghost_wft"])
        with patch("src.core.redis_client.redis_client.hgetall", return_value={}):
            result = await tool._handle_scan_placeholders(ctx, {})

        assert result["success"] is False
        assert "文件不存在" in result["error"]

    @pytest.mark.asyncio
    async def test_execute_scan_then_merge_results(self):
        """execute 走 scan_placeholders 路由：_merge_results 合并 placeholders/variable_names"""
        from unittest.mock import AsyncMock

        from src.tools.word.word_lib import WordFileHandler
        from src.tools.word.word_process_tool import WordProcessTool

        doc = Document()
        doc.add_paragraph("【客户名称】")
        save = WordFileHandler.save_temp(doc, file_name="wft_scan_exec.docx")

        tool = WordProcessTool()
        mock_router = AsyncMock()
        mock_router.route.return_value = {"task": "scan_placeholders", "params": {}}
        tool._router = mock_router

        result = await tool.execute(context="查看模板占位符", file_paths=[save["file_path"]])

        assert result["success"] is True
        assert result["placeholders"] == [
            {"name": "客户名称", "syntax": "fullwidth_bracket", "count": 1}
        ]
        assert result["variable_names"] == ["客户名称"]

    @pytest.mark.asyncio
    async def test_execute_fill_template_merges_new_fields(self):
        """execute 走 fill_template 路由：_merge_results 透传新反馈字段"""
        from unittest.mock import AsyncMock

        from src.tools.word.word_lib import WordFileHandler
        from src.tools.word.word_process_tool import WordProcessTool

        doc = Document()
        doc.add_paragraph("{{甲方}}与{{乙方}}")
        save = WordFileHandler.save_temp(doc, file_name="wft_fill_exec.docx")

        tool = WordProcessTool()
        mock_router = AsyncMock()
        mock_router.route.return_value = {
            "task": "fill_template",
            "params": {"variables": {"甲方": "A公司"}},
        }
        tool._router = mock_router

        result = await tool.execute(context="填充模板", file_paths=[save["file_path"]])

        assert result["success"] is True
        assert result["variables_replaced"] == result["total"] == 1
        assert result["per_variable"] == {"甲方": 1}
        assert result["unmatched_variables"] == []
        assert result["remaining_placeholders"] == ["乙方"]
