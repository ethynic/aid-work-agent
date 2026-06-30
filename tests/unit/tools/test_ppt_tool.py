"""PPT 工具 Phase 0 回归测试。"""

from pathlib import Path

import pytest
from pptx import Presentation

pytestmark = pytest.mark.tools


class FakePlanner:
    def __init__(self, plan):
        self.plan = plan
        self.topic_calls = []
        self.content_calls = []

    async def plan_from_topic(self, topic, slide_count=None, theme_id=None):
        self.topic_calls.append(topic)
        return self.plan

    async def plan_from_content(self, content, theme_id=None):
        self.content_calls.append(content)
        return self.plan


def _sample_plan(title="测试演示文稿"):
    return {
        "title": title,
        "theme_id": 14,
        "style": "soft",
        "slides": [
            {
                "type": "cover",
                "title": title,
                "subtitle": "Phase 0 回归",
                "presenter": "AID",
                "date": "2026-06-30",
            },
            {
                "type": "content",
                "layout": "bullets",
                "title": "核心要点",
                "points": ["稳定生成", "可编辑 PPTX", "保持旧路径"],
            },
            {
                "type": "summary",
                "title": "总结",
                "takeaways": ["回归测试覆盖旧能力"],
                "next_steps": ["继续 Phase 1"],
            },
        ],
    }


def _assert_pptx(path, expected_slides):
    pptx_path = Path(path)
    assert pptx_path.exists()
    assert pptx_path.stat().st_size > 0
    prs = Presentation(str(pptx_path))
    assert len(prs.slides) == expected_slides


def _slide_texts(path):
    prs = Presentation(str(path))
    return [
        "\n".join(
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text
        )
        for slide in prs.slides
    ]


@pytest.mark.asyncio
async def test_topic_generates_ppt_with_python_pptx(tmp_path, monkeypatch):
    from src.tools.ppt.generator import PPTGenerator
    from src.tools.ppt.ppt_process_tool import PptProcessTool

    monkeypatch.setenv("PPT_RENDERER", "python_pptx")
    monkeypatch.setattr(PPTGenerator, "_get_output_dir", lambda self: tmp_path)

    tool = PptProcessTool()
    planner = FakePlanner(_sample_plan("企业 AI 应用"))
    tool._planner = planner

    result = await tool.execute(context="企业 AI 应用")

    assert result["success"] is True
    assert result["renderer"] == "python_pptx"
    assert result["slide_count"] == 3
    assert planner.topic_calls == ["企业 AI 应用"]
    _assert_pptx(result["file_path"], 3)
    assert "企业 AI 应用" in _slide_texts(result["file_path"])[0]


@pytest.mark.asyncio
async def test_markdown_outline_generates_ppt_with_python_pptx(tmp_path, monkeypatch):
    from src.tools.ppt.generator import PPTGenerator
    from src.tools.ppt.ppt_process_tool import PptProcessTool

    monkeypatch.setenv("PPT_RENDERER", "python_pptx")
    monkeypatch.setattr(PPTGenerator, "_get_output_dir", lambda self: tmp_path)

    markdown = "# 数据治理方案\n\n## 背景\n\n- 数据分散\n- 口径不一\n\n## 路线\n\n1. 标准化\n2. 自动化"
    tool = PptProcessTool()
    planner = FakePlanner(_sample_plan("数据治理方案"))
    tool._planner = planner

    result = await tool.execute(context=markdown)

    assert result["success"] is True
    assert planner.content_calls == [markdown]
    _assert_pptx(result["file_path"], 3)
    assert "核心要点" in _slide_texts(result["file_path"])[1]


@pytest.mark.asyncio
async def test_template_pptx_with_content_generates_ppt(tmp_path, monkeypatch):
    from src.tools.ppt.ppt_process_tool import PptProcessTool
    from src.tools.ppt.template_analyzer import TemplateAnalyzer

    monkeypatch.setenv("PPT_RENDERER", "python_pptx")
    monkeypatch.setattr(TemplateAnalyzer, "_get_output_dir", lambda self: tmp_path)

    template_path = tmp_path / "template.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "模板标题"
    prs.save(str(template_path))

    tool = PptProcessTool()
    planner = FakePlanner(_sample_plan("模板生成"))
    tool._planner = planner

    result = await tool.execute(
        context="# 模板生成\n\n- 使用模板布局\n- 填充内容",
        file_paths=[str(template_path)],
    )

    assert result["success"] is True
    assert "基于模板" in result["message"]
    _assert_pptx(result["file_path"], 3)
    slide_texts = _slide_texts(result["file_path"])
    assert "模板生成" in slide_texts[0]
    assert "稳定生成" in slide_texts[1]


@pytest.mark.asyncio
async def test_empty_input_returns_clear_error():
    from src.tools.ppt.ppt_process_tool import PptProcessTool

    result = await PptProcessTool().execute(context="  ")

    assert result["success"] is False
    assert "请提供主题或内容" in result["error"]


@pytest.mark.asyncio
async def test_blank_file_paths_are_treated_as_empty_input():
    from src.tools.ppt.ppt_process_tool import PptProcessTool

    result = await PptProcessTool().execute(context=None, file_paths=["", "  "])

    assert result["success"] is False
    assert "请提供主题或内容" in result["error"]


def test_ppt_config_reads_switches(monkeypatch):
    from src.tools.ppt.ppt_config import get_ppt_config

    monkeypatch.setenv("PPT_RENDERER", "pptxgenjs")
    monkeypatch.setenv("PPT_ENABLE_HTML_EXPORT", "true")
    monkeypatch.setenv("PPT_QA_STRICT", "1")

    config = get_ppt_config()

    assert config.renderer == "pptxgenjs"
    assert config.enable_html_export is True
    assert config.qa_strict is True


def test_ppt_config_invalid_renderer_falls_back(monkeypatch):
    from src.tools.ppt.ppt_config import get_ppt_config

    monkeypatch.setenv("PPT_RENDERER", "unknown")

    assert get_ppt_config().renderer == "python_pptx"


def test_ppt_capabilities_reports_missing_node(monkeypatch, tmp_path):
    from src.tools.ppt import ppt_capabilities

    monkeypatch.setattr(ppt_capabilities.shutil, "which", lambda command: None)

    caps = ppt_capabilities.get_capabilities(renderer_dir=tmp_path / "renderer-node")

    assert caps["node_renderer"]["available"] is False
    assert "Node.js" in caps["node_renderer"]["reason"]
    assert caps["render_validation"]["available"] is False


def test_ppt_capabilities_reports_missing_playwright_browser(monkeypatch):
    from src.tools.ppt import ppt_capabilities

    monkeypatch.setattr(
        ppt_capabilities,
        "probe_playwright",
        lambda package_available: {
            "available": False,
            "browser": "chromium",
            "reason": "Playwright Chromium 未安装或不可用",
        },
    )

    caps = ppt_capabilities.get_capabilities()

    assert caps["html_export"]["available"] is False
    assert "Chromium" in caps["html_export"]["reason"]


@pytest.mark.asyncio
async def test_pptxgenjs_renderer_missing_node_falls_back_without_stack(tmp_path, monkeypatch):
    from src.tools.ppt import ppt_capabilities
    from src.tools.ppt.generator import PPTGenerator
    from src.tools.ppt.ppt_process_tool import PptProcessTool

    monkeypatch.setenv("PPT_RENDERER", "pptxgenjs")
    monkeypatch.setattr(ppt_capabilities.shutil, "which", lambda command: None)
    monkeypatch.setattr(PPTGenerator, "_get_output_dir", lambda self: tmp_path)

    tool = PptProcessTool()
    tool._planner = FakePlanner(_sample_plan("Node 缺失回退"))

    result = await tool.execute(context="Node 缺失回退")

    assert result["success"] is True
    assert result["renderer"] == "python_pptx"
    assert "PptxGenJS 渲染器依赖不可用" in result["warnings"][0]
    assert "Traceback" not in str(result)
    _assert_pptx(result["file_path"], 3)


@pytest.mark.asyncio
async def test_unexpected_value_error_does_not_leak_details(monkeypatch):
    from src.tools.ppt.ppt_process_tool import PptProcessTool

    tool = PptProcessTool()

    async def raise_sensitive_error(context, file_paths):
        raise ValueError("api_key=secret-value C:\\tenant\\private.pptx")

    monkeypatch.setattr(tool, "_handle_auto", raise_sensitive_error)

    result = await tool.execute(context="测试异常")

    assert result["success"] is False
    assert result["error"] == "PPT生成失败，请检查输入内容或稍后重试"
    assert "secret-value" not in str(result)
    assert "private.pptx" not in str(result)
