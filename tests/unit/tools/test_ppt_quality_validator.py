"""Phase 7 unified PPTX quality validation tests."""

import json
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from src.tools.ppt.quality_validator import PPTQualityValidator

pytestmark = pytest.mark.tools


def _presentation(path: Path, *, blank: bool = False, image: bool = False) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if not blank:
        slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5), Inches(5), Inches(1)
        ).text = "可交付内容"
    if image:
        image_path = path.with_suffix(".png")
        Image.new("RGB", (40, 30), "#1677ff").save(image_path)
        slide.shapes.add_picture(
            str(image_path), Inches(1), Inches(2), Inches(2), Inches(1.5)
        )
    prs.save(path)
    return path


def _layout(*, raster_readable: bool = True) -> dict:
    return {
        "version": "1.0",
        "slide_count": 1,
        "node_count": 1,
        "slides": [
            {
                "id": "s1",
                "node_count": 1,
                "node_types": {"raster": 1},
                "slide_size": {"width": 10, "height": 7.5},
                "objects": [
                    {
                        "type": "raster",
                        "bounds": {"x": 0, "y": 0, "w": 5, "h": 7.5},
                        "readable": raster_readable,
                        "source_size_bytes": 100 if raster_readable else 0,
                    }
                ],
            }
        ],
    }


def test_real_pptx_writes_layout_and_qa_report_with_stable_fields(
    tmp_path, monkeypatch
):
    path = _presentation(tmp_path / "valid.pptx", image=True)
    monkeypatch.setattr("src.tools.ppt.quality_validator.shutil.which", lambda _: None)

    report = PPTQualityValidator().validate(
        path, expected_slide_count=1, layout=_layout()
    )

    summary = report["summary"]
    assert summary["file_valid"] is True
    assert summary["openable"] is True
    assert summary["slide_count"] == 1
    assert summary["nonempty_slide_count"] == 1
    assert summary["image_count"] == 1
    assert summary["invalid_image_count"] == 0
    assert summary["raster_layer_count"] == 1
    assert summary["raster_readable_count"] == 1
    assert summary["editable_ratio"] == 0.5
    assert summary["preview_status"] == "unavailable"
    assert summary["deliverable"] is True
    layout_path = tmp_path / "valid.layout.json"
    report_path = tmp_path / "valid.qa-report.json"
    assert layout_path.is_file()
    assert report_path.is_file()
    serialized = report_path.read_text(encoding="utf-8")
    assert str(path) not in serialized
    assert json.loads(serialized)["summary"] == summary


def test_corrupt_pptx_is_sanitized_and_strict_failure(tmp_path):
    path = tmp_path / "customer-secret.pptx"
    path.write_bytes(b"not a pptx package")

    report = PPTQualityValidator().validate(path, strict=True)

    assert report["status"] == "failed"
    assert report["summary"]["openable"] is False
    assert report["summary"]["deliverable"] is False
    assert "损坏或无法打开" in report["errors"][0]
    assert str(path) not in json.dumps(report, ensure_ascii=False)


def test_empty_page_count_mismatch_and_out_of_bounds_are_reported(tmp_path):
    path = _presentation(tmp_path / "empty.pptx", blank=True)
    prs = Presentation(path)
    prs.slides[0].shapes.add_textbox(
        Inches(12.5), Inches(1), Inches(2), Inches(1)
    )
    prs.save(path)

    report = PPTQualityValidator().validate(
        path, expected_slide_count=2, strict=False
    )

    assert report["summary"]["page_count_matches"] is False
    assert report["summary"]["empty_slide_count"] == 1
    assert report["summary"]["out_of_bounds_count"] == 1
    assert report["summary"]["qa_passed"] is False
    assert report["summary"]["deliverable"] is True


def test_zero_slide_deck_and_preview_page_mismatch_are_errors(
    tmp_path, monkeypatch
):
    zero_slide = tmp_path / "zero.pptx"
    Presentation().save(zero_slide)
    monkeypatch.setattr("src.tools.ppt.quality_validator.shutil.which", lambda _: None)
    zero_report = PPTQualityValidator().validate(zero_slide)
    assert any("不包含幻灯片" in item for item in zero_report["errors"])

    path = _presentation(tmp_path / "preview-mismatch.pptx")
    validator = PPTQualityValidator()
    monkeypatch.setattr(
        validator,
        "_render_preview",
        lambda _: {
            "summary": {"preview_status": "png", "preview_page_count": 2},
            "warnings": [],
        },
    )
    mismatch_report = validator.validate(path)
    assert any("预览页数" in item for item in mismatch_report["errors"])


def test_corrupt_embedded_image_and_unreadable_raster_are_errors(tmp_path):
    source = _presentation(tmp_path / "source.pptx", image=True)
    corrupt = tmp_path / "corrupt-image.pptx"
    with zipfile.ZipFile(source, "r") as incoming, zipfile.ZipFile(
        corrupt, "w"
    ) as outgoing:
        for info in incoming.infolist():
            content = incoming.read(info.filename)
            if info.filename.startswith("ppt/media/"):
                content = b"broken-image"
            outgoing.writestr(info, content)

    report = PPTQualityValidator().validate(
        corrupt, layout=_layout(raster_readable=False)
    )

    assert report["summary"]["invalid_image_count"] == 1
    assert report["summary"]["unreadable_raster_count"] == 1
    assert any("无效图片" in error for error in report["errors"])
    assert any("raster" in error for error in report["errors"])


def test_libreoffice_conversion_failure_degrades_to_warning(tmp_path, monkeypatch):
    path = _presentation(tmp_path / "preview.pptx")
    monkeypatch.setattr(
        "src.tools.ppt.quality_validator.shutil.which",
        lambda name: "soffice" if name == "soffice" else None,
    )

    class Completed:
        returncode = 1

    monkeypatch.setattr(
        "src.tools.ppt.quality_validator.subprocess.run",
        lambda *args, **kwargs: Completed(),
    )
    report = PPTQualityValidator().validate(path)

    assert report["summary"]["preview_status"] == "failed"
    assert report["summary"]["qa_passed"] is True
    assert any("已降级" in warning for warning in report["warnings"])


def test_tool_strict_blocks_and_non_strict_preserves_delivery(tmp_path, monkeypatch):
    from src.tools.ppt.ppt_process_tool import PptProcessTool

    path = tmp_path / "broken.pptx"
    path.write_bytes(b"broken")
    monkeypatch.setattr("src.tools.ppt.quality_validator.shutil.which", lambda _: None)
    tool = PptProcessTool()

    monkeypatch.setenv("PPT_QA_STRICT", "false")
    non_strict = tool._apply_quality_validation(
        {"success": True, "file_path": str(path), "slide_count": 1}
    )
    assert non_strict["success"] is True
    assert non_strict["file_path"] == str(path)
    assert non_strict["qa_summary"]["deliverable"] is True

    monkeypatch.setenv("PPT_QA_STRICT", "true")
    strict = tool._apply_quality_validation(
        {"success": True, "file_path": str(path), "slide_count": 1}
    )
    assert strict["success"] is False
    assert "file_path" not in strict
    assert strict["qa_summary"]["deliverable"] is False
    assert str(path) not in strict["error"]
    assert not path.exists()
    assert not path.with_suffix(".layout.json").exists()
    assert not path.with_suffix(".qa-report.json").exists()


def test_tool_validates_both_output_files_independently(tmp_path, monkeypatch):
    from src.tools.ppt.ppt_process_tool import PptProcessTool

    primary = _presentation(tmp_path / "editable.pptx")
    alternate = _presentation(tmp_path / "high-fidelity.pptx")
    monkeypatch.setenv("PPT_QA_STRICT", "true")
    monkeypatch.setattr("src.tools.ppt.quality_validator.shutil.which", lambda _: None)

    result = PptProcessTool()._apply_quality_validation(
        {
            "success": True,
            "file_path": str(primary),
            "alternate_file_path": str(alternate),
            "slide_count": 1,
            "qa_summary": {"editable_ratio": 0.8},
            "_layout_qa": _layout(),
            "_alternate_layout_qa": _layout(),
        }
    )

    assert result["success"] is True
    assert result["qa_summary"]["editable_ratio"] == 0.8
    assert result["alternate_qa_summary"]["qa_passed"] is True
    assert result["alternate_qa_summary"]["raster_layer_count"] == 1
    assert "_layout_qa" not in result
    assert "_alternate_layout_qa" not in result


def test_strict_alternate_failure_removes_both_outputs(tmp_path, monkeypatch):
    from src.tools.ppt.ppt_process_tool import PptProcessTool

    primary = _presentation(tmp_path / "editable.pptx")
    alternate = tmp_path / "high-fidelity.pptx"
    alternate.write_bytes(b"broken")
    monkeypatch.setenv("PPT_QA_STRICT", "true")
    monkeypatch.setattr("src.tools.ppt.quality_validator.shutil.which", lambda _: None)

    result = PptProcessTool()._apply_quality_validation(
        {
            "success": True,
            "file_path": str(primary),
            "alternate_file_path": str(alternate),
            "slide_count": 1,
        }
    )

    assert result["success"] is False
    assert not primary.exists()
    assert not alternate.exists()
    assert not primary.with_suffix(".layout.json").exists()
    assert not alternate.with_suffix(".qa-report.json").exists()
