"""
Word 文档处理工具单元测试

测试 WordProcessTool 及其子模块（mock python-docx 和 markitdown）
"""

import pytest
import re
from unittest.mock import AsyncMock, MagicMock, patch
from pathlib import Path

pytestmark = pytest.mark.tools


# =============================================================================
# 工具定义和参数验证
# =============================================================================

class TestWordProcessToolDefinition:
    """WordProcessTool 工具定义测试"""

    def test_tool_name(self):
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()
        assert tool.name == "word_process"

    def test_tool_display_name(self):
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()
        assert tool.display_name == "Word文档处理"

    def test_tool_definition_has_schema(self):
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()
        defn = tool.to_tool_definition()
        assert defn["name"] == "word_process"
        assert "input_schema" in defn
        schema = defn["input_schema"]
        properties = schema.get("properties", {})
        assert "instruction" in properties
        assert "content" in properties
        assert "content_type" in properties
        assert "output_name" in properties
        assert "context" in properties
        assert "file_paths" in properties
        # task 和 params 不再暴露给 Agent
        assert "task" not in properties
        assert "params" not in properties

    def test_tool_has_description(self):
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()
        assert len(tool.description) > 50
        assert "Word" in tool.description
        assert "触发规则" in tool.description


# =============================================================================
# Pipeline 执行测试（通过 mock router 直接测试各 handler）
# =============================================================================

class TestWordProcessToolExecution:
    """工具执行测试"""

    @pytest.mark.asyncio
    async def test_execute_no_context_no_files(self):
        """无 context 无 file_paths 时路由返回空"""
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()

        mock_router = AsyncMock()
        mock_router.route.return_value = {"task": "", "params": {}, "error": "缺少 context 和 file_paths"}
        tool._router = mock_router

        result = await tool.execute()
        assert result["success"] is False

    @pytest.mark.asyncio
    async def test_execute_router_returns_invalid_task(self):
        """路由返回无效操作时返回错误"""
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()

        mock_router = AsyncMock()
        mock_router.route.return_value = {"task": "unknown_op", "params": {}, "error": ""}
        tool._router = mock_router

        result = await tool.execute(context="测试")
        assert result["success"] is False
        assert "无效操作" in result["error"]

    @pytest.mark.asyncio
    async def test_execute_router_exception(self):
        """路由抛异常时返回错误"""
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()

        mock_router = AsyncMock()
        mock_router.route.side_effect = Exception("LLM down")
        tool._router = mock_router

        result = await tool.execute(context="读取文件", file_paths=["/tmp/test.docx"])
        assert result["success"] is False

    @pytest.mark.asyncio
    async def test_execute_list_templates(self):
        """路由返回 list_templates 时正常执行"""
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()

        mock_router = AsyncMock()
        mock_router.route.return_value = {"task": "list_templates", "params": {}}
        tool._router = mock_router

        result = await tool.execute(context="查看模板")
        assert result["success"] is True
        assert "templates" in result

    @pytest.mark.asyncio
    async def test_execute_read_missing_file(self):
        """路由返回 read 但文件不存在时失败"""
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()

        mock_router = AsyncMock()
        mock_router.route.return_value = {"task": "read", "params": {}}
        tool._router = mock_router

        result = await tool.execute(context="读取文件", file_paths=["/nonexistent.docx"])
        assert result["success"] is False

    @pytest.mark.asyncio
    async def test_execute_md_to_word_no_context(self):
        """路由返回 md_to_word 但无 context 且无 md 文件时失败"""
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()

        mock_router = AsyncMock()
        mock_router.route.return_value = {"task": "md_to_word", "params": {}}
        tool._router = mock_router

        # 无 context 且 file_paths 指向不存在的 .md 文件
        result = await tool.execute(file_paths=["/nonexistent.md"])
        assert result["success"] is False

    @pytest.mark.asyncio
    async def test_resolve_markdown_context_directly_to_word(self):
        """纯 Markdown 正文已足够明确时，不依赖内部 LLM 路由。"""
        from src.tools.word.word_process_tool import WordProcessTool
        tool = WordProcessTool()

        mock_router = AsyncMock()
        tool._router = mock_router
        context = (
            "## 安顺坝陵河大桥3天2晚行程\n\n"
            "**出发日期**：2026年7月11日（周六）\n\n"
            "| 天数 | 时段 | 行程安排 |\n"
            "|------|------|----------|\n"
            "| D1 | 上午 | 贵阳接站 |"
        )

        result = await tool._resolve_task(context, None)

        assert result["task"] == "md_to_word"
        assert result["params"]["template"] == "default"
        mock_router.route.assert_not_called()

    def test_extract_markdown_body_removes_tool_instruction(self):
        """生成指令只用于路由，不应写入 Word 正文。"""
        from src.tools.word.word_process_tool import WordProcessTool

        context = (
            "生成一份贵州安顺坝陵河3天2晚行程Word文档，格式为Markdown表格。\n\n"
            "## 安顺坝陵河大桥3天2晚行程\n\n"
            "| 天数 | 时段 | 行程安排 |\n"
            "|------|------|----------|\n"
            "| D1 | 上午 | 贵阳接站 |"
        )

        body = WordProcessTool._extract_markdown_body(context)

        assert body.startswith("## 安顺坝陵河大桥3天2晚行程")
        assert "生成一份" not in body

    def test_normalize_legacy_context_splits_instruction_and_content(self):
        """旧 context 调用会拆分为 instruction + content。"""
        from src.tools.word.word_process_tool import WordProcessTool

        tool = WordProcessTool()
        context = (
            "生成一份贵州安顺坝陵河3天2晚行程Word文档，格式为Markdown表格。\n\n"
            "## 安顺坝陵河大桥3天2晚行程\n\n"
            "| 天数 | 时段 | 行程安排 |\n"
            "|------|------|----------|\n"
            "| D1 | 上午 | 贵阳接站 |"
        )

        normalized = tool._normalize_input(
            context=context,
            instruction=None,
            content=None,
            content_type=None,
            output_name=None,
        )

        assert normalized["instruction"].startswith("生成一份贵州安顺坝陵河")
        assert normalized["content"].startswith("## 安顺坝陵河大桥3天2晚行程")
        assert "生成一份" not in normalized["content"]
        assert normalized["content_type"] == "auto"

    @pytest.mark.asyncio
    async def test_execute_structured_instruction_content_to_word(self):
        """新入参 instruction + content 路径不污染正文，并优先使用 output_name。"""
        from src.tools.word.word_process_tool import WordProcessTool

        tool = WordProcessTool()
        mock_router = AsyncMock()
        tool._router = mock_router
        content = (
            "## 安顺坝陵河大桥3天2晚行程\n\n"
            "**出发日期**：2026年7月11日（周六）\n\n"
            "| 天数 | 时段 | 行程安排 |\n"
            "|------|------|----------|\n"
            "| D1 | 上午 | 贵阳接站 |"
        )

        with patch("src.tools.word.md_to_word.convert_async", new_callable=AsyncMock) as mock_convert, \
             patch("src.tools.word.md_to_word.save_as") as mock_save:
            mock_convert.return_value = MagicMock()
            mock_save.return_value = {"file_path": "/tmp/custom.docx", "file_size": 123}

            result = await tool.execute(
                instruction="生成贵州安顺坝陵河3天2晚行程Word文档",
                content=content,
                content_type="markdown",
                output_name="安顺坝陵河3天2晚行程.docx",
            )

        converted_text = mock_convert.call_args.args[0]
        assert result["success"] is True
        assert converted_text.startswith("## 安顺坝陵河大桥3天2晚行程")
        assert "生成贵州安顺坝陵河" not in converted_text
        assert mock_save.call_args.kwargs["file_name"] == "安顺坝陵河3天2晚行程.docx"
        mock_router.route.assert_not_called()

    @pytest.mark.asyncio
    async def test_execute_structured_markdown_simple_document_bypasses_router(self):
        """content_type=markdown 且 instruction 明确生成 Word 时，简单 Markdown 也走规则路由。"""
        from src.tools.word.word_process_tool import WordProcessTool

        tool = WordProcessTool()
        mock_router = AsyncMock()
        tool._router = mock_router
        content = "# 项目说明\n\n这是一个没有表格和列表的简单 Markdown 文档。"

        with patch("src.tools.word.md_to_word.convert_async", new_callable=AsyncMock) as mock_convert, \
             patch("src.tools.word.md_to_word.save_as") as mock_save:
            mock_convert.return_value = MagicMock()
            mock_save.return_value = {"file_path": "/tmp/project.docx", "file_size": 123}

            result = await tool.execute(
                instruction="生成Word文档",
                content=content,
                content_type="markdown",
            )

        assert result["success"] is True
        assert mock_convert.call_args.args[0] == content
        mock_router.route.assert_not_called()

    @pytest.mark.asyncio
    async def test_execute_output_name_is_sanitized(self):
        """显式 output_name 只作为文件名使用，不允许携带路径片段。"""
        from src.tools.word.word_process_tool import WordProcessTool

        tool = WordProcessTool()
        content = "# 项目说明\n\n这是一个简单 Markdown 文档。"

        with patch("src.tools.word.md_to_word.convert_async", new_callable=AsyncMock) as mock_convert, \
             patch("src.tools.word.md_to_word.save_as") as mock_save:
            mock_convert.return_value = MagicMock()
            mock_save.return_value = {"file_path": "/tmp/project.docx", "file_size": 123}

            result = await tool.execute(
                instruction="生成Word文档",
                content=content,
                content_type="markdown",
                output_name="../阶段一:验证?.docx",
            )

        assert result["success"] is True
        assert mock_save.call_args.kwargs["file_name"] == "阶段一验证.docx"

    @pytest.mark.asyncio
    async def test_handle_md_to_word_uses_clean_markdown_body(self):
        """md_to_word 执行前清理 context，避免指令行进入 docx。"""
        from src.tools.word.word_process_tool import PipelineContext, WordProcessTool

        tool = WordProcessTool()
        context = (
            "生成一份贵州安顺坝陵河3天2晚行程Word文档，格式为Markdown表格。\n\n"
            "## 安顺坝陵河大桥3天2晚行程\n\n"
            "| 天数 | 时段 | 行程安排 |\n"
            "|------|------|----------|\n"
            "| D1 | 上午 | 贵阳接站 |"
        )

        with patch("src.tools.word.md_to_word.convert_async", new_callable=AsyncMock) as mock_convert, \
             patch("src.tools.word.md_to_word.save_as") as mock_save:
            mock_convert.return_value = MagicMock()
            mock_save.return_value = {"file_path": "/tmp/安顺坝陵河大桥3天2晚行程.docx", "file_size": 123}

            result = await tool._handle_md_to_word(PipelineContext(context=context), {})

        converted_text = mock_convert.call_args.args[0]
        assert result["success"] is True
        assert converted_text.startswith("## 安顺坝陵河大桥3天2晚行程")
        assert "生成一份" not in converted_text
        assert mock_save.call_args.kwargs["file_name"] == "安顺坝陵河大桥3天2晚行程.docx"


# =============================================================================
# WordReader 测试
# =============================================================================

class TestWordReader:
    """WordReader 模块测试"""

    def test_read_content_file_not_found(self):
        from src.tools.word.word_reader import read_content
        result = read_content("/nonexistent.docx")
        assert result["success"] is False
        assert "不存在" in result["error"]

    def test_analyze_structure_file_not_found(self):
        from src.tools.word.word_reader import analyze_structure
        result = analyze_structure("/nonexistent.docx")
        assert result["success"] is False

    def test_extract_metadata_file_not_found(self):
        from src.tools.word.word_reader import extract_metadata
        result = extract_metadata("/nonexistent.docx")
        assert result["success"] is False


# =============================================================================
# TemplateManager 测试
# =============================================================================

class TestTemplateManager:
    """TemplateManager 模块测试"""

    def test_list_templates(self):
        from src.tools.word.template_manager import list_templates
        templates = list_templates()
        assert len(templates) == 4
        names = [t["name"] for t in templates]
        assert "default" in names
        assert "formal" in names
        assert "modern" in names
        assert "report" in names

    def test_get_template_default(self):
        from src.tools.word.template_manager import get_template
        template = get_template("default")
        assert template["name"] == "default"
        assert "body" in template

    def test_get_template_not_found(self):
        from src.tools.word.template_manager import get_template
        with pytest.raises(ValueError, match="未知模板"):
            get_template("nonexistent")

    def test_fill_template_replaces_variables(self):
        from src.tools.word.template_manager import fill_template
        from docx import Document

        doc = Document()
        doc.add_paragraph("{{甲方}}和{{乙方}}签订合同")
        doc.add_paragraph("金额：[金额]")

        result = fill_template(doc, {
            "甲方": "XX公司",
            "乙方": "YY公司",
            "金额": "50万",
        })
        # fill_template 返回替换反馈 Dict（total/per_variable/...）
        assert result["total"] >= 3
        assert "XX公司" in doc.paragraphs[0].text
        assert "YY公司" in doc.paragraphs[0].text
        assert "50万" in doc.paragraphs[1].text


# =============================================================================
# WordModifier 测试
# =============================================================================

class TestWordModifier:
    """WordModifier 模块测试"""

    def test_replace_text(self):
        from src.tools.word.word_modifier import batch_modify
        from docx import Document

        doc = Document()
        doc.add_paragraph("Hello World")
        result = batch_modify(doc, [
            {"type": "replace_text", "target": "World", "replacement": "Python"}
        ])
        assert result["success"] is True
        assert result["operations_applied"] == 1
        assert "Python" in doc.paragraphs[0].text

    def test_insert_paragraph(self):
        from src.tools.word.word_modifier import batch_modify
        from docx import Document

        doc = Document()
        doc.add_paragraph("First")
        result = batch_modify(doc, [
            {"type": "insert_paragraph", "after_index": 0, "text": "Inserted"}
        ])
        assert result["success"] is True

    def test_delete_paragraph(self):
        from src.tools.word.word_modifier import batch_modify
        from docx import Document

        doc = Document()
        doc.add_paragraph("First")
        doc.add_paragraph("Second")
        result = batch_modify(doc, [
            {"type": "delete_paragraph", "index": 0}
        ])
        assert result["success"] is True

    def test_unknown_modify_op(self):
        from src.tools.word.word_modifier import execute_modify_op
        from docx import Document

        doc = Document()
        result = execute_modify_op(doc, {"type": "unknown_op"})
        assert result["success"] is False

    def test_invalid_paragraph_index(self):
        from src.tools.word.word_modifier import execute_modify_op
        from docx import Document

        doc = Document()
        doc.add_paragraph("Test")
        result = execute_modify_op(doc, {"type": "delete_paragraph", "index": 99})
        assert result["success"] is False


# =============================================================================
# WordFormatter 测试
# =============================================================================

class TestWordFormatter:
    """WordFormatter 模块测试"""

    def test_set_font(self):
        from src.tools.word.word_formatter import batch_format
        from docx import Document

        doc = Document()
        doc.add_paragraph("Hello World")
        result = batch_format(doc, [
            {"type": "set_font", "scope": "all", "font_name": "FangSong", "font_size": 16}
        ])
        assert result["success"] is True
        assert result["operations_applied"] == 1

    def test_set_page_margins(self):
        from src.tools.word.word_formatter import batch_format
        from docx import Document

        doc = Document()
        result = batch_format(doc, [
            {"type": "set_page_margins", "top": 3.7, "bottom": 3.5, "left": 2.8, "right": 2.6}
        ])
        assert result["success"] is True

    def test_set_page_size_invalid(self):
        from src.tools.word.word_formatter import execute_format_op
        from docx import Document

        doc = Document()
        result = execute_format_op(doc, {"type": "set_page_size", "size": "Invalid"})
        assert result["success"] is False

    def test_unknown_format_op(self):
        from src.tools.word.word_formatter import execute_format_op
        from docx import Document

        doc = Document()
        result = execute_format_op(doc, {"type": "unknown_op"})
        assert result["success"] is False

    def test_set_header(self):
        from src.tools.word.word_formatter import execute_format_op
        from docx import Document

        doc = Document()
        result = execute_format_op(doc, {"type": "set_header", "text": "机密文件"})
        assert result["success"] is True

    def test_set_footer_with_page_number(self):
        from src.tools.word.word_formatter import execute_format_op
        from docx import Document

        doc = Document()
        result = execute_format_op(doc, {"type": "set_footer", "text": "第", "include_page_number": True})
        assert result["success"] is True


# =============================================================================
# WordDiffer 测试
# =============================================================================

class TestWordDiffer:
    """WordDiffer 模块测试"""

    def test_diff_identical_docs(self):
        from src.tools.word.word_differ import diff
        from src.tools.word.word_lib import WordFileHandler
        from docx import Document

        doc1 = Document()
        doc1.add_paragraph("Same content")
        save1 = WordFileHandler.save_temp(doc1, file_name="same1.docx")
        save2 = WordFileHandler.save_temp(doc1, file_name="same2.docx")

        result = diff(save1["file_path"], save2["file_path"])
        assert result["success"] is True
        assert result["change_count"] == 0

    def test_diff_different_docs(self):
        from src.tools.word.word_differ import diff
        from src.tools.word.word_lib import WordFileHandler
        from docx import Document

        doc1 = Document()
        doc1.add_paragraph("Old text")
        save1 = WordFileHandler.save_temp(doc1, file_name="old.docx")

        doc2 = Document()
        doc2.add_paragraph("New text")
        doc2.add_paragraph("Extra paragraph")
        save2 = WordFileHandler.save_temp(doc2, file_name="new.docx")

        result = diff(save1["file_path"], save2["file_path"])
        assert result["success"] is True
        assert result["change_count"] > 0

    def test_diff_markdown_format(self):
        from src.tools.word.word_differ import diff
        from src.tools.word.word_lib import WordFileHandler
        from docx import Document

        doc1 = Document()
        doc1.add_paragraph("Version A")
        save1 = WordFileHandler.save_temp(doc1, file_name="a.docx")

        doc2 = Document()
        doc2.add_paragraph("Version B")
        save2 = WordFileHandler.save_temp(doc2, file_name="b.docx")

        result = diff(save1["file_path"], save2["file_path"], output_format="markdown")
        assert result["success"] is True
        assert "##" in result["diff_text"]

    def test_diff_file_not_found(self):
        from src.tools.word.word_differ import diff
        result = diff("/nonexistent1.docx", "/nonexistent2.docx")
        assert result["success"] is False


# =============================================================================
# WordToMd 测试
# =============================================================================

class TestWordToMd:
    """WordToMd 模块测试"""

    def test_convert_file_not_found(self):
        from src.tools.word.word_to_md import convert
        with pytest.raises(FileNotFoundError):
            convert("/nonexistent.docx")

    def test_convert_success(self):
        from src.tools.word.word_to_md import convert
        from src.tools.word.md_to_word import convert as md_convert, save_as

        doc = md_convert("# Test\n\nHello world", template="default")
        result = save_as(doc, file_name="w2md_test.docx")

        md_text = convert(result["file_path"])
        assert len(md_text) > 0
        assert "Test" in md_text or "test" in md_text.lower()


# =============================================================================
# MdToWord 测试
# =============================================================================

class TestMdToWord:
    """MdToWord 模块测试"""

    def test_convert_with_headings(self):
        from src.tools.word.md_to_word import convert, save_as

        md = "# Title\n\n## Subtitle\n\nParagraph text"
        doc = convert(md, template="default", title="Test")
        result = save_as(doc, file_name="md_test.docx")

        assert Path(result["file_path"]).exists()
        assert result["file_size"] > 0

    def test_convert_with_table(self):
        from src.tools.word.md_to_word import convert, save_as

        md = "| A | B |\n|---|---|\n| 1 | 2 |"
        doc = convert(md, template="formal")
        result = save_as(doc, file_name="table_test.docx")
        assert Path(result["file_path"]).exists()

    def test_convert_with_code_block(self):
        from src.tools.word.md_to_word import convert, save_as

        md = "```\nprint('hello')\n```"
        doc = convert(md)
        result = save_as(doc, file_name="code_test.docx")
        assert Path(result["file_path"]).exists()

    def test_convert_with_list(self):
        from src.tools.word.md_to_word import convert, save_as

        md = "- Item 1\n- Item 2\n\n1. First\n2. Second"
        doc = convert(md)
        result = save_as(doc, file_name="list_test.docx")
        assert Path(result["file_path"]).exists()

    def test_convert_with_ignored_template_name(self):
        """传入旧模板名称时静默忽略，使用 Pandoc 默认样式"""
        from src.tools.word.md_to_word import convert, save_as

        doc = convert("# Test", template="nonexistent")
        result = save_as(doc, file_name="ignored_template.docx")
        assert Path(result["file_path"]).exists()
        assert result["file_size"] > 0


# =============================================================================
# _find_pandoc 测试
# =============================================================================

class TestFindPandoc:
    """Pandoc 可执行文件定位测试"""

    def test_find_pandoc_in_path(self):
        """PATH 中存在 pandoc 时返回该路径"""
        from src.tools.word.md_to_word import _find_pandoc

        with patch("shutil.which", return_value="/usr/bin/pandoc"):
            assert _find_pandoc() == "/usr/bin/pandoc"

    def test_find_pandoc_not_installed_raises_with_guide(self):
        """Pandoc 未安装时抛 FileNotFoundError，并附带安装指引（Fail loud）。

        WHY：避免 subprocess 抛出难以理解的 [WinError 2]，
        让 Agent/用户直接看到缺失依赖与解决路径，而不是陷入无意义重试。
        """
        from src.tools.word.md_to_word import _find_pandoc

        with patch("shutil.which", return_value=None), \
             patch("pathlib.Path.exists", return_value=False):
            with pytest.raises(FileNotFoundError) as exc_info:
                _find_pandoc()

        msg = str(exc_info.value)
        assert "Pandoc" in msg
        assert "pandoc-install-guide" in msg


class TestNormalizeMarkdown:
    """Markdown 规范化测试"""

    def test_normalize_literal_newlines(self):
        from src.tools.word.md_to_word import normalize_markdown

        text = "Line1\\nLine2\\nLine3"
        result = normalize_markdown(text)
        assert result == "Line1\nLine2\nLine3"

    def test_normalize_literal_newlines_preserve_code_block(self):
        from src.tools.word.md_to_word import normalize_markdown

        text = "text\\n```python\\nprint('a\\nb')\\n```\\nmore"
        result = normalize_markdown(text)
        # 代码块内的 \\n 不替换
        assert "print('a\\nb')" in result

    def test_normalize_unclosed_code_block(self):
        from src.tools.word.md_to_word import normalize_markdown

        text = "```\ncode without closing"
        result = normalize_markdown(text)
        assert result.endswith("```")

    def test_normalize_missing_table_separator(self):
        from src.tools.word.md_to_word import normalize_markdown

        text = "| A | B |\n| 1 | 2 |"
        result = normalize_markdown(text)
        lines = result.split('\n')
        # 第二行应该是分隔行
        assert re.match(r'^\|[\s\-:|]+$', lines[1])

    def test_normalize_block_spacing(self):
        from src.tools.word.md_to_word import normalize_markdown

        text = "paragraph\n# Heading\nparagraph"
        result = normalize_markdown(text)
        # 标题前后应该有空行
        assert "\n\n# Heading\n\n" in result

    def test_normalize_hr_to_newpage(self):
        from src.tools.word.md_to_word import normalize_markdown

        text = "above\n---\nbelow"
        result = normalize_markdown(text)
        assert "\\newpage" in result
        assert "---" not in result.split('\n')


# =============================================================================
# WordLib 核心库测试
# =============================================================================

class TestWordLibCore:
    """WordLib 核心功能测试"""

    def test_resolve_font_name(self):
        from src.tools.word.word_lib import resolve_font_name
        assert resolve_font_name("宋体") == "SimSun"
        assert resolve_font_name("黑体") == "SimHei"
        assert resolve_font_name("Arial") == "Arial"
        assert resolve_font_name("") == ""

    def test_resolve_alignment(self):
        from src.tools.word.word_lib import resolve_alignment
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        assert resolve_alignment("CENTER") == WD_ALIGN_PARAGRAPH.CENTER
        assert resolve_alignment("JUSTIFY") == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert resolve_alignment("unknown") is None

    def test_parse_color(self):
        from src.tools.word.word_lib import parse_color
        color = parse_color("#FF0000")
        assert color is not None
        assert color[0] == 255
        assert color[1] == 0
        assert color[2] == 0

        assert parse_color("") is None
        assert parse_color("invalid") is None

    def test_replace_text_cross_run(self):
        from src.tools.word.word_lib import replace_text_cross_run
        from docx import Document

        doc = Document()
        para = doc.add_paragraph()
        para.add_run("Hello ")
        para.add_run("World")

        count = replace_text_cross_run(para.runs, "Hello World", "Hi Earth")
        assert count == 1
        assert para.text == "Hi Earth"

    def test_replace_text_cross_run_single_run(self):
        from src.tools.word.word_lib import replace_text_cross_run
        from docx import Document

        doc = Document()
        doc.add_paragraph("Hello World")

        count = replace_text_cross_run(doc.paragraphs[0].runs, "World", "Python")
        assert count == 1
        assert "Python" in doc.paragraphs[0].text

    def test_replace_text_cross_run_no_match(self):
        from src.tools.word.word_lib import replace_text_cross_run
        from docx import Document

        doc = Document()
        doc.add_paragraph("Hello World")

        count = replace_text_cross_run(doc.paragraphs[0].runs, "NotFound", "X")
        assert count == 0

    def test_file_handler_save_temp(self):
        from src.tools.word.word_lib import WordFileHandler
        from docx import Document

        doc = Document()
        doc.add_paragraph("Test")
        result = WordFileHandler.save_temp(doc, file_name="handler_test.docx")
        assert Path(result["file_path"]).exists()
        assert result["file_size"] > 0

    def test_file_handler_copy_and_open(self):
        from src.tools.word.word_lib import WordFileHandler
        from docx import Document

        doc = Document()
        doc.add_paragraph("Original")
        save = WordFileHandler.save_temp(doc, file_name="original.docx")

        doc2, info = WordFileHandler.copy_and_open(save["file_path"])
        assert info["original_name"] == "original.docx"
        assert "Original" in doc2.paragraphs[0].text

    def test_file_handler_copy_not_found(self):
        from src.tools.word.word_lib import WordFileHandler
        with pytest.raises(FileNotFoundError):
            WordFileHandler.copy_and_open("/nonexistent.docx")

    def test_page_sizes(self):
        from src.tools.word.word_lib import PAGE_SIZES
        assert "A4" in PAGE_SIZES
        assert PAGE_SIZES["A4"] == (21.0, 29.7)


# =============================================================================
# WordRouter 测试
# =============================================================================

class TestWordRouter:
    """WordRouter 内部 LLM 路由器测试"""

    def test_build_prompt_with_context(self):
        from src.tools.word.word_router import WordRouter
        router = WordRouter()
        prompt = router._build_prompt("用户要求生成Word", None)
        assert "用户要求生成Word" in prompt
        assert "(无附件)" in prompt

    def test_build_prompt_with_files(self):
        from src.tools.word.word_router import WordRouter
        router = WordRouter()
        prompt = router._build_prompt("读取文件", ["/tmp/test.docx"])
        assert "读取文件" in prompt
        assert "/tmp/test.docx" in prompt

    def test_parse_response_valid_json(self):
        from src.tools.word.word_router import WordRouter
        router = WordRouter()
        result = router._parse_response('{"task": "md_to_word", "params": {"template": "default"}, "reason": "test"}')
        assert result["task"] == "md_to_word"
        assert result["params"]["template"] == "default"

    def test_parse_response_json_in_code_block(self):
        from src.tools.word.word_router import WordRouter
        router = WordRouter()
        content = '```json\n{"task": "read", "params": {}, "reason": "test"}\n```'
        result = router._parse_response(content)
        assert result["task"] == "read"

    def test_parse_response_invalid_json(self):
        from src.tools.word.word_router import WordRouter
        router = WordRouter()
        result = router._parse_response("not json at all")
        assert result["task"] == ""
        assert "error" in result

    def test_parse_response_invalid_task(self):
        from src.tools.word.word_router import WordRouter
        router = WordRouter()
        result = router._parse_response('{"task": "invalid_op", "params": {}, "reason": "test"}')
        assert result["task"] == ""
        assert "error" in result

    def test_parse_response_multi_task(self):
        from src.tools.word.word_router import WordRouter
        router = WordRouter()
        result = router._parse_response('{"task": "read,modify", "params": {"operations": []}, "reason": "test"}')
        assert result["task"] == "read,modify"

    @pytest.mark.asyncio
    async def test_route_llm_success(self):
        from src.tools.word.word_router import WordRouter
        router = WordRouter()

        mock_gateway = AsyncMock()
        mock_gateway.get_model_name = lambda: "test-model"
        mock_gateway.chat_no_thinking.return_value = {
            "content": '{"task": "md_to_word", "params": {"template": "default"}, "reason": "用户要求生成Word"}'
        }
        router._gateway = mock_gateway

        result = await router.route("用户要求将行程生成Word文档", None)
        assert result["task"] == "md_to_word"
        assert result["params"]["template"] == "default"

    @pytest.mark.asyncio
    async def test_route_llm_failure(self):
        from src.tools.word.word_router import WordRouter
        router = WordRouter()

        mock_gateway = AsyncMock()
        mock_gateway.get_model_name = lambda: "test-model"
        mock_gateway.chat_no_thinking.side_effect = Exception("LLM unavailable")
        router._gateway = mock_gateway

        result = await router.route("用户要求生成Word", None)
        assert result["task"] == ""
        assert "error" in result

    @pytest.mark.asyncio
    async def test_route_llm_empty_response(self):
        from src.tools.word.word_router import WordRouter
        router = WordRouter()

        mock_gateway = AsyncMock()
        mock_gateway.get_model_name = lambda: "test-model"
        mock_gateway.chat_no_thinking.return_value = {"content": ""}
        router._gateway = mock_gateway

        result = await router.route("some context", None)
        assert result["task"] == ""
        assert "error" in result
