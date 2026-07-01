"""Phase 6 template-following tests using real PPTX packages."""

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
from PIL import Image

from src.tools.ppt.template_analyzer import (
    TemplateAnalysisError,
    TemplateAnalyzer,
)

pytestmark = pytest.mark.tools


def _create_template(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "源模板标题"
    slide.placeholders[1].text = "源模板正文"
    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    title_run.font.name = "Arial"
    title_run.font.size = Pt(30)
    title_run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    brand = slide.shapes.add_textbox(
        Inches(0.4), Inches(6.9), Inches(4), Inches(0.3)
    )
    brand.name = "Brand Footer"
    brand.text = "ACME CONFIDENTIAL"
    prs.save(path)
    return path


def _plan(**slide_overrides):
    slide = {
        "type": "content",
        "title": "季度经营复盘",
        "points": ["收入增长 20%", "客户留存改善"],
        **slide_overrides,
    }
    return {"title": "模板跟随输出", "theme_id": 14, "style": "soft", "slides": [slide]}


def test_template_audit_extracts_dimensions_masters_layouts_and_source_slides(tmp_path):
    template = _create_template(tmp_path / "brand-template.pptx")
    artifacts = tmp_path / "artifacts"

    audit = TemplateAnalyzer().analyze(template, artifact_dir=artifacts)

    assert audit["slide_size"]["width_inches"] == pytest.approx(13.333, abs=0.001)
    assert audit["slide_size"]["height_inches"] == 7.5
    assert audit["masters"]
    assert audit["layouts"]
    assert any(layout["placeholders"] for layout in audit["layouts"])
    assert audit["slides"][0]["layout_name"]
    assert "源模板标题" in audit["slides"][0]["text_summary"]
    assert "Arial" in audit["slides"][0]["fonts"]
    assert "123456" in audit["slides"][0]["colors"]
    persisted = json.loads(
        (artifacts / "template-audit.json").read_text(encoding="utf-8")
    )
    assert persisted == audit
    assert str(template) not in json.dumps(persisted, ensure_ascii=False)


def test_frame_map_and_deviation_log_are_written_before_generation(tmp_path):
    template = _create_template(tmp_path / "template.pptx")
    artifacts = tmp_path / "artifacts"
    analyzer = TemplateAnalyzer()
    audit = analyzer.analyze(template, artifacts)

    matches = analyzer.match_content_to_layouts(_plan(), audit, artifacts)

    frame_map = json.loads(
        (artifacts / "template-frame-map.json").read_text(encoding="utf-8")
    )
    deviations = json.loads(
        (artifacts / "deviation-log.json").read_text(encoding="utf-8")
    )
    assert frame_map["frames"][0]["strategy"] == "layout"
    assert frame_map["frames"][0]["layout_index"] is not None
    assert {target["role"] for target in frame_map["frames"][0]["editTargets"]} == {
        "title",
        "body",
    }
    assert deviations == {"version": "1.0", "deviations": []}
    assert matches[0]["frame"] == frame_map["frames"][0]


def test_explicit_source_slide_clone_preserves_brand_and_replaces_placeholders(
    tmp_path, monkeypatch
):
    template = _create_template(tmp_path / "template.pptx")
    artifacts = tmp_path / "artifacts"
    output_dir = tmp_path / "output"
    analyzer = TemplateAnalyzer()
    monkeypatch.setattr(analyzer, "_get_output_dir", lambda: output_dir)
    plan = _plan(template_slide_index=0)
    audit = analyzer.analyze(template, artifacts)
    matches = analyzer.match_content_to_layouts(plan, audit, artifacts)

    output = analyzer.generate_from_template(template, matches, plan, artifacts)

    generated = Presentation(output)
    assert len(generated.slides) == 1
    texts = [
        shape.text
        for shape in generated.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    ]
    assert "季度经营复盘" in texts
    assert any("收入增长 20%" in text for text in texts)
    assert "ACME CONFIDENTIAL" in texts
    assert generated.slide_width == Inches(13.333)
    assert matches[0]["frame"]["strategy"] == "clone"


def test_missing_body_placeholder_declares_addition_and_deviation(tmp_path, monkeypatch):
    template = tmp_path / "blank-template.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(template)
    artifacts = tmp_path / "artifacts"
    analyzer = TemplateAnalyzer()
    monkeypatch.setattr(analyzer, "_get_output_dir", lambda: tmp_path / "output")
    audit = analyzer.analyze(template, artifacts)
    matches = analyzer.match_content_to_layouts(
        _plan(layout_index=6), audit, artifacts
    )

    frame = matches[0]["frame"]
    assert {item["role"] for item in frame["allowedAdditions"]} == {
        "title",
        "body",
    }
    assert "object_addition_required" in frame["deviation_codes"]
    output = analyzer.generate_from_template(template, matches, _plan(), artifacts)
    generated = Presentation(output)
    generated_text = [
        shape.text
        for shape in generated.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    ]
    assert "季度经营复盘" in generated_text
    assert any(
        "收入增长 20%" in shape.text
        for shape in generated.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )


def test_clone_declares_deleted_source_placeholder_and_preserves_image(
    tmp_path, monkeypatch
):
    image_path = tmp_path / "brand.png"
    Image.new("RGB", (16, 16), (18, 52, 86)).save(image_path)
    template = tmp_path / "template.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Source"
    body = slide.placeholders[1]
    slide.shapes._spTree.remove(body.element)
    slide.shapes.add_picture(str(image_path), Inches(11), Inches(6), Inches(1), Inches(1))
    prs.save(template)

    artifacts = tmp_path / "artifacts"
    analyzer = TemplateAnalyzer()
    monkeypatch.setattr(analyzer, "_get_output_dir", lambda: tmp_path / "output")
    plan = _plan(template_slide_index=0)
    audit = analyzer.analyze(template, artifacts)
    matches = analyzer.match_content_to_layouts(plan, audit, artifacts)

    body_addition = next(
        item
        for item in matches[0]["frame"]["allowedAdditions"]
        if item["role"] == "body"
    )
    assert body_addition["reason"] == "source_placeholder_missing"

    output = analyzer.generate_from_template(template, matches, plan, artifacts)
    generated = Presentation(output)
    assert any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        for shape in generated.slides[0].shapes
    )
    assert any(
        "收入增长 20%" in shape.text
        for shape in generated.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )


def test_generation_requires_complete_audit_artifacts(tmp_path):
    template = _create_template(tmp_path / "template.pptx")
    artifacts = tmp_path / "artifacts"
    analyzer = TemplateAnalyzer()
    audit = analyzer.analyze(template, artifacts)
    matches = analyzer.match_content_to_layouts(_plan(), audit, artifacts)
    (artifacts / "template-frame-map.json").unlink()

    with pytest.raises(TemplateAnalysisError, match="审计产物不完整"):
        analyzer.generate_from_template(template, matches, _plan(), artifacts)


def test_unavailable_requested_layout_records_safe_fallback(tmp_path):
    template = _create_template(tmp_path / "template.pptx")
    artifacts = tmp_path / "artifacts"
    analyzer = TemplateAnalyzer()
    audit = analyzer.analyze(template, artifacts)

    matches = analyzer.match_content_to_layouts(
        _plan(layout_index=999), audit, artifacts
    )

    assert matches[0]["frame"]["layout_index"] is not None
    assert "requested_frame_unavailable" in matches[0]["frame"]["deviation_codes"]


def test_clone_failure_updates_persisted_frame_map(tmp_path, monkeypatch):
    template = _create_template(tmp_path / "template.pptx")
    artifacts = tmp_path / "artifacts"
    analyzer = TemplateAnalyzer()
    monkeypatch.setattr(analyzer, "_get_output_dir", lambda: tmp_path / "output")
    monkeypatch.setattr(
        analyzer,
        "_clone_slide",
        lambda *_: (_ for _ in ()).throw(ValueError("simulated clone failure")),
    )
    plan = _plan(template_slide_index=0)
    audit = analyzer.analyze(template, artifacts)
    matches = analyzer.match_content_to_layouts(plan, audit, artifacts)

    output = analyzer.generate_from_template(template, matches, plan, artifacts)

    Presentation(output)
    persisted = json.loads(
        (artifacts / "template-frame-map.json").read_text(encoding="utf-8")
    )
    frame = persisted["frames"][0]
    assert frame["strategy"] == "layout"
    assert frame["source_slide_index"] is None
    assert "clone_fallback" in frame["deviation_codes"]
    deviation_log = json.loads(
        (artifacts / "deviation-log.json").read_text(encoding="utf-8")
    )
    assert any(
        item["code"] == "clone_fallback"
        for item in deviation_log["deviations"]
    )


def test_added_textboxes_stay_inside_custom_slide_size(tmp_path, monkeypatch):
    template = tmp_path / "small-template.pptx"
    prs = Presentation()
    prs.slide_width = Inches(5)
    prs.slide_height = Inches(3)
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(template)
    artifacts = tmp_path / "artifacts"
    analyzer = TemplateAnalyzer()
    monkeypatch.setattr(analyzer, "_get_output_dir", lambda: tmp_path / "output")
    audit = analyzer.analyze(template, artifacts)
    matches = analyzer.match_content_to_layouts(
        _plan(layout_index=6), audit, artifacts
    )

    output = analyzer.generate_from_template(template, matches, _plan(), artifacts)

    generated = Presentation(output)
    assert generated.slide_width == Inches(5)
    assert generated.slide_height == Inches(3)
    for shape in generated.slides[0].shapes:
        assert shape.left >= 0
        assert shape.top >= 0
        assert shape.left + shape.width <= generated.slide_width
        assert shape.top + shape.height <= generated.slide_height


def test_invalid_template_error_does_not_expose_path(tmp_path):
    missing = tmp_path / "sensitive-customer-name.pptx"

    with pytest.raises(TemplateAnalysisError) as error:
        TemplateAnalyzer().analyze(missing)

    assert str(missing) not in str(error.value)
    assert "不存在或不可访问" in str(error.value)


def test_corrupt_template_error_is_sanitized(tmp_path):
    corrupt = tmp_path / "sensitive-customer-name.pptx"
    corrupt.write_bytes(b"not-a-valid-pptx")

    with pytest.raises(TemplateAnalysisError) as error:
        TemplateAnalyzer().analyze(corrupt)

    assert str(corrupt) not in str(error.value)
    assert "无效或无法读取" in str(error.value)
