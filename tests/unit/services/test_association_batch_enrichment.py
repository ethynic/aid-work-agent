import csv
import asyncio
import json

import pytest
from openpyxl import Workbook, load_workbook

from src.services.association_batch_enrichment import (
    AssociationBatchEnricher,
    parse_association_input,
    write_enrichment_workbook,
)
from src.services.association_enrichment_providers import ProjectAssociationProviders
from src.services.association_profile_extractor import PROFILE_FIELDS


pytestmark = pytest.mark.unit


class _SearchStub:
    def __init__(self, responses):
        self.responses = list(responses)
        self.keywords = []

    async def execute(self, *, keyword, **_kwargs):
        self.keywords.append(keyword)
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response


class _GatewayStub:
    def __init__(self, contents):
        self.contents = list(contents)
        self.calls = 0

    async def chat(self, **_kwargs):
        self.calls += 1
        return {"content": self.contents.pop(0)}


def test_text_csv_and_excel_inputs_produce_one_ordered_deduplicated_list(tmp_path):
    csv_path = tmp_path / "input.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["协会名称", "备注"])
        writer.writerow(["中国日用玻璃协会", "重复"])
        writer.writerow(["中国缝制机械协会", "新增"])

    assert parse_association_input(
        text_values=["中国日用玻璃协会，中国  日用玻璃协会"],
        input_path=csv_path,
    ) == ["中国日用玻璃协会", "中国缝制机械协会"]

    xlsx_path = tmp_path / "input.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["单位名称"])
    worksheet.append(["中国物资再生协会"])
    workbook.save(xlsx_path)
    assert parse_association_input(input_path=xlsx_path) == ["中国物资再生协会"]


def test_input_without_association_column_fails_loud(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("姓名\n张三\n", encoding="utf-8")
    with pytest.raises(ValueError, match="INPUT_ASSOCIATION_COLUMN_NOT_FOUND"):
        parse_association_input(input_path=path)


def test_empty_and_corrupt_input_files_fail_loud(tmp_path):
    empty_csv = tmp_path / "empty.csv"
    empty_csv.write_text("", encoding="utf-8")
    with pytest.raises(ValueError, match="INPUT_ASSOCIATION_LIST_EMPTY"):
        parse_association_input(input_path=empty_csv)

    corrupt_xlsx = tmp_path / "corrupt.xlsx"
    corrupt_xlsx.write_bytes(b"not an xlsx archive")
    with pytest.raises(Exception):
        parse_association_input(input_path=corrupt_xlsx)


@pytest.mark.asyncio
async def test_batch_falls_back_from_https_to_http_and_then_enriches_wechat():
    attempts = []

    async def resolve(_name):
        return "https://legacy.example.org/"

    async def collect(url, headless, **kwargs):
        attempts.append((url, headless))
        if url.startswith("https://"):
            raise ConnectionError("old site")
        return {
            "secretary_general_name": "王秘书",
            "official_website": url,
        }

    async def fallback(_name):
        raise AssertionError("official HTTP fallback succeeded")

    async def wechat(association, person, role):
        assert (association, person, role) == ("测试协会", "王秘书", "秘书长")
        return "18612345678"

    enricher = AssociationBatchEnricher(
        official_site_resolver=resolve,
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
        headless=False,
    )
    row = await enricher.enrich_one("测试协会")

    assert attempts == [
        ("https://legacy.example.org/", False),
        ("http://legacy.example.org/", False),
    ]
    assert row.values["secretary_general_mobile"] == "18612345678"
    assert row.processing_status == "partial"
    assert "profile:member_director_not_found" in row.errors


@pytest.mark.asyncio
async def test_one_failed_association_does_not_stop_later_associations():
    async def resolve(name):
        if name == "坏协会":
            raise RuntimeError("resolver down")
        return None

    async def collect(_url, _headless, **kwargs):
        raise AssertionError("no official URL")

    async def fallback(name):
        if name == "坏协会":
            raise RuntimeError("search down")
        return {"address": "北京市"}

    async def wechat(_association, _person, _role):
        return None

    enricher = AssociationBatchEnricher(
        official_site_resolver=resolve,
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
    )
    rows = await enricher.enrich_many(["坏协会", "好协会"])

    assert [row.association_name for row in rows] == ["坏协会", "好协会"]
    assert rows[0].processing_status == "failed"
    assert rows[1].values["address"] == "北京市"


@pytest.mark.asyncio
async def test_cleanup_failure_keeps_verified_mobile_and_stops_batch():
    from src.services.association_batch_enrichment import WechatRpaError

    calls = []

    async def resolve(_name):
        return "https://association.example.cn/"

    async def collect(_url, _headless, **kwargs):
        return {
            "secretary_general_name": "陈戟",
        }

    async def fallback(_name):
        raise AssertionError("official collection succeeded")

    async def wechat(_association, person, _role):
        calls.append(person)
        if person == "陈戟":
            raise WechatRpaError(
                "PLUGIN_CLOSE_TIMEOUT",
                stage="cleanup",
                session_fatal=True,
                recovered_mobile="18512345678",
            )
        return "18612345678"

    enricher = AssociationBatchEnricher(
        official_site_resolver=resolve,
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
    )

    rows = await enricher.enrich_many(["测试协会", "不得继续协会"])

    assert calls == ["陈戟"]
    assert rows.aborted is True
    assert rows.abort_error_code == "PLUGIN_CLOSE_TIMEOUT"
    assert [row.association_name for row in rows] == ["测试协会"]
    assert rows[0].values["secretary_general_mobile"] == "18512345678"
    assert "wechat:秘书长" in rows[0].sources
    assert any(
        error == "wechat:秘书长:PLUGIN_CLOSE_TIMEOUT:cleanup"
        for error in rows[0].errors
    )


@pytest.mark.asyncio
async def test_nonfatal_wechat_failure_allows_next_person():
    calls = []

    async def resolve(_name):
        return "https://example.cn"

    async def collect(_url, _headless, **kwargs):
        return {
            "secretary_general_name": "秘书长乙",
            "member_director_name": "会员丙",
            "office_director_name": "办公室丁",
        }

    async def fallback(_name):
        return {}

    async def wechat(_association, person, _role):
        calls.append(person)
        if person == "秘书长乙":
            raise RuntimeError("CARD_LOCATE_FAILED")
        return None

    enricher = AssociationBatchEnricher(
        official_site_resolver=resolve,
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
    )

    await enricher.enrich_one("测试协会")

    assert calls == ["秘书长乙", "会员丙", "办公室丁"]


@pytest.mark.asyncio
async def test_last_association_session_failure_is_explicitly_aborted():
    from src.services.association_batch_enrichment import WechatRpaError

    async def resolve(_name):
        return "https://example.cn"

    async def collect(_url, _headless, **kwargs):
        return {"secretary_general_name": "秘书甲"}

    async def fallback(_name):
        return {}

    async def wechat(association, _person, _role):
        if association == "最后协会":
            raise WechatRpaError(
                "SESSION_CLEANUP_FAILED",
                stage="cleanup",
                session_fatal=True,
            )
        return None

    enricher = AssociationBatchEnricher(
        official_site_resolver=resolve,
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
    )

    rows = await enricher.enrich_many(["正常协会", "最后协会"])

    assert len(rows) == 2
    assert rows.aborted is True
    assert rows.abort_error_code == "SESSION_CLEANUP_FAILED"


def test_excel_contains_business_and_audit_columns_but_error_is_redacted(tmp_path):
    from src.services.association_batch_enrichment import AssociationEnrichmentRow

    row = AssociationEnrichmentRow(
        association_name="测试协会",
        processing_status="partial",
        errors=["provider returned 18612345678"],
        processed_at="2026-07-29T00:00:00+00:00",
    )
    row.values["secretary_general_mobile"] = "18612345678"
    path = write_enrichment_workbook([row], tmp_path / "result.xlsx")
    workbook = load_workbook(path, data_only=True)
    try:
        worksheet = workbook.active
        headers = [cell.value for cell in worksheet[1]]
        values = dict(zip(headers, [cell.value for cell in worksheet[2]]))
    finally:
        workbook.close()

    # 表头输出中文标签（3 个目标联系角色列）。
    from src.services.association_batch_enrichment import _EXCEL_HEADERS
    assert headers == [label for label, _ in _EXCEL_HEADERS]
    assert values["秘书长\n手机"] == "18612345678"
    # 未命中的联系角色列存在且为空
    assert "会员服务负责人\n姓名" in values
    assert values["会员服务负责人\n姓名"] is None


def test_excel_treats_external_formula_prefixes_as_text(tmp_path):
    from src.services.association_batch_enrichment import AssociationEnrichmentRow

    row = AssociationEnrichmentRow(
        association_name="=HYPERLINK(\"https://evil.invalid\")",
        processing_status="partial",
        sources=["+cmd"],
        errors=["@payload"],
        processed_at="-1",
    )
    row.values["address"] = "=1+1"
    path = write_enrichment_workbook([row], tmp_path / "safe.xlsx")
    workbook = load_workbook(path, data_only=False)
    try:
        worksheet = workbook.active
        headers = [cell.value for cell in worksheet[1]]
        cells = dict(zip(headers, [cell.value for cell in worksheet[2]]))
        types = dict(zip(headers, [cell.data_type for cell in worksheet[2]]))
    finally:
        workbook.close()

    for header_name in (
        "客户名称",
        "单位地址",
        "处理时间",
    ):
        assert cells[header_name].startswith("'")
        assert types[header_name] == "s"


def test_web_fallback_accepts_model_values_without_verbatim_quote_check():
    """删除逐字校验后，fallback 路径只做 schema/类型/手机号禁止校验；
    模型给出的 value 一律接受，不再要求 value 逐字出现在 quote 中。"""
    parsed = {
        name: {"value": None, "evidence_quote": None, "source_url": None}
        for name in PROFILE_FIELDS
    }
    parsed["secretary_general_name"] = {
        "value": "潘华",
        "evidence_quote": "秘书长：潘  华",
        "source_url": "https://example.org/profile",
    }
    result = ProjectAssociationProviders._validate_fallback_evidence(
        parsed,
        [{
            "url": "https://example.org/profile",
            "title": "协会简介",
            "content": "第九届理事会秘书长：潘  华",
        }],
    )
    assert result["secretary_general_name"] == "潘华"

    parsed["secretary_general_name"]["value"] = "虚构姓名"
    rejected = []
    result = ProjectAssociationProviders._validate_fallback_evidence(
        parsed,
        [{
            "url": "https://example.org/profile",
            "title": "协会简介",
            "content": "第九届理事会秘书长：潘  华",
        }],
        rejected,
    )
    assert result["secretary_general_name"] == "虚构姓名"
    assert rejected == []


@pytest.mark.asyncio
async def test_timed_out_child_process_is_killed_and_reaped():
    class HangingProcess:
        def __init__(self):
            self.killed = False
            self.waited = False

        async def communicate(self):
            await asyncio.Event().wait()

        def kill(self):
            self.killed = True

        async def wait(self):
            self.waited = True

    process = HangingProcess()
    with pytest.raises(RuntimeError, match="CHILD_TIMEOUT"):
        await ProjectAssociationProviders._communicate_with_timeout(
            process,
            timeout_seconds=0.001,
            error_code="CHILD_TIMEOUT",
        )
    assert process.killed is True
    assert process.waited is True


@pytest.mark.asyncio
async def test_timed_out_child_preserves_code_when_it_exits_before_kill():
    class ExitedProcess:
        def __init__(self):
            self.waited = False

        async def communicate(self):
            await asyncio.Event().wait()

        def kill(self):
            raise ProcessLookupError

        async def wait(self):
            self.waited = True

    process = ExitedProcess()
    with pytest.raises(RuntimeError, match="WECHAT_RPA_TIMEOUT"):
        await ProjectAssociationProviders._communicate_with_timeout(
            process,
            timeout_seconds=0.001,
            error_code="WECHAT_RPA_TIMEOUT",
        )
    assert process.waited is True


@pytest.mark.asyncio
async def test_wechat_audit_failure_does_not_discard_found_mobile(
    monkeypatch, tmp_path
):
    import src.services.association_enrichment_providers as module

    class CompletedProcess:
        def __init__(self, stdout):
            self.returncode = 0
            self._stdout = stdout

        async def communicate(self):
            return self._stdout, b""

    processes = iter(
        [
            CompletedProcess(
                json.dumps(
                    {
                        "ok": True,
                        "session_closed": True,
                        "status": "found",
                        "artifact_ref": "artifact.dpapi",
                    }
                ).encode("utf-8")
            ),
            CompletedProcess(
                json.dumps(
                    {"matched": True, "mobile": "13912345678"}
                ).encode("utf-8")
            ),
        ]
    )

    async def fake_subprocess(*_args, **_kwargs):
        return next(processes)

    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def broken_audit(*_args, **_kwargs):
        raise RuntimeError("audit unavailable")

    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)
    monkeypatch.setattr(providers, "_audit_wechat_artifact", broken_audit)
    monkeypatch.setattr(
        providers,
        "_validated_wechat_artifact_path",
        lambda _ref: tmp_path / "artifact.dpapi",
    )

    assert (
        await providers.wechat_mobile("测试协会", "张三", "会长")
        == "13912345678"
    )
    assert providers._wechat_handoff_ready is True


@pytest.mark.asyncio
async def test_wechat_mobile_retries_when_souyisou_empty(monkeypatch, tmp_path):
    """搜一搜结果为空（inconclusive：没读到列表）时重发组合键重搜一次；第二轮搜到 → 返回手机号。

    业务意图：用户要「搜一搜空才重试，有结果不重试」。inconclusive = 搜一搜没读到列表
    （焦点丢失/没加载），该重试；not_found = 有列表但没匹配手机号，算成功不重试。
    """
    import src.services.association_enrichment_providers as module

    class CompletedProcess:
        def __init__(self, stdout):
            self.returncode = 0
            self._stdout = stdout

        async def communicate(self):
            return self._stdout, b""

    processes = iter([
        # 第一轮 collect：没搜到 → 触发重试
        CompletedProcess(
            json.dumps({"ok": True, "session_closed": True, "status": "inconclusive", "checked": 0}).encode("utf-8")
        ),
        # 第二轮 collect：搜到
        CompletedProcess(
            json.dumps({
                "ok": True, "session_closed": True, "status": "found",
                "artifact_ref": "artifact.dpapi",
            }).encode("utf-8")
        ),
        # extract-mobile.ps1
        CompletedProcess(
            json.dumps({"matched": True, "mobile": "13912345678"}).encode("utf-8")
        ),
    ])

    async def fake_subprocess(*_args, **_kwargs):
        return next(processes)

    async def noop(*_args, **_kwargs):
        return None

    providers = ProjectAssociationProviders(repository_root=tmp_path)
    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)
    monkeypatch.setattr(providers, "_audit_wechat_artifact", noop)
    monkeypatch.setattr(
        providers,
        "_validated_wechat_artifact_path",
        lambda _ref: tmp_path / "artifact.dpapi",
    )

    assert await providers.wechat_mobile("测试协会", "张三", "会长") == "13912345678"


@pytest.mark.asyncio
async def test_wechat_mobile_returns_none_after_retry_exhausted(monkeypatch, tmp_path):
    """两轮搜一搜都为空（inconclusive）→ 返回 None（重试一次仍空，不无限重试）。"""
    import src.services.association_enrichment_providers as module

    class CompletedProcess:
        def __init__(self, stdout):
            self.returncode = 0
            self._stdout = stdout

        async def communicate(self):
            return self._stdout, b""

    processes = iter([
        CompletedProcess(
            json.dumps({"ok": True, "session_closed": True, "status": "inconclusive", "checked": 0}).encode("utf-8")
        ),
        CompletedProcess(
            json.dumps({"ok": True, "session_closed": True, "status": "inconclusive", "checked": 0}).encode("utf-8")
        ),
    ])

    async def fake_subprocess(*_args, **_kwargs):
        return next(processes)

    providers = ProjectAssociationProviders(repository_root=tmp_path)
    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)

    assert await providers.wechat_mobile("测试协会", "张三", "会长") is None


@pytest.mark.asyncio
async def test_wechat_mobile_no_retry_when_not_found(monkeypatch, tmp_path):
    """搜一搜有列表但没匹配到手机号（not_found）→ 直接返回 None，不重试。

    业务意图：not_found 算「搜一搜成功」（有结果），不该重发组合键。回归保护用户报告的
    「有结果也搜两遍」bug——若重试条件误用 status != found，not_found 也会搜两次。
    """
    import src.services.association_enrichment_providers as module

    class CompletedProcess:
        def __init__(self, stdout):
            self.returncode = 0
            self._stdout = stdout

        async def communicate(self):
            return self._stdout, b""

    call_count = 0

    async def fake_subprocess(*_args, **_kwargs):
        nonlocal call_count
        call_count += 1
        return CompletedProcess(
            json.dumps({"ok": True, "session_closed": True, "status": "not_found"}).encode("utf-8")
        )

    providers = ProjectAssociationProviders(repository_root=tmp_path)
    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)

    assert await providers.wechat_mobile("测试协会", "张三", "会长") is None
    assert call_count == 1  # not_found 不重试，只搜一次


@pytest.mark.asyncio
async def test_wechat_mobile_no_retry_after_detail_checked(monkeypatch, tmp_path):
    """进过详情页（inconclusive + checked>0）→ 不重试，直接返回 None。

    业务意图：checked>0 说明已经打开过详情尽力搜过，重搜结果一样。回归保护用户报告的
    「秘书长手机搜两次」——首版重试条件 status==inconclusive 没看 checked，导致进过详情
    也重搜一遍。
    """
    import src.services.association_enrichment_providers as module

    class CompletedProcess:
        def __init__(self, stdout):
            self.returncode = 0
            self._stdout = stdout

        async def communicate(self):
            return self._stdout, b""

    call_count = 0

    async def fake_subprocess(*_args, **_kwargs):
        nonlocal call_count
        call_count += 1
        return CompletedProcess(
            json.dumps({"ok": True, "session_closed": True, "status": "inconclusive", "checked": 1}).encode("utf-8")
        )

    providers = ProjectAssociationProviders(repository_root=tmp_path)
    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)

    assert await providers.wechat_mobile("测试协会", "张三", "会长") is None
    assert call_count == 1  # 进过详情(checked>0)不重试，只搜一次


@pytest.mark.asyncio
async def test_wechat_inter_query_handoff_runs_once_only_after_confirmed_close(
    monkeypatch, tmp_path
):
    providers = ProjectAssociationProviders(repository_root=tmp_path)
    events = []

    monkeypatch.setattr(
        providers, "_send_alt_tab_once", lambda: events.append("alt_tab")
    )

    async def fake_sleep(seconds):
        events.append(("sleep", seconds))

    providers._wechat_handoff_sleep = fake_sleep

    assert await providers._prepare_wechat_query_handoff() is False
    assert events == []

    providers._wechat_handoff_ready = True
    assert await providers._prepare_wechat_query_handoff() is True
    assert events == ["alt_tab", ("sleep", 1.0)]

    assert await providers._prepare_wechat_query_handoff() is False
    assert events == ["alt_tab", ("sleep", 1.0)]


@pytest.mark.asyncio
async def test_wechat_inter_query_handoff_failure_consumes_authorization(
    monkeypatch, tmp_path
):
    providers = ProjectAssociationProviders(repository_root=tmp_path)
    providers._wechat_handoff_ready = True

    def fail_alt_tab():
        raise RuntimeError("injection failed")

    monkeypatch.setattr(providers, "_send_alt_tab_once", fail_alt_tab)
    with pytest.raises(RuntimeError, match="injection failed"):
        await providers._prepare_wechat_query_handoff()

    assert providers._wechat_handoff_ready is False


def test_wechat_artifact_ref_is_scoped_to_local_uuid_directory(
    monkeypatch, tmp_path
):
    artifact_root = (
        tmp_path / "AidWorkAgent" / "wechat-souyisou-rpa" / "artifacts"
    )
    artifact_root.mkdir(parents=True)
    valid = artifact_root / f"{'a' * 32}.dpapi"
    valid.write_bytes(b"encrypted")
    outside = tmp_path / f"{'b' * 32}.dpapi"
    outside.write_bytes(b"encrypted")
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))

    assert (
        ProjectAssociationProviders._validated_wechat_artifact_path(str(valid))
        == valid.resolve()
    )
    with pytest.raises(ValueError, match="WECHAT_ARTIFACT_REF_INVALID"):
        ProjectAssociationProviders._validated_wechat_artifact_path(str(outside))
    with pytest.raises(ValueError, match="WECHAT_ARTIFACT_REF_INVALID"):
        ProjectAssociationProviders._validated_wechat_artifact_path(
            str(artifact_root / ".." / valid.name)
        )


@pytest.mark.asyncio
@pytest.mark.asyncio
async def test_website_only_official_profile_uses_fallback_then_wechat():
    async def resolve(_name):
        return "https://association.example.cn/"

    async def collect(url, _headless, **kwargs):
        return {"official_website": url}

    async def fallback(_name):
        return {"secretary_general_name": "潘华"}

    async def wechat(association, person, role):
        assert (association, person, role) == ("测试协会", "潘华", "秘书长")
        return "18612345678"

    enricher = AssociationBatchEnricher(
        official_site_resolver=resolve,
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
    )

    row = await enricher.enrich_one("测试协会")

    assert row.values["official_website"] == "https://association.example.cn/"
    assert row.values["secretary_general_name"] == "潘华"
    assert row.values["secretary_general_mobile"] == "18612345678"
    assert "wenxin_search" in row.sources


@pytest.mark.asyncio
async def test_search_profile_returns_basic_info(monkeypatch, tmp_path):
    """文心采集失败时降级 llm_gateway 直出（DeepSeek 不联网）。

    DeepSeek 直出不联网，秘书长/会员/办公室负责人姓名及手机号即使模型返回也
    强制 null（避免过期/幻觉，交给官网组织领导页或微信搜一搜精确取证）。
    显式 mock _collect_wenxin 返回 None，避免依赖 runtime import / 9222 状态。
    """
    import src.services.association_enrichment_providers as module

    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_wenxin(_name):
        return None  # 文心采集失败 → fallback DeepSeek 直出
    monkeypatch.setattr(providers, "_collect_wenxin", fake_wenxin)

    mock_result = {name: None for name in PROFILE_FIELDS}
    mock_result["address"] = "北京市测试路1号"
    mock_result["official_website"] = "https://example.cn"
    # 模型瞎给的姓名/手机号——fallback 分支必须强制清掉
    mock_result["secretary_general_name"] = "李四"
    mock_result["member_director_name"] = "王五"
    mock_result["office_director_name"] = "赵六"
    mock_result["secretary_general_mobile"] = "13800000000"

    class MockGateway:
        async def chat(self, **kwargs):
            return {"content": json.dumps(mock_result, ensure_ascii=False)}

    monkeypatch.setattr(module, "llm_gateway", MockGateway())

    result = await providers.search_profile("测试协会")
    assert result["address"] == "北京市测试路1号"
    assert result["secretary_general_name"] is None
    assert result["member_director_name"] is None
    assert result["office_director_name"] is None
    assert result["secretary_general_mobile"] is None


class _RecordingGateway:
    """记录 DeepSeek 收到的 messages，供验证文心/降级分支用。"""

    def __init__(self, content: str):
        self.content = content
        self.messages = None

    async def chat(self, **kwargs):
        self.messages = kwargs.get("messages")
        return {"content": self.content}


@pytest.mark.asyncio
async def test_search_profile_parses_wenxin_raw_text(monkeypatch, tmp_path):
    """文心联网采集到原文时，DeepSeek 从原文提取——官网/三个联系人姓名均来自原文。"""
    import src.services.association_enrichment_providers as module

    providers = ProjectAssociationProviders(repository_root=tmp_path)
    wenxin_answer = (
        "地址：北京市测试路1号\n"
        "官网网址：http://www.zgct.org.cn\n"
        "主管单位：工业和信息化部\n"
        "现任秘书长：李四\n"
        "会员服务部主任：王五\n"
        "办公室主任：赵六"
    )

    async def fake_wenxin(name):
        assert name == "测试协会"
        return {"ok": True, "answer": wenxin_answer, "note": ""}

    monkeypatch.setattr(providers, "_collect_wenxin", fake_wenxin)

    parsed = {name: None for name in PROFILE_FIELDS}
    parsed["address"] = "北京市测试路1号"
    parsed["official_website"] = "http://www.zgct.org.cn"
    parsed["secretary_general_name"] = "李四"
    parsed["member_director_name"] = "王五"
    parsed["office_director_name"] = "赵六"
    gateway = _RecordingGateway(json.dumps(parsed, ensure_ascii=False))
    monkeypatch.setattr(module, "llm_gateway", gateway)

    result = await providers.search_profile("测试协会")

    # DeepSeek 收到的必须是文心原文，而非泛泛提问（降级路径）
    assert "原文" in gateway.messages[0]["content"]
    assert wenxin_answer in gateway.messages[1]["content"]
    assert "给出测试协会的以下信息" not in gateway.messages[1]["content"]
    # 官网/三个联系人姓名从原文提取（文心联网，可靠）；手机号强制 null（由微信取证）
    assert result["official_website"] == "http://www.zgct.org.cn"
    assert result["address"] == "北京市测试路1号"
    assert result["secretary_general_name"] == "李四"
    assert result["member_director_name"] == "王五"
    assert result["office_director_name"] == "赵六"
    assert result["secretary_general_mobile"] is None


@pytest.mark.asyncio
async def test_search_profile_falls_back_when_wenxin_returns_none(monkeypatch, tmp_path):
    """文心 spawn 失败/返回 None 时，降级原 DeepSeek 直出（保底，不至整步空）。"""
    import src.services.association_enrichment_providers as module

    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_wenxin(name):
        return None

    monkeypatch.setattr(providers, "_collect_wenxin", fake_wenxin)

    parsed = {name: None for name in PROFILE_FIELDS}
    parsed["official_website"] = "https://example.cn"
    gateway = _RecordingGateway(json.dumps(parsed, ensure_ascii=False))
    monkeypatch.setattr(module, "llm_gateway", gateway)

    result = await providers.search_profile("测试协会")

    # 降级：user content 是泛泛提问，不含原文
    assert "给出测试协会的以下信息" in gateway.messages[1]["content"]
    assert result["official_website"] == "https://example.cn"


@pytest.mark.asyncio
async def test_search_profile_falls_back_on_captcha(monkeypatch, tmp_path):
    """文心返回验证码时，降级 DeepSeek 直出。"""
    import src.services.association_enrichment_providers as module

    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_wenxin(name):
        return {"ok": False, "answer": "", "note": "captcha"}

    monkeypatch.setattr(providers, "_collect_wenxin", fake_wenxin)

    parsed = {name: None for name in PROFILE_FIELDS}
    gateway = _RecordingGateway(json.dumps(parsed, ensure_ascii=False))
    monkeypatch.setattr(module, "llm_gateway", gateway)

    result = await providers.search_profile("测试协会")

    assert gateway.messages is not None
    assert all(v is None for v in result.values())


@pytest.mark.asyncio
async def test_progress_reports_stages_and_redacts_mobile():
    messages = []

    async def resolve(_name):
        return "https://association.example.cn/"

    async def collect(_url, _headless, **kwargs):
        return {"member_director_name": "张三"}

    async def fallback(_name):
        return {"member_director_name": "张三"}

    async def wechat(_association, _person, _role):
        return "18612345678"

    enricher = AssociationBatchEnricher(
        official_site_resolver=resolve,
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
        progress_reporter=messages.append,
    )
    row = await enricher.enrich_one("测试协会 18612345678")

    assert row.values["member_director_mobile"] == "18612345678"
    assert any("DeepSeek" in m or "基础信息" in m for m in messages)
    assert any("官网" in m for m in messages)
    assert any("微信检索" in m or "手机号" in m for m in messages)
    assert all("18612345678" not in message for message in messages)
    assert any("186****5678" in message for message in messages)


@pytest.mark.asyncio
async def test_progress_reporter_failure_does_not_interrupt_enrichment():
    async def resolve(_name):
        return None

    async def fallback(_name):
        return {"address": "北京市"}

    async def wechat(_association, _person, _role):
        return None

    def broken_reporter(_message):
        raise RuntimeError("terminal closed")

    enricher = AssociationBatchEnricher(
        official_site_resolver=resolve,
        official_profile_collector=lambda *_args: None,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
        progress_reporter=broken_reporter,
    )

    row = await enricher.enrich_one("测试协会")

    assert row.values["address"] == "北京市"
    assert row.processing_status == "partial"


@pytest.mark.asyncio
async def test_official_profile_passes_hostname_as_verified_domain(
    monkeypatch, tmp_path
):
    import src.services.association_enrichment_providers as module
    from src.services.association_profile_extractor import (
        AssociationProfile,
        ExtractionResult,
    )

    providers = ProjectAssociationProviders(repository_root=tmp_path)
    captured = {}

    async def collect(entry_url, domain, *, headless, **budgets):
        captured["collector"] = (entry_url, domain, headless, budgets)
        return []

    async def extract(pages, verified_domain, gateway=None):
        captured["extractor"] = (pages, verified_domain)
        return ExtractionResult(
            status="success",
            profile=AssociationProfile(
                **{name: None for name in PROFILE_FIELDS}
            ),
        )

    monkeypatch.setattr(module, "collect_official_pages_with_playwright", collect)
    monkeypatch.setattr(module, "extract_association_profile", extract)

    await providers.collect_official_profile(
        "https://www.example.cn/about/index.html", False
    )

    assert captured["collector"][1] == "www.example.cn"
    assert captured["extractor"][1] == "www.example.cn"


@pytest.mark.asyncio
async def test_official_profile_retries_invalid_llm_evidence_once(
    monkeypatch, tmp_path
):
    import src.services.association_enrichment_providers as module
    from src.services.association_profile_extractor import (
        AssociationProfile,
        ExtractionResult,
    )

    providers = ProjectAssociationProviders(repository_root=tmp_path)
    calls = []

    async def collect(_entry_url, _domain, *, headless, **_budgets):
        assert headless is False
        return []

    async def extract(pages, verified_domain, gateway=None):
        calls.append((pages, verified_domain))
        if len(calls) == 1:
            return ExtractionResult(
                status="inconclusive",
                reason_code="INVALID_EVIDENCE",
            )
        return ExtractionResult(
            status="success",
            profile=AssociationProfile(
                **{name: None for name in PROFILE_FIELDS}
            ),
        )

    monkeypatch.setattr(module, "collect_official_pages_with_playwright", collect)
    monkeypatch.setattr(module, "extract_association_profile", extract)

    result = await providers.collect_official_profile(
        "https://www.example.cn/about/index.html", False
    )

    assert len(calls) == 2
    assert calls[0][1] == calls[1][1] == "www.example.cn"
    assert set(result) == set(PROFILE_FIELDS)


@pytest.mark.asyncio
async def test_official_main_extraction_failure_still_runs_leadership_fallback(
    monkeypatch, tmp_path
):
    import src.services.association_enrichment_providers as module
    from src.services.association_profile_extractor import (
        ExtractionResult,
        VerifiedOfficialPage,
    )

    providers = ProjectAssociationProviders(repository_root=tmp_path)
    leadership_page = VerifiedOfficialPage(
        url="https://example.cn/leaders",
        title="领导机构",
        content="现任秘书长张秘书",
        verified_official=True,
    )

    async def collect(*_args, **_kwargs):
        return [leadership_page]

    async def extract(*_args, **_kwargs):
        return ExtractionResult(
            status="inconclusive",
            reason_code="STRICT_JSON_INVALID",
        )

    async def focused(pages, domain):
        assert pages == [leadership_page]
        assert domain == "example.cn"
        return {
            "secretary_general_name": "张秘书",
            "member_director_name": None,
            "office_director_name": None,
        }

    monkeypatch.setattr(module, "collect_official_pages_with_playwright", collect)
    monkeypatch.setattr(module, "extract_association_profile", extract)
    monkeypatch.setattr(providers, "_extract_leadership", focused)

    result = await providers.collect_official_profile("https://example.cn", False)

    assert result["secretary_general_name"] == "张秘书"
    assert set(result) == set(PROFILE_FIELDS)


@pytest.mark.asyncio
async def test_official_missing_secretary_triggers_llm_leadership_search(
    monkeypatch, tmp_path
):
    """词表采集没拿到秘书长 → 触发 LLM 逐层导航搜索，找到后纠错（覆盖文心错值）。

    这是官网纠正文心错秘书长姓名的主链路：搜索找到的页面交给 _extract_leadership，
    提取结果经 enricher override=True 覆盖文心值。
    """
    import src.services.association_enrichment_providers as module
    from src.services.association_profile_extractor import (
        ExtractionResult,
        VerifiedOfficialPage,
    )

    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def collect(*_args, **_kwargs):
        return []  # 词表采集一页没拿到

    async def extract(*_args, **_kwargs):
        return ExtractionResult(status="inconclusive", reason_code="STRICT_JSON_INVALID")

    search_calls = []

    async def fake_search(entry_url, domain, gateway, *, association_name="", **_kw):
        search_calls.append((entry_url, association_name))
        return [
            VerifiedOfficialPage(
                url="https://example.cn/about/leaders",
                title="现任领导",
                content="会长：张会长\n秘书长：王建琪",
                verified_official=True,
            )
        ]

    async def focused(pages, _domain):
        if not pages:
            # 第一次调用：词表采集没页面，leadership_pages 为空 → 返回空
            return {}
        # 第二次调用：LLM 搜索找到的现任领导页
        assert len(pages) == 1
        assert "王建琪" in pages[0].content
        return {
            "secretary_general_name": "王建琪",
            "member_director_name": None,
            "office_director_name": None,
        }

    monkeypatch.setattr(module, "collect_official_pages_with_playwright", collect)
    monkeypatch.setattr(module, "extract_association_profile", extract)
    monkeypatch.setattr(
        "src.services.official_site_leadership_search.search_leadership_pages_with_playwright",
        fake_search,
    )
    monkeypatch.setattr(providers, "_extract_leadership", focused)

    result = await providers.collect_official_profile(
        "https://example.cn", False, association_name="测试学会"
    )

    assert search_calls == [("https://example.cn", "测试学会")]
    assert result["secretary_general_name"] == "王建琪"


@pytest.mark.asyncio
async def test_official_secretary_found_skips_llm_leadership_search(
    monkeypatch, tmp_path
):
    """词表采集/领导兜底已拿到秘书长 → 不再触发 LLM 搜索（省 LLM 调用）。"""
    import src.services.association_enrichment_providers as module
    from src.services.association_profile_extractor import (
        AssociationProfile,
        ExtractionResult,
    )

    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def collect(*_args, **_kwargs):
        return []

    profile_kwargs = {name: None for name in PROFILE_FIELDS}
    profile_kwargs["secretary_general_name"] = "王建琪"
    profile = AssociationProfile(**profile_kwargs)

    async def extract(*_args, **_kwargs):
        return ExtractionResult(status="success", profile=profile)

    async def fake_search(*_args, **_kwargs):
        raise AssertionError("secretary already found; search must not run")

    monkeypatch.setattr(module, "collect_official_pages_with_playwright", collect)
    monkeypatch.setattr(module, "extract_association_profile", extract)
    monkeypatch.setattr(
        "src.services.official_site_leadership_search.search_leadership_pages_with_playwright",
        fake_search,
    )

    result = await providers.collect_official_profile("https://example.cn", False)

    assert result["secretary_general_name"] == "王建琪"


@pytest.mark.asyncio
async def test_credit_guard_stops_batch_before_next_association():
    """余额守卫抛异常 → 当前协会标 aborted、批次熔断，后续协会不再处理。

    真机教训：中途 402（NoCreditError）被当单协会失败吞掉后，后续协会仍会
    白跑文心/微信等不计费步骤且无提醒——守卫把透支/白跑范围压到一个协会。
    """

    async def collect(_url, _headless, **_kwargs):
        raise AssertionError("no official URL")

    async def fallback(_name):
        return {"address": "北京市"}

    async def wechat(_association, _person, _role):
        return None

    guard_calls = []

    class _NoCredit(RuntimeError):
        error_code = "NO_CREDIT"

    async def guard():
        guard_calls.append(1)
        if len(guard_calls) >= 2:
            raise _NoCredit("积分余额不足（0.00）")

    enricher = AssociationBatchEnricher(
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
        credit_guard=guard,
    )
    rows = await enricher.enrich_many(["协会一", "协会二", "协会三"])

    assert guard_calls and len(guard_calls) == 2  # 第二个协会开始前拦停
    assert [r.association_name for r in rows] == ["协会一", "协会二"]
    assert rows[0].processing_status == "partial"
    assert rows[1].processing_status == "aborted"
    assert "credit_guard:NO_CREDIT" in rows[1].errors
    assert rows.aborted is True
    assert rows.abort_error_code == "NO_CREDIT"
    # 协会三未处理


@pytest.mark.asyncio
async def test_credit_guard_pass_keeps_batch_running():
    """守卫通过（余额充足）→ 批次正常跑完全部协会。"""

    async def collect(_url, _headless, **_kwargs):
        raise AssertionError("no official URL")

    async def fallback(_name):
        return {"address": "北京市"}

    async def wechat(_association, _person, _role):
        return None

    async def guard():
        return None

    enricher = AssociationBatchEnricher(
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
        credit_guard=guard,
    )
    rows = await enricher.enrich_many(["协会一", "协会二"])

    assert [r.association_name for r in rows] == ["协会一", "协会二"]
    assert rows.aborted is False


@pytest.mark.asyncio
async def test_official_secretary_miss_emits_audit_event(monkeypatch, tmp_path):
    """官网跑完整链路（词表+LLM搜索）仍无秘书长 → 记 official_secretary_miss
    审计事件（协会名+官网 URL），供后续程序优化统计。"""
    import src.services.association_enrichment_providers as module
    from src.services.association_profile_extractor import (
        AssociationProfile,
        ExtractionResult,
    )

    providers = ProjectAssociationProviders(repository_root=tmp_path)
    events = []
    providers._audit = lambda **event: events.append(event)

    async def collect(*_args, **_kwargs):
        return []

    async def extract(*_args, **_kwargs):
        return ExtractionResult(
            status="success",
            profile=AssociationProfile(**{name: None for name in PROFILE_FIELDS}),
        )

    async def fake_search(*_args, **_kwargs):
        return []

    monkeypatch.setattr(module, "collect_official_pages_with_playwright", collect)
    monkeypatch.setattr(module, "extract_association_profile", extract)
    monkeypatch.setattr(
        "src.services.official_site_leadership_search.search_leadership_pages_with_playwright",
        fake_search,
    )

    result = await providers.collect_official_profile(
        "https://example.cn", False, association_name="测试学会"
    )

    assert result["secretary_general_name"] is None
    miss_events = [e for e in events if e.get("kind") == "official_secretary_miss"]
    assert len(miss_events) == 1
    assert miss_events[0]["detail"] == {
        "official_url": "https://example.cn",
        "association_name": "测试学会",
    }


@pytest.mark.asyncio
async def test_official_secretary_found_no_miss_audit(monkeypatch, tmp_path):
    """拿到秘书长 → 不发 official_secretary_miss（避免噪声）。"""
    import src.services.association_enrichment_providers as module
    from src.services.association_profile_extractor import (
        AssociationProfile,
        ExtractionResult,
    )

    providers = ProjectAssociationProviders(repository_root=tmp_path)
    events = []
    providers._audit = lambda **event: events.append(event)

    async def collect(*_args, **_kwargs):
        return []

    profile_kwargs = {name: None for name in PROFILE_FIELDS}
    profile_kwargs["secretary_general_name"] = "王建琪"

    async def extract(*_args, **_kwargs):
        return ExtractionResult(
            status="success", profile=AssociationProfile(**profile_kwargs)
        )

    monkeypatch.setattr(module, "collect_official_pages_with_playwright", collect)
    monkeypatch.setattr(module, "extract_association_profile", extract)

    await providers.collect_official_profile(
        "https://example.cn", False, association_name="测试学会"
    )

    assert not [e for e in events if e.get("kind") == "official_secretary_miss"]


@pytest.mark.asyncio
async def test_wechat_nonzero_exit_audits_failure_and_available_artifact(
    monkeypatch, tmp_path
):
    import src.services.association_enrichment_providers as module

    class FailedProcess:
        returncode = 1

        async def communicate(self):
            return (
                json.dumps(
                    {
                        "ok": False,
                        "session_closed": False,
                        "status": "failed",
                        "artifact_ref": "failure.dpapi",
                    }
                ).encode("utf-8"),
                b"",
            )

        def kill(self):
            pass

        async def wait(self):
            pass

    async def fake_subprocess(*_args, **_kwargs):
        return FailedProcess()

    events = []
    providers = ProjectAssociationProviders(
        repository_root=tmp_path,
        audit_callback=lambda **event: events.append(event),
    )
    artifact_reads = []

    async def fake_audit_artifact(ref, **_kwargs):
        artifact_reads.append(ref)

    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)
    monkeypatch.setattr(providers, "_audit_wechat_artifact", fake_audit_artifact)

    with pytest.raises(RuntimeError, match="WECHAT_RPA_FAILED"):
        await providers.wechat_mobile("测试协会", "张三", "会长")

    assert artifact_reads == ["failure.dpapi"]
    assert [event["kind"] for event in events] == [
        "wechat_start",
        "wechat_failed",
    ]


def test_workbook_token_usage_uses_separate_sheet_without_changing_main_contract(tmp_path):
    from src.services.association_batch_enrichment import AssociationEnrichmentRow

    path = write_enrichment_workbook(
        [AssociationEnrichmentRow(association_name="协会一")],
        tmp_path / "tokens.xlsx",
        token_usage={
            "associations": {
                "协会一": {
                    "input_tokens": 100,
                    "cached_input_tokens": 40,
                    "output_tokens": 20,
                    "total_tokens": 120,
                    "call_count": 1,
                }
            },
            "batch_overhead": {
                "input_tokens": 10,
                "cached_input_tokens": 2,
                "output_tokens": 3,
                "total_tokens": 13,
                "call_count": 1,
            },
            "batch": {
                "input_tokens": 110,
                "cached_input_tokens": 42,
                "output_tokens": 23,
                "total_tokens": 133,
                "call_count": 2,
                "average_input_tokens": 110,
                "average_cached_input_tokens": 42,
                "average_output_tokens": 23,
                "average_total_tokens": 133,
            },
        },
    )
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        from src.services.association_batch_enrichment import _EXCEL_HEADERS
        assert workbook["协会信息"].max_column == len(_EXCEL_HEADERS)
        token_sheet = workbook["Token用量"]
        assert token_sheet.cell(2, 1).value == "协会一"
        assert token_sheet.cell(2, 2).value == 100
        assert token_sheet.cell(2, 3).value == 40
        assert token_sheet.cell(3, 1).value == "批次开销（清单解析）"
        assert token_sheet.cell(3, 5).value == 13
        assert token_sheet.cell(7, 1).value == "cached_input_tokens"
    finally:
        workbook.close()


@pytest.mark.asyncio
async def test_wechat_session_cleanup_failure_keeps_original_error_after_audit(
    monkeypatch, tmp_path
):
    import src.services.association_enrichment_providers as module

    class CompletedProcess:
        returncode = 1

        async def communicate(self):
            return (
                json.dumps(
                    {
                        "ok": False,
                        "session_closed": False,
                        "status": "failed",
                        "error_code": "PLUGIN_CLOSE_TIMEOUT",
                        "stage": "cleanup",
                        "result_status": "found",
                        "artifact_ref": "cleanup.dpapi",
                    }
                ).encode("utf-8"),
                "可能包含手机号 18512345678".encode("utf-8"),
            )

    async def fake_subprocess(*_args, **_kwargs):
        return CompletedProcess()

    events = []
    providers = ProjectAssociationProviders(
        repository_root=tmp_path,
        audit_callback=lambda **event: events.append(event),
    )

    async def broken_audit_artifact(*_args, **_kwargs):
        raise RuntimeError("audit unavailable")

    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)
    monkeypatch.setattr(providers, "_audit_wechat_artifact", broken_audit_artifact)
    async def recovered_mobile(*_args, **_kwargs):
        return "18512345678"

    monkeypatch.setattr(providers, "_extract_wechat_mobile", recovered_mobile)

    with pytest.raises(RuntimeError, match="PLUGIN_CLOSE_TIMEOUT") as caught:
        await providers.wechat_mobile("测试协会", "张三", "会长")

    assert caught.value.stage == "cleanup"
    assert caught.value.session_fatal is True
    assert caught.value.recovered_mobile == "18512345678"

    assert [event["kind"] for event in events] == [
        "wechat_start",
        "wechat_audit_failed",
        "wechat_failed",
    ]
    assert events[-1]["summary"] == "PLUGIN_CLOSE_TIMEOUT"
    assert events[-1]["detail"]["stderr_bytes"] > 0
    assert len(events[-1]["detail"]["stderr_sha256"]) == 64
    assert "18512345678" not in json.dumps(events[-1], ensure_ascii=False)


@pytest.mark.asyncio
async def test_wechat_unsubmitted_input_failure_is_nonfatal_after_clean_close(
    monkeypatch, tmp_path
):
    import src.services.association_enrichment_providers as module

    class CompletedProcess:
        returncode = 1

        async def communicate(self):
            return (
                json.dumps(
                    {
                        "ok": False,
                        "error_code": "SEARCH_INPUT_READBACK_MISMATCH",
                        "stage": "input_verify",
                        "session_closed": True,
                    }
                ).encode("utf-8"),
                b"",
            )

    async def fake_subprocess(*_args, **_kwargs):
        return CompletedProcess()

    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    with pytest.raises(RuntimeError, match="SEARCH_INPUT_READBACK_MISMATCH") as caught:
        await providers.wechat_mobile("测试协会", "张三", "会长")

    assert caught.value.stage == "input_verify"
    assert caught.value.session_fatal is False
    assert caught.value.recovered_mobile is None


@pytest.mark.asyncio
async def test_wechat_input_failure_is_fatal_when_cleanup_is_not_confirmed(
    monkeypatch, tmp_path
):
    import src.services.association_enrichment_providers as module

    class CompletedProcess:
        returncode = 1

        async def communicate(self):
            return (
                json.dumps(
                    {
                        "ok": False,
                        "error_code": "SEARCH_INPUT_LOCATOR_FAILED",
                        "stage": "input_verify",
                        "session_closed": False,
                    }
                ).encode("utf-8"),
                b"",
            )

    async def fake_subprocess(*_args, **_kwargs):
        return CompletedProcess()

    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    with pytest.raises(RuntimeError, match="SEARCH_INPUT_LOCATOR_FAILED") as caught:
        await providers.wechat_mobile("测试协会", "张三", "会长")

    assert caught.value.session_fatal is True


@pytest.mark.asyncio
async def test_wechat_timeout_audits_stable_code_and_reraises(monkeypatch, tmp_path):
    import src.services.association_enrichment_providers as module

    class HangingProcess:
        returncode = None

    async def fake_subprocess(*_args, **_kwargs):
        return HangingProcess()

    timeout_calls = []

    async def fake_communicate(*_args, **kwargs):
        timeout_calls.append(kwargs)
        raise RuntimeError("WECHAT_RPA_TIMEOUT")

    events = []
    providers = ProjectAssociationProviders(
        repository_root=tmp_path,
        audit_callback=lambda **event: events.append(event),
    )
    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)
    monkeypatch.setattr(
        providers,
        "_communicate_with_timeout",
        fake_communicate,
    )

    with pytest.raises(RuntimeError, match="WECHAT_RPA_TIMEOUT"):
        await providers.wechat_mobile("测试协会", "张三", "会长")

    assert [event["kind"] for event in events] == [
        "wechat_start",
        "wechat_failed",
    ]
    assert events[-1]["summary"] == "WECHAT_RPA_TIMEOUT"
    assert timeout_calls == [{
        "timeout_seconds": 600,
        "error_code": "WECHAT_RPA_TIMEOUT",
    }]


@pytest.mark.asyncio
async def test_wechat_invalid_json_audits_response_error(monkeypatch, tmp_path):
    import src.services.association_enrichment_providers as module

    class CompletedProcess:
        returncode = 0

        async def communicate(self):
            return b"not-json", b""

    async def fake_subprocess(*_args, **_kwargs):
        return CompletedProcess()

    events = []
    providers = ProjectAssociationProviders(
        repository_root=tmp_path,
        audit_callback=lambda **event: events.append(event),
    )
    monkeypatch.setattr(module.asyncio, "create_subprocess_exec", fake_subprocess)

    with pytest.raises(RuntimeError, match="WECHAT_RPA_RESPONSE_INVALID"):
        await providers.wechat_mobile("测试协会", "张三", "会长")

    assert events[-1]["kind"] == "wechat_failed"
    assert events[-1]["summary"] == "WECHAT_RPA_RESPONSE_INVALID"


# ---- 秘书长手机号「文心快速路径」enricher 级测试 ----


@pytest.mark.asyncio
async def test_secretary_wenxin_hit_skips_wechat():
    """秘书长文心命中手机号 → 填值并跳过微信；会员/办公室负责人仍走微信。"""
    async def collect(url, _headless, **kwargs):
        return {}

    async def fallback(_name):
        return {
            "secretary_general_name": "秘书长乙",
            "member_director_name": "会员丙",
            "office_director_name": "办公室丁",
        }

    wenxin_calls = []
    async def wenxin_secretary_mobile(association, name):
        wenxin_calls.append((association, name))
        return "13812345678"  # 秘书长文心命中

    wechat_calls = []
    async def wechat(association, person, role):
        wechat_calls.append((association, person, role))
        return "13900000000"

    enricher = AssociationBatchEnricher(
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
        wenxin_secretary_mobile_provider=wenxin_secretary_mobile,
    )
    row = await enricher.enrich_one("测试协会")

    assert row.values["secretary_general_mobile"] == "13812345678"
    assert row.values["member_director_mobile"] == "13900000000"
    assert row.values["office_director_mobile"] == "13900000000"
    assert "wenxin_mobile:秘书长" in row.sources
    assert "wechat:会员部主任" in row.sources
    assert "wechat:办公室主任" in row.sources
    assert "wechat:秘书长" not in row.sources
    # 秘书长走文心（命中跳过微信），会员/办公室负责人走微信
    assert wechat_calls == [
        ("测试协会", "会员丙", "会员部主任"),
        ("测试协会", "办公室丁", "办公室主任"),
    ]
    assert wenxin_calls == [("测试协会", "秘书长乙")]


@pytest.mark.asyncio
async def test_secretary_wenxin_miss_falls_back_to_wechat():
    """秘书长文心未命中(返None) → 落回微信搜一搜。"""
    async def collect(url, _headless, **kwargs):
        return {}

    async def fallback(_name):
        return {"secretary_general_name": "秘书长乙"}

    async def wenxin_secretary_mobile(_association, _name):
        return None  # 文心未查到

    async def wechat(_association, _person, role):
        return "13900000001" if role == "秘书长" else "13900000000"

    enricher = AssociationBatchEnricher(
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
        wenxin_secretary_mobile_provider=wenxin_secretary_mobile,
    )
    row = await enricher.enrich_one("测试协会")

    assert row.values["secretary_general_mobile"] == "13900000001"
    assert "wechat:秘书长" in row.sources
    assert "wenxin_mobile:秘书长" not in row.sources


@pytest.mark.asyncio
async def test_secretary_wenxin_exception_falls_back_to_wechat():
    """秘书长文心 provider 抛异常 → 不崩，落回微信兜底。"""
    async def collect(url, _headless, **kwargs):
        return {}

    async def fallback(_name):
        return {"secretary_general_name": "秘书长乙"}

    async def wenxin_secretary_mobile(_association, _name):
        raise RuntimeError("wenxin boom")

    async def wechat(_association, _person, role):
        return "13900000002" if role == "秘书长" else "13900000000"

    enricher = AssociationBatchEnricher(
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
        wenxin_secretary_mobile_provider=wenxin_secretary_mobile,
    )
    row = await enricher.enrich_one("测试协会")

    assert row.values["secretary_general_mobile"] == "13900000002"
    assert "wechat:秘书长" in row.sources


@pytest.mark.asyncio
async def test_non_secretary_roles_never_use_wenxin_mobile():
    """会员/办公室负责人手机号永远不走文心（只秘书长走文心快速路径）。"""
    async def collect(url, _headless, **kwargs):
        return {}

    async def fallback(_name):
        return {
            "secretary_general_name": "秘书长乙",
            "member_director_name": "会员丙",
            "office_director_name": "办公室丁",
        }

    wenxin_calls = []
    async def wenxin_secretary_mobile(_association, name):
        wenxin_calls.append(name)
        return None

    async def wechat(_association, _person, _role):
        return None

    enricher = AssociationBatchEnricher(
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
        wenxin_secretary_mobile_provider=wenxin_secretary_mobile,
    )
    await enricher.enrich_one("测试协会")

    # 文心只被秘书长调用一次，会员/办公室负责人从不调用
    assert wenxin_calls == ["秘书长乙"]


@pytest.mark.asyncio
async def test_secretary_wenxin_skipped_when_no_secretary_name():
    """无秘书长姓名时，文心 provider 不被调用（Step4 守卫先 continue）。"""
    async def collect(url, _headless, **kwargs):
        return {}

    async def fallback(_name):
        return {"member_director_name": "会员丙"}  # 无秘书长姓名

    wenxin_calls = []
    async def wenxin_secretary_mobile(_association, name):
        wenxin_calls.append(name)
        return None

    async def wechat(_association, _person, _role):
        return None

    enricher = AssociationBatchEnricher(
        official_profile_collector=collect,
        fallback_profile_provider=fallback,
        wechat_mobile_provider=wechat,
        wenxin_secretary_mobile_provider=wenxin_secretary_mobile,
    )
    await enricher.enrich_one("测试协会")

    assert wenxin_calls == []


# ---- wenxin_search_secretary_mobile providers 级测试（正则圈候选 + LLM 判别归属）----


class _MobileJudgeGateway:
    """按号码返回固定 mobile 判别结果的假网关。"""

    def __init__(self, mobile: str | None):
        self.mobile = mobile
        self.messages = None

    async def chat(self, **kwargs):
        self.messages = kwargs.get("messages")
        return {"content": json.dumps({"mobile": self.mobile}, ensure_ascii=False)}


@pytest.mark.asyncio
async def test_wenxin_secretary_mobile_extracts_valid_number(monkeypatch, tmp_path):
    """文心回答明确绑定秘书长本人 + LLM 确认 → 返回该号码。"""
    import src.services.association_enrichment_providers as module
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_query(_query, **kwargs):
        return {"ok": True, "answer": "秘书长张三，手机 13812345678", "note": ""}
    monkeypatch.setattr(providers, "_collect_wenxin_query", fake_query)
    gateway = _MobileJudgeGateway("13812345678")
    monkeypatch.setattr(module, "llm_gateway", gateway)

    mobile = await providers.wenxin_search_secretary_mobile("测试协会", "张三")
    assert mobile == "13812345678"


@pytest.mark.asyncio
async def test_wenxin_secretary_mobile_rejects_landline_fax(monkeypatch, tmp_path):
    """回答含座机/传真 + 手机 → 候选只含手机号（座机0开头/传真被 1[3-9] 正则排除）。"""
    import src.services.association_enrichment_providers as module
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_query(_query, **kwargs):
        return {
            "ok": True,
            "answer": "电话：010-12345678 传真：010-87654321 手机：13812345678",
            "note": "",
        }
    monkeypatch.setattr(providers, "_collect_wenxin_query", fake_query)
    gateway = _MobileJudgeGateway("13812345678")
    monkeypatch.setattr(module, "llm_gateway", gateway)

    mobile = await providers.wenxin_search_secretary_mobile("测试协会", "张三")
    assert mobile == "13812345678"


@pytest.mark.asyncio
async def test_wenxin_mobile_llm_picks_right_person_when_multiple(monkeypatch, tmp_path):
    """多个手机号（协会其他人号码在前）→ LLM 判别出秘书长本人的号码，不再取第一个。

    这是纯正则取第一个方案的真实事故场景：文心回答常含办公室/其他负责人的
    号码且排在前面，取第一个会张冠李戴。
    """
    import src.services.association_enrichment_providers as module
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_query(_query, **kwargs):
        return {
            "ok": True,
            "answer": (
                "办公室李四：13811112222。"
                "秘书长张三：13833334444。"
                "会员部王五：13855556666。"
            ),
            "note": "",
        }
    monkeypatch.setattr(providers, "_collect_wenxin_query", fake_query)
    gateway = _MobileJudgeGateway("13833334444")  # LLM 选出张三的号码（非第一个）
    monkeypatch.setattr(module, "llm_gateway", gateway)

    mobile = await providers.wenxin_search_secretary_mobile("测试协会", "张三")
    assert mobile == "13833334444"
    # 判别提示必须带上全部候选 + 原文，且要求明确绑定才返回
    assert "13811112222" in gateway.messages[1]["content"]
    assert "张三" in gateway.messages[0]["content"]
    assert "严禁猜测" in gateway.messages[0]["content"]


@pytest.mark.asyncio
async def test_wenxin_mobile_llm_unconfirmed_returns_none(monkeypatch, tmp_path):
    """LLM 无法确认候选中有秘书长本人的号码 → None（落回微信兜底，宁缺毋错）。"""
    import src.services.association_enrichment_providers as module
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_query(_query, **kwargs):
        return {"ok": True, "answer": "13811112222 13833334444", "note": ""}
    monkeypatch.setattr(providers, "_collect_wenxin_query", fake_query)
    monkeypatch.setattr(
        module, "llm_gateway", _MobileJudgeGateway(None)  # LLM 判定都不是本人
    )

    mobile = await providers.wenxin_search_secretary_mobile("测试协会", "张三")
    assert mobile is None


@pytest.mark.asyncio
async def test_wenxin_mobile_llm_hallucinated_number_rejected(monkeypatch, tmp_path):
    """LLM 返回不在候选中的号码（幻觉/编造）→ 拒绝返回 None。"""
    import src.services.association_enrichment_providers as module
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_query(_query, **kwargs):
        return {"ok": True, "answer": "秘书长张三：13812345678", "note": ""}
    monkeypatch.setattr(providers, "_collect_wenxin_query", fake_query)
    monkeypatch.setattr(
        module, "llm_gateway", _MobileJudgeGateway("13900000000")  # 不在候选中
    )

    mobile = await providers.wenxin_search_secretary_mobile("测试协会", "张三")
    assert mobile is None


@pytest.mark.asyncio
async def test_wenxin_secretary_mobile_none_on_no_mobile(monkeypatch, tmp_path):
    """回答只有座机、无手机号 → None（不进 LLM 判别）。"""
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_query(_query, **kwargs):
        return {"ok": True, "answer": "办公电话 010-12345678，无手机号", "note": ""}
    monkeypatch.setattr(providers, "_collect_wenxin_query", fake_query)

    mobile = await providers.wenxin_search_secretary_mobile("测试协会", "张三")
    assert mobile is None


@pytest.mark.asyncio
async def test_wenxin_secretary_mobile_none_on_wenxin_failure(monkeypatch, tmp_path):
    """文心采集失败(返回None) → None。"""
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_query(_query, **kwargs):
        return None
    monkeypatch.setattr(providers, "_collect_wenxin_query", fake_query)

    mobile = await providers.wenxin_search_secretary_mobile("测试协会", "张三")
    assert mobile is None


@pytest.mark.asyncio
async def test_wenxin_secretary_mobile_none_on_captcha(monkeypatch, tmp_path):
    """文心返回验证码 → None。"""
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_query(_query, **kwargs):
        return {"ok": False, "note": "captcha"}
    monkeypatch.setattr(providers, "_collect_wenxin_query", fake_query)

    mobile = await providers.wenxin_search_secretary_mobile("测试协会", "张三")
    assert mobile is None


@pytest.mark.asyncio
async def test_wenxin_secretary_mobile_none_on_llm_failure(monkeypatch, tmp_path):
    """LLM 判别调用失败（非 JSON/超时）→ 不崩，None 落回微信兜底。"""
    import src.services.association_enrichment_providers as module
    providers = ProjectAssociationProviders(repository_root=tmp_path)

    async def fake_query(_query, **kwargs):
        return {"ok": True, "answer": "秘书长张三：13812345678", "note": ""}
    monkeypatch.setattr(providers, "_collect_wenxin_query", fake_query)

    class BoomGateway:
        async def chat(self, **kwargs):
            raise RuntimeError("llm down")
    monkeypatch.setattr(module, "llm_gateway", BoomGateway())

    mobile = await providers.wenxin_search_secretary_mobile("测试协会", "张三")
    assert mobile is None
