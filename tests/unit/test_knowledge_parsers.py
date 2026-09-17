"""
文档解析器测试
"""

import pytest

pytestmark = pytest.mark.skills

from src.knowledge.parsers import DocumentParseError
from src.knowledge.parsers.word_parser import WordParser
from src.knowledge.parsers.excel_parser import ExcelParser
from src.knowledge.parsers.ppt_parser import PPTParser
from src.knowledge.parsers.pdf_parser import PDFParser
from src.knowledge.parsers.parser_factory import parser_factory

# OLE 复合文件魔数：加密 docx/xlsx/pptx 和老版 .doc/.xls/.ppt 都是这种格式（非 zip）
OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 512


class TestParserFactory:
    """解析器工厂测试"""

    def test_get_word_parser(self):
        parser = parser_factory.get_parser("test.docx")
        assert isinstance(parser, WordParser)

    def test_get_excel_parser(self):
        parser = parser_factory.get_parser("test.xlsx")
        assert isinstance(parser, ExcelParser)

    def test_get_ppt_parser(self):
        parser = parser_factory.get_parser("test.pptx")
        assert isinstance(parser, PPTParser)

    def test_get_pdf_parser(self):
        parser = parser_factory.get_parser("test.pdf")
        assert isinstance(parser, PDFParser)

    def test_unsupported_format_returns_none(self):
        parser = parser_factory.get_parser("test.xyz")
        assert parser is None


class TestWordParser:
    """Word 解析器测试"""

    def test_supported_extensions(self):
        parser = WordParser()
        assert ".docx" in parser.supported_extensions()

    @pytest.mark.asyncio
    async def test_parse_missing_file(self):
        parser = WordParser()
        with pytest.raises(Exception):
            await parser.parse("non_existent_file.docx")

    @pytest.mark.asyncio
    async def test_parse_encrypted_or_fake_docx_raises_friendly_error(self, tmp_path):
        """加密 docx / .doc 改后缀（OLE 复合文件，非 zip）应抛带指引的 DocumentParseError"""
        file_path = tmp_path / "encrypted.docx"
        file_path.write_bytes(OLE_MAGIC)

        parser = WordParser()
        with pytest.raises(DocumentParseError) as exc_info:
            await parser.parse(str(file_path))
        assert "另存为" in str(exc_info.value)

    @pytest.mark.asyncio
    async def test_parse_wrong_content_type_docx_raises_friendly_error(self, tmp_path):
        """合法 zip 但主部件 content type 不是 Word（其他 OOXML 改后缀）应抛 DocumentParseError"""
        from pptx import Presentation

        file_path = tmp_path / "fake.docx"
        Presentation().save(file_path)

        parser = WordParser()
        with pytest.raises(DocumentParseError) as exc_info:
            await parser.parse(str(file_path))
        assert "另存为" in str(exc_info.value)


class TestExcelParser:
    """Excel 解析器测试"""

    def test_supported_extensions(self):
        parser = ExcelParser()
        assert ".xlsx" in parser.supported_extensions()

    @pytest.mark.asyncio
    async def test_excel_parse_data_rows_are_own_paragraphs(self, tmp_path):
        """每个数据行用空行（\\n\\n）分隔成独立段落，保证 TextChunker 能按行粒度切块"""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "SKU"
        ws.append(["SKU编号", "名称", "描述"])
        ws.append(["YD0001", "防晒冰袖", "冰丝凉感面料，UPF50+防晒，透气吸汗，高弹贴合，不卷边不滑落，户外骑行、跑步、开车、日常通勤均可，男女同款多色可选" * 5])
        ws.append(["YD0002", "跑步鞋", "网面透气鞋面，橡胶软底大底，轻便舒适，透气不闷脚，日常穿搭、通勤、散步、短途出行均可，百搭显瘦" * 5])
        file_path = tmp_path / "test.xlsx"
        wb.save(file_path)

        parser = ExcelParser()
        result = await parser.parse(str(file_path))
        text = result.text

        # 段落之间用空行分隔
        assert "\n\n" in text

        # 每个数据行是独立段落（不在同一段里）
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        assert paragraphs[0].startswith("## 工作表: SKU")
        assert any(p.startswith("YD0001 | 防晒冰袖") for p in paragraphs)
        assert any(p.startswith("YD0002 | 跑步鞋") for p in paragraphs)

        # 用 TextChunker 验证：数据行足够长时能切成多个 chunk
        # （旧实现 \n join 会把行合并成大段落，超长后按 6000 字符硬切，单个商品被稀释）
        from src.knowledge.chunker import TextChunker
        chunker = TextChunker(chunk_size=512, overlap=64)
        chunks = chunker.chunk(text)
        assert len(chunks) >= 2

    @pytest.mark.asyncio
    async def test_parse_encrypted_or_fake_xlsx_raises_friendly_error(self, tmp_path):
        """加密 xlsx / .xls 改后缀（OLE 复合文件，非 zip）应抛带指引的 DocumentParseError"""
        file_path = tmp_path / "encrypted.xlsx"
        file_path.write_bytes(OLE_MAGIC)

        parser = ExcelParser()
        with pytest.raises(DocumentParseError) as exc_info:
            await parser.parse(str(file_path))
        assert "另存为" in str(exc_info.value)

    @pytest.mark.asyncio
    async def test_parse_zip_without_workbook_part_xlsx_raises_friendly_error(self, tmp_path):
        """合法 zip 但缺少工作簿部件（其他 OOXML 改后缀）应抛带指引的 DocumentParseError"""
        from docx import Document

        file_path = tmp_path / "fake.xlsx"
        Document().save(file_path)

        parser = ExcelParser()
        with pytest.raises(DocumentParseError) as exc_info:
            await parser.parse(str(file_path))
        assert "另存为" in str(exc_info.value)


class TestPPTParser:
    """PPT 解析器测试"""

    def test_supported_extensions(self):
        parser = PPTParser()
        assert ".pptx" in parser.supported_extensions()

    @pytest.mark.asyncio
    async def test_parse_encrypted_or_fake_pptx_raises_friendly_error(self, tmp_path):
        """加密 pptx / .ppt 改后缀（OLE 复合文件，非 zip）应抛带指引的 DocumentParseError"""
        file_path = tmp_path / "encrypted.pptx"
        file_path.write_bytes(OLE_MAGIC)

        parser = PPTParser()
        with pytest.raises(DocumentParseError) as exc_info:
            await parser.parse(str(file_path))
        assert "另存为" in str(exc_info.value)

    @pytest.mark.asyncio
    async def test_parse_wrong_content_type_pptx_raises_friendly_error(self, tmp_path):
        """合法 zip 但主部件 content type 不是 PPT（其他 OOXML 改后缀）应抛 DocumentParseError"""
        from docx import Document

        file_path = tmp_path / "fake.pptx"
        Document().save(file_path)

        parser = PPTParser()
        with pytest.raises(DocumentParseError) as exc_info:
            await parser.parse(str(file_path))
        assert "另存为" in str(exc_info.value)


class TestPDFParser:
    """PDF 解析器测试"""

    def test_supported_extensions(self):
        parser = PDFParser()
        assert ".pdf" in parser.supported_extensions()
