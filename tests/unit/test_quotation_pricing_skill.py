"""quotation-pricing 技能移植单测：注册发现、roundtrip 全链路、生成器列序回归、非法输入。"""

import json
import subprocess
import sys
from pathlib import Path

import pytest

pytestmark = pytest.mark.skills

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SKILL_DIR = PROJECT_ROOT / "src" / "skills" / "quotation-pricing-1.0.0"
SAMPLE_XLSX = SKILL_DIR / "sample" / "HappyRoom_家具软装.xlsx"


def _load_registry():
    from src.core.skill_registry import SkillRegistry

    registry = SkillRegistry()
    registry.load_from_directory(SKILL_DIR.parent)
    return registry


def _run(cmd, cwd):
    return subprocess.run(
        cmd, cwd=str(cwd), capture_output=True, text=True, timeout=120,
    )


class TestQuotationPricingSkillRegistration:
    def test_skill_registered(self):
        registry = _load_registry()
        skill = registry.get("quotation-pricing")
        assert skill is not None
        assert skill.version == "1.0.0"

    def test_requires_tools_declared(self):
        registry = _load_registry()
        skill = registry.get("quotation-pricing")
        assert list(getattr(skill, "requires_tools", [])) == [
            "knowledge_base_search", "knowledge_file_search",
        ]

    def test_scripts_and_references_present(self):
        for name in ("parse_quotation.py", "fill_quotation.py", "review.py"):
            assert (SKILL_DIR / "scripts" / name).is_file()
        for name in ("kb_search.md", "workflow.md"):
            assert (SKILL_DIR / "references" / name).is_file()
        assert SAMPLE_XLSX.is_file()


class TestRoundtrip:
    @pytest.fixture()
    def workdir(self, tmp_path):
        wd = tmp_path / "rt"
        wd.mkdir()
        return wd

    def test_parse_extracts_upstream_layout(self, workdir):
        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "parse_quotation.py"),
             str(SAMPLE_XLSX), "--json", "items_extracted.json", "--project", "Demo Hotel"],
            workdir,
        )
        assert result.returncode == 0, result.stderr
        data = json.loads((workdir / "items_extracted.json").read_text(encoding="utf-8"))
        assert data["project"] == "Demo Hotel"
        assert data["item_count"] == 3
        codes = {it["code"] for it in data["items"]}
        assert codes == {"HR-1601", "HR-1602", "HR-1701"}
        first = data["items"][0]
        assert first["row"] == 8
        assert first["area"] == "HappyRoom"
        assert "单人沙发" in first["title_cn_en"]
        assert any("velvet" in b for b in first["spec_bullets"])

    def test_fill_and_review_full_chain(self, workdir):
        # parse
        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "parse_quotation.py"),
             str(SAMPLE_XLSX), "--json", "items_extracted.json"],
            workdir,
        )
        assert result.returncode == 0, result.stderr

        # 模拟 agent 填写：全范围 + 1 直配 + 1 替换 + 1 AI 建议价
        data = json.loads((workdir / "items_extracted.json").read_text(encoding="utf-8"))
        fills = {
            "HR-1601": {"unit_price": 1200, "source": "知识库-KB",
                        "kb_match": "单人沙发 KF-201（价目表 2026.xlsx）", "spec_consistent": True},
            "HR-1602": {"unit_price": 2600, "source": "知识库-KB(替换)",
                        "kb_match": "双人沙发 KF-202", "spec_consistent": False,
                        "replacement_desc": "原：布艺 W1800 → 替换：真皮 W1600"},
            "HR-1701": {"unit_price": 680, "source": "AI建议参考价"},
        }
        for it in data["items"]:
            it["in_scope"] = True
            it.update(fills[it["code"]])
        (workdir / "items_extracted.json").write_text(
            json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

        # fill（复制原清单到工作目录，验证不破坏原文件）
        import shutil
        shutil.copy(str(SAMPLE_XLSX), str(workdir / SAMPLE_XLSX.name))
        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "fill_quotation.py"),
             "--json", "items_extracted.json"],
            workdir,
        )
        assert result.returncode == 0, result.stderr
        priced = workdir / "HappyRoom_家具软装_priced.xlsx"
        assert priced.is_file()

        from openpyxl import load_workbook
        wb = load_workbook(str(priced))
        ws = wb.worksheets[0]
        assert ws.cell(8, 9).value == 1200           # I 列单价
        assert ws.cell(8, 13).value == "知识库-KB"    # M 列来源
        assert ws.cell(9, 14).value == "原：布艺 W1800 → 替换：真皮 W1600" or "替换" in str(ws.cell(9, 14).value)
        assert ws.cell(10, 13).value == "AI建议参考价"
        assert "替换与差异 Replacements" in wb.sheetnames
        repl = wb["替换与差异 Replacements"]
        assert repl.cell(2, 1).value == "HR-1602"    # 替换 sheet 首行编号
        # 原文件未被破坏
        wb_orig = load_workbook(str(workdir / SAMPLE_XLSX.name))
        assert "替换与差异 Replacements" not in wb_orig.sheetnames
        assert wb_orig.worksheets[0].cell(8, 9).value in (None, "")

        # review
        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "review.py"),
             "--json", "items_extracted.json"],
            workdir,
        )
        assert result.returncode == 0, result.stderr
        report = json.loads((workdir / "review_report.json").read_text(encoding="utf-8"))
        assert report["in_scope_total"] == 3
        assert report["quoted"] == 3
        assert report["unquoted"] == 0
        assert report["mismatched"] == 1
        assert report["kb_sourced"] == 2
        assert report["ai_suggested"] == 1
        assert report["complete"] is True

    def test_review_flags_unquoted(self, workdir):
        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "parse_quotation.py"),
             str(SAMPLE_XLSX), "--json", "items_extracted.json"],
            workdir,
        )
        assert result.returncode == 0, result.stderr
        data = json.loads((workdir / "items_extracted.json").read_text(encoding="utf-8"))
        for it in data["items"]:
            it["in_scope"] = True
            if it["code"] == "HR-1601":
                it["unit_price"] = 1200
                it["source"] = "知识库-KB"
            # 其余留空
        (workdir / "items_extracted.json").write_text(
            json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "review.py"),
             "--json", "items_extracted.json"],
            workdir,
        )
        assert result.returncode == 0, result.stderr
        report = json.loads((workdir / "review_report.json").read_text(encoding="utf-8"))
        assert report["complete"] is False
        assert report["unquoted"] == 2
        assert {u["code"] for u in report["unquoted_items"]} == {"HR-1602", "HR-1701"}


class TestGeneratorLayoutRegression:
    """spec-to-quotation-list 2.0.1 列序对齐上游（B 描述 / C:D 区域合并）后，
    上游 parse_quotation.py 必须能正确抽取生成产物——两技能衔接的关键契约。"""

    def test_generated_xlsx_parseable_by_upstream(self, tmp_path):
        pytest.importorskip("fitz")
        wd = tmp_path / "gen"
        wd.mkdir()
        items = {
            "project": {"title": "T", "project": "冒烟", "project_no": "T-1", "issue_date": "27 OCT 2026"},
            "parameters": {"客房总数": {"value": 100, "remark": "test"}},
            "items": [
                {"code": "HR-1601", "cn": "单人沙发", "en": "Armchair velvet",
                 "space": "HappyRoom", "theme": "家具软装", "category": "家具 Furniture",
                 "uom": "Nos", "qty_rule": {"mode": "fixed", "value": 3},
                 "basis": "平面布置图", "vendor": "TO BID", "ref": "2.1 p18"},
            ],
        }
        (wd / "items.json").write_text(json.dumps(items, ensure_ascii=False), encoding="utf-8")
        import fitz
        doc = fitz.open()
        doc.new_page()
        doc.save(str(wd / "dummy.pdf"))
        doc.close()

        gen = _run(
            [sys.executable,
             str(PROJECT_ROOT / "src" / "skills" / "spec-to-quotation-list-2.0.0" / "scripts" / "generate_quotation_lists.py"),
             "--pdf", "dummy.pdf", "--items", "items.json", "--out", "out", "--project-name", "冒烟"],
            wd,
        )
        assert gen.returncode == 0, gen.stderr
        out_xlsx = wd / "out" / "HappyRoom_家具软装.xlsx"
        assert out_xlsx.is_file()

        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "parse_quotation.py"),
             str(out_xlsx), "--json", "parsed.json"],
            wd,
        )
        assert result.returncode == 0, result.stderr
        data = json.loads((wd / "parsed.json").read_text(encoding="utf-8"))
        assert data["item_count"] == 1
        it = data["items"][0]
        assert it["code"] == "HR-1601"
        assert it["area"] == "HappyRoom"                    # C:D 合并区域取到空间名
        assert "单人沙发" in it["title_cn_en"]               # B 列描述拆出标题
        assert it["qty_expr"] == 3                          # E 列数量
        assert it["uom"] == "Nos"
        assert it["brand"] == "TO BID"


class TestInvalidInputs:
    def test_parse_file_without_header_falls_back_and_skips(self, tmp_path):
        from openpyxl import Workbook
        bad = tmp_path / "bad.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.cell(1, 1, "随便什么")
        ws.cell(2, 1, "HR-1")
        ws.cell(2, 2, "某产品")
        wb.save(str(bad))

        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "parse_quotation.py"),
             str(bad), "--json", "out.json"],
            tmp_path,
        )
        assert result.returncode == 0, result.stderr   # 不崩溃，fallback 表头行继续解析
        data = json.loads((tmp_path / "out.json").read_text(encoding="utf-8"))
        assert data["item_count"] >= 0

    def test_fill_with_missing_file_skips(self, tmp_path):
        data = {"project": "x", "items": [
            {"file": "not_exist.xlsx", "row": 8, "code": "HR-1", "in_scope": True,
             "unit_price": 100, "source": "知识库-KB", "title_cn_en": "x"},
        ]}
        (tmp_path / "items_extracted.json").write_text(
            json.dumps(data, ensure_ascii=False), encoding="utf-8")
        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "fill_quotation.py"),
             "--json", "items_extracted.json"],
            tmp_path,
        )
        assert result.returncode == 0, result.stderr
        assert "找不到文件" in (result.stdout or "")

    def test_fill_with_non_numeric_price_writes_verbatim(self, tmp_path):
        import shutil
        from openpyxl import load_workbook
        shutil.copy(str(SAMPLE_XLSX), str(tmp_path / SAMPLE_XLSX.name))
        data = {"project": "x", "items": [
            {"file": SAMPLE_XLSX.name, "row": 8, "code": "HR-1601", "in_scope": True,
             "unit_price": "面议", "source": "待确认", "title_cn_en": "单人沙发"},
        ]}
        (tmp_path / "items_extracted.json").write_text(
            json.dumps(data, ensure_ascii=False), encoding="utf-8")
        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "fill_quotation.py"),
             "--json", "items_extracted.json"],
            tmp_path,
        )
        assert result.returncode == 0, result.stderr
        wb = load_workbook(str(tmp_path / "HappyRoom_家具软装_priced.xlsx"))
        assert wb.worksheets[0].cell(8, 9).value == "面议"
        assert wb.worksheets[0].cell(8, 13).value == "待确认"

    def test_fill_out_of_scope_left_blank(self, tmp_path):
        import shutil
        from openpyxl import load_workbook
        shutil.copy(str(SAMPLE_XLSX), str(tmp_path / SAMPLE_XLSX.name))
        data = {"project": "x", "items": [
            {"file": SAMPLE_XLSX.name, "row": 8, "code": "HR-1601", "in_scope": False,
             "unit_price": 999, "source": "知识库-KB", "title_cn_en": "单人沙发"},
        ]}
        (tmp_path / "items_extracted.json").write_text(
            json.dumps(data, ensure_ascii=False), encoding="utf-8")
        result = _run(
            [sys.executable, str(SKILL_DIR / "scripts" / "fill_quotation.py"),
             "--json", "items_extracted.json"],
            tmp_path,
        )
        assert result.returncode == 0, result.stderr
        wb = load_workbook(str(tmp_path / "HappyRoom_家具软装_priced.xlsx"))
        assert wb.worksheets[0].cell(8, 9).value in (None, "")   # 范围外不填
