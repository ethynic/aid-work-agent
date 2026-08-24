"""
Excel ETL 黄金夹具生成脚本（确定性，可重复运行）

按 docs/tools/excel/excel-etl-gap-analysis.md §2 的脏点表合成 5 个来源报表 xlsx，
与原样本 `ExcelAI测试.zip` 等效（原样本不在本地）：

1. 202608德勤派单增减人员.xlsx —— 换行复合表头、重复列名（派单日期 x2）、
   int 序列日期 + 日期格式、约 30 行全空行、参保地三级串、公积金比例 "5%:5%"
2. 万宝盛华人员变动通知.xlsx —— 同 Sheet 多段（数据行 -> 0806补充 分段行 -> 重复表头）、
   datetime 日期、企/个比例拆两列 int、始缴年月 "2026/08"
3. 活悦安徽分增减表.xlsx —— 两行表头（R1 说明 + R2 真表头）、脱敏身份证、最后缴纳月 "202608"
4. 外管离职导出（上海中企）.xlsx —— 标题/说明/T0 三行前置 + 合并单元格、
   8 个险种停止缴纳日期列（字符串日期）、下岗原因离职文案
5. 上海信息数据（社保&公积金）.xlsx —— 表头嵌枚举说明、备注页脚行、
   公积金比例 float 0.05 + "0%" 格式、日期混排（datetime 与字符串同列）
6. 生成标准模板格式.xlsx —— 单 Sheet1 仅 18 列标准表头（取自 schema 资产），无数据行；
   供 schema 指纹匹配测试（确定性生成，已 .gitignore 排除）

身份证号均按 ISO 7064 MOD 11-2 计算了正确校验位（Phase 2 校验层不误报）。

运行：.venv/bin/python tests/fixtures/excel_etl/build_fixtures.py
（expected_render.json / golden_records.json 为手工推演的黄金文件，不在本脚本生成）
"""

import datetime
import sys
from pathlib import Path

import openpyxl

FIXTURE_DIR = Path(__file__).resolve().parent
_REPO_ROOT = FIXTURE_DIR.parents[2]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

# 身份证校验位（ISO 7064 MOD 11-2）
_ID_WEIGHTS = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
_ID_CHECK_CHARS = "10X98765432"


def _id(area: str, year: int, month: int, day: int, seq: int) -> str:
    """生成带正确校验位的 18 位身份证号"""
    body = f"{area}{year:04d}{month:02d}{day:02d}{seq:03d}"
    check = _ID_CHECK_CHARS[sum(int(c) * w for c, w in zip(body, _ID_WEIGHTS)) % 11]
    return body + check


# ---- 人员数据（与 golden_records.json 逐字段对应）----
P_ZHANGWEI = _id("340203", 1990, 3, 7, 425)     # 340203199003074259 德勤新增
P_LIJUAN = _id("340203", 1995, 11, 12, 36)      # 340203199511120362 德勤新增
P_WANGQIANG = _id("340222", 1988, 6, 21, 517)   # 340222198806215179 德勤减少
P_CHENJING = _id("340203", 1992, 7, 18, 628)    # 340203199207186287 万宝增员
P_LIUYANG = _id("340203", 2000, 1, 5, 219)      # 340203200001052190 万宝增员
P_ZHAOMIN = _id("340203", 1989, 9, 30, 840)     # 340203198909308408 万宝减员
P_SUNLEI = _id("340203", 1992, 12, 3, 451)      # 34020319921203451X 万宝减员（校验位 X）
P_ZHENGHUA = _id("310101", 1985, 4, 17, 233)    # 310101198504172334 外管
P_FENGLI = _id("310101", 1990, 8, 25, 146)      # 310101199008251468 外管
P_HEJUN = _id("310101", 1991, 6, 11, 372)       # 310101199106113720 上海增员
P_XUTING = _id("310101", 1995, 2, 27, 184)      # 310101199502271843 上海增员
P_SHENXIN = _id("310101", 1988, 3, 19, 525)     # 310101198803195253 上海减员

# Excel 序列日期（数字格式为日期，裸值是 int——渲染层的首要脏点）
SERIAL_PAIDAN_1 = 46239    # 2026-08-05
SERIAL_PAIDAN_2 = 46246    # 2026-08-12

DATE_FMT = "yyyy-mm-dd"


def _fill_rows(ws, start_row: int, rows) -> None:
    """从 start_row 行起逐行写入二维数组（None 跳过）"""
    for r_off, row_vals in enumerate(rows):
        for c_off, val in enumerate(row_vals, start=1):
            if val is None:
                continue
            ws.cell(row=start_row + r_off, column=c_off, value=val)


def _serial_date(ws, row: int, col: int, serial: int) -> None:
    """写入 Excel 序列日期：int 值 + 日期数字格式（LLM 裸读只能看到 46239）"""
    cell = ws.cell(row=row, column=col, value=serial)
    cell.number_format = DATE_FMT


def build_deliqidian() -> str:
    """1. 德勤派单增减人员：换行表头/重复列名/序列日期/全空行/空列"""
    wb = openpyxl.Workbook()

    header = [
        "序号", "姓名", "身份证号", "手机号码", "派单日期", "参保地",
        "社保缴至月份\n（如6月继续缴纳则填）", "公积金比例", "派单日期", "备注",
    ]
    region = "安徽省,芜湖市,镜湖区"

    # 新增 sheet：表头 + 30 行全空行 + 2 行有效数据（模拟 max_row=114 仅 3 行有效）
    ws_add = wb.active
    ws_add.title = "新增"
    _fill_rows(ws_add, 1, [header])
    add_rows = [
        [1, "张伟", P_ZHANGWEI, "13805512366", None, region, "202607", "5%:5%", None, None],
        [2, "李娟", P_LIJUAN, "13905514578", None, region, None, "5%:5%", None, "8月新增"],
    ]
    _fill_rows(ws_add, 32, add_rows)
    for r in (32, 33):
        _serial_date(ws_add, r, 5, SERIAL_PAIDAN_1)
        _serial_date(ws_add, r, 9, SERIAL_PAIDAN_2)
    # 第 11 列：仅设数字格式的全空列（验证空列保留/列对齐）
    # 注："General" 是默认格式不会序列化，用 "@" 文本格式才能让该列真实存在于文件中
    ws_add.cell(row=1, column=11).number_format = "@"

    # 减少 sheet：表头 + 20 行全空行 + 1 行有效数据
    ws_dec = wb.create_sheet("减少")
    _fill_rows(ws_dec, 1, [header])
    _fill_rows(ws_dec, 22, [
        [1, "王强", P_WANGQIANG, "13705516789", None, region, "202606", "5%:5%", None, "7月离职"],
    ])
    _serial_date(ws_dec, 22, 5, SERIAL_PAIDAN_1)
    _serial_date(ws_dec, 22, 9, SERIAL_PAIDAN_2)

    path = FIXTURE_DIR / "202608德勤派单增减人员.xlsx"
    wb.save(path)
    return path.name


def build_wanbaoshenghua() -> str:
    """2. 万宝盛华人员变动通知：datetime/企个比例拆列/始缴年月字符串/多段结构"""
    wb = openpyxl.Workbook()

    add_header = ["雇员姓名", "身份证号码", "手机号码", "入职日期", "始缴年月",
                  "公积金比例-企业", "公积金比例-个人", "参保地", "用工形式"]
    ws_add = wb.active
    ws_add.title = "增员"
    _fill_rows(ws_add, 1, [add_header])
    _fill_rows(ws_add, 2, [
        ["陈静", P_CHENJING, "13605513456", datetime.datetime(2026, 8, 1), "2026/08",
         5, 5, "安徽省,芜湖市,镜湖区", "派遣"],
        ["刘洋", P_LIUYANG, "13505517890", datetime.datetime(2026, 8, 3), "2026/08",
         5, 5, "安徽省,芜湖市,镜湖区", "派遣"],
    ])

    dec_header = ["雇员姓名", "身份证号码", "离职日期", "停缴年月",
                  "公积金比例-企业", "公积金比例-个人", "离职原因", "参保地"]
    ws_dec = wb.create_sheet("减员")
    # 多段结构：表头 -> 1 行数据 -> 0806补充 分段行 -> 重复表头 -> 1 行数据
    _fill_rows(ws_dec, 1, [dec_header])
    _fill_rows(ws_dec, 2, [
        ["赵敏", P_ZHAOMIN, datetime.datetime(2026, 7, 31), "2026/07", 5, 5,
         "个人原因离职", "安徽省,芜湖市,镜湖区"],
        ["0806补充"],
        dec_header,
        ["孙磊", P_SUNLEI, datetime.datetime(2026, 7, 15), "2026/07", 5, 5,
         "合同期满", "安徽省,芜湖市,镜湖区"],
    ])

    path = FIXTURE_DIR / "万宝盛华人员变动通知.xlsx"
    wb.save(path)
    return path.name


def build_huoyue() -> str:
    """3. 活悦安徽分增减表：两行表头/脱敏身份证/最后缴纳月 202608"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "减员"
    _fill_rows(ws, 1, [
        ["2026年8月安徽分公司社保减员名单"],
        ["序号", "姓名", "身份证号", "社保最后缴纳月", "公积金最后缴纳月", "参保地", "备注"],
        [1, "周霞", "340203********1234", "202608", "202608", "安徽省,芜湖市,镜湖区", "8月减员"],
        [2, "吴涛", "340203********5678", "202608", "202608", "安徽省,芜湖市,镜湖区", None],
    ])
    path = FIXTURE_DIR / "活悦安徽分增减表.xlsx"
    wb.save(path)
    return path.name


def build_waiguan() -> str:
    """4. 外管离职导出（上海中企）：标题/说明/T0 前置 + 合并单元格 + 8 险种停止缴纳日期列"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "减员"
    _fill_rows(ws, 1, [
        ["上海中企人力资源服务有限公司离职人员名单"],
        ["说明：本表数据截至2026年8月，停止缴纳时间以当月最后一天为准"],
        ["T0"],
    ])
    ws.merge_cells("A1:E1")  # 标题跨列合并
    stop_cols = [f"{kind}停止缴纳时间" for kind in
                 ("养老保险", "失业保险", "医疗保险", "工伤保险", "生育保险", "档案", "补充工伤", "住房公积金")]
    header = ["序号", "姓名", "身份证号", "用工信息结束时间", "参保地", *stop_cols, "下岗原因"]
    _fill_rows(ws, 4, [header])
    _fill_rows(ws, 5, [
        [1, "郑华", P_ZHENGHUA, "2026-08-15", "上海市", *["2026-08-31"] * 8, "个人原因离职"],
        [2, "冯丽", P_FENGLI, "2026-08-20", "上海市", *["2026-08-31"] * 8, "合同到期"],
    ])
    path = FIXTURE_DIR / "外管离职导出（上海中企）.xlsx"
    wb.save(path)
    return path.name


def build_shanghai() -> str:
    """5. 上海信息数据：表头嵌枚举/备注页脚/float 0.05 + 0% 格式/日期混排"""
    wb = openpyxl.Workbook()

    add_header = ["员工姓名", "身份证号", "性别\n1男\n2女", "手机号码", "参保地", "参保年月",
                  "基本工资", "公积金基数", "公积金比例", "劳动合同起始时间", "劳动合同终止时间",
                  "岗位", "学历", "备注"]
    ws_add = wb.active
    ws_add.title = "增员表"
    _fill_rows(ws_add, 1, [add_header])
    _fill_rows(ws_add, 2, [
        ["何军", P_HEJUN, 1, "13805519999", "上海市", "2026/09", 6500, 6500, 0.05,
         datetime.datetime(2026, 9, 1), datetime.datetime(2029, 8, 31), "操作工", "大专", None],
        # 日期混排：同列既有 datetime 又有字符串
        ["许婷", P_XUTING, 2, "13705518888", "上海市", "2026/09", 5800, 5800, 0.05,
         "2026/9/15", "2029/9/14", "质检员", "本科", "9月中旬到岗"],
    ])
    # 公积金比例列：float 0.05 + 百分比数字格式（渲染层按格式输出 5%）
    for r in (2, 3):
        ws_add.cell(row=r, column=9).number_format = "0%"
    # 数据区后 2 行备注页脚（首列非空，渲染层须原样保留）
    _fill_rows(ws_add, 4, [
        ["备注：以上人员自2026年9月起缴纳社保"],
        ["数据来源：上海信息（社保&公积金）"],
    ])

    dec_header = ["员工姓名", "身份证号", "手机号码", "参保地", "减员年月",
                  "用工信息结束时间", "社保最后缴纳月", "公积金最后缴纳月", "离职原因"]
    ws_dec = wb.create_sheet("减员表")
    _fill_rows(ws_dec, 1, [dec_header])
    _fill_rows(ws_dec, 2, [
        ["沈鑫", P_SHENXIN, "13905517777", "上海市", "2026/08",
         datetime.datetime(2026, 8, 10), "202607", "202607", "员工主动提出辞职"],
    ])

    path = FIXTURE_DIR / "上海信息数据（社保&公积金）.xlsx"
    wb.save(path)
    return path.name


def build_standard_template() -> str:
    """6. 生成标准模板格式.xlsx：单 Sheet 名 Sheet1，仅 18 列表头行（schema.template_header[0] 顺序），无数据行

    表头取 src/tools/excel/assets/schema.json 的固化定义（单一事实来源）；
    模板指纹 = sha256("Sheet1\\n" + 18 个去空白表头 join("\\n"))，与
    excel_extract.template_fingerprint_from_headers 同规则（详见该函数 docstring）。
    """
    from src.tools.excel.excel_extract import load_default_schema

    schema = load_default_schema()
    headers = [f["template_header"][0] for f in schema["fields"]]
    assert len(headers) == 18, f"schema 字段数异常: {len(headers)}"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = schema["sheet_name"]  # "Sheet1"
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    path = FIXTURE_DIR / "生成标准模板格式.xlsx"
    wb.save(path)
    return path.name


def build_zip() -> str:
    """7. ExcelAI测试样本.zip：5 个来源 + 标准模板打成一个包（Phase 3 文件入口 e2e 夹具）

    等效原样本 `ExcelAI测试.zip`；成员固定时间戳保证确定性（*.zip 全局 gitignore，
    需要时运行本脚本重建）。
    """
    import zipfile

    members = [
        "202608德勤派单增减人员.xlsx",
        "万宝盛华人员变动通知.xlsx",
        "活悦安徽分增减表.xlsx",
        "外管离职导出（上海中企）.xlsx",
        "上海信息数据（社保&公积金）.xlsx",
        "生成标准模板格式.xlsx",
    ]
    # 依赖前面的构建函数产物（含 .gitignore 排除的模板）
    main()
    zip_path = FIXTURE_DIR / "ExcelAI测试样本.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in members:
            src = FIXTURE_DIR / name
            assert src.exists(), f"夹具缺失，先运行 build: {src}"
            info = zipfile.ZipInfo(name, date_time=(2026, 8, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zf.writestr(info, src.read_bytes())
    return zip_path.name


def main() -> None:
    built = [
        build_deliqidian(),
        build_wanbaoshenghua(),
        build_huoyue(),
        build_waiguan(),
        build_shanghai(),
        build_standard_template(),
    ]
    for name in built:
        print(f"生成: {FIXTURE_DIR / name}")


if __name__ == "__main__":
    main()
