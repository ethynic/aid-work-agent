"""真实工作簿验证宽表完整性，异步入口验证 Runner 心跳的调度前提。"""

import asyncio
import threading
import time
from unittest.mock import patch

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from src.tools.file.excel_reader import ExcelReader
from src.tools.file.read_tool import ReadTool

pytestmark = pytest.mark.integration


@pytest.fixture(scope="session", autouse=True)
def init_db_pool():
    """本文件仅组合文件解析与事件循环，覆盖目录级数据库fixture。"""
    yield


def test_styled_xfc_column_does_not_make_small_workbook_exceed_heartbeat_budget(tmp_path):
    """复现线上 A1:XFC28，完整保留工作表范围及真实值，不能靠裁掉远列提速。"""
    path = tmp_path / "sky_bridge.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "项目"
    sheet.append(["项目", "价格", "可用"])
    sheet.append(["桥梁科普馆", 0, False])
    sheet["J27"] = "最后的业务值"
    sheet["XFC28"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    workbook.save(path)
    workbook.close()

    started = time.perf_counter()
    result = ExcelReader().read_excel_document(str(path))
    elapsed = time.perf_counter() - started

    assert result["success"], result
    assert elapsed < 10, f"小型样式宽表读取耗时 {elapsed:.2f}s，接近生产30秒租约"
    data = result["sheet_data"]
    assert (data["max_row"], data["max_column"]) == (28, 16383)
    assert len(data["rows"]) == 28
    assert all(len(row) == 16383 for row in data["rows"])
    assert data["rows"][1][:3] == ["桥梁科普馆", 0, False]
    assert data["rows"][26][9] == "最后的业务值"
    assert "最后的业务值" in result["content"]
    assert "False" in result["content"]
    assert result["sheets_info"][0]["max_column"] == 16383


def test_selected_sheet_and_full_content_preserve_multiple_sheets_and_long_values(tmp_path):
    """指定工作表只影响结构化rows，全文仍包含全部工作表，长内容不得被列宽截断。"""
    path = tmp_path / "sheets.xlsx"
    workbook = Workbook()
    workbook.active.title = "摘要"
    workbook.active.append(["天空之桥"])
    sheet = workbook.create_sheet("项目")
    sheet.append(["项目", None, "说明"])
    long_value = "完整说明" * 40
    sheet.append(["桥模DIY", None, long_value])
    workbook.create_sheet("空表")
    workbook.save(path)
    workbook.close()

    result = ExcelReader().read_excel_document(str(path), sheet_name="项目")

    assert result["success"]
    assert result["current_sheet"] == "项目"
    assert result["sheet_data"]["rows"][1] == ["桥模DIY", None, long_value]
    assert result["sheet_data"]["headers"] == ["项目", "", "说明"]
    assert all(f"=== 工作表: {name} ===" in result["content"] for name in ["摘要", "项目", "空表"])
    assert long_value in result["content"]


@pytest.mark.asyncio
async def test_excel_parse_keeps_event_loop_available_for_heartbeat(tmp_path):
    """解析线程停留在IO/计算期间，事件循环必须能继续执行心跳而非等解析返回。"""
    path = tmp_path / "blocked.xlsx"
    path.write_bytes(b"reader stub")
    tool = ReadTool()
    entered, release = threading.Event(), threading.Event()

    def parse(_path):
        entered.set()
        if not release.wait(3):
            raise TimeoutError("解析先于心跳恢复")
        return {"content": "1\t天空之桥", "read_lines": 1}

    with patch.object(tool, "_read_excel_document", side_effect=parse):
        task = asyncio.create_task(tool.execute(file_path=str(path)))
        try:
            assert await asyncio.to_thread(entered.wait, 2)
            heartbeat_seen = False

            async def heartbeat():
                nonlocal heartbeat_seen
                heartbeat_seen = True

            await asyncio.wait_for(heartbeat(), timeout=1)
            assert heartbeat_seen
            assert not task.done(), "事件循环直到同步解析结束后才恢复调度"
        finally:
            release.set()
            result = await task
    assert result["content"] == "1\t天空之桥"


@pytest.mark.asyncio
async def test_cancelling_excel_read_propagates_without_publishing_success(tmp_path):
    """取消await会立即传播；只读解析线程可自行完成，不能把晚到结果当成功回复。"""
    path = tmp_path / "cancel.xlsx"
    path.write_bytes(b"reader stub")
    entered, release, finished = threading.Event(), threading.Event(), threading.Event()
    tool = ReadTool()

    def parse(_path):
        entered.set()
        try:
            release.wait(3)
            return {"content": "晚到结果"}
        finally:
            finished.set()

    with patch.object(tool, "_read_excel_document", side_effect=parse):
        task = asyncio.create_task(tool.execute(file_path=str(path)))
        try:
            assert await asyncio.to_thread(entered.wait, 2)
            task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await asyncio.wait_for(task, timeout=1)
        finally:
            release.set()
            assert await asyncio.to_thread(finished.wait, 2)
