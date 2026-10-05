"""Real generated files under trusted skill-child identity and shared storage."""
import json
from pathlib import Path
import shlex

import pytest

from .provider import Reply, tool_reply
from .test_worker import workers, api_pair, prices, accept, terminal

pytestmark = pytest.mark.integration


@pytest.mark.parametrize("actor_name", ["a", "global"])
def test_actual_trusted_skill_child_document_owners_use_shared_root_and_override_stale_saas_identity(workers, actors, service_database, actor_name):
    actor = actors[actor_name]
    script, proof = workers.root / "document-child.py", workers.root / "document-proof.json"
    script.write_text("""import json, sys
from pathlib import Path
from docx import Document
import fitz, openpyxl
from src.saas.context import set_tenant_context
from src.tools.context import current_tool_execution_context
from src.tools.word.word_lib import WordFileHandler
from src.tools.pdf.pdf_lib import PdfFileHandler
from src.tools.pdf.pdf_process_tool import PdfProcessTool
from src.tools.excel.excel_lib import ExcelFileHandler
from src.tools.excel.excel_process_tool import ExcelProcessTool
from src.tools.ppt.generator import PPTGenerator
set_tenant_context('fictional-ambient-foreign-tenant', 'fictional-ambient-foreign-user')
word = Document()
word.add_paragraph('fixture-word-readable')
word_path = WordFileHandler.save_temp(word, file_name='fixture-word.docx')['file_path']
pdf = fitz.open()
page = pdf.new_page()
page.insert_text((72,72), 'fixture-pdf-readable')
source = Path.cwd()/'source.pdf'
pdf.save(str(source)); pdf.close()
pdf_path = PdfFileHandler.save_temp(str(source), file_name='fixture-pdf.pdf')['file_path']
excel = openpyxl.Workbook(); excel.active['A1']='fixture-excel-readable'
excel_path = ExcelFileHandler.save_temp(excel, file_name='fixture-excel.xlsx')['file_path']
ppt_path = PPTGenerator().generate({'title':'fixture-ppt', 'slides':[{'type':'cover','title':'fixture-ppt-readable'}]})
pdf_tool, excel_tool = PdfProcessTool(), ExcelProcessTool()
for tool in (pdf_tool, excel_tool):
    tool.set_tenant_id('fictional-stale-tenant')
    tool.set_user_id('fictional-stale-user')
context = current_tool_execution_context()
assert context is not None
proof={'paths':[word_path,pdf_path,excel_path,ppt_path], 'tenant_id':context.tenant_id,
       'user_id':context.user_id,'session_id':context.session_id,
       'tool_call_id':context.tool_call_id,'pdf_identity':pdf_tool._resolve_tenant_user(),
       'excel_identity':excel_tool._resolve_tenant_user(), 'main_loaded':'src.main' in sys.modules}
Path(""" + repr(str(proof)) + ").write_text(json.dumps(proof))\n")
    marker = workers.provider.register(tool_reply("use_skill", {"skill":"excel-to-template"}, call_id="document-load"),
        tool_reply("skill_execute", {"skill":"excel-to-template", "command":"python "+shlex.quote(str(script))}, call_id="document-child-execute"),
        Reply(content="documents-generated-by-child"))
    accepted = accept(workers.api, actor, marker)
    child, _ = workers.start()
    workers.assert_clean_exit(child)
    assert terminal(service_database, accepted["runner_id"])["status"] == "completed"
    observed = json.loads(proof.read_text())
    assert observed["tenant_id"] == actor.tenant_id and observed["user_id"] == actor.user_id
    assert observed["session_id"] == actor.session_id and observed["tool_call_id"] == "document-child-execute"
    assert observed["pdf_identity"] == observed["excel_identity"] == [actor.tenant_id, actor.user_id]
    assert not observed["main_loaded"]
    destination = workers.storage / "tenants" / (actor.tenant_id or "_anonymous") / "conversation"
    paths = [Path(path) for path in observed["paths"]]
    assert all(path.parent == destination and path.exists() for path in paths)
    from docx import Document
    import fitz, openpyxl
    from pptx import Presentation
    assert Document(paths[0]).paragraphs[0].text == "fixture-word-readable"
    with fitz.open(paths[1]) as pdf:
        assert "fixture-pdf-readable" in pdf[0].get_text()
    workbook = openpyxl.load_workbook(paths[2])
    try:
        assert workbook.active['A1'].value == "fixture-excel-readable"
    finally:
        workbook.close()
    assert len(Presentation(paths[3]).slides) == 1
    assert not (workers.storage / "tenants" / "fictional-ambient-foreign-tenant").exists()
