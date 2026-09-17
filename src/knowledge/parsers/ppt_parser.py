"""
PPT 文档解析器
"""

from typing import List
from loguru import logger
from . import BaseParser, DocumentParseError, ParseResult


class PPTParser(BaseParser):
    """PPT (.pptx) 文档解析器"""

    def supported_extensions(self) -> List[str]:
        return [".pptx"]

    async def parse(self, file_path: str) -> ParseResult:
        logger.info(f"后端日志：开始解析 PPT 文档: {file_path}")
        try:
            from pptx import Presentation
            from pptx.exc import PackageNotFoundError

            try:
                prs = Presentation(file_path)
            except PackageNotFoundError:
                # python-pptx 对「非 zip 格式文件」统一抛此异常：加密 pptx（OLE 复合文件）、
                # 老版 .ppt 改后缀、损坏文件等
                raise DocumentParseError(
                    "文件不是标准 PPT 文档（可能已设置打开密码，或是老版 .ppt 改了后缀），"
                    "请用 Office/WPS 打开后另存为未加密的 .pptx 再上传"
                )
            except ValueError as e:
                # python-pptx 1.x 对「合法 zip 但主部件 content type 不是 PPT」
                # （其他 OOXML 文件改后缀）抛 ValueError
                raise DocumentParseError(
                    "文件不是标准 PPT 文档（内容格式与 .pptx 不符），"
                    "请用 Office/WPS 打开后另存为标准 .pptx 再上传"
                ) from e

            paragraphs = []
            total_slides = len(prs.slides)

            for i, slide in enumerate(prs.slides, 1):
                paragraphs.append(f"## 幻灯片 {i}")

                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            text = paragraph.text.strip()
                            if text:
                                paragraphs.append(text)

                    # 提取表格
                    if shape.has_table:
                        table = shape.table
                        for row in table.rows:
                            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_text:
                                paragraphs.append(" | ".join(row_text))

                paragraphs.append("")  # 幻灯片之间空行

            text = "\n".join(paragraphs)
            metadata = {
                "slide_count": total_slides,
                "char_count": len(text),
            }

            logger.info(f"后端日志：PPT 文档解析完成: {file_path}, 幻灯片数={total_slides}")
            return ParseResult(text=text, metadata=metadata)

        except DocumentParseError as e:
            logger.warning(f"后端日志：PPT 文档内容不合法被拒绝: {file_path}, 原因: {e}")
            raise
        except Exception as e:
            logger.opt(exception=True).error(f"后端日志：PPT 文档解析失败: {file_path}, error: {e}")
            raise
