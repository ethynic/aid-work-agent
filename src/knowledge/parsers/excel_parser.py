"""
Excel 文档解析器

二维表 sheet 走「一行数据 = 一个 chunk」的行级分块旁路（表头字段名注入每个
数据行），非二维表 sheet 回退普通文本分块。格式判定为启发式（零成本），
可选 LLM 兜底（knowledge.excel_llm_layout_fallback，默认关闭）。
"""

import re
from collections import Counter
from typing import Any, Dict, List, Optional, Tuple
from loguru import logger
from . import BaseParser, ParseResult, ParsedChunk

# 与 TextChunker.MAX_EMBEDDING_CHUNK_CHARS 一致（embedding API 输入上限），
# 独立常量避免解析器反向依赖 chunker 内部实现
MAX_EXCEL_CHUNK_CHARS = 6000
# 预留给 service 侧注入的「文档标题：{filename}\n」前缀，避免最终 chunk 突破上限
_PREFIX_HEADROOM_CHARS = 100
_LONG_CHUNK_OVERLAP = 64

_SENTENCE_END_CHARS = set('。！？')


def _sheet_matrix(worksheet) -> List[Tuple[int, List[str]]]:
    """sheet 转矩阵：返回 (Excel 实际行号, 单元格文本列表)，跳过全空行"""
    result = []
    for row_no, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        cells = [str(cell).strip() if cell is not None else "" for cell in row]
        if any(cells):
            result.append((row_no, cells))
    return result


def _row_span(row: List[str]) -> int:
    """行跨度 = 最后一个非空单元格的下标 + 1（中间/尾部空单元格不影响跨度）"""
    span = 0
    for i, c in enumerate(row):
        if c:
            span = i + 1
    return span


def _mode_column_count(rows: List[List[str]]) -> int:
    counter = Counter(_row_span(r) for r in rows)
    return counter.most_common(1)[0][0]


def _header_score(row: List[str]) -> float:
    """首行「像表头」得分：逐格判定（短文本 / 少数字 / 无句末标点），
    返回通过的单元格占非空单元格比例"""
    cells = [c for c in row if c]
    if not cells:
        return 0.0
    passed = 0
    for c in cells:
        if len(c) > 30:
            continue
        if any(ch in _SENTENCE_END_CHARS for ch in c):
            continue
        digits = sum(1 for ch in c if ch.isdigit())
        if digits > len(c) * 0.5 or re.search(r'\d{5,}', c):
            continue
        passed += 1
    return passed / len(cells)


def _looks_like_header(row: List[str]) -> bool:
    """像表头 = 得分 >= 0.6 且列名基本唯一"""
    cells = [c for c in row if c]
    if not cells:
        return False
    if len(set(cells)) / len(cells) < 0.9:
        return False
    return _header_score(row) >= 0.6


def _detect_layout(rows: List[List[str]]) -> Dict[str, Any]:
    """单个 sheet 的格式判定（启发式，docs/system/knowledge-base/excel-row-chunking-design.md §3.3）

    Returns:
        {"layout": "table"|"freeform", "has_header": bool, "rule": int, "score": float}
    """
    if len(rows) < 3:
        # 少于 3 行无法区分「表头 + 数据」与散文行（2 行散文的首行也常"像表头"）
        return {"layout": "freeform", "has_header": False, "rule": 1, "score": 0.0}

    first, data_rows = rows[0], rows[1:]
    mode = _mode_column_count(data_rows)
    score = _header_score(first)

    # 列结构不规则（某行跨度远小于众数=备注行/合并行，或明显超出）或单列内容（列表/散文）
    if mode < 2 or any(_row_span(r) * 2 < mode or _row_span(r) > mode + 2 for r in data_rows):
        return {"layout": "freeform", "has_header": False, "rule": 2, "score": score}
    # 首行跨度与数据行众数不一致 -> 首行可能是标题/合并说明行
    if _row_span(first) != mode:
        return {"layout": "freeform", "has_header": False, "rule": 3, "score": score}

    if _looks_like_header(first) and _row_text(first) in {_row_text(r) for r in data_rows}:
        # 无表头数据表（首行与数据行同构），首行即数据
        return {"layout": "table", "has_header": False, "rule": 4, "score": score}
    if _looks_like_header(first):
        return {"layout": "table", "has_header": True, "rule": 5, "score": score}

    # 其余：保守按表处理，首行也当数据行（不注入表头）。
    # 灰色地带（0.4~0.6）留给 LLM 兜底重判（默认关闭）
    return {"layout": "table", "has_header": False, "rule": 6, "score": score}


def _row_text(row: List[str]) -> str:
    return " | ".join(c for c in row if c)


def _split_long_text(text: str, width: int = MAX_EXCEL_CHUNK_CHARS) -> List[str]:
    """超长 chunk 定长切割（等价 TextChunker._split_long_text，带重叠）"""
    if len(text) <= width:
        return [text]
    chunks = []
    start = 0
    while start < len(text):
        chunks.append(text[start:start + width])
        start += width - _LONG_CHUNK_OVERLAP
    return chunks


class ExcelParser(BaseParser):
    """Excel (.xlsx) 文档解析器"""

    def supported_extensions(self) -> List[str]:
        return [".xlsx"]

    async def parse(self, file_path: str) -> ParseResult:
        logger.info(f"后端日志：开始解析 Excel 文档: {file_path}")
        from openpyxl import load_workbook

        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            return await self._parse_workbook(wb, file_path)
        except Exception as e:
            logger.opt(exception=True).error(f"后端日志：Excel 文档解析失败: {file_path}, error: {e}")
            raise
        finally:
            wb.close()

    async def _parse_workbook(self, wb, file_path: str) -> ParseResult:
        total_sheets = len(wb.sheetnames)
        multi_sheet = total_sheets > 1

        layout_report = []
        table_chunks: List[ParsedChunk] = []
        sheet_texts: Dict[str, str] = {}    # 全部 sheet 的老格式文本（供 text/raw_text/摘要）
        freeform_in_mixed: List[Tuple[str, str]] = []  # (sheet, 文本)，混合文件时单独分块

        for sheet_name in wb.sheetnames:
            sheet_rows = _sheet_matrix(wb[sheet_name])
            cell_rows = [cells for _, cells in sheet_rows]

            sheet_text_lines = [f"## 工作表: {sheet_name}"]
            for cells in cell_rows:
                row_text = _row_text(cells)
                if row_text:
                    sheet_text_lines.append(row_text)
            sheet_text = "\n\n".join(sheet_text_lines)
            sheet_texts[sheet_name] = sheet_text

            verdict = _detect_layout(cell_rows)
            if verdict["layout"] == "table" and verdict["rule"] == 6 \
                    and 0.4 <= verdict["score"] < 0.6:
                llm_layout = await self._llm_judge_layout(cell_rows)
                if llm_layout == "freeform":
                    verdict = {"layout": "freeform", "has_header": False,
                               "rule": 7, "score": verdict["score"]}
                elif llm_layout == "table":
                    verdict = {**verdict, "rule": 8}

            layout_report.append({"sheet": sheet_name, **verdict})

            if verdict["layout"] != "table":
                freeform_in_mixed.append((sheet_name, sheet_text))
                continue

            headers = None
            data_entries = list(zip(sheet_rows, cell_rows))
            if verdict["has_header"]:
                headers = self._build_headers(cell_rows[0])
                data_entries = list(zip(sheet_rows[1:], cell_rows[1:]))

            for (row_no, _), cells in data_entries:
                body = " | ".join(self._format_row_pairs(cells, headers))
                if not body:
                    continue
                prefix = f"[工作表: {sheet_name}] " if multi_sheet else ""
                full = prefix + body
                # 分段宽度预留前缀空间，service 侧还会注入文档标题前缀
                segments = _split_long_text(full, MAX_EXCEL_CHUNK_CHARS - _PREFIX_HEADROOM_CHARS)
                for seg_i, seg in enumerate(segments):
                    metadata = {
                        "chunk_type": "excel_row",
                        "sheet_name": sheet_name,
                        "row_number": row_no,
                    }
                    if len(segments) > 1:
                        metadata["segment"] = seg_i
                    table_chunks.append(ParsedChunk(text=seg, metadata=metadata))

        # ParseResult.text 保持老格式（供摘要 / documents.raw_text），全部 sheet 拼接
        text = "\n\n".join(sheet_texts[name] for name in wb.sheetnames)

        # 混合文件（含 table + freeform sheet）：freeform sheet 文本用 TextChunker
        # 就近分块，统一走 precomputed 旁路；纯 freeform 文件返回 None 完全走老路径。
        # 空 sheet（只有标题行、无内容）不产生 chunk
        precomputed: List[ParsedChunk] = list(table_chunks)
        if table_chunks:
            from src.knowledge.chunker import TextChunker
            mini_chunker = TextChunker()
            for sheet_name, sheet_text in freeform_in_mixed:
                if sheet_text == f"## 工作表: {sheet_name}":
                    continue
                for c in mini_chunker.chunk(sheet_text):
                    precomputed.append(ParsedChunk(
                        text=c["text"],
                        metadata={"chunk_type": "excel_freeform", "sheet_name": sheet_name},
                    ))

        logger.info(
            f"后端日志：Excel 文档解析完成: {file_path}, 工作表数={total_sheets}, "
            f"layout_report={layout_report}, chunks={len(precomputed)}"
        )
        metadata = {
            "sheet_count": total_sheets,
            "sheets": list(wb.sheetnames),
            "char_count": len(text),
            "layout_report": layout_report,
        }
        return ParseResult(text=text, metadata=metadata,
                           precomputed_chunks=precomputed or None)

    def _build_headers(self, first_row: List[str]) -> List[str]:
        """表头：空表头单元格用 列N 占位"""
        return [c if c else f"列{i + 1}" for i, c in enumerate(first_row)]

    def _format_row_pairs(self, cells: List[str], headers: Optional[List[str]]) -> List[str]:
        """数据行转「字段: 值」对；无表头时输出原值；空值保留字段名"""
        pairs = []
        for col_idx, value in enumerate(cells):
            if headers is None:
                if value:
                    pairs.append(value)
                continue
            key = headers[col_idx] if col_idx < len(headers) else f"列{col_idx + 1}"
            pairs.append(f"{key}: {value}")
        return pairs

    async def _llm_judge_layout(self, cell_rows: List[List[str]]) -> Optional[str]:
        """灰色地带的 LLM 兜底判定，返回 "table"|"freeform"，失败/未启用返回 None。

        仅在 knowledge.excel_llm_layout_fallback=true 时调用（默认关闭）。
        """
        try:
            from src.config.settings import settings
            knowledge_cfg = getattr(settings, "knowledge", None)
            if not getattr(knowledge_cfg, "excel_llm_layout_fallback", False):
                return None

            from src.llm.gateway import LLMGateway
            from src.services.session_record import record_background_llm_usage
            from src.tools.context import resolve_llm_gateway

            preview = "\n".join(_row_text(r) for r in cell_rows[:20])
            # 关思考调用（不切模型）：主链路思考 token 会烧穿 max_tokens 导致 content 为空
            gateway = resolve_llm_gateway() or LLMGateway()
            response = await gateway.chat_no_thinking(
                messages=[{"role": "user", "content":
                    "判断以下 Excel 工作表内容是否为「第一行表头 + 后续数据行」的二维数据表。"
                    "如果是回答 table，否则回答 freeform，只回答这一个单词。\n\n" + preview}],
                temperature=0,
                max_tokens=16,
            )
            record_background_llm_usage(
                response.get("usage") if isinstance(response, dict) else None,
                source="excel_layout_detect",
                model=gateway.get_model_name(),  # chat_no_thinking 沿用主链路模型，按主模型单价计费
            )
            answer = (response.get("content", "") or "").strip().lower()
            if "table" in answer:
                return "table"
            if "freeform" in answer:
                return "freeform"
            return None
        except Exception as e:
            logger.warning(f"后端日志：Excel 格式 LLM 兜底判定失败，沿用启发式结果: {e}")
            return None
