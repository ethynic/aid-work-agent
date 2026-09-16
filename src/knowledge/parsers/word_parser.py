"""
Word 文档解析器
"""

from typing import List
from loguru import logger
from . import BaseParser, DocumentParseError, ParseResult


class WordParser(BaseParser):
    """Word (.docx) 文档解析器"""

    def supported_extensions(self) -> List[str]:
        return [".docx"]

    async def parse(self, file_path: str) -> ParseResult:
        logger.info(f"后端日志：开始解析 Word 文档: {file_path}")
        try:
            from docx import Document
            from docx.opc.exceptions import PackageNotFoundError

            try:
                doc = Document(file_path)
            except PackageNotFoundError:
                # python-docx 对「非 zip 格式文件」统一抛此异常：加密 docx（OLE 复合文件）、
                # 老版 .doc 改后缀、损坏文件等
                raise DocumentParseError(
                    "文件不是标准 Word 文档（可能已设置打开密码，或是老版 .doc 改了后缀），"
                    "请用 Office/WPS 打开后另存为未加密的 .docx 再上传"
                )

            # 提取所有段落文本
            paragraphs = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)

            # 提取表格文本
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        paragraphs.append(" | ".join(row_text))

            text = "\n".join(paragraphs)

            if not text:
                logger.warning(f"后端日志：Word 文档内容为空: {file_path}")

            metadata = {
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "char_count": len(text),
            }

            logger.info(f"后端日志：Word 文档解析完成: {file_path}, 字符数={len(text)}")
            return ParseResult(text=text, metadata=metadata)

        except DocumentParseError as e:
            logger.warning(f"后端日志：Word 文档内容不合法被拒绝: {file_path}, 原因: {e}")
            raise
        except Exception as e:
            logger.opt(exception=True).error(f"后端日志：Word 文档解析失败: {file_path}, error: {e}")
            raise
