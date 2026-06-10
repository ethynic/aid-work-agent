from __future__ import annotations

import io
import json
import re
import uuid
from dataclasses import dataclass, field
from typing import List, Optional, Tuple
from urllib.parse import quote

import openpyxl
from fastapi.responses import StreamingResponse
from loguru import logger

from src.db.database import get_db_connection


_IDENTIFIER_RE = re.compile(r"^[a-zA-Z_][a-zA-Z0-9_]*$")


@dataclass
class ExportTableConfig:
    table: str
    columns: List[str]
    filename: str
    where_clause: str = ""
    where_params: Tuple = field(default_factory=tuple)
    order_by: str = "id"
    db_label: Optional[str] = None


def _validate_identifier(value: str, kind: str) -> None:
    if not isinstance(value, str) or not _IDENTIFIER_RE.match(value):
        raise ValueError(f"非法{kind}: {value}")


def _validate_config(cfg: ExportTableConfig) -> None:
    _validate_identifier(cfg.table, "表名")
    _validate_identifier(cfg.order_by, "排序字段")
    for col_name in cfg.columns:
        _validate_identifier(col_name, "列名")


def _get_uuid_column_type(cursor, table: str) -> Optional[str]:
    cursor.execute(
        """
        SELECT data_type
        FROM information_schema.columns
        WHERE table_schema = current_schema()
          AND table_name = %s
          AND column_name = 'uuid'
        LIMIT 1
        """,
        (table,),
    )
    row = cursor.fetchone()
    if not row:
        return None
    return row["data_type"] if isinstance(row, dict) else row[0]


def _fill_missing_uuid(cursor, table: str, tenant_id: str) -> int:
    uuid_column_type = _get_uuid_column_type(cursor, table)
    if not uuid_column_type:
        return 0

    missing_condition = "uuid IS NULL"
    params = [tenant_id]
    if uuid_column_type in {"text", "character varying", "character"}:
        missing_condition = "uuid IS NULL OR uuid = %s"
        params.append("")

    cursor.execute(
        f'SELECT ctid FROM {table} WHERE tenant_id = %s AND ({missing_condition})',
        params,
    )
    rows = cursor.fetchall()
    row_ctids = [row["ctid"] if isinstance(row, dict) else row[0] for row in rows]
    if not row_ctids:
        return 0

    updates = [(str(uuid.uuid4()), row_ctid) for row_ctid in row_ctids]
    cursor.executemany(f'UPDATE {table} SET uuid = %s WHERE ctid = %s::tid', updates)
    return len(updates)


def export_table_to_excel(cfg: ExportTableConfig, tenant_id: str) -> StreamingResponse:
    _validate_config(cfg)

    col_str = ", ".join(f'"{c}"' for c in cfg.columns)
    params = [tenant_id, *cfg.where_params]
    sql = (
        f'SELECT {col_str} FROM {cfg.table} '
        f'WHERE tenant_id = %s {cfg.where_clause} ORDER BY "{cfg.order_by}"'
    )
    logger.info(f"导出表 {cfg.table}, tenant={tenant_id}, columns={len(cfg.columns)}")

    conn_fn = get_db_connection(cfg.db_label) if cfg.db_label else get_db_connection()
    with conn_fn as conn:
        cursor = conn.cursor()
        filled_count = _fill_missing_uuid(cursor, cfg.table, tenant_id)
        if filled_count:
            conn.commit()
            logger.info(f"导出表 {cfg.table} 前自动补齐 uuid: {filled_count} 条")
        cursor.execute(sql, params)
        rows = cursor.fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cfg.filename[:31]

    for col_idx, col_name in enumerate(cfg.columns, 1):
        ws.cell(row=1, column=col_idx, value=col_name)

    for row_idx, row in enumerate(rows, 2):
        for col_idx, col_name in enumerate(cfg.columns, 1):
            value = row[col_name] if isinstance(row, dict) else row[col_idx - 1]
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            ws.cell(row=row_idx, column=col_idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    encoded_filename = quote(f"{cfg.filename}.xlsx")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )
