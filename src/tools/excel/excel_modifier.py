"""
Excel 内容修改

批量修改 Excel 内容：写入单元格、插入/删除行列、合并单元格等。
"""

from typing import Any, Dict, List

import openpyxl
from loguru import logger


def batch_modify(wb: openpyxl.Workbook, operations: List[Dict]) -> Dict[str, Any]:
    """
    批量修改 Excel 内容。

    operations 中每项需包含 type 字段，支持的类型：
    write_cell, write_range, insert_rows, delete_rows, insert_columns,
    delete_columns, rename_sheet, add_sheet, delete_sheet, merge_cells,
    unmerge_cells, sort_data
    """
    if not operations:
        return {"success": False, "error": "操作列表为空"}

    applied = 0
    errors = []

    for op in operations:
        op_type = op.get("type", "")
        try:
            ws = _get_sheet(wb, op.get("sheet"))
            if ws is None and op_type not in ("add_sheet", "rename_sheet", "delete_sheet"):
                errors.append(f"Sheet '{op.get('sheet')}' 不存在")
                continue

            if op_type == "write_cell":
                _op_write_cell(ws, op)
            elif op_type == "write_range":
                _op_write_range(ws, op)
            elif op_type == "insert_rows":
                _op_insert_rows(ws, op)
            elif op_type == "delete_rows":
                _op_delete_rows(ws, op)
            elif op_type == "insert_columns":
                _op_insert_columns(ws, op)
            elif op_type == "delete_columns":
                _op_delete_columns(ws, op)
            elif op_type == "rename_sheet":
                old_name = op.get("old_name", "")
                new_name = op.get("new_name", "")
                if old_name in wb.sheetnames:
                    wb[old_name].title = new_name
                else:
                    errors.append(f"Sheet '{old_name}' 不存在")
                    continue
            elif op_type == "add_sheet":
                name = op.get("name", "NewSheet")
                ws = wb.create_sheet(title=name)
                data = op.get("data", [])
                if data:
                    for row_idx, row in enumerate(data, 1):
                        for col_idx, val in enumerate(row, 1):
                            ws.cell(row=row_idx, column=col_idx, value=val)
            elif op_type == "delete_sheet":
                name = op.get("name", "")
                if name in wb.sheetnames and len(wb.sheetnames) > 1:
                    wb.remove(wb[name])
                else:
                    errors.append(f"无法删除 Sheet '{name}'")
                    continue
            elif op_type == "merge_cells":
                ws.merge_cells(op.get("range", ""))
            elif op_type == "unmerge_cells":
                ws.unmerge_cells(op.get("range", ""))
            elif op_type == "sort_data":
                _op_sort_data(ws, op)
            else:
                errors.append(f"不支持的操作类型: {op_type}")
                continue

            applied += 1
        except Exception as e:
            errors.append(f"{op_type} 操作失败: {e}")
            logger.warning(f"[ExcelModifier] {op_type} 失败: {e}")

    return {
        "success": applied > 0,
        "operations_applied": applied,
        "errors": errors if errors else None,
    }


def _get_sheet(wb: openpyxl.Workbook, sheet_name: str = None):
    if sheet_name and sheet_name in wb.sheetnames:
        return wb[sheet_name]
    elif not sheet_name:
        return wb.active
    return None


def _op_write_cell(ws, op):
    cell_ref = op.get("cell", "A1")
    value = op.get("value", "")
    cell = ws[cell_ref]
    cell.value = value


def _op_write_range(ws, op):
    start = op.get("start", "A1")
    data = op.get("data", [])

    from openpyxl.utils import range_boundaries
    col_start, row_start, _, _ = range_boundaries(start)

    for row_idx, row_data in enumerate(data):
        if isinstance(row_data, list):
            for col_idx, val in enumerate(row_data):
                ws.cell(row=row_start + row_idx, column=col_start + col_idx, value=val)


def _op_insert_rows(ws, op):
    position = op.get("position", 1)
    count = op.get("count", 1)
    ws.insert_rows(position, amount=count)


def _op_delete_rows(ws, op):
    start = op.get("start", 1)
    count = op.get("count", 1)
    ws.delete_rows(start, amount=count)


def _op_insert_columns(ws, op):
    position = op.get("position", 1)
    count = op.get("count", 1)
    ws.insert_cols(position, amount=count)


def _op_delete_columns(ws, op):
    start = op.get("start", 1)
    count = op.get("count", 1)
    ws.delete_cols(start, amount=count)


def _op_sort_data(ws, op):
    column = op.get("column", "A")
    order = op.get("order", "asc")
    has_header = op.get("has_header", True)

    from openpyxl.utils import column_index_from_string
    col_idx = column_index_from_string(column)

    # 收集数据行
    start_row = 2 if has_header else 1
    rows_data = []
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row,
                            min_col=1, max_col=ws.max_column):
        rows_data.append([cell.value for cell in row])

    # 排序
    reverse = order == "desc"
    rows_data.sort(key=lambda r: r[col_idx - 1] if r[col_idx - 1] is not None else "", reverse=reverse)

    # 写回
    for row_idx, row_data in enumerate(rows_data, start_row):
        for col_idx, val in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
