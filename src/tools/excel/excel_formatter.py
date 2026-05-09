"""
Excel 格式化

批量设置 Excel 格式样式：字体、边框、背景色、对齐、数字格式、列宽行高等。
"""

from copy import copy
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from loguru import logger

from src.tools.excel.excel_lib import (
    DEFAULT_BORDER,
    auto_column_width,
    parse_color,
    parse_range,
    resolve_font_name,
)


def batch_format(wb: openpyxl.Workbook, format_operations: List[Dict]) -> Dict[str, Any]:
    """
    批量设置 Excel 格式。

    format_operations 中每项需包含 type 字段，支持的类型：
    set_font, set_fill, set_border, set_alignment, set_number_format,
    set_column_width, set_row_height, freeze_panes, auto_filter
    """
    if not format_operations:
        return {"success": False, "error": "格式化操作列表为空"}

    applied = 0
    errors = []

    for op in format_operations:
        op_type = op.get("type", "")
        try:
            ws = _get_sheet(wb, op.get("sheet"))
            if ws is None:
                errors.append(f"Sheet '{op.get('sheet')}' 不存在")
                continue

            if op_type == "set_font":
                _op_set_font(ws, op)
            elif op_type == "set_fill":
                _op_set_fill(ws, op)
            elif op_type == "set_border":
                _op_set_border(ws, op)
            elif op_type == "set_alignment":
                _op_set_alignment(ws, op)
            elif op_type == "set_number_format":
                _op_set_number_format(ws, op)
            elif op_type == "set_column_width":
                _op_set_column_width(ws, op)
            elif op_type == "set_row_height":
                _op_set_row_height(ws, op)
            elif op_type == "freeze_panes":
                cell = op.get("cell", "A2")
                ws.freeze_panes = cell
            elif op_type == "auto_filter":
                rng = op.get("range")
                if rng:
                    ws.auto_filter.ref = rng
            else:
                errors.append(f"不支持的操作类型: {op_type}")
                continue

            applied += 1
        except Exception as e:
            errors.append(f"{op_type} 失败: {e}")
            logger.warning(f"[ExcelFormatter] {op_type} 失败: {e}")

    return {
        "success": applied > 0,
        "operations_applied": applied,
        "errors": errors if errors else None,
    }


def _get_sheet(wb: openpyxl.Workbook, sheet_name: str = None):
    if sheet_name and sheet_name in wb.sheetnames:
        return wb[sheet_name]
    elif not sheet_name:
        return wb.active
    return None


def _apply_to_range(ws, range_str: str, apply_fn):
    """对范围内的每个单元格应用函数"""
    if not range_str:
        return

    (min_col, min_row), (max_col, max_row) = parse_range(range_str)
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            apply_fn(ws.cell(row=row, column=col))


def _op_set_font(ws, op):
    font_name = resolve_font_name(op.get("font_name", ""))
    size = op.get("size") or op.get("font_size")
    bold = op.get("bold")
    italic = op.get("italic")
    color = parse_color(op.get("color", ""))

    font_kwargs = {}
    if font_name:
        font_kwargs["name"] = font_name
    if size is not None:
        font_kwargs["size"] = float(size)
    if bold is not None:
        font_kwargs["bold"] = bool(bold)
    if italic is not None:
        font_kwargs["italic"] = bool(italic)
    if color:
        font_kwargs["color"] = color

    font = Font(**font_kwargs)

    def apply(cell):
        cell.font = font

    _apply_to_range(ws, op.get("range", ""), apply)


def _op_set_fill(ws, op):
    color = parse_color(op.get("color", ""))
    if not color:
        return
    fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    def apply(cell):
        cell.fill = fill

    _apply_to_range(ws, op.get("range", ""), apply)


def _op_set_border(ws, op):
    style = op.get("style", "thin")
    color = parse_color(op.get("color", "D9D9D9")) or "D9D9D9"

    side = Side(style=style, color=color)
    border = Border(left=side, right=side, top=side, bottom=side)

    def apply(cell):
        cell.border = border

    _apply_to_range(ws, op.get("range", ""), apply)


def _op_set_alignment(ws, op):
    horizontal = op.get("horizontal", "left")
    vertical = op.get("vertical", "center")
    wrap_text = op.get("wrap_text", False)

    alignment = Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap_text)

    def apply(cell):
        cell.alignment = alignment

    _apply_to_range(ws, op.get("range", ""), apply)


def _op_set_number_format(ws, op):
    fmt = op.get("format", "@")
    # 支持中文名称映射
    from src.tools.excel.excel_lib import NUMBER_FORMATS
    fmt = NUMBER_FORMATS.get(fmt, fmt)

    def apply(cell):
        cell.number_format = fmt

    _apply_to_range(ws, op.get("range", ""), apply)


def _op_set_column_width(ws, op):
    columns = op.get("columns", [])
    widths = op.get("widths", [])

    if not columns:
        return

    if widths == "auto" or (isinstance(widths, list) and "auto" in widths):
        auto_column_width(ws)
        return

    for col_name, width in zip(columns, widths):
        col_idx = column_index_from_string(col_name)
        ws.column_dimensions[get_column_letter(col_idx)].width = float(width)


def _op_set_row_height(ws, op):
    rows = op.get("rows", [])
    height = op.get("height", 20)

    for row in rows:
        ws.row_dimensions[row].height = float(height)
