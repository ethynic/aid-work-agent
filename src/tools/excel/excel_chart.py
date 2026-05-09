"""
Excel 图表生成

在 Excel 中创建图表：柱状图、折线图、饼图、散点图、面积图。
"""

from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart
from openpyxl.utils import column_index_from_string
from loguru import logger


def create_chart(wb: openpyxl.Workbook, chart_type: str = "bar",
                 x_column: str = "", y_columns: Optional[List[str]] = None,
                 title: Optional[str] = None,
                 sheet_name: Optional[str] = None,
                 **kwargs) -> Dict[str, Any]:
    """
    在 Excel 中创建图表并嵌入到工作表中。

    chart_type: bar | line | pie | scatter | area
    """
    if not y_columns:
        return {"success": False, "error": "需要指定 y_columns（数据列）"}

    ws = _get_sheet(wb, sheet_name)
    if ws is None:
        return {"success": False, "error": f"Sheet '{sheet_name}' 不存在"}

    try:
        max_row = ws.max_row
        max_col = ws.max_column

        if max_row < 2 or max_col < 1:
            return {"success": False, "error": "数据不足，无法生成图表"}

        # 检测数据行范围（跳过表头）
        header_row = 1
        data_start = 2

        # 构建 x 轴数据引用
        x_col_idx = column_index_from_string(x_column) if x_column else 1
        x_data = Reference(ws, min_col=x_col_idx, min_row=data_start, max_row=max_row)

        # 构建 y 轴数据引用
        chart = _create_chart_obj(chart_type, title, **kwargs)

        for y_col_name in y_columns:
            y_col_idx = column_index_from_string(y_col_name)
            y_data = Reference(ws, min_col=y_col_idx, min_row=header_row, max_row=max_row)

            if chart_type == "pie":
                chart.add_data(y_data, titles_from_data=True)
            elif chart_type == "scatter":
                from openpyxl.chart import Series
                x_values = Reference(ws, min_col=x_col_idx, min_row=data_start, max_row=max_row)
                y_values = Reference(ws, min_col=y_col_idx, min_row=data_start, max_row=max_row)
                series = Series(y_values, x_values, title_from_data=True)
                chart.series.append(series)
            else:
                chart.add_data(y_data, titles_from_data=True)
                if chart_type == "bar" and x_column:
                    chart.set_categories(x_data)

        # 设置图表尺寸
        chart.width = kwargs.get("width", 20)
        chart.height = kwargs.get("height", 12)

        # 图例和标签
        if "show_legend" in kwargs and not kwargs["show_legend"]:
            chart.legend = None
        if kwargs.get("show_labels"):
            from openpyxl.chart.label import DataLabelList
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True

        # 轴标题
        if kwargs.get("x_title"):
            chart.x_axis.title = kwargs["x_title"]
        if kwargs.get("y_title"):
            chart.y_axis.title = kwargs["y_title"]

        # 嵌入到工作表
        ws.add_chart(chart, _find_chart_position(ws))

        return {
            "success": True,
            "chart_type": chart_type,
            "chart_title": title or "",
            "data_range": f"{ws.title}!A1:{openpyxl.utils.get_column_letter(max_col)}{max_row}",
        }
    except Exception as e:
        logger.error(f"[ExcelChart] 创建图表失败: {e}", exc_info=True)
        return {"success": False, "error": f"创建图表失败: {e}"}


def _get_sheet(wb: openpyxl.Workbook, sheet_name: str = None):
    if sheet_name and sheet_name in wb.sheetnames:
        return wb[sheet_name]
    elif not sheet_name:
        return wb.active
    return None


def _create_chart_obj(chart_type: str, title: str = None, **kwargs):
    """创建对应类型的图表对象"""
    style = kwargs.get("style", 10)

    if chart_type == "bar":
        chart = BarChart()
        chart.type = kwargs.get("bar_type", "col")  # col=柱状, bar=条形
        chart.style = style
    elif chart_type == "line":
        chart = LineChart()
        chart.style = style
    elif chart_type == "pie":
        chart = PieChart()
        chart.style = style
    elif chart_type == "scatter":
        chart = ScatterChart()
        chart.style = style
    elif chart_type == "area":
        from openpyxl.chart import AreaChart
        chart = AreaChart()
        chart.style = style
    else:
        chart = BarChart()
        chart.style = style

    if title:
        chart.title = title

    return chart


def _find_chart_position(ws):
    """找到一个合适的位置放置图表（不覆盖数据）"""
    max_col = ws.max_column
    # 放在数据区域右侧，留一列间隔
    col_letter = openpyxl.utils.get_column_letter(max_col + 2)
    return f"{col_letter}1"
