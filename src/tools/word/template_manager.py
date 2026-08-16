"""
Word 模板管理器

从 word_lib.py 迁移模板加载/列表功能，提供变量占位符引擎：
- 多占位符语法识别与替换：{{x}} / {x} / 【x】 / [x] / %x%
- 占位符扫描：scan_placeholders 列出模板中的占位符及出现次数
- 表格行循环展开：{{#列表名}} ... {{/列表名}} 区间行按列表条数复制
"""

import copy
import json
import re
from pathlib import Path
from typing import Any, Dict, Iterator, List, Tuple

from loguru import logger

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.tools.word.word_lib import (
    iter_all_paragraphs,
    paragraph_text_runs,
    replace_text_cross_run,
)


TEMPLATES_DIR = Path(__file__).parent / "assets" / "templates"

# =============================================================================
# 占位符语法定义
# =============================================================================

# 占位符语法集：(语法名, 编译正则)，按优先级排序——高优先级语法命中的区间
# 不再被低优先级语法重复匹配（如 {{x}} 内部的 {x} 不重复识别）。
# 注意：%x% 语法因共享分隔符易跨匹配（"10%，折扣：%折扣%" 会先匹配出 "%，折扣：%"
# 吞掉真占位符），不在此列，由 _iter_percent_matches 逐起点单独处理
PLACEHOLDER_SYNTAXES: List[Tuple[str, "re.Pattern[str]"]] = [
    # {{变量名}}：允许内侧空格，{{ name }} 与 {{name}} 等价
    ("double_brace", re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")),
    # {变量名}：排除内部再含花括号和 % 的串（避免误吞 {{x}} 和 100% 附近的文本）
    ("single_brace", re.compile(r"\{([^{}%]+?)\}")),
    # 【变量名】（中文全角方括号）
    ("fullwidth_bracket", re.compile(r"【([^【】]+?)】")),
    # [变量名]（半角方括号，正文引用 [1] 等易误报，配合启发式过滤）
    ("square_bracket", re.compile(r"\[([^\[\]]+?)\]")),
]

# 误报风险高的语法（正文方括号引用、百分数），扫描/填充时应用启发式过滤
LOOSE_SYNTAXES = {"square_bracket", "percent"}

# 表格行循环标记：{{#列表名}} 开始 / {{/列表名}} 结束（同一行内或跨行均可）
ROW_LOOP_START_RE = re.compile(r"\{\{\s*#\s*([^{}\s]+?)\s*\}\}")
ROW_LOOP_END_RE = re.compile(r"\{\{\s*/\s*([^{}\s]+?)\s*\}\}")

# 启发式过滤：变量名须为标识符样式（字母/数字/下划线/中文），长度 >= 2 且至少含一个
# 字母/中文/下划线（剔除 [1] 纯数字、%，折扣：% 这类含标点的跨匹配误报）
_VAR_NAME_TOKEN_RE = re.compile(r"[A-Za-z0-9_\u4e00-\u9fff]+")
_VAR_NAME_HINT_RE = re.compile(r"[A-Za-z_\u4e00-\u9fff]")


def _looks_like_var_name(name: str) -> bool:
    """[x] / %x% 语法的启发式变量名过滤，降低正文误报"""
    return (
        len(name) >= 2
        and bool(_VAR_NAME_HINT_RE.search(name))
        and _VAR_NAME_TOKEN_RE.fullmatch(name) is not None
    )


def _iter_placeholder_matches(text: str) -> Iterator[Tuple[str, str, str]]:
    """
    按语法优先级产出 (实际匹配串, 变量名, 语法名)。

    已被高优先级语法命中的字符区间不再被低优先级语法重复匹配
    （如 {{name}} 只按 double_brace 识别一次，内部 {name} 不重复识别）。
    高误报语法的启发式不通过的匹配（如 50%-60%、[1]）既不产出也不占用区间，
    避免吞掉后面真正的 %变量% / [变量]。
    """
    claimed_spans: List[Tuple[int, int]] = []
    for syntax, regex in PLACEHOLDER_SYNTAXES:
        for m in regex.finditer(text):
            start, end = m.span()
            # 与已命中区间重叠则跳过
            if any(start < e and end > s for s, e in claimed_spans):
                continue
            name = m.group(1).strip()
            if syntax in LOOSE_SYNTAXES and not _looks_like_var_name(name):
                continue
            claimed_spans.append((start, end))
            yield m.group(0), name, syntax
    yield from _iter_percent_matches(text, claimed_spans)


def _iter_percent_matches(text: str, claimed_spans: List[Tuple[int, int]]
                          ) -> Iterator[Tuple[str, str, str]]:
    """%变量名% 逐起点匹配（优先级最低，最后处理）。

    不用 finditer：其对共享分隔符的非重叠推进会先吞出 "%，折扣：% 这类误报"
    并跳过其后的真占位符。这里逐个 % 起点（升序）尝试配对到下一个 %，
    启发式不通过的起点不占用区间，后续起点仍可命中真占位符。
    """
    percent_positions = [i for i, ch in enumerate(text) if ch == "%"]
    for i in percent_positions:
        j = text.find("%", i + 1)
        if j == -1:
            break  # 该起点之后再无闭合 %，更靠后的起点也不可能配对
        start, end = i, j + 1
        if any(start < e and end > s for s, e in claimed_spans):
            continue
        name = text[i + 1:j].strip()
        if not _looks_like_var_name(name):
            continue
        claimed_spans.append((start, end))
        yield text[start:end], name, "percent"


def _is_scalar(value: Any) -> bool:
    """普通替换只接受标量值；列表/字典变量留给行循环展开处理"""
    return isinstance(value, (str, int, float, bool))


# =============================================================================
# 占位符扫描
# =============================================================================

def scan_placeholders(doc: Document) -> List[Dict[str, Any]]:
    """
    扫描全文档（正文/嵌套表格/文本框/内容控件/页眉页脚/超链接）中的占位符。

    Returns:
        [{"name": 变量名, "syntax": 语法名, "count": 出现次数}]，按 count 降序。
        [x] / %x% 语法经启发式过滤（变量名须为标识符样式且长度 >= 2，
        剔除正文方括号引用 [1]、百分数等误报，见 _iter_placeholder_matches）。
        去重合并后的变量名列表可用 unique_variable_names() 从结果派生。
    """
    counters: Dict[Tuple[str, str], int] = {}
    for para in iter_all_paragraphs(doc):
        text = "".join(r.text for r in paragraph_text_runs(para))
        if not text:
            continue
        for _match_text, name, syntax in _iter_placeholder_matches(text):
            key = (name, syntax)
            counters[key] = counters.get(key, 0) + 1

    return [
        {"name": name, "syntax": syntax, "count": count}
        for (name, syntax), count in sorted(counters.items(), key=lambda kv: -kv[1])
    ]


def unique_variable_names(placeholders: List[Dict[str, Any]]) -> List[str]:
    """从 scan_placeholders 结果提取去重后的变量名列表（合并不同语法的同名变量）"""
    names: List[str] = []
    for p in placeholders:
        if p["name"] not in names:
            names.append(p["name"])
    return names


# =============================================================================
# 占位符填充
# =============================================================================

def fill_template(doc: Document, variables: Dict[str, Any]) -> Dict[str, Any]:
    """
    替换文档中所有匹配的变量占位符，并返回替换反馈。

    支持格式：{{x}} / {x} / 【x】 / [x] / %x%（{{ x }} 带内侧空格同样识别，
    按原文实际出现形式逐串替换，不误伤正文普通方括号/百分号文本）。
    搜索范围：正文（含嵌套表格/文本框/内容控件/超链接）+ 页眉页脚。
    列表变量配合 {{#列表名}}...{{/列表名}} 行循环标记做表格行展开。

    Returns:
        {
            "total": 替换总数,
            "per_variable": {变量名: 替换次数},
            "unmatched_variables": [提供了但全文没匹配到的变量],
            "remaining_placeholders": [替换后仍留在文档中的占位符名（去重）]
        }
    """
    per_variable: Dict[str, int] = {}

    # 先做表格行循环展开（列表变量），再做全文档普通替换
    consumed_lists, loop_counts = _expand_row_loops(doc, variables)
    for name, cnt in loop_counts.items():
        per_variable[name] = per_variable.get(name, 0) + cnt

    for para in iter_all_paragraphs(doc):
        _replace_placeholders_in_paragraph(para, variables, per_variable)

    matched_names = set(per_variable.keys()) | set(consumed_lists)
    return {
        "total": sum(per_variable.values()),
        "per_variable": per_variable,
        "unmatched_variables": [k for k in variables if k not in matched_names],
        "remaining_placeholders": unique_variable_names(scan_placeholders(doc)),
    }


def _replace_placeholders_in_paragraph(para: Paragraph, variables: Dict[str, Any],
                                       per_variable: Dict[str, int]) -> None:
    """替换单个段落中命中的占位符（跨 run 匹配，含超链接内 run）"""
    runs = paragraph_text_runs(para)
    if not runs:
        return
    text = "".join(r.text for r in runs)

    # 实际匹配串 -> 变量名（同串多处出现由 replace_text_cross_run 一次全替换）
    matched: Dict[str, str] = {}
    for match_text, name, _syntax in _iter_placeholder_matches(text):
        if name in variables and _is_scalar(variables[name]) and match_text not in matched:
            matched[match_text] = name

    for match_text, name in matched.items():
        count = replace_text_cross_run(runs, match_text, str(variables[name]))
        if count:
            per_variable[name] = per_variable.get(name, 0) + count


def _replace_placeholders_in_element(el, scope: Dict[str, Any],
                                     per_variable: Dict[str, int]) -> None:
    """替换指定 XML 元素（如复制的表格行）内所有段落的占位符，并清除行循环标记"""
    for p_el in el.iter(qn('w:p')):
        para = Paragraph(p_el, None)
        runs = paragraph_text_runs(para)
        if not runs:
            continue
        # 清除行内 {{#x}} / {{/x}} 循环标记文本（含跨 run），必须清除干净
        text = "".join(r.text for r in runs)
        for regex in (ROW_LOOP_START_RE, ROW_LOOP_END_RE):
            for marker_text in {m.group(0) for m in regex.finditer(text)}:
                replace_text_cross_run(runs, marker_text, "")
        _replace_placeholders_in_paragraph(para, scope, per_variable)


# =============================================================================
# 表格行循环展开
# =============================================================================

def _expand_row_loops(doc: Document, variables: Dict[str, Any]
                      ) -> Tuple[List[str], Dict[str, int]]:
    """
    表格行循环展开：{{#列表名}} 开始、{{/列表名}} 结束，行区间 = 两个标记所在的行。

    行为：
    - variables[列表名] 为 List[Dict[str, str]]，区间内模板行按条数 deepcopy 复制，
      每份用对应条目的字段值替换行内占位符（条目字段 > 全局 variables 兜底），
      复制完成后删除原模板行
    - 列表为空时删除整个区间行
    - deepcopy 整行天然保留格式与行内合并单元格样式；行内嵌套表格允许（替换按行内段落处理）
    - 标记文本本身清除干净；多个循环按文档序选取、倒序展开（避免先行展开
      增删行使后续区间行号失效）；同名循环只处理第一个（其余 warning 跳过）

    已知限制：
    - 合并单元格（vMerge）跨行的循环区间明确不支持（跨行 vMerge 复制后结构错乱）
    - 循环标记不查找页眉页脚中的表格（页眉表格通常不承载明细行循环）

    Returns:
        (成功处理的列表变量名, 行内字段替换计数 {字段名: 次数})
    """
    consumed: List[str] = []
    counts: Dict[str, int] = {}
    processed_names: set = set()

    # 遍历所有表格（含嵌套表格）；先物化列表，避免展开时增删行影响 lxml 迭代器
    for tbl_el in list(doc.element.body.iter(qn('w:tbl'))):
        table = Table(tbl_el, doc)
        loops = _pair_row_loop_markers(table)
        # 同名循环只处理文档序第一个（跨表同样适用）
        first_loops = []
        for start_idx, end_idx, name in loops:
            if name in processed_names:
                logger.warning(f"[TemplateManager] 同名表格行循环 {{{{{name}}}}} 只处理第一个，其余跳过")
                continue
            processed_names.add(name)
            first_loops.append((start_idx, end_idx, name))
        # 倒序（开始行号降序）展开：先处理靠后的行区间。前面的展开会增删行使
        # 后面区间的行号索引失效（错删/错复制无关行），倒序则前面的索引始终不变
        for start_idx, end_idx, name in sorted(first_loops, key=lambda t: -t[0]):
            entries = variables.get(name)
            if not isinstance(entries, list):
                # 未提供列表数据：不展开，保留原行（占位符进入 remaining 反馈）
                continue
            _expand_single_loop(table, start_idx, end_idx, name, entries, variables, counts)
            consumed.append(name)

    return consumed, counts


def _pair_row_loop_markers(table: Table) -> List[Tuple[int, int, str]]:
    """配对表格内的行循环标记，返回 [(开始行号, 结束行号, 列表名)]（文档序）"""
    markers: List[Tuple[int, str, str]] = []  # (行号, kind, 列表名)
    for row_idx, row in enumerate(table.rows):
        seen_tc: set = set()
        for cell in row.cells:
            if id(cell._tc) in seen_tc:
                # 水平合并单元格在 row.cells 中重复出现，去重避免标记重复配对
                continue
            seen_tc.add(id(cell._tc))
            for para in cell.paragraphs:
                text = para.text
                for m in ROW_LOOP_START_RE.finditer(text):
                    markers.append((row_idx, "start", m.group(1)))
                for m in ROW_LOOP_END_RE.finditer(text):
                    markers.append((row_idx, "end", m.group(1)))

    loops: List[Tuple[int, int, str]] = []
    pending: Dict[str, int] = {}
    for row_idx, kind, name in markers:
        if kind == "start":
            if name in pending:
                logger.warning(f"[TemplateManager] 同名表格行循环 {{{{{name}}}}} 重复开始标记，跳过")
                continue
            pending[name] = row_idx
        else:
            if name in pending:
                loops.append((pending.pop(name), row_idx, name))
            else:
                logger.warning(f"[TemplateManager] 未配对的行循环结束标记 {{{{/{name}}}}}，跳过")
    return loops


def _expand_single_loop(table: Table, start_idx: int, end_idx: int, name: str,
                        entries: List[Dict[str, Any]], variables: Dict[str, Any],
                        counts: Dict[str, int]) -> None:
    """展开单个行循环：复制区间行并按条目替换，最后删除原模板行"""
    tr_elems = [row._tr for row in table.rows[start_idx:end_idx + 1]]

    if entries:
        anchor = tr_elems[-1]
        for entry in entries:
            # 条目字段优先，全局 variables 兜底（非列表变量也可在行内使用）
            scope = dict(variables)
            if isinstance(entry, dict):
                scope.update(entry)
            for tr_el in tr_elems:
                new_el = copy.deepcopy(tr_el)
                anchor.addnext(new_el)
                anchor = new_el
                _replace_placeholders_in_element(new_el, scope, counts)

    # 删除原模板行
    for tr_el in tr_elems:
        tr_el.getparent().remove(tr_el)


# =============================================================================
# 命名模板（JSON）
# =============================================================================

def get_template(template_name: str) -> Dict[str, Any]:
    """加载命名模板。"""
    template_path = TEMPLATES_DIR / f"{template_name}.json"
    if not template_path.exists():
        available = [p.stem for p in sorted(TEMPLATES_DIR.glob("*.json"))] if TEMPLATES_DIR.exists() else []
        raise ValueError(f"未知模板 '{template_name}'。可用模板: {available}")
    return json.loads(template_path.read_text(encoding="utf-8"))


def list_templates() -> List[Dict[str, str]]:
    """列出所有可用模板的名称和描述。"""
    if not TEMPLATES_DIR.exists():
        return []
    result = []
    for p in sorted(TEMPLATES_DIR.glob("*.json")):
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
            result.append({
                "name": data.get("name", p.stem),
                "display_name": data.get("display_name", p.stem),
                "description": data.get("description", ""),
            })
        except (json.JSONDecodeError, KeyError):
            result.append({"name": p.stem, "display_name": p.stem, "description": ""})
    return result
