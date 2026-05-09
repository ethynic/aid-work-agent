"""
Word 文档读取与分析

合并 word_lib.py 的 document_to_dict 和 word_reader.py 的简化读取能力。
"""

from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.oxml.ns import qn


def read_content(file_path: str) -> Dict[str, Any]:
    """读取 Word 文档的文本内容和表格"""
    path = Path(file_path)
    if not path.exists():
        return {"success": False, "error": f"文件不存在: {file_path}"}

    try:
        doc = Document(str(path))
    except Exception as e:
        return {"success": False, "error": f"无法打开文档: {e}"}

    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text.strip())

    tables = []
    for table in doc.tables:
        table_text = []
        if table.rows:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            table_text.append(" | ".join(headers))
            table_text.append("-" * len(" | ".join(headers)))
            for row in table.rows[1:]:
                row_data = [cell.text.strip() for cell in row.cells]
                table_text.append(" | ".join(row_data))
        tables.append('\n'.join(table_text))

    content_parts = []
    content_parts.append('\n'.join(paragraphs))
    if tables:
        for i, table in enumerate(tables, 1):
            content_parts.append(f"\n\n表格 {i}:\n{table}")

    doc_info = _extract_properties(doc)

    return {
        "success": True,
        "content": ''.join(content_parts),
        "paragraphs": paragraphs,
        "tables": tables,
        "document_info": doc_info,
        "file_path": str(path.absolute()),
    }


def analyze_structure(file_path: str, detailed: bool = False) -> Dict[str, Any]:
    """分析 Word 文档的完整结构（段落样式、字体、表格布局、页面设置）"""
    path = Path(file_path)
    if not path.exists():
        return {"success": False, "error": f"文件不存在: {file_path}"}

    try:
        doc = Document(str(path))
    except Exception as e:
        return {"success": False, "error": f"无法打开文档: {e}"}

    result = document_to_dict(doc, detailed=detailed)
    result["file_path"] = str(path.absolute())
    result["file_name"] = path.name
    return result


def extract_metadata(file_path: str) -> Dict[str, Any]:
    """提取文档元信息（标题、作者、创建时间等）"""
    path = Path(file_path)
    if not path.exists():
        return {"success": False, "error": f"文件不存在: {file_path}"}

    try:
        doc = Document(str(path))
    except Exception as e:
        return {"success": False, "error": f"无法打开文档: {e}"}

    return {
        "success": True,
        "metadata": _extract_properties(doc),
        "file_path": str(path.absolute()),
    }


def document_to_dict(doc: Document, detailed: bool = False) -> Dict[str, Any]:
    """将 python-docx Document 转为结构化字典"""
    result = {
        "success": True,
        "document_properties": _extract_properties(doc),
        "styles_found": _extract_style_names(doc),
        "sections": [],
        "summary": {
            "paragraph_count": 0,
            "table_count": 0,
            "section_count": len(doc.sections),
        },
    }

    total_chars = 0

    for section_idx, section in enumerate(doc.sections):
        section_data = {
            "section_index": section_idx,
            "paragraphs": [],
            "tables": [],
            "page_setup": _extract_page_setup(section),
        }

        section_data["header"] = _extract_header_footer(section, "header")
        section_data["footer"] = _extract_header_footer(section, "footer")

        body = doc.element.body
        para_idx = 0
        table_idx = 0

        for element in body:
            if element.tag == qn("w:p"):
                if para_idx < len(doc.paragraphs):
                    para = doc.paragraphs[para_idx]
                    text = para.text
                    if text.strip() or detailed:
                        para_data = _extract_paragraph(para, para_idx, detailed)
                        section_data["paragraphs"].append(para_data)
                        total_chars += len(text)
                    else:
                        section_data["paragraphs"].append({
                            "index": para_idx,
                            "text": "",
                            "style": para.style.name if para.style else "Normal",
                        })
                    para_idx += 1
            elif element.tag == qn("w:tbl"):
                if table_idx < len(doc.tables):
                    table = doc.tables[table_idx]
                    table_data = _extract_table(table, table_idx)
                    section_data["tables"].append(table_data)
                    total_chars += len(table_data.get("text_content", ""))
                    table_idx += 1

        result["sections"].append(section_data)
        result["summary"]["paragraph_count"] += para_idx
        result["summary"]["table_count"] += table_idx

    result["summary"]["total_chars"] = total_chars
    return result


def _extract_properties(doc: Document) -> Dict[str, Any]:
    props = doc.core_properties
    return {
        "title": props.title or "",
        "author": props.author or "",
        "subject": props.subject or "",
        "keywords": props.keywords or "",
        "created": str(props.created) if props.created else "",
        "modified": str(props.modified) if props.modified else "",
        "last_modified_by": props.last_modified_by or "",
        "revision": props.revision or 0,
        "category": props.category or "",
    }


def _extract_style_names(doc: Document) -> List[str]:
    names = set()
    for para in doc.paragraphs:
        if para.style:
            names.add(para.style.name)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if para.style:
                        names.add(para.style.name)
    return sorted(names)


def _extract_page_setup(section) -> Dict[str, Any]:
    try:
        return {
            "page_width_cm": round(section.page_width.cm, 2) if section.page_width else None,
            "page_height_cm": round(section.page_height.cm, 2) if section.page_height else None,
            "top_margin_cm": round(section.top_margin.cm, 2) if section.top_margin else None,
            "bottom_margin_cm": round(section.bottom_margin.cm, 2) if section.bottom_margin else None,
            "left_margin_cm": round(section.left_margin.cm, 2) if section.left_margin else None,
            "right_margin_cm": round(section.right_margin.cm, 2) if section.right_margin else None,
            "orientation": str(section.orientation) if section.orientation else "PORTRAIT",
        }
    except Exception:
        return {}


def _extract_header_footer(section, which: str) -> Dict[str, Any]:
    try:
        hf = section.header if which == "header" else section.footer

        if hf.is_linked_to_previous:
            return {"linked_to_previous": True, "text": ""}

        paragraphs = [para.text.strip() for para in hf.paragraphs if para.text.strip()]
        return {
            "linked_to_previous": False,
            "text": "\n".join(paragraphs),
            "paragraph_count": len(paragraphs),
        }
    except Exception:
        return {"text": ""}


def _extract_paragraph(para, index: int, detailed: bool = False) -> Dict[str, Any]:
    data = {
        "index": index,
        "text": para.text,
        "style": para.style.name if para.style else "Normal",
    }

    pf = para.paragraph_format
    if pf.alignment is not None:
        data["alignment"] = str(pf.alignment).replace("WD_ALIGN_PARAGRAPH.", "")
    if pf.line_spacing is not None:
        data["line_spacing"] = pf.line_spacing
    if pf.space_before is not None:
        data["space_before_pt"] = pf.space_before.pt
    if pf.space_after is not None:
        data["space_after_pt"] = pf.space_after.pt

    if para.runs:
        first_run = para.runs[0]
        font = first_run.font
        data["font_name"] = font.name or ""
        data["font_size"] = font.size.pt if font.size else None
        data["bold"] = font.bold
        data["italic"] = font.italic

        if detailed:
            data["runs"] = []
            for run in para.runs:
                run_data = {"text": run.text}
                rf = run.font
                if rf.name:
                    run_data["font_name"] = rf.name
                if rf.size:
                    run_data["font_size"] = rf.size.pt
                if rf.bold is not None:
                    run_data["bold"] = rf.bold
                if rf.italic is not None:
                    run_data["italic"] = rf.italic
                if rf.underline is not None:
                    run_data["underline"] = rf.underline
                if rf.color and rf.color.rgb:
                    run_data["color"] = str(rf.color.rgb)
                data["runs"].append(run_data)

    return data


def _extract_table(table, index: int) -> Dict[str, Any]:
    rows_data = []
    for row in table.rows:
        row_data = [cell.text.strip() for cell in row.cells]
        rows_data.append(row_data)

    text_content = " | ".join(" | ".join(row) for row in rows_data)

    return {
        "index": index,
        "rows": len(table.rows),
        "cols": len(table.columns) if table.columns else 0,
        "data": rows_data,
        "text_content": text_content,
        "style": table.style.name if table.style else "",
    }
