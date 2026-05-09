"""
Word 文档对比

基于 difflib 实现段落级文本对比。
"""

import difflib
from pathlib import Path
from typing import Any, Dict, List

from docx import Document


def diff(file_path_old: str, file_path_new: str,
         output_format: str = "text", include_style: bool = False) -> Dict[str, Any]:
    """对比两个 Word 文档的差异"""
    path_old = Path(file_path_old)
    path_new = Path(file_path_new)

    if not path_old.exists():
        return {"success": False, "error": f"旧版本文件不存在: {file_path_old}"}
    if not path_new.exists():
        return {"success": False, "error": f"新版本文件不存在: {file_path_new}"}

    try:
        doc_old = Document(str(path_old))
        doc_new = Document(str(path_new))
    except Exception as e:
        return {"success": False, "error": f"无法打开文档: {e}"}

    # 提取段落文本
    old_paragraphs = [p.text for p in doc_old.paragraphs]
    new_paragraphs = [p.text for p in doc_new.paragraphs]

    # 段落级对比
    matcher = difflib.SequenceMatcher(None, old_paragraphs, new_paragraphs)
    changes = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            continue
        elif tag == "replace":
            for oi, ni in zip(range(i1, i2), range(j1, j2)):
                changes.append({
                    "type": "modified",
                    "location": "paragraph",
                    "index": ni,
                    "old_text": old_paragraphs[oi] if oi < len(old_paragraphs) else "",
                    "new_text": new_paragraphs[ni] if ni < len(new_paragraphs) else "",
                })
            # 处理不等长替换
            if i2 - i1 > j2 - j1:
                for oi in range(i1 + (j2 - j1), i2):
                    changes.append({
                        "type": "deleted",
                        "location": "paragraph",
                        "index": oi,
                        "old_text": old_paragraphs[oi],
                        "new_text": "",
                    })
            elif j2 - j1 > i2 - i1:
                for ni in range(j1 + (i2 - i1), j2):
                    changes.append({
                        "type": "added",
                        "location": "paragraph",
                        "index": ni,
                        "old_text": "",
                        "new_text": new_paragraphs[ni],
                    })
        elif tag == "insert":
            for ni in range(j1, j2):
                changes.append({
                    "type": "added",
                    "location": "paragraph",
                    "index": ni,
                    "old_text": "",
                    "new_text": new_paragraphs[ni],
                })
        elif tag == "delete":
            for oi in range(i1, i2):
                changes.append({
                    "type": "deleted",
                    "location": "paragraph",
                    "index": oi,
                    "old_text": old_paragraphs[oi],
                    "new_text": "",
                })

    # 表格对比
    old_tables = _extract_tables_text(doc_old)
    new_tables = _extract_tables_text(doc_new)
    for t_idx in range(max(len(old_tables), len(new_tables))):
        if t_idx >= len(old_tables):
            changes.append({
                "type": "added",
                "location": "table",
                "index": t_idx,
                "old_text": "",
                "new_text": f"新增表格 ({len(doc_new.tables[t_idx].rows)}行)",
            })
        elif t_idx >= len(new_tables):
            changes.append({
                "type": "deleted",
                "location": "table",
                "index": t_idx,
                "old_text": f"删除表格 ({len(doc_old.tables[t_idx].rows)}行)",
                "new_text": "",
            })
        else:
            table_changes = _diff_tables(doc_old.tables[t_idx], doc_new.tables[t_idx], t_idx)
            changes.extend(table_changes)

    # 生成报告
    summary = _generate_summary(changes)
    diff_text = _format_diff(changes, path_old.name, path_new.name, output_format)

    return {
        "success": True,
        "diff_text": diff_text,
        "changes": changes,
        "change_count": len(changes),
        "summary": summary,
    }


def _extract_tables_text(doc: Document) -> List[str]:
    return [" | ".join(cell.text.strip() for cell in row.cells)
            for table in doc.tables for row in table.rows]


def _diff_tables(table_old, table_new, table_idx: int) -> List[Dict]:
    """逐单元格对比两个表格"""
    changes = []
    max_rows = max(len(table_old.rows), len(table_new.rows))
    max_cols = max(
        len(table_old.columns) if table_old.columns else 0,
        len(table_new.columns) if table_new.columns else 0,
    )

    for r in range(max_rows):
        for c in range(max_cols):
            old_text = _safe_cell_text(table_old, r, c)
            new_text = _safe_cell_text(table_new, r, c)
            if old_text != new_text:
                changes.append({
                    "type": "modified",
                    "location": "table_cell",
                    "table_index": table_idx,
                    "row": r,
                    "col": c,
                    "old_text": old_text,
                    "new_text": new_text,
                })
    return changes


def _safe_cell_text(table, row: int, col: int) -> str:
    try:
        return table.rows[row].cells[col].text.strip()
    except (IndexError, AttributeError):
        return ""


def _generate_summary(changes: List[Dict]) -> str:
    modified = sum(1 for c in changes if c["type"] == "modified")
    added = sum(1 for c in changes if c["type"] == "added")
    deleted = sum(1 for c in changes if c["type"] == "deleted")

    parts = []
    if modified:
        parts.append(f"{modified}处修改")
    if added:
        parts.append(f"{added}处新增")
    if deleted:
        parts.append(f"{deleted}处删除")

    if not parts:
        return "两个文档完全相同，无差异"
    return f"共发现{len(changes)}处差异：" + "，".join(parts)


def _format_diff(changes: List[Dict], old_name: str, new_name: str,
                 output_format: str) -> str:
    if output_format == "markdown":
        return _format_markdown(changes, old_name, new_name)
    return _format_text(changes, old_name, new_name)


def _format_text(changes: List[Dict], old_name: str, new_name: str) -> str:
    lines = [
        "文档差异报告",
        f"旧文件：{old_name} | 新文件：{new_name}",
        "",
    ]

    modified = [c for c in changes if c["type"] == "modified"]
    added = [c for c in changes if c["type"] == "added"]
    deleted = [c for c in changes if c["type"] == "deleted"]

    if modified:
        lines.append(f"━━━ 修改（{len(modified)}处）━━━")
        for c in modified:
            if c["location"] == "paragraph":
                lines.append(f"\n第{c['index']}段：")
                lines.append(f"  -「{c['old_text'][:100]}」")
                lines.append(f"  +「{c['new_text'][:100]}」")
            elif c["location"] == "table_cell":
                lines.append(f"\n表{c['table_index']+1}，第{c['row']+1}行第{c['col']+1}列：")
                lines.append(f"  -「{c['old_text'][:100]}」")
                lines.append(f"  +「{c['new_text'][:100]}」")

    if added:
        lines.append(f"\n━━━ 新增（{len(added)}处）━━━")
        for c in added:
            if c["location"] == "paragraph":
                lines.append(f"\n第{c['index']}段（新增）：")
                lines.append(f"  +「{c['new_text'][:100]}」")

    if deleted:
        lines.append(f"\n━━━ 删除（{len(deleted)}处）━━━")
        for c in deleted:
            if c["location"] == "paragraph":
                lines.append(f"\n第{c['index']}段（删除）：")
                lines.append(f"  -「{c['old_text'][:100]}」")

    return '\n'.join(lines)


def _format_markdown(changes: List[Dict], old_name: str, new_name: str) -> str:
    lines = [
        "# 文档差异报告",
        "",
        f"**旧文件**：{old_name} | **新文件**：{new_name}",
        "",
    ]

    modified = [c for c in changes if c["type"] == "modified"]
    added = [c for c in changes if c["type"] == "added"]
    deleted = [c for c in changes if c["type"] == "deleted"]

    if modified:
        lines.append(f"## 修改（{len(modified)}处）")
        for c in modified:
            if c["location"] == "paragraph":
                lines.append(f"\n**第{c['index']}段**：")
                lines.append(f"- ❌ {c['old_text'][:100]}")
                lines.append(f"- ✅ {c['new_text'][:100]}")
            elif c["location"] == "table_cell":
                lines.append(f"\n**表{c['table_index']+1}，第{c['row']+1}行第{c['col']+1}列**：")
                lines.append(f"- ❌ {c['old_text'][:100]}")
                lines.append(f"- ✅ {c['new_text'][:100]}")

    if added:
        lines.append(f"\n## 新增（{len(added)}处）")
        for c in added:
            if c["location"] == "paragraph":
                lines.append(f"\n**第{c['index']}段**：")
                lines.append(f"- ✅ {c['new_text'][:100]}")

    if deleted:
        lines.append(f"\n## 删除（{len(deleted)}处）")
        for c in deleted:
            if c["location"] == "paragraph":
                lines.append(f"\n**第{c['index']}段**：")
                lines.append(f"- ❌ {c['old_text'][:100]}")

    return '\n'.join(lines)
