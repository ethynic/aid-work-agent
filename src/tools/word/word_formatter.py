"""
Word 文档格式化

从 word_ops.py 迁移 format 操作逻辑。
"""

from typing import Any, Dict

from docx import Document
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.tools.word.word_lib import StyleManager, resolve_alignment, PAGE_SIZES


def batch_format(doc: Document, operations: list) -> Dict[str, Any]:
    """批量执行多个格式化操作"""
    results = []
    for op in operations:
        result = execute_format_op(doc, op)
        results.append(result)
    return {
        "success": True,
        "operations_applied": len(results),
        "results": results,
    }


def execute_format_op(doc: Document, op: Dict[str, Any]) -> Dict[str, Any]:
    """执行单个格式操作"""
    op_type = op.get("type", "") or op.get("op", "")

    try:
        handlers = {
            "set_font": _fmt_set_font,
            "set_paragraph_format": _fmt_set_paragraph_format,
            "set_page_margins": _fmt_set_page_margins,
            "set_page_size": _fmt_set_page_size,
            "set_header": _fmt_set_header,
            "set_footer": _fmt_set_footer,
            "set_table_style": _fmt_set_table_style,
        }
        handler = handlers.get(op_type)
        if not handler:
            return {"op": op_type, "success": False, "error": f"未知格式操作: {op_type}"}
        return handler(doc, op)
    except Exception as e:
        return {"op": op_type, "success": False, "error": str(e)}


def _get_scope_paragraphs(doc: Document, scope: str):
    """根据 scope 获取匹配的段落列表"""
    if scope in ("all", "paragraphs"):
        return doc.paragraphs
    elif scope.startswith("heading"):
        try:
            level = int(scope.replace("heading", ""))
            return [p for p in doc.paragraphs if p.style and p.style.name == f"Heading {level}"]
        except ValueError:
            return []
    elif scope == "table":
        paras = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paras.extend(cell.paragraphs)
        return paras
    else:
        return [p for p in doc.paragraphs if p.style and p.style.name == scope]


def _fmt_set_font(doc: Document, op: Dict[str, Any]) -> Dict[str, Any]:
    """设置字体"""
    scope = op.get("scope", "all")
    paragraphs = _get_scope_paragraphs(doc, scope)
    count = 0

    font_spec = {k: v for k, v in op.items() if k in
                 ("font_name", "font_size", "bold", "italic", "underline", "color", "color_hex")}

    if not font_spec:
        return {"op": "set_font", "success": False, "error": "未指定任何字体属性"}

    for para in paragraphs:
        for run in para.runs:
            StyleManager.apply_font(run, font_spec)
            count += 1
        if not para.runs and para.text:
            run = para.add_run(para.text)
            para.text = ""
            StyleManager.apply_font(run, font_spec)
            count += 1

    return {"op": "set_font", "success": True, "scope": scope, "runs_affected": count}


def _fmt_set_paragraph_format(doc: Document, op: Dict[str, Any]) -> Dict[str, Any]:
    """设置段落格式"""
    scope = op.get("scope", "all")
    paragraphs = _get_scope_paragraphs(doc, scope)

    para_spec = {k: v for k, v in op.items() if k in
                 ("alignment", "line_spacing", "space_before", "space_after",
                  "first_line_indent")}

    if not para_spec:
        return {"op": "set_paragraph_format", "success": False, "error": "未指定任何段落格式属性"}

    count = 0
    for para in paragraphs:
        StyleManager.apply_paragraph_format(para, para_spec)
        count += 1

    return {"op": "set_paragraph_format", "success": True, "scope": scope, "paragraphs_affected": count}


def _fmt_set_page_margins(doc: Document, op: Dict[str, Any]) -> Dict[str, Any]:
    """设置页面边距"""
    for section in doc.sections:
        if "top" in op:
            section.top_margin = Cm(op["top"])
        if "bottom" in op:
            section.bottom_margin = Cm(op["bottom"])
        if "left" in op:
            section.left_margin = Cm(op["left"])
        if "right" in op:
            section.right_margin = Cm(op["right"])

    return {"op": "set_page_margins", "success": True, "sections_affected": len(doc.sections)}


def _fmt_set_page_size(doc: Document, op: Dict[str, Any]) -> Dict[str, Any]:
    """设置页面大小和方向"""
    size = op.get("size", "A4")
    orientation = op.get("orientation", "portrait")

    if size.upper() not in PAGE_SIZES:
        return {"op": "set_page_size", "success": False,
                "error": f"不支持的页面大小: {size}。支持: {', '.join(PAGE_SIZES.keys())}"}

    w, h = PAGE_SIZES[size.upper()]
    for section in doc.sections:
        if orientation.lower() == "landscape":
            section.page_width = Cm(h)
            section.page_height = Cm(w)
        else:
            section.page_width = Cm(w)
            section.page_height = Cm(h)

    return {"op": "set_page_size", "success": True, "size": size, "orientation": orientation}


def _fmt_set_header(doc: Document, op: Dict[str, Any]) -> Dict[str, Any]:
    """设置页眉"""
    if not doc.sections:
        return {"op": "set_header", "success": False, "error": "文档没有 section"}

    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False

    for para in header.paragraphs:
        para.text = ""

    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    para.text = op.get("text", "")

    font_spec = {k: v for k, v in op.items() if k in
                 ("font_name", "font_size", "bold", "italic", "color")}
    if font_spec and para.runs:
        StyleManager.apply_font(para.runs[0], font_spec)

    align = op.get("alignment")
    if align:
        resolved = resolve_alignment(align)
        if resolved is not None:
            para.paragraph_format.alignment = resolved

    return {"op": "set_header", "success": True}


def _fmt_set_footer(doc: Document, op: Dict[str, Any]) -> Dict[str, Any]:
    """设置页脚（支持页码）"""
    if not doc.sections:
        return {"op": "set_footer", "success": False, "error": "文档没有 section"}

    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False

    for para in footer.paragraphs:
        para.text = ""

    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.text = op.get("text", "")

    if op.get("include_page_number"):
        run = para.add_run()
        fldChar1 = OxmlElement("w:fldChar")
        fldChar1.set(qn("w:fldCharType"), "begin")
        run._element.append(fldChar1)

        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")
        instrText.text = " PAGE "
        run._element.append(instrText)

        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "end")
        run._element.append(fldChar2)

    font_spec = {k: v for k, v in op.items() if k in
                 ("font_name", "font_size", "bold", "italic", "color")}
    if font_spec and para.runs:
        StyleManager.apply_font(para.runs[0], font_spec)

    align = op.get("alignment")
    if align:
        resolved = resolve_alignment(align)
        if resolved is not None:
            para.paragraph_format.alignment = resolved

    return {"op": "set_footer", "success": True}


def _fmt_set_table_style(doc: Document, op: Dict[str, Any]) -> Dict[str, Any]:
    """设置表格样式"""
    table_index = int(op.get("table_index", 0))
    style_name = op.get("style", "Table Grid")

    if table_index < 0 or table_index >= len(doc.tables):
        return {"op": "set_table_style", "success": False, "error": f"无效表格索引: {table_index}"}

    table = doc.tables[table_index]
    try:
        table.style = style_name
    except KeyError:
        return {"op": "set_table_style", "success": False, "error": f"无效表格样式: {style_name}"}

    if op.get("header_row_bold") and table.rows:
        for cell in table.rows[0].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.bold = True

    return {"op": "set_table_style", "success": True, "table_index": table_index, "style": style_name}
