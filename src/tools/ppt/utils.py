"""
PPT 工具通用工具函数

颜色解析、单位转换、形状创建等共享功能。
"""

import re
from typing import Optional, Tuple

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


def parse_color(color_str: str) -> Optional[RGBColor]:
    """解析颜色字符串为 RGBColor（支持 #RRGGBB 和纯 RRGGBB 格式）。"""
    if not color_str:
        return None
    color_str = color_str.lstrip("#")
    if len(color_str) == 6:
        try:
            r = int(color_str[:2], 16)
            g = int(color_str[2:4], 16)
            b = int(color_str[4:6], 16)
            return RGBColor(r, g, b)
        except ValueError:
            return None
    return None


def inches(val: float) -> int:
    """浮点英寸 → Emu 整数。"""
    return Inches(val)


def pt(val: float) -> int:
    """浮点 pt → Emu 整数。"""
    return Pt(val)


def alignment(align_str: str) -> int:
    """字符串对齐方式 → PP_ALIGN 枚举值。"""
    mapping = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }
    return mapping.get(align_str.lower(), PP_ALIGN.LEFT)


def set_slide_bg(slide, color_hex: str):
    """设置幻灯片背景色。"""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = parse_color(color_hex) or RGBColor(0xFF, 0xFF, 0xFF)


def add_textbox(
    slide,
    text: str,
    x: float, y: float, w: float, h: float,
    font_size: float = 16,
    bold: bool = False,
    color: str = "000000",
    font_name_cn: str = "微软雅黑",
    font_name_en: str = "Arial",
    align: str = "left",
    line_spacing: float = 1.2,
):
    """添加文本框到幻灯片。

    x/y/w/h 单位为英寸。font_size 单位为 pt。color 为 RRGGBB 格式。
    """
    from pptx.util import Pt as Ppt, Inches as In
    from pptx.dml.color import RGBColor as RGB
    from pptx.enum.text import PP_ALIGN as PA

    txBox = slide.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Ppt(font_size)
    p.font.bold = bold
    p.font.color.rgb = parse_color(color) or RGB(0, 0, 0)

    # 设置中英文字体
    p.font.name = font_name_en
    run_elem = p.runs[0]._r if p.runs else p._r
    from pptx.oxml.ns import qn
    rPr = run_elem.find(qn("a:rPr"))
    if rPr is None:
        rPr = run_elem.makeelement(qn("a:rPr"), {})
        run_elem.insert(0, rPr)
    rPr.set("lang", "zh-CN")
    rPr.set("altLang", "en-US")
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = run_elem.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font_name_cn)

    align_map = {"left": PA.LEFT, "center": PA.CENTER, "right": PA.RIGHT}
    p.alignment = align_map.get(align, PA.LEFT)

    p.space_after = Pt(0)
    p.space_before = Pt(0)

    return txBox


def add_shape_rect(
    slide,
    x: float, y: float, w: float, h: float,
    fill_color: str = "000000",
    border_color: Optional[str] = None,
    border_width: float = 0,
):
    """添加矩形形状。"""
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = parse_color(fill_color) or RGBColor(0, 0, 0)

    if border_color:
        shape.line.color.rgb = parse_color(border_color)
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()

    return shape


def add_shape_rounded_rect(
    slide,
    x: float, y: float, w: float, h: float,
    fill_color: str = "000000",
    border_color: Optional[str] = None,
    border_width: float = 0,
    corner_radius: float = 0.1,
):
    """添加圆角矩形形状。"""
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = parse_color(fill_color) or RGBColor(0, 0, 0)

    if border_color:
        shape.line.color.rgb = parse_color(border_color)
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()

    # 调整圆角大小
    try:
        shape.adjustments[0] = min(corner_radius / max(w, h), 0.5)
    except (AttributeError, IndexError):
        pass

    return shape


def add_line(
    slide,
    x1: float, y1: float, x2: float, y2: float,
    color: str = "000000",
    width: float = 2,
):
    """添加线条。"""
    from pptx.enum.shapes import MSO_SHAPE

    connector = slide.shapes.add_connector(
        1,  # MSO_CONNECTOR.STRAIGHT
        Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = parse_color(color) or RGBColor(0, 0, 0)
    connector.line.width = Pt(width)

    return connector


def add_page_number(slide, index: int, total: int, color: str = "888888"):
    """在幻灯片右下角添加页码。"""
    from src.tools.ppt.theme import PAGE_NUMBER_X, PAGE_NUMBER_Y

    text = f"{index} / {total}"
    add_textbox(
        slide, text,
        x=PAGE_NUMBER_X, y=PAGE_NUMBER_Y, w=1.0, h=0.4,
        font_size=10, color=color, align="right",
    )


def truncate_text(text: str, max_chars: int = 30) -> str:
    """截断文本到指定最大字符数。"""
    if len(text) <= max_chars:
        return text
    return text[: max_chars - 1] + "…"
