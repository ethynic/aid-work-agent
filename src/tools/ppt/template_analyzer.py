"""Auditable PowerPoint template analysis and template-following generation."""

from __future__ import annotations

import copy
import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


class TemplateAnalysisError(RuntimeError):
    """Sanitized template processing failure."""


class TemplateAnalyzer:
    """Analyze and generate decks without introducing an unrelated theme."""

    MAX_TEMPLATE_BYTES = 50 * 1024 * 1024
    AUDIT_NAME = "template-audit.json"
    FRAME_MAP_NAME = "template-frame-map.json"
    DEVIATION_NAME = "deviation-log.json"

    def __init__(self):
        self.deviations: list[dict[str, Any]] = []
        self.frame_map: dict[str, Any] = {"version": "1.0", "frames": []}

    def analyze(
        self, template_path: str | Path, artifact_dir: str | Path | None = None
    ) -> dict:
        """Return a serializable audit and optionally persist it."""
        source = self._validate_template_path(template_path)
        try:
            prs = Presentation(str(source))
        except (OSError, ValueError, KeyError, PackageNotFoundError) as exc:
            raise TemplateAnalysisError("PPTX 模板无效或无法读取") from exc

        master_indexes = {
            str(master.part.partname): index
            for index, master in enumerate(prs.slide_masters)
        }
        layout_indexes = {
            str(layout.part.partname): index
            for index, layout in enumerate(prs.slide_layouts)
        }
        audit = {
            "version": "1.0",
            "slide_size": {
                "width_emu": prs.slide_width,
                "height_emu": prs.slide_height,
                "width_inches": round(prs.slide_width / 914400, 4),
                "height_inches": round(prs.slide_height / 914400, 4),
            },
            # Retained for callers from Phase 0-5.
            "slide_width": prs.slide_width,
            "slide_height": prs.slide_height,
            "theme_colors": self._extract_theme_colors(prs),
            "theme_fonts": self._extract_theme_fonts(prs),
            "masters": [
                {
                    "index": index,
                    "name": master.name,
                    "shape_count": len(master.shapes),
                    "layout_indexes": [
                        layout_indexes[str(layout.part.partname)]
                        for layout in master.slide_layouts
                        if str(layout.part.partname) in layout_indexes
                    ],
                }
                for index, master in enumerate(prs.slide_masters)
            ],
            "layouts": [
                self._analyze_layout(
                    layout,
                    index,
                    master_indexes.get(str(layout.slide_master.part.partname), 0),
                )
                for index, layout in enumerate(prs.slide_layouts)
            ],
            "slides": [
                self._analyze_slide(slide, index, layout_indexes)
                for index, slide in enumerate(prs.slides)
            ],
        }
        if artifact_dir is not None:
            self._write_json(artifact_dir, self.AUDIT_NAME, audit)
        return audit

    def _extract_theme_colors(self, prs: Presentation) -> dict[str, str]:
        colors: dict[str, str] = {}
        try:
            theme = prs.slide_masters[0].part.theme_part.element
            scheme = theme.find(".//" + qn("a:clrScheme"))
            if scheme is None:
                return colors
            for item in scheme:
                if not len(item):
                    continue
                value = item[0].get("val") or item[0].get("lastClr")
                if value and len(value) == 6:
                    colors[item.tag.rsplit("}", 1)[-1]] = value.upper()
        except (AttributeError, IndexError):
            pass
        return colors

    def _extract_theme_fonts(self, prs: Presentation) -> dict[str, str]:
        fonts: dict[str, str] = {}
        try:
            theme = prs.slide_masters[0].part.theme_part.element
            for key, xpath in {
                "major_latin": ".//a:majorFont/a:latin",
                "minor_latin": ".//a:minorFont/a:latin",
                "major_east_asian": ".//a:majorFont/a:ea",
                "minor_east_asian": ".//a:minorFont/a:ea",
            }.items():
                nodes = theme.xpath(xpath)
                if nodes and nodes[0].get("typeface"):
                    fonts[key] = nodes[0].get("typeface")
        except (AttributeError, IndexError):
            pass
        return fonts

    def _analyze_layout(self, layout, index: int, master_index: int) -> dict:
        return {
            "index": index,
            "name": layout.name,
            "master_index": master_index,
            "inferred_type": self._infer_type(layout),
            "placeholders": [
                self._placeholder_summary(placeholder)
                for placeholder in layout.placeholders
            ],
        }

    def _analyze_slide(self, slide, index: int, layout_indexes: dict[int, int]) -> dict:
        texts: list[str] = []
        fonts: set[str] = set()
        colors: set[str] = set()
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if shape.text.strip():
                texts.append(shape.text.strip()[:200])
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        fonts.add(run.font.name)
                    color = self._run_color(run)
                    if color:
                        colors.add(color)
        return {
            "index": index,
            "layout_index": layout_indexes.get(
                str(slide.slide_layout.part.partname), 0
            ),
            "layout_name": slide.slide_layout.name,
            "shape_count": len(slide.shapes),
            "placeholder_count": len(slide.placeholders),
            "placeholders": [
                self._placeholder_summary(placeholder)
                for placeholder in slide.placeholders
            ],
            "text_summary": texts,
            "fonts": sorted(fonts),
            "colors": sorted(colors),
            "has_notes": bool(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()),
        }

    @staticmethod
    def _run_color(run) -> str | None:
        try:
            rgb = run.font.color.rgb
            return str(rgb).upper() if rgb is not None else None
        except AttributeError:
            return None

    @staticmethod
    def _placeholder_summary(placeholder) -> dict:
        placeholder_type = placeholder.placeholder_format.type
        return {
            "idx": placeholder.placeholder_format.idx,
            "type": str(placeholder_type),
            "type_code": int(placeholder_type),
            "name": placeholder.name,
            "position": {
                "left": placeholder.left,
                "top": placeholder.top,
                "width": placeholder.width,
                "height": placeholder.height,
            },
        }

    def _infer_type(self, layout) -> str:
        name = layout.name.lower()
        if "title" in name and "content" not in name and "and" not in name:
            return "cover"
        if "section" in name or "header" in name:
            return "section"
        return "content"

    def match_content_to_layouts(
        self,
        plan: dict,
        analysis: dict,
        artifact_dir: str | Path | None = None,
    ) -> list[dict]:
        """Build an auditable output-page to template-frame mapping."""
        layouts = analysis.get("layouts", [])
        source_slides = analysis.get("slides", [])
        self.deviations = []
        matches: list[dict] = []
        frames: list[dict] = []

        for output_index, slide_data in enumerate(plan.get("slides", [])):
            requested_slide_value = slide_data.get("template_slide_index")
            requested_layout_value = slide_data.get("layout_index")
            requested_slide = self._valid_index(
                requested_slide_value, len(source_slides)
            )
            requested_layout = self._valid_index(
                requested_layout_value, len(layouts)
            )
            strategy = "clone" if requested_slide is not None else "layout"
            if requested_slide is not None:
                source_info = source_slides[requested_slide]
                layout_index = source_info["layout_index"]
            else:
                layout_index = requested_layout
                if layout_index is None:
                    layout_index = self._best_layout_index(slide_data, layouts)
            layout = layouts[layout_index] if layout_index is not None else None
            available_placeholder_indexes = None
            if requested_slide is not None:
                available_placeholder_indexes = {
                    placeholder["idx"]
                    for placeholder in source_slides[requested_slide].get(
                        "placeholders", []
                    )
                }
            edit_targets, additions = self._build_edit_targets(
                slide_data, layout, available_placeholder_indexes
            )
            deviations = self._frame_deviations(
                output_index, slide_data, layout, additions
            )
            if (
                requested_slide_value is not None and requested_slide is None
            ) or (
                requested_layout_value is not None and requested_layout is None
            ):
                deviations.append(
                    self._deviation(output_index, "requested_frame_unavailable")
                )
            if (
                layout
                and slide_data.get("type")
                and layout["inferred_type"] != slide_data["type"]
            ):
                deviations.append(
                    self._deviation(output_index, "layout_type_fallback")
                )
            self.deviations.extend(deviations)
            frame = {
                "output_slide_index": output_index,
                "source_slide_index": requested_slide,
                "layout_index": layout_index,
                "layout_name": layout["name"] if layout else None,
                "strategy": strategy,
                "editTargets": edit_targets,
                "allowedAdditions": additions,
                "deviation_codes": [item["code"] for item in deviations],
            }
            frames.append(frame)
            matches.append({"slide": slide_data, "layout": layout, "frame": frame})

        self.frame_map = {"version": "1.0", "frames": frames}
        if artifact_dir is not None:
            self._write_json(artifact_dir, self.FRAME_MAP_NAME, self.frame_map)
            self._write_json(
                artifact_dir,
                self.DEVIATION_NAME,
                {"version": "1.0", "deviations": self.deviations},
            )
        return matches

    @staticmethod
    def _valid_index(value: Any, length: int) -> int | None:
        return value if isinstance(value, int) and 0 <= value < length else None

    def _best_layout_index(self, slide_data: dict, layouts: list[dict]) -> int | None:
        if not layouts:
            return None
        slide_type = slide_data.get("type", "content")
        candidates = [
            layout for layout in layouts if layout["inferred_type"] == slide_type
        ]
        if not candidates:
            candidates = [
                layout for layout in layouts if layout["inferred_type"] == "content"
            ]
        if not candidates:
            return layouts[0]["index"]
        content_length = sum(len(str(item)) for item in slide_data.get("points", []))
        return max(
            candidates,
            key=lambda layout: self._body_capacity(layout) - min(
                self._body_capacity(layout), content_length
            ),
        )["index"]

    @staticmethod
    def _body_capacity(layout: dict) -> int:
        body = [
            item
            for item in layout.get("placeholders", [])
            if any(token in item["type"] for token in ("BODY", "OBJECT", "SUBTITLE"))
        ]
        return max(
            (
                item["position"]["width"] * item["position"]["height"] // 10**11
                for item in body
            ),
            default=0,
        )

    def _build_edit_targets(
        self,
        slide_data: dict,
        layout: dict | None,
        available_placeholder_indexes: set[int] | None = None,
    ) -> tuple[list[dict], list[dict]]:
        placeholders = layout.get("placeholders", []) if layout else []
        targets: list[dict] = []
        additions: list[dict] = []
        title = next(
            (
                item
                for item in placeholders
                if "TITLE" in item["type"] and "SUBTITLE" not in item["type"]
            ),
            None,
        )
        body = next(
            (
                item
                for item in placeholders
                if any(token in item["type"] for token in ("BODY", "OBJECT", "SUBTITLE"))
            ),
            None,
        )
        if (
            title
            and slide_data.get("title")
            and (
                available_placeholder_indexes is None
                or title["idx"] in available_placeholder_indexes
            )
        ):
            targets.append(
                {"placeholder_idx": title["idx"], "role": "title", "source": "title"}
            )
        elif slide_data.get("title"):
            additions.append(
                {
                    "type": "text",
                    "role": "title",
                    "reason": (
                        "source_placeholder_missing"
                        if title
                        else "missing_title_placeholder"
                    ),
                }
            )
        if (
            body
            and self._body_text(slide_data)
            and (
                available_placeholder_indexes is None
                or body["idx"] in available_placeholder_indexes
            )
        ):
            targets.append(
                {"placeholder_idx": body["idx"], "role": "body", "source": "content"}
            )
        elif self._body_text(slide_data):
            additions.append(
                {
                    "type": "text",
                    "role": "body",
                    "reason": (
                        "source_placeholder_missing"
                        if body
                        else "missing_body_placeholder"
                    ),
                }
            )
        return targets, additions

    def _frame_deviations(
        self,
        output_index: int,
        slide_data: dict,
        layout: dict | None,
        additions: list[dict],
    ) -> list[dict]:
        items: list[dict] = []
        if layout is None:
            items.append(self._deviation(output_index, "no_matching_layout"))
        if additions:
            items.append(self._deviation(output_index, "object_addition_required"))
        content_length = len(self._body_text(slide_data))
        if layout and content_length > max(300, self._body_capacity(layout)):
            items.append(self._deviation(output_index, "content_too_dense"))
        return items

    @staticmethod
    def _deviation(output_index: int, code: str) -> dict:
        messages = {
            "no_matching_layout": "未找到可安全匹配的模板版式",
            "object_addition_required": "模板缺少内容占位符，已声明新增文本对象",
            "content_too_dense": "内容密度可能超过模板占位符容量",
            "clone_fallback": "源页面克隆失败，已降级为版式新建页",
            "requested_frame_unavailable": "指定的模板页面或版式不可用，已选择最接近版式",
            "layout_type_fallback": "未找到同类型模板版式，已选择最接近版式",
        }
        return {
            "output_slide_index": output_index,
            "code": code,
            "message": messages[code],
        }

    def generate_from_template(
        self,
        template_path: str | Path,
        matches: list[dict],
        plan: dict,
        artifact_dir: str | Path | None = None,
    ) -> str:
        """Generate slides from the template's own layouts and source frames."""
        source = self._validate_template_path(template_path)
        if artifact_dir is not None:
            self._require_artifacts(artifact_dir)
        try:
            prs = Presentation(str(source))
        except (OSError, ValueError, KeyError, PackageNotFoundError) as exc:
            raise TemplateAnalysisError("PPTX 模板无效或无法读取") from exc
        source_slides = list(prs.slides)
        self._remove_all_slides(prs)

        for output_index, match in enumerate(matches):
            frame = match["frame"]
            layout_index = frame.get("layout_index")
            if layout_index is None or not 0 <= layout_index < len(prs.slide_layouts):
                raise TemplateAnalysisError("模板缺少可用版式，无法安全生成")
            slide = None
            source_index = frame.get("source_slide_index")
            if frame.get("strategy") == "clone" and source_index is not None:
                try:
                    slide = self._clone_slide(prs, source_slides[source_index])
                except (AttributeError, KeyError, ValueError):
                    self.deviations.append(
                        self._deviation(output_index, "clone_fallback")
                    )
            if slide is None:
                slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
            self._fill_frame(slide, match["slide"], frame)

        output_dir = self._get_output_dir()
        output_dir.mkdir(parents=True, exist_ok=True)
        safe_name = self._safe_output_name(plan.get("title", "presentation"))
        output_path = output_dir / f"{safe_name}.pptx"
        try:
            prs.save(str(output_path))
        except OSError as exc:
            raise TemplateAnalysisError("模板 PPTX 输出失败") from exc
        if artifact_dir is not None and self.deviations:
            self._write_json(
                artifact_dir,
                self.DEVIATION_NAME,
                {"version": "1.0", "deviations": self.deviations},
            )
        return str(output_path.resolve())

    @staticmethod
    def _remove_all_slides(prs: Presentation) -> None:
        while prs.slides:
            slide_id = prs.slides._sldIdLst[0]
            prs.part.drop_rel(slide_id.get(qn("r:id")))
            del prs.slides._sldIdLst[0]

    def _clone_slide(self, prs: Presentation, source_slide):
        target = prs.slides.add_slide(source_slide.slide_layout)
        shape_tree = target.shapes._spTree
        for shape in list(target.shapes):
            shape_tree.remove(shape.element)
        relationship_map: dict[str, str] = {}
        for relationship in source_slide.part.rels.values():
            if relationship.reltype.endswith(("/slideLayout", "/notesSlide")):
                continue
            if relationship.is_external:
                new_id = target.part.rels.add_relationship(
                    relationship.reltype, relationship.target_ref, is_external=True
                )
            else:
                new_id = target.part.relate_to(
                    relationship.target_part, relationship.reltype
                )
            relationship_map[relationship.rId] = new_id
        for source_shape in source_slide.shapes:
            element = copy.deepcopy(source_shape.element)
            for node in element.iter():
                for attribute, value in list(node.attrib.items()):
                    if value in relationship_map:
                        node.set(attribute, relationship_map[value])
            shape_tree.insert_element_before(element, "p:extLst")
        return target

    def _fill_frame(self, slide, data: dict, frame: dict) -> None:
        placeholders = {
            placeholder.placeholder_format.idx: placeholder
            for placeholder in slide.placeholders
        }
        for target in frame.get("editTargets", []):
            placeholder = placeholders.get(target["placeholder_idx"])
            if placeholder is None or not getattr(placeholder, "has_text_frame", False):
                continue
            placeholder.text = (
                str(data.get("title", ""))
                if target["role"] == "title"
                else self._body_text(data)
            )
        for addition in frame.get("allowedAdditions", []):
            if addition["type"] != "text":
                continue
            is_title = addition["role"] == "title"
            textbox = slide.shapes.add_textbox(
                Inches(0.8),
                Inches(0.5 if is_title else 1.8),
                Inches(11.7 if is_title else 8.4),
                Inches(0.8 if is_title else 4.5),
            )
            textbox.text_frame.text = (
                str(data.get("title", "")) if is_title else self._body_text(data)
            )
            for paragraph in textbox.text_frame.paragraphs:
                paragraph.font.size = Pt(28 if is_title else 20)

    @staticmethod
    def _body_text(data: dict) -> str:
        points = data.get("points") or data.get("takeaways") or []
        if points:
            return "\n".join(str(item) for item in points)
        return str(data.get("subtitle") or data.get("intro") or "")

    def _require_artifacts(self, artifact_dir: str | Path) -> None:
        directory = Path(artifact_dir)
        required = [self.AUDIT_NAME, self.FRAME_MAP_NAME, self.DEVIATION_NAME]
        if any(not (directory / name).is_file() for name in required):
            raise TemplateAnalysisError("模板审计产物不完整，已停止生成")

    def _validate_template_path(self, template_path: str | Path) -> Path:
        try:
            source = Path(template_path).expanduser().resolve(strict=True)
        except (OSError, RuntimeError) as exc:
            raise TemplateAnalysisError("PPTX 模板不存在或不可访问") from exc
        if (
            not source.is_file()
            or source.suffix.lower() != ".pptx"
            or source.stat().st_size > self.MAX_TEMPLATE_BYTES
        ):
            raise TemplateAnalysisError("PPTX 模板无效或超过 50MB 限制")
        return source

    @staticmethod
    def _safe_output_name(title: Any) -> str:
        value = "".join(
            character if character.isalnum() or character in "._- " else "_"
            for character in str(title)
        ).strip(" .")
        return value or "presentation"

    @staticmethod
    def _write_json(
        artifact_dir: str | Path, filename: str, payload: dict
    ) -> Path:
        directory = Path(artifact_dir).resolve()
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / filename
        path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return path

    def _get_output_dir(self) -> Path:
        try:
            from src.main import _get_tenant_upload_dir

            return _get_tenant_upload_dir()
        except (ImportError, AttributeError):
            return Path("storage/ppt")
