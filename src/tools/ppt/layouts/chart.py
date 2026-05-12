"""
图表页渲染器

使用 python-pptx 原生图表 API 创建柱状图、折线图、饼图。
"""

from src.tools.ppt.theme import PPTTheme, CANVAS_WIDTH
from src.tools.ppt.utils import (
    set_slide_bg, add_textbox, add_shape_rect, add_page_number,
)


def render_chart(slide, data: dict, theme: PPTTheme, index: int, total: int):
    """渲染图表页。"""
    set_slide_bg(slide, theme.bg)

    # 页面标题
    add_textbox(
        slide, data.get("title", ""),
        x=0.8, y=0.4, w=11.7, h=0.9,
        font_size=28, bold=True, color=theme.primary,
    )

    # 装饰线
    add_shape_rect(slide, x=0.8, y=1.2, w=1.5, h=0.04, fill_color=theme.accent)

    chart_data = data.get("chart", {})
    chart_type = chart_data.get("type", "bar")

    try:
        _add_chart(slide, chart_data, chart_type, theme)
    except Exception:
        # 图表创建失败时显示占位信息
        add_textbox(
            slide, f"[{chart_type} 图表]",
            x=2.0, y=2.5, w=9.0, h=3.0,
            font_size=24, color=theme.secondary,
            align="center",
        )

    # takeaway 注释
    takeaway = data.get("takeaway", "")
    if takeaway:
        add_textbox(
            slide, f"💡 {takeaway}",
            x=0.8, y=6.0, w=10.0, h=0.5,
            font_size=12, color=theme.accent,
        )

    add_page_number(slide, index, total, color=theme.secondary)


def _add_chart(slide, chart_data: dict, chart_type: str, theme: PPTTheme):
    """使用 python-pptx 图表 API 创建图表。"""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from src.tools.ppt.utils import parse_color

    labels = chart_data.get("labels", [])
    series_list = chart_data.get("series", [])

    if not labels or not series_list:
        return

    # 构建图表数据
    chart_type_enum = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }.get(chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)

    cd = CategoryChartData()
    cd.categories = labels
    for s in series_list:
        cd.add_series(s.get("name", ""), s.get("values", []))

    chart_frame = slide.shapes.add_chart(
        chart_type_enum,
        Inches(1.5), Inches(1.6), Inches(10.0), Inches(4.2),
        cd,
    )

    chart = chart_frame.chart
    chart.has_legend = len(series_list) > 1

    # 样式调整
    try:
        plot = chart.plots[0]
        plot.gap_width = 120

        primary_color = parse_color(theme.primary)
        accent_color = parse_color(theme.accent)

        for idx, series in enumerate(plot.series):
            color = primary_color if idx == 0 else accent_color
            if chart_type == "bar":
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color
            elif chart_type == "line":
                series.format.line.color.rgb = color
                series.format.line.width = Pt(2.5)
                series.smooth = True
    except (AttributeError, IndexError):
        pass

    # 图例位置
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
