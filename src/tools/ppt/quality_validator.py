"""Unified structural and optional rendered-preview validation for PPTX files."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
from zipfile import BadZipFile
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError


class PPTQualityValidator:
    """Validate PPTX delivery quality without exposing internal paths."""

    VERSION = "1.0"
    PREVIEW_TIMEOUT_SECONDS = 90

    def validate(
        self,
        pptx_path: str | Path,
        *,
        expected_slide_count: int | None = None,
        layout: dict | None = None,
        editable_ratio: float | None = None,
        warnings: list[str] | None = None,
        strict: bool = False,
    ) -> dict[str, Any]:
        path = Path(pptx_path)
        report_warnings = list(dict.fromkeys(warnings or []))
        errors: list[str] = []
        structural_layout = self._empty_layout()
        summary = self._base_summary(expected_slide_count)

        if not path.is_file() or path.stat().st_size == 0:
            errors.append("PPTX 文件不存在或为空")
        else:
            summary["file_valid"] = True
            summary["pptx_size_bytes"] = path.stat().st_size
            try:
                prs = Presentation(str(path))
            except (OSError, ValueError, KeyError, BadZipFile, PackageNotFoundError):
                errors.append("PPTX 文件损坏或无法打开")
            else:
                summary["openable"] = True
                structural_layout, structure = self._inspect_presentation(prs)
                summary.update(structure)
                if structure["slide_count"] == 0:
                    errors.append("PPTX 不包含幻灯片")
                if (
                    expected_slide_count is not None
                    and structure["slide_count"] != expected_slide_count
                ):
                    summary["page_count_matches"] = False
                    errors.append("PPTX 页数与预期不一致")
                if structure["empty_slide_count"]:
                    errors.append("PPTX 包含空白页")
                if structure["out_of_bounds_count"]:
                    errors.append("PPTX 包含越界对象")
                if structure["invalid_image_count"]:
                    errors.append("PPTX 包含无效图片")

        effective_layout = self._merge_layout(structural_layout, layout)
        raster = self._raster_summary(effective_layout)
        summary.update(raster)
        if raster["unreadable_raster_count"]:
            errors.append("PPTX 包含不可读取的 raster 图层")
        summary["editable_ratio"] = (
            round(max(0.0, min(1.0, editable_ratio)), 4)
            if editable_ratio is not None
            else self._editable_ratio(effective_layout)
        )

        if summary["openable"]:
            preview = self._render_preview(path)
            summary.update(preview["summary"])
            report_warnings.extend(preview["warnings"])
            if (
                summary["preview_status"] == "png"
                and summary["preview_page_count"] != summary["slide_count"]
            ):
                errors.append("PPTX 预览页数与结构页数不一致")
        else:
            summary.update(
                {"preview_status": "skipped", "preview_page_count": 0}
            )

        report_warnings = list(dict.fromkeys(report_warnings))
        errors = list(dict.fromkeys(errors))
        summary.update(
            {
                "warning_count": len(report_warnings),
                "error_count": len(errors),
                "qa_passed": not errors,
                "strict_mode": strict,
                "deliverable": not errors or not strict,
            }
        )
        report = {
            "version": self.VERSION,
            "status": "passed" if not errors else ("failed" if strict else "warning"),
            "summary": summary,
            "warnings": report_warnings,
            "errors": errors,
        }
        self._write_artifacts(path, effective_layout, report)
        return report

    def _inspect_presentation(self, prs: Presentation) -> tuple[dict, dict]:
        slides: list[dict] = []
        object_count = 0
        empty_count = 0
        out_of_bounds_count = 0
        image_count = 0
        invalid_image_count = 0
        width = int(prs.slide_width)
        height = int(prs.slide_height)

        for slide_index, slide in enumerate(prs.slides):
            objects: list[dict] = []
            slide_nonempty = False
            for shape in slide.shapes:
                object_count += 1
                bounds = {
                    "x": int(shape.left),
                    "y": int(shape.top),
                    "w": int(shape.width),
                    "h": int(shape.height),
                }
                out_of_bounds = (
                    bounds["x"] < 0
                    or bounds["y"] < 0
                    or bounds["x"] + bounds["w"] > width
                    or bounds["y"] + bounds["h"] > height
                )
                out_of_bounds_count += int(out_of_bounds)
                shape_type = self._shape_type(shape)
                readable = True
                if shape.shape_type in {
                    MSO_SHAPE_TYPE.PICTURE,
                    MSO_SHAPE_TYPE.LINKED_PICTURE,
                }:
                    image_count += 1
                    readable = self._image_blob_readable(shape)
                    invalid_image_count += int(not readable)
                    slide_nonempty = slide_nonempty or readable
                elif getattr(shape, "has_text_frame", False):
                    slide_nonempty = slide_nonempty or bool(shape.text.strip())
                elif not getattr(shape, "is_placeholder", False):
                    slide_nonempty = True
                objects.append(
                    {
                        "type": shape_type,
                        "bounds_emu": bounds,
                        "out_of_bounds": out_of_bounds,
                        "readable": readable,
                    }
                )
            empty_count += int(not slide_nonempty)
            slides.append(
                {
                    "index": slide_index,
                    "object_count": len(objects),
                    "empty": not slide_nonempty,
                    "objects": objects,
                }
            )

        layout = {
            "version": self.VERSION,
            "slide_count": len(slides),
            "slide_size_emu": {"width": width, "height": height},
            "slides": slides,
        }
        return layout, {
            "slide_count": len(slides),
            "page_count_matches": True,
            "nonempty_slide_count": len(slides) - empty_count,
            "empty_slide_count": empty_count,
            "object_count": object_count,
            "out_of_bounds_count": out_of_bounds_count,
            "image_count": image_count,
            "invalid_image_count": invalid_image_count,
        }

    @staticmethod
    def _shape_type(shape) -> str:
        if shape.shape_type in {
            MSO_SHAPE_TYPE.PICTURE,
            MSO_SHAPE_TYPE.LINKED_PICTURE,
        }:
            return "image"
        if getattr(shape, "has_table", False):
            return "table"
        if getattr(shape, "has_chart", False):
            return "chart"
        if getattr(shape, "has_text_frame", False):
            return "text"
        return "shape"

    @staticmethod
    def _image_blob_readable(shape) -> bool:
        try:
            blob = shape.image.blob
            if not blob:
                return False
            with Image.open(BytesIO(blob)) as image:
                image.verify()
            return True
        except (AttributeError, OSError, ValueError):
            return False

    def _render_preview(self, path: Path) -> dict:
        libreoffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not libreoffice:
            return {
                "summary": {
                    "preview_status": "unavailable",
                    "preview_page_count": 0,
                },
                "warnings": ["LibreOffice 不可用，已降级为 PPTX 结构检查"],
            }
        with tempfile.TemporaryDirectory(prefix="ppt-qa-preview-") as temp_dir:
            directory = Path(temp_dir)
            try:
                completed = subprocess.run(
                    [
                        libreoffice,
                        "--headless",
                        f"-env:UserInstallation={(directory / 'lo-profile').as_uri()}",
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        str(directory),
                        str(path.resolve()),
                    ],
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=self.PREVIEW_TIMEOUT_SECONDS,
                    check=False,
                )
            except (OSError, subprocess.TimeoutExpired):
                return self._preview_fallback("LibreOffice 预览失败，已降级为结构检查")
            pdf_path = directory / f"{path.stem}.pdf"
            if completed.returncode != 0 or not pdf_path.is_file() or pdf_path.stat().st_size == 0:
                return self._preview_fallback("LibreOffice 预览失败，已降级为结构检查")
            pdftoppm = shutil.which("pdftoppm")
            if not pdftoppm:
                return {
                    "summary": {
                        "preview_status": "pdf",
                        "preview_page_count": 0,
                    },
                    "warnings": ["PNG 预览工具不可用，已完成 PDF 预览检查"],
                }
            try:
                rendered = subprocess.run(
                    [
                        pdftoppm,
                        "-png",
                        "-r",
                        "96",
                        str(pdf_path),
                        str(directory / "slide"),
                    ],
                    capture_output=True,
                    timeout=self.PREVIEW_TIMEOUT_SECONDS,
                    check=False,
                )
            except (OSError, subprocess.TimeoutExpired):
                return self._preview_fallback("PNG 预览失败，已保留结构检查结果")
            pages = sorted(directory.glob("slide-*.png"))
            valid_pages = sum(self._preview_image_readable(item) for item in pages)
            if rendered.returncode != 0 or not pages or valid_pages != len(pages):
                return self._preview_fallback("PNG 预览失败，已保留结构检查结果")
            return {
                "summary": {
                    "preview_status": "png",
                    "preview_page_count": valid_pages,
                },
                "warnings": [],
            }

    @staticmethod
    def _preview_image_readable(path: Path) -> bool:
        try:
            with Image.open(path) as image:
                image.verify()
            return path.stat().st_size > 0
        except (OSError, ValueError):
            return False

    @staticmethod
    def _preview_fallback(message: str) -> dict:
        return {
            "summary": {
                "preview_status": "failed",
                "preview_page_count": 0,
            },
            "warnings": [message],
        }

    @staticmethod
    def _merge_layout(structural: dict, supplied: dict | None) -> dict:
        if not supplied or not isinstance(supplied.get("slides"), list):
            return structural
        merged = dict(structural)
        merged["renderer_layout"] = {
            "version": supplied.get("version"),
            "slide_count": supplied.get("slide_count"),
            "node_count": supplied.get("node_count"),
            "slides": supplied["slides"],
        }
        return merged

    @staticmethod
    def _raster_summary(layout: dict) -> dict:
        renderer = layout.get("renderer_layout", {})
        raster_count = 0
        readable_count = 0
        unreadable_count = 0
        raster_area = 0.0
        slide_area = 0.0
        for slide in renderer.get("slides", []):
            for item in slide.get("objects", []):
                if item.get("type") != "raster":
                    continue
                raster_count += 1
                readable = item.get("readable", False)
                readable_count += int(readable)
                unreadable_count += int(not readable)
                bounds = item.get("bounds", {})
                raster_area += max(0.0, bounds.get("w", 0)) * max(
                    0.0, bounds.get("h", 0)
                )
            size = slide.get("slide_size", {})
            slide_area += max(0.0, size.get("width", 0)) * max(
                0.0, size.get("height", 0)
            )
        return {
            "raster_layer_count": raster_count,
            "raster_readable_count": readable_count,
            "unreadable_raster_count": unreadable_count,
            "raster_area_ratio": round(
                min(1.0, raster_area / slide_area), 4
            )
            if slide_area
            else 0.0,
        }

    def _editable_ratio(self, layout: dict) -> float:
        return round(1.0 - self._raster_summary(layout)["raster_area_ratio"], 4)

    @staticmethod
    def _base_summary(expected_slide_count: int | None) -> dict:
        return {
            "file_valid": False,
            "openable": False,
            "pptx_size_bytes": 0,
            "slide_count": 0,
            "expected_slide_count": expected_slide_count,
            "page_count_matches": expected_slide_count is None,
            "nonempty_slide_count": 0,
            "empty_slide_count": 0,
            "object_count": 0,
            "out_of_bounds_count": 0,
            "image_count": 0,
            "invalid_image_count": 0,
        }

    @staticmethod
    def _empty_layout() -> dict:
        return {
            "version": "1.0",
            "slide_count": 0,
            "slide_size_emu": {"width": 0, "height": 0},
            "slides": [],
        }

    @staticmethod
    def _write_artifacts(path: Path, layout: dict, report: dict) -> None:
        if not path.parent.is_dir():
            return
        for suffix, payload in (
            (".layout.json", layout),
            (".qa-report.json", report),
        ):
            artifact = path.with_suffix(suffix)
            temp_path: Path | None = None
            try:
                with tempfile.NamedTemporaryFile(
                    mode="w",
                    encoding="utf-8",
                    dir=artifact.parent,
                    prefix=f".{artifact.name}.",
                    suffix=".tmp",
                    delete=False,
                ) as stream:
                    json.dump(payload, stream, ensure_ascii=False, indent=2)
                    stream.flush()
                    os.fsync(stream.fileno())
                    temp_path = Path(stream.name)
                os.replace(temp_path, artifact)
            finally:
                if temp_path is not None:
                    temp_path.unlink(missing_ok=True)
