"""
PPT 生成引擎

遍历大纲 JSON，调用各页面类型的布局渲染器，生成最终的 .pptx 文件。
"""

import uuid
from pathlib import Path
from typing import Any, Dict, Optional

from loguru import logger
from pptx import Presentation
from pptx.util import Inches

from src.tools.ppt.theme import PPTTheme, get_theme, CANVAS_WIDTH, CANVAS_HEIGHT
from src.tools.ppt.layouts.cover import render_cover
from src.tools.ppt.layouts.toc import render_toc
from src.tools.ppt.layouts.section import render_section
from src.tools.ppt.layouts.content import render_content
from src.tools.ppt.layouts.chart import render_chart
from src.tools.ppt.layouts.summary import render_summary


class PPTGenerator:
    """PPT 生成核心引擎。"""

    def __init__(self, theme: Optional[PPTTheme] = None):
        self.theme = theme or get_theme()
        self.prs = Presentation()
        self.prs.slide_width = Inches(CANVAS_WIDTH)
        self.prs.slide_height = Inches(CANVAS_HEIGHT)

    def generate(self, plan: dict) -> str:
        """根据大纲生成 PPT，返回文件路径。"""
        slides = plan.get("slides", [])
        total = len(slides)

        if total == 0:
            raise ValueError("PPT 大纲中没有任何页面")

        for i, slide_data in enumerate(slides):
            slide_type = slide_data.get("type", "content")
            self._create_slide(slide_type, slide_data, i + 1, total)

        output_path = self._save(plan.get("title", "presentation"))
        logger.info(f"[PPTGenerator] 生成完成: {output_path}, 共 {total} 页")
        return output_path

    def _create_slide(self, slide_type: str, data: dict, index: int, total: int):
        """根据类型调用对应的布局渲染器。"""
        # 使用空白布局
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

        if slide_type == "cover":
            render_cover(slide, data, self.theme)
        elif slide_type == "toc":
            render_toc(slide, data, self.theme, index, total)
        elif slide_type == "section":
            render_section(slide, data, self.theme)
        elif slide_type == "summary":
            render_summary(slide, data, self.theme)
        elif slide_type == "content":
            layout = data.get("layout", "bullets")
            if layout == "chart":
                render_chart(slide, data, self.theme, index, total)
            else:
                render_content(slide, data, self.theme, index, total)
        else:
            # 未知类型按内容页处理
            render_content(slide, data, self.theme, index, total)

    def _save(self, title: str) -> str:
        """保存到文件，返回绝对路径。"""
        save_dir = self._get_output_dir()
        save_dir.mkdir(parents=True, exist_ok=True)

        # 清理文件名
        safe_name = "".join(c if c.isalnum() or c in "._- " else "_" for c in title)
        file_name = f"{safe_name}.pptx"
        output_path = save_dir / file_name

        self.prs.save(str(output_path))
        return str(output_path.absolute())

    def _get_output_dir(self) -> Path:
        """获取输出目录。"""
        try:
            from src.main import _get_tenant_upload_dir
            return _get_tenant_upload_dir()
        except (ImportError, AttributeError):
            return Path("storage/ppt")
