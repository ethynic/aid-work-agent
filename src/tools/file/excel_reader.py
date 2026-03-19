"""
Excel文档处理工具

提供Excel文档的读取功能，支持提取工作表、单元格、表格等数据
"""

from pathlib import Path
from typing import Any, Dict, List, Optional

from loguru import logger


class ExcelReader:
    """Excel文档读取器"""

    def __init__(self):
        """初始化Excel读取器"""
        self.openpyxl = None
        self._init_openpyxl()

    def _init_openpyxl(self):
        """初始化openpyxl库"""
        try:
            import openpyxl
            self.openpyxl = openpyxl
            logger.info("openpyxl库初始化成功")
        except ImportError:
            logger.error("openpyxl库未安装，请运行: pip install openpyxl")
            self.openpyxl = None

    def read_excel_document(self, file_path: str, sheet_name: Optional[str] = None) -> Dict[str, Any]:
        """
        读取Excel文档内容

        Args:
            file_path: Excel文档路径
            sheet_name: 工作表名称（可选），如果不指定则读取第一个工作表

        Returns:
            包含文档内容的字典
        """
        if self.openpyxl is None:
            return {
                "success": False,
                "error": "openpyxl库未安装，无法读取Excel文档"
            }

        try:
            path = Path(file_path)
            if not path.exists():
                return {
                    "success": False,
                    "error": f"文件不存在: {file_path}"
                }

            # 检查文件扩展名
            if path.suffix.lower() not in ['.xlsx', '.xls']:
                return {
                    "success": False,
                    "error": f"不支持的文件格式: {path.suffix}，仅支持.xlsx和.xls格式"
                }

            # 检查.xls格式（需要转换）
            if path.suffix.lower() == '.xls':
                return {
                    "success": False,
                    "error": "目前仅支持.xlsx格式，请将.xls文件转换为.xlsx格式"
                }

            # 读取Excel文档
            workbook = self.openpyxl.load_workbook(path, data_only=True)

            # 提取所有工作表信息
            sheet_names = workbook.sheetnames
            sheet_count = len(sheet_names)

            # 获取工作表信息
            sheets_info = []
            for sheet_name_item in sheet_names:
                sheet = workbook[sheet_name_item]
                sheets_info.append({
                    "name": sheet_name_item,
                    "title": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "row_count": sheet.max_row,
                    "column_count": sheet.max_column
                })

            # 读取指定工作表或第一个工作表
            target_sheet_name = sheet_name if sheet_name else sheet_names[0]
            if target_sheet_name not in sheet_names:
                return {
                    "success": False,
                    "error": f"工作表不存在: {target_sheet_name}，可用工作表: {', '.join(sheet_names)}"
                }

            sheet_data = self._extract_sheet_data(workbook[target_sheet_name])

            # 提取文档信息
            doc_info = self._extract_document_info(workbook, path)

            # 合并所有工作表内容（用于完整内容输出）
            all_sheets_content = []
            for sheet_name_item in sheet_names:
                sheet = workbook[sheet_name_item]
                sheet_text = self._format_sheet_to_text(sheet)
                all_sheets_content.append(f"=== 工作表: {sheet_name_item} ===\n{sheet_text}")

            full_content = '\n\n'.join(all_sheets_content)
            total_lines = len(full_content.splitlines())

            workbook.close()

            logger.info(f"成功读取Excel文档: {path} (工作表: {sheet_count})")

            return {
                "success": True,
                "message": "成功读取Excel文档",
                "file_path": str(path),
                "file_type": "xlsx",
                "sheet_count": sheet_count,
                "sheet_names": sheet_names,
                "sheets_info": sheets_info,
                "current_sheet": target_sheet_name,
                "sheet_data": sheet_data,
                "document_info": doc_info,
                "content": full_content,
                "total_lines": total_lines
            }

        except Exception as e:
            logger.error(f"读取Excel文档失败: {e}")
            return {
                "success": False,
                "error": f"读取Excel文档失败: {str(e)}"
            }

    def _extract_sheet_data(self, sheet) -> Dict[str, Any]:
        """
        提取工作表数据

        Args:
            sheet: openpyxl的Worksheet对象

        Returns:
            工作表数据字典
        """
        data = {
            "name": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "rows": [],
            "headers": [],
            "data": []
        }

        # 获取第一行作为表头
        if sheet.max_row >= 1:
            headers = []
            for col in range(1, sheet.max_column + 1):
                cell_value = sheet.cell(1, col).value
                headers.append(str(cell_value) if cell_value is not None else "")
            data["headers"] = headers

        # 获取所有数据行
        rows = []
        for row in range(1, sheet.max_row + 1):
            row_data = []
            for col in range(1, sheet.max_column + 1):
                cell_value = sheet.cell(row, col).value
                row_data.append(cell_value)
            rows.append(row_data)

        data["rows"] = rows

        # 提取数据（跳过表头）
        if sheet.max_row > 1:
            data["data"] = rows[1:]
        else:
            data["data"] = []

        return data

    def _format_sheet_to_text(self, sheet) -> str:
        """
        将工作表格式化为文本

        Args:
            sheet: openpyxl的Worksheet对象

        Returns:
            格式化的文本内容
        """
        text_lines = []

        if sheet.max_row == 0:
            return "(空工作表)"

        # 获取每列的最大宽度（用于格式化）
        col_widths = []
        for col in range(1, sheet.max_column + 1):
            max_width = 0
            for row in range(1, sheet.max_row + 1):
                cell_value = sheet.cell(row, col).value
                cell_str = str(cell_value) if cell_value is not None else ""
                max_width = max(max_width, len(cell_str))
            col_widths.append(min(max_width + 2, 50))  # 限制最大宽度

        # 格式化每一行
        for row in range(1, sheet.max_row + 1):
            row_parts = []
            for col in range(1, sheet.max_column + 1):
                cell_value = sheet.cell(row, col).value
                cell_str = str(cell_value) if cell_value is not None else ""
                width = col_widths[col - 1]
                row_parts.append(cell_str.ljust(width))
            text_lines.append(" | ".join(row_parts))

        return '\n'.join(text_lines)

    def _extract_document_info(self, workbook, path: Path) -> Dict[str, Any]:
        """
        提取文档元信息

        Args:
            workbook: openpyxl的Workbook对象
            path: 文件路径

        Returns:
            文档信息字典
        """
        info = {
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "sheet_count": len(workbook.sheetnames),
            "active_sheet": workbook.active.title if workbook.active else None,
        }

        try:
            # 尝试获取工作簿属性（如果有的话）
            if hasattr(workbook, 'properties'):
                props = workbook.properties
                if props:
                    info.update({
                        "title": props.title or "",
                        "author": props.creator or "",
                        "created": str(props.created) if props.created else "",
                        "modified": str(props.modified) if props.modified else "",
                    })
        except Exception as e:
            logger.warning(f"提取文档信息失败: {e}")

        return info


def is_excel_document(file_path: str) -> bool:
    """
    判断文件是否为Excel文档

    Args:
        file_path: 文件路径

    Returns:
        是否为Excel文档
    """
    path = Path(file_path)
    return path.suffix.lower() in ['.xlsx', '.xls']
