"""验证：拉丁/数字→Noto Sans、CJK→思源黑体的双轨字体改造。

覆盖：字体文件存在且可加载；remap_typeface 按标签分流；native deck 的
latin/cs=Noto Sans、ea=思源黑体；混合文本（44.63万）测量与防溢出仍准确；
embed 同时嵌入两款字体；预览 PNG 无报错；路线一仍可生成。
"""
import os
import sys
import zipfile
import tempfile

_HERE = os.path.dirname(os.path.abspath(__file__))


def _bases():
    """候选根目录：脚本目录及其上级 + CWD。

    覆盖两种布局：workspace（pptx_gen/verify_font_split.py，资源在
    pptx_gen/assets/）与打包/安装后的技能库（verify/verify_font_split.py，
    模板在 assets/、demo 内容与配图在 examples/）。
    """
    out = [_HERE, os.path.dirname(_HERE), os.getcwd()]
    seen, uniq = set(), []
    for b in out:
        if b and b not in seen:
            seen.add(b)
            uniq.append(b)
    return uniq


def _find(name):
    """在候选根的各资源子目录下找文件，返回首个命中（找不到返回 None）。"""
    for b in _bases():
        for sub in ("", "assets", "examples/assets", "examples"):
            p = os.path.join(b, sub, name)
            if os.path.exists(p):
                return p
    return None


# sys.path：指向包含 pptx_gen 包的目录（必须先于 import pptx_gen）
for _b in _bases():
    if os.path.isdir(os.path.join(_b, "pptx_gen")):
        if _b not in sys.path:
            sys.path.insert(0, _b)
        break

from pptx import Presentation
from pptx.oxml.ns import qn

import pptx_gen
from pptx_gen import fonts, native, generate_from_vi, preview_pngs, generate
from pptx_gen.fonts import (FONT_NAME, FONT_NAME_LATIN, font_file,
                            latin_font_file, is_cjk, remap_typeface,
                            remap_presentation, embed_fonts)

TPL = _find("demo_template.pptx")
CONTENT = _find("demo_content.json")
CHART = _find("chart.png")
LINE_CHART = _find("chart_line.png")
PANELS = [_find("panel_%s.png" % s) for s in "abc"]
# generate_from_vi / preview_pngs 用 base_dir 解析 content 里的相对图片路径
IMG_DIR = os.path.dirname(CHART) if CHART else None
for _n, _p in (("demo_template.pptx", TPL), ("demo_content.json", CONTENT),
               ("chart.png", CHART)):
    if not _p:
        print("FAIL missing demo asset: %s" % _n)
        sys.exit(1)


def _content():
    """复用 demo 内容，但保证含一个 stats 页（数字+万）用于校验。"""
    import json
    with open(CONTENT, encoding="utf-8") as fh:
        c = json.load(fh)
    has_stats = any(s.get("role") == "stats" for s in c["slides"])
    if not has_stats:
        c["slides"].insert(1, {
            "role": "stats", "title": "经营概览",
            "stats": [
                {"num": "44.63万", "label": "活跃客户", "desc": "同比增长 12.4%"},
                {"num": "1,234.56", "label": "成交额(亿)", "desc": "环比 +3.1%"},
                {"num": "88%", "label": "续约率", "desc": "目标 85%"},
            ],
        })
    for s in c["slides"]:
        if s.get("role") == "image_text":
            s["image"] = "chart_line.png"
        if s.get("role") == "image_grid":
            s["images"] = [{"src": "panel_a.png"}, {"src": "panel_b.png"},
                           {"src": "panel_c.png"}]
        if s.get("role") == "image":
            s["image"] = "chart.png"
    return c


fails = []


def check(cond, msg):
    print(("PASS" if cond else "FAIL"), msg)
    if not cond:
        fails.append(msg)


# 1) 字体文件存在且可加载
for p in [font_file(False), font_file(True), latin_font_file(False),
          latin_font_file(True)]:
    check(p and os.path.exists(p), "bundled font exists: %s" % p)
    if p:
        try:
            from PIL import ImageFont
            ImageFont.truetype(p, 40)
            ok = True
        except Exception as e:
            ok = False
            print("   load error:", e)
        check(ok, "PIL can load: %s" % os.path.basename(p))

# 2) is_cjk 分类
check(is_cjk("万") and is_cjk("中") and not is_cjk("4") and not is_cjk("A"),
      "is_cjk classifies CJK vs Latin")

# 3) remap_typeface 按标签分流
chk_lat = remap_typeface("Calibri", tag="a:latin")
chk_ea = remap_typeface("微软雅黑", tag="a:ea")
check(chk_lat[0] == FONT_NAME_LATIN, "latin 'Calibri' -> %s (got %s)" %
      (FONT_NAME_LATIN, chk_lat[0]))
check(chk_ea[0] == FONT_NAME, "ea '微软雅黑' -> %s (got %s)" %
      (FONT_NAME, chk_ea[0]))
# 免费拉丁字体保留
check(remap_typeface("Noto Sans", tag="a:latin")[0] is None,
      "free 'Noto Sans' kept in latin slot")
check(remap_typeface("Source Han Sans CN", tag="a:ea")[0] is None,
      "free 'Source Han Sans CN' kept in ea slot")

# 4) 混合文本测量：数字用 Noto Sans，CJK 用思源
w = native.measure("44.63万", 104, True)
print("   measure('44.63万',104,bold)=", round(w, 2), "px")
check(200 < w < 460, "mixed '44.63万' width in sane range")

# 5) 生成 native deck，检查 typeface 写法
tmpdir = tempfile.mkdtemp()
out = os.path.join(tmpdir, "test.pptx")
vi = pptx_gen.extract_vi(TPL)
res = generate_from_vi(TPL, _content(), out,
                       vi_override={"brand": "#14395F", "accent": "#FFB545",
                                    "dark": "#0A1E33"},
                       base_dir=IMG_DIR)
check(os.path.exists(out), "native deck generated: %s" % out)

prs = Presentation(out)
latins, eas, css = set(), set(), set()
for slide in prs.slides:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                rPr = r._r.find(qn("a:rPr"))
                if rPr is not None:
                    for tag, coll in (("a:latin", latins), ("a:ea", eas),
                                      ("a:cs", css)):
                        e = rPr.find(qn(tag))
                        if e is not None:
                            coll.add(e.get("typeface"))
check(latins <= {FONT_NAME_LATIN}, "all latin slots = %s (got %s)" %
      (FONT_NAME_LATIN, sorted(latins)))
check(eas <= {FONT_NAME}, "all ea slots = %s (got %s)" % (FONT_NAME, sorted(eas)))
check(css <= {FONT_NAME_LATIN}, "all cs slots = %s (got %s)" %
      (FONT_NAME_LATIN, sorted(css)))

# 6) 防溢出：stats 页大数字单行应被 fit 到框内（est_w <= 框宽*0.92 预算）
#    直接复用 native 的 fit_size 逻辑验证 44.63万 在 cw-52 宽内单行
cw = 420
fit = native.fit_size("44.63万", 104, cw - 52, h=(cw - 52) * 0.92,
                      max_lines=1, bold=True, safety=0.92)
fit_w = native.est_w("44.63万", fit, bold=True)
check(fit_w <= (cw - 52) * 0.92 + 1,
      "fit '44.63万' single-line within budget (size=%d, w=%.1f<=%.1f)"
      % (fit, fit_w, (cw - 52) * 0.92))

# 7) embed 同时嵌入两款字体
emb = embed_fonts(out, out)
print("   embed:", emb)
check(emb.get("embedded"), "fonts embedded")
zf = zipfile.ZipFile(out)
fnt = [n for n in zf.namelist() if n.endswith(".fntdata")]
check(len(fnt) >= 2, "embedded >=2 font parts (got %d)" % len(fnt))
pres = zf.read("ppt/presentation.xml").decode("utf-8")
check(FONT_NAME in pres and FONT_NAME_LATIN in pres,
      "embeddedFontLst lists both %s and %s" % (FONT_NAME, FONT_NAME_LATIN))

# 8) 预览 PNG 无报错
prev = os.path.join(tmpdir, "preview")
pngs = preview_pngs(TPL, _content(), prev, base_dir=IMG_DIR)
check(len(pngs) > 0 and all(os.path.exists(p) for p in pngs),
      "preview_pngs produced %d pngs" % len(pngs))

# 9) 路线一仍可生成
spec = {"slides": [{"mode": "copy", "template_slide_index": 0,
                    "replacements": {}}], "keep_template_slides": True}
r1 = generate(TPL, spec, os.path.join(tmpdir, "r1.pptx"))
check(os.path.exists(r1["out"]), "route-1 generate works")

print("\n==== %s ====" % ("ALL PASS" if not fails else "FAILURES: %d" % len(fails)))
for f in fails:
    print("  -", f)
sys.exit(1 if fails else 0)
