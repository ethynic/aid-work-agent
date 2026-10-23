#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""商用免费字体策略验证（verify_font_policy.py）。

断言（双轨字体版），缺一不可：
  1. 包内字体文件可解析（font_file→思源 OTF、latin_font_file→Noto Sans TTF，
     必须返回 assets/fonts/ 内文件——防止静默回退系统字体导致"测得准≠渲染
     一致"的分叉）；
  2. measure() 与"逐字符用对应脚本字体"的直接测量一致（字体活性证明）；
  3. 路线二 deck：latin/cs 全是 FONT_NAME_LATIN、ea 全是 FONT_NAME；
  4. 路线一 deck（模板含 Calibri/主题引用）经 remap 后同样全免费——
     含主题 a:font script=... 回退、a:buFont、空 a:ea，全 deck 任何
     typeface 属性都不得落在免费集合之外；
  5. embed_fonts 嵌入后：两款字体各有 embeddedFont 条目、fntdata 部件齐全、
     rels 可解析、python-pptx 能重开、幂等。

本脚本位置无关：可从 workspace 根目录跑，也可从打包后的技能 verify/ 目录跑
（自动定位 pptx_gen 包与 demo 资源）。
"""
import json
import os
import sys
import tempfile
import zipfile

_HERE = os.path.dirname(os.path.abspath(__file__))


def _bases():
    out = [_HERE, os.path.dirname(_HERE), os.getcwd()]
    seen, uniq = set(), []
    for b in out:
        if b and b not in seen:
            seen.add(b)
            uniq.append(b)
    return uniq


def _find(name):
    for b in _bases():
        for sub in ("", "assets", "examples/assets", "examples",
                    "pptx_gen/assets"):
            p = os.path.join(b, sub, name)
            if os.path.exists(p):
                return p
    return None


for _b in _bases():
    if os.path.isdir(os.path.join(_b, "pptx_gen")):
        if _b not in sys.path:
            sys.path.insert(0, _b)
        break

from lxml import etree
from pptx import Presentation

from pptx_gen import (FONT_NAME, FONT_NAME_LATIN, analyze, embed_fonts,
                      extract_vi, font_file, generate, generate_from_vi,
                      latin_font_file, remap_presentation)
from pptx_gen.fonts import is_cjk
from pptx_gen.native import measure
from PIL import ImageFont

TPL = _find("demo_template.pptx")
if not TPL:
    print("FAIL missing demo_template.pptx")
    sys.exit(1)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
fails = []


def check(name, cond, detail=""):
    print(("[OK ] " if cond else "[FAIL] ") + name + ((" | " + detail) if detail else ""))
    if not cond:
        fails.append(name)


# ---- 1) 包内字体文件 -----------------------------------------------------
reg, bold = font_file(False), font_file(True)
lat_reg, lat_bold = latin_font_file(False), latin_font_file(True)
check("bundled regular OTF resolves", bool(reg) and reg.endswith("SourceHanSansCN-Regular.otf"), str(reg))
check("bundled bold OTF resolves", bool(bold) and bold.endswith("SourceHanSansCN-Bold.otf"), str(bold))
check("bundled latin regular TTF resolves", bool(lat_reg) and lat_reg.endswith("NotoSans-Regular.ttf"), str(lat_reg))
check("bundled latin bold TTF resolves", bool(lat_bold) and lat_bold.endswith("NotoSans-Bold.ttf"), str(lat_bold))
check("OFL license shipped", os.path.exists(os.path.join(os.path.dirname(reg), "OFL-1.1.txt")))
check("FONT_NAME is free family", FONT_NAME in ("Source Han Sans CN",), FONT_NAME)
check("FONT_NAME_LATIN is free family", FONT_NAME_LATIN in ("Noto Sans",), FONT_NAME_LATIN)

# ---- 2) measure() 字体活性（按字符脚本分流）--------------------------------
def _mixed_expect(s, size, use_bold):
    cjk = ImageFont.truetype(bold if use_bold else reg, size)
    lat = ImageFont.truetype(lat_bold if use_bold else lat_reg, size)
    return sum((cjk if is_cjk(c) else lat).getlength(c) for c in s)


exp = _mixed_expect("44.63万", 82, False)
check("measure() == per-char mixed measurement",
      abs(measure("44.63万", 82) - exp) < 0.01,
      "measure=%s expect=%s" % (measure("44.63万", 82), exp))
check("bold measured wider than regular", measure("44.63万", 82, True) > measure("44.63万", 82))
lat_only = ImageFont.truetype(lat_reg, 82).getlength("44.63")
check("pure digits measured with Noto Sans",
      abs(measure("44.63", 82) - lat_only) < 0.01,
      "measure=%s noto=%s" % (measure("44.63", 82), lat_only))

# ---- 3) 路线二 deck 全免费 ------------------------------------------------
tmp = tempfile.mkdtemp(prefix="font_policy_")
tpl = TPL
CONTENT = _find("demo_content.json")
content = json.load(open(CONTENT, encoding="utf-8"))
o2 = os.path.join(tmp, "r2.pptx")
rep2 = generate_from_vi(tpl, content, o2)
fonts_in_deck = set()
with zipfile.ZipFile(o2) as z:
    for n in z.namelist():
        if n.endswith(".xml") and ("slides/" in n or "theme" in n or "Master" in n or "Layout" in n):
            root = etree.fromstring(z.read(n))
            for el in root.iter():
                if el.tag in ("{%s}latin" % A, "{%s}ea" % A, "{%s}cs" % A):
                    tf = el.get("typeface")
                    if tf and not tf.startswith("+"):
                        fonts_in_deck.add(tf)
check("route2 deck: only free dual-track fonts",
      fonts_in_deck == {FONT_NAME, FONT_NAME_LATIN}, str(sorted(fonts_in_deck)))
# 全 deck 任何 typeface 属性（含主题里 a:font script=... 的按脚本回退字体）
# 都不得出现商业字体——宋体/ＭＳ Ｐゴシック 这类模板默认值同样是授权风险
all_tf = set()
with zipfile.ZipFile(o2) as z:
    for n in z.namelist():
        if n.endswith(".xml"):
            for el in etree.fromstring(z.read(n)).iter():
                tf = el.get("typeface")
                if tf and not tf.startswith("+"):
                    all_tf.add(tf)
check("route2 NO typeface outside free set (incl. script fallbacks)",
      all_tf <= {FONT_NAME, FONT_NAME_LATIN}, str(sorted(all_tf)))
# 逐标签核对：latin/cs → Noto Sans，ea → 思源
lat_in_deck, ea_in_deck = set(), set()
with zipfile.ZipFile(o2) as z:
    for n in z.namelist():
        if n.endswith(".xml") and "slides/" in n:
            root = etree.fromstring(z.read(n))
            for el in root.iter():
                if el.tag == "{%s}latin" % A and el.get("typeface") and not el.get("typeface").startswith("+"):
                    lat_in_deck.add(el.get("typeface"))
                if el.tag == "{%s}ea" % A and el.get("typeface") and not el.get("typeface").startswith("+"):
                    ea_in_deck.add(el.get("typeface"))
check("route2 latin slots == FONT_NAME_LATIN", lat_in_deck == {FONT_NAME_LATIN}, str(sorted(lat_in_deck)))
check("route2 ea slots == FONT_NAME", ea_in_deck == {FONT_NAME}, str(sorted(ea_in_deck)))
check("route2 font_report present", isinstance(rep2.get("fonts"), dict))

# ---- 4) 路线一 deck remap 后全免费 ----------------------------------------
spec = {"slides": [
    {"mode": "copy", "template_slide_index": 0,
     "replacements": {"客户名称 · 项目汇报": "字体合规 · 路线一"}},
    {"mode": "placeholder", "layout_index": 1,
     "fields": {"0": "商用免费字体", "1": ["思源黑体 CN（OFL）", "零授权风险"]}},
], "keep_template_slides": False}
o1 = os.path.join(tmp, "r1.pptx")
rep1 = generate(tpl, spec, o1)
fonts_r1 = set()
with zipfile.ZipFile(o1) as z:
    for n in z.namelist():
        if n.endswith(".xml"):
            root = etree.fromstring(z.read(n))
            for el in root.iter():
                if el.tag in ("{%s}latin" % A, "{%s}ea" % A, "{%s}cs" % A):
                    tf = el.get("typeface")
                    if tf and not tf.startswith("+"):
                        fonts_r1.add(tf)
check("route1 deck: only free dual-track fonts",
      fonts_r1 <= {FONT_NAME, FONT_NAME_LATIN}, str(sorted(fonts_r1)))
# 全 deck 任何 typeface 属性都不得出现商业字体（含主题 a:font script=... 回退）
all_tf_r1 = set()
with zipfile.ZipFile(o1) as z:
    for n in z.namelist():
        if n.endswith(".xml"):
            for el in etree.fromstring(z.read(n)).iter():
                tf = el.get("typeface")
                if tf and not tf.startswith("+"):
                    all_tf_r1.add(tf)
check("route1 NO typeface outside free set (incl. script fallbacks)",
      all_tf_r1 <= {FONT_NAME, FONT_NAME_LATIN}, str(sorted(all_tf_r1)))
check("route1 remap report recorded", bool(rep1.get("fonts", {}).get("replaced")),
      json.dumps(rep1.get("fonts", {}).get("replaced", {}), ensure_ascii=False))
# 按脚本标签分流：模板主题里 a:latin 的 Calibri 必须映射到 Noto Sans（不是
# 中文字体）；a:ea 的空值必须被补成思源黑体（否则继承主题的中文落查看器默认字体）
# 注意：+mj-lt/+mn-ea 这类是主题引用，解析到主题字体（现已合规），不算违规
latin_slots, ea_slots = set(), set()
with zipfile.ZipFile(o1) as z:
    for n in z.namelist():
        if n.endswith(".xml"):
            for el in etree.fromstring(z.read(n)).iter():
                tf = el.get("typeface")
                if not tf or tf.startswith("+"):
                    continue
                if el.tag == "{%s}latin" % A:
                    latin_slots.add(tf)
                elif el.tag == "{%s}ea" % A:
                    ea_slots.add(tf)
check("route1 latin slots all Noto Sans", latin_slots <= {FONT_NAME_LATIN},
      str(sorted(latin_slots)))
check("route1 ea slots all Source Han", ea_slots <= {FONT_NAME},
      str(sorted(ea_slots)))
cal = rep1.get("fonts", {}).get("replaced", {}).get("Calibri")
check("route1 theme Calibri(a:latin) -> Noto Sans",
      cal is not None and cal[0] == FONT_NAME_LATIN, str(cal))
check("route1 empty a:ea filled", bool(rep1.get("fonts", {}).get("filled")),
      json.dumps(rep1.get("fonts", {}).get("filled", {}), ensure_ascii=False))
# remap_fonts=False 时保留模板原字体（客户自有授权场景）
o1k = os.path.join(tmp, "r1_keep.pptx")
generate(tpl, spec, o1k, remap_fonts=False)
prs = Presentation(o1k)
check("route1 remap_fonts=False still opens", len(prs.slides._sldIdLst) == 2)

# ---- 5) embed_fonts 嵌入 ---------------------------------------------------
o2e = os.path.join(tmp, "r2_embed.pptx")
import shutil
shutil.copy2(o2, o2e)
emb = embed_fonts(o2e)
check("embed_fonts reports embedded", emb.get("embedded") is True, str(emb))
with zipfile.ZipFile(o2e) as z:
    names = set(z.namelist())
    ct = z.read("[Content_Types].xml").decode("utf-8")
    pres = etree.fromstring(z.read("ppt/presentation.xml"))
    rels = etree.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
check("fntdata parts present (both fonts x regular/bold)",
      {"ppt/fonts/font1.fntdata", "ppt/fonts/font2.fntdata",
       "ppt/fonts/font3.fntdata", "ppt/fonts/font4.fntdata"} <= names,
      str(sorted(n for n in names if n.endswith(".fntdata"))))
check("content-types has fntdata", 'Extension="fntdata"' in ct)
eflst = pres.find("{%s}embeddedFontLst" % P)
check("presentation has embeddedFontLst", eflst is not None)
tfs = {ef.find("{%s}font" % P).get("typeface")
       for ef in eflst} if eflst is not None else set()
check("embeddedFontLst lists both fonts",
      tfs == {FONT_NAME, FONT_NAME_LATIN}, str(sorted(tfs)))
check("presentation embedTrueTypeFonts=1", pres.get("embedTrueTypeFonts") == "1")
rids = set()
for ef in (eflst or []):
    for child in ef:
        rid = child.get("{%s}id" % R)
        if rid:
            rids.add(rid)
rel_ids = {r.get("Id") for r in rels}
check("embedded font rels resolve", rids and rids <= rel_ids, str(rids - rel_ids))
check("font rels target fntdata", all(("fonts/font" in (r.get("Target") or "")) for r in rels if r.get("Id") in rids))
prs2 = Presentation(o2e)
check("embedded deck reopens with python-pptx", len(prs2.slides._sldIdLst) == 15)
# 幂等
embed_fonts(o2e)
check("embed_fonts idempotent", zipfile.ZipFile(o2e).namelist().count("ppt/fonts/font1.fntdata") == 1)

print()
if fails:
    print("FAILURES:", len(fails))
    for x in fails:
        print(" -", x)
    sys.exit(1)
print("ALL FONT POLICY CHECKS PASSED")
