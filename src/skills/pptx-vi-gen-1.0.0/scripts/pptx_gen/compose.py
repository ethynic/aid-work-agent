"""Compose per-page PNGs into a 16:9 .pptx (full-bleed image slides)."""
import os

from pptx import Presentation
from pptx.util import Emu

SLIDE_W = Emu(12192000)  # 13.333 in  (16:9)
SLIDE_H = Emu(6858000)   # 7.5 in


def compose(png_paths, out_path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    for png in png_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(png, 0, 0, width=SLIDE_W, height=SLIDE_H)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    return out_path
