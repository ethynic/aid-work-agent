"""Extract a customer's visual identity from their .pptx template.

Pulls theme colors + theme fonts + a logo image from the template, filling
the rest from a sane default so the renderer always has a complete VI.

Font policy: fonts are NEVER taken from the customer template for rendering —
only the commercially-free bundled font (see fonts.py) is used, so generated
decks carry zero font-licensing risk. Template theme fonts are still read but
remapped to their free equivalent for the audit trail.
"""
import glob
import io

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

from . import fonts as _fonts

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def default_vi():
    return {
        "brand": "#1F4E79",
        "accent": "#2E9CCA",
        "dark": "#10263D",
        "light": "#F5F7FA",
        "ink": "#1A1A1A",
        "muted": "#6B7280",
        # 商用免费双轨字体（OFL）：CJK→思源黑体，拉丁→Noto Sans
        "font_title": _fonts.font_file(bold=True),
        "font_body": _fonts.font_file(bold=False),
        "font_latin_title": _fonts.latin_font_file(bold=True),
        "font_latin_body": _fonts.latin_font_file(bold=False),
        "logo_img": None,
    }


def _first_picture(container):
    try:
        for shp in container.shapes:
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                return shp
    except Exception:
        return None
    return None


def _logo_image(prs):
    for container in [prs.slide_masters[0]] + list(prs.slides):
        pic = _first_picture(container)
        if pic is not None:
            try:
                return Image.open(io.BytesIO(pic.image.blob)).convert("RGBA")
            except Exception:
                continue
    return None


def extract_vi(template_path):
    vi = default_vi()
    prs = Presentation(template_path)

    try:
        theme_part = prs.slide_masters[0].part.part_related_by(RT.THEME)
        root = etree.fromstring(theme_part.blob)

        scheme = root.find(".//a:clrScheme", NS)
        if scheme is not None:
            def col(tag):
                el = scheme.find(f"a:{tag}/a:srgbClr", NS)
                if el is not None and el.get("val"):
                    return "#" + el.get("val")
                # sysClr 的实际颜色在 lastClr（val 只是系统色名如 windowText）
                el = scheme.find(f"a:{tag}/a:sysClr", NS)
                return ("#" + el.get("lastClr")) if (el is not None and el.get("lastClr")) else None
            vi["brand"] = col("accent1") or vi["brand"]
            vi["accent"] = col("accent2") or col("accent1") or vi["accent"]
            vi["dark"] = col("dk2") or col("dk1") or vi["dark"]

        fs = root.find(".//a:fontScheme", NS)
        if fs is not None:
            mj = fs.find(".//a:majorFont/a:latin", NS)
            mn = fs.find(".//a:minorFont/a:latin", NS)
            # 模板主题字体只作审计记录（theme_latin_*），绝不写入
            # font_latin_* —— 那两个键必须是"字体文件路径"（预览渲染用），
            # 曾经把映射后的字体"名"写进去导致 PIL 回退成微型点阵字体。
            for el, key in ((mj, "theme_latin_title"),
                            (mn, "theme_latin_body")):
                if el is not None:
                    tf = el.get("typeface")
                    # 主题字体是 latin 标签：映射为免费拉丁字体（Noto Sans）
                    new, _why = _fonts.remap_typeface(tf, tag="a:latin",
                                                      strict=True)
                    vi[key] = new or _fonts.FONT_NAME_LATIN
    except Exception:
        pass

    vi["logo_img"] = _logo_image(prs)
    return vi


if __name__ == "__main__":
    import json
    import sys
    v = extract_vi(sys.argv[1])
    v.pop("logo_img", None)
    print(json.dumps({k: (os.path.basename(val) if k.startswith("font_") and val else val)
                      for k, val in v.items()}, ensure_ascii=False, indent=2))
