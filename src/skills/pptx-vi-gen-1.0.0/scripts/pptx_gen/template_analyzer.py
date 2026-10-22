"""Introspect a template so an upper-layer agent/LLM can decide routing.

Produces two JSON-serializable views:
  - analyze_layouts(): every slide layout with its placeholders (idx/type)
  - analyze_slides(): every template slide with its text shapes + pictures

The per-slide text map doubles as the replacement-key dictionary for the
deep-copy route, and the placeholder list drives the placeholder route.
"""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _ph_type(ph):
    try:
        return str(ph.placeholder_format.type).split(" ")[0].split(".")[-1]
    except Exception:
        return "?"


def analyze_layouts(template_path):
    prs = Presentation(template_path)
    layouts = []
    for i, layout in enumerate(prs.slide_layouts):
        phs = []
        for ph in layout.placeholders:
            phs.append({
                "idx": ph.placeholder_format.idx,
                "type": _ph_type(ph),
                "name": ph.name,
                "has_text": bool(ph.has_text_frame),
            })
        layouts.append({"index": i, "name": layout.name, "placeholders": phs})
    return layouts


def analyze_slides(template_path):
    prs = Presentation(template_path)
    slides = []
    for i, slide in enumerate(prs.slides):
        texts, has_picture = [], False
        for shp in slide.shapes:
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                has_picture = True
            if shp.has_text_frame and shp.text_frame.text.strip():
                texts.append({"shape": shp.name, "text": shp.text_frame.text})
        slides.append({
            "index": i,
            "layout": slide.slide_layout.name,
            "has_picture": has_picture,
            "texts": texts,
        })
    return slides


def analyze(template_path):
    return {
        "layouts": analyze_layouts(template_path),
        "slides": analyze_slides(template_path),
    }


if __name__ == "__main__":
    import json
    import sys
    print(json.dumps(analyze(sys.argv[1]), ensure_ascii=False, indent=2))
