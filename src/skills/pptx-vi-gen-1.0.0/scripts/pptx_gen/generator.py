"""Unified entry point: template + content spec -> filled .pptx.

deck_spec (JSON-serializable):
{
  "slides": [
    {"mode": "copy",        "template_slide_index": 0, "replacements": {"old": "new", ...}},
    {"mode": "placeholder", "layout_index": 1,          "fields": {"0": "title", "1": ["b1","b2"]}}
  ],
  "keep_template_slides": false
}

"copy"        -> deep-copy a template slide, replace its text (keeps every
                 picture/shape/background of that template page).
"placeholder" -> add a fresh slide from a layout and fill its placeholders.

The original template slides are dropped by default so the output contains
only the generated pages.
"""
import os

from pptx import Presentation
from pptx.oxml.ns import qn

from .content_filler import fill_placeholders, replace_text
from .slide_copy import copy_slide


def _drop_first_n_slides(prs, n):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst)[:n]:
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def generate(template_path, deck_spec, out_path):
    prs = Presentation(template_path)
    original_count = len(prs.slides)
    log = []

    for spec in deck_spec.get("slides", []):
        mode = spec.get("mode", "copy")
        if mode == "copy":
            src = prs.slides[spec["template_slide_index"]]
            new = copy_slide(prs, src)
            touched = replace_text(new, spec.get("replacements", {}))
            log.append({"mode": "copy", "shapes_changed": touched})
        elif mode == "placeholder":
            layout = prs.slide_layouts[spec["layout_index"]]
            new = prs.slides.add_slide(layout)
            filled = fill_placeholders(new, spec.get("fields", {}))
            log.append({"mode": "placeholder", "placeholders_filled": filled})
        else:
            raise ValueError(f"unknown mode: {mode!r}")

    if not deck_spec.get("keep_template_slides", False):
        _drop_first_n_slides(prs, original_count)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    return {"out": out_path, "slides": log, "generated": len(log)}


if __name__ == "__main__":
    import json
    import sys
    tpl, spec_path, out = sys.argv[1], sys.argv[2], sys.argv[3]
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    print(json.dumps(generate(tpl, spec, out), ensure_ascii=False, indent=2))
