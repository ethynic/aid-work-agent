"""Route 3: native editable PPTX generation.

Same content.json schema and VI palette as the Pillow renderer, but every
page is built from native PPTX shapes (text boxes / rounded rects / ovals /
pictures) laid out on a 1920x1080 grid (1 px = 6350 EMU, 1 px = 0.5 pt), so
all text stays fully editable in PowerPoint / WPS.

Public API:
    generate_native(template_path, content, out_path,
                    vi_override=None, base_dir=None)
        -> {"out": path, "slides": n}
"""
import os
import sys

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .fonts import (FONT_NAME, FONT_NAME_LATIN, font_file, latin_font_file,
                     is_cjk, remap_presentation)
from .fonts import embed_fonts as _embed_fonts
from .vi import default_vi, extract_vi

EMU_PX = 6350                     # 12192000 EMU / 1920 px  (144 px per inch)
W, H = 1920, 1080
MARGIN = 120
FOOTER_Y = 1000
# 商用免费双轨字体（见 fonts.py）：CJK→思源黑体(FONT_NAME)，拉丁→Noto Sans
FONT = FONT_NAME
FONT_LATIN = FONT_NAME_LATIN

INK, BODY, MUTED = "FFFFFF", "E8F0F8", "A9C2DA"
DARKTXT, CYAN = "0A1E33", "35B8D9"


# ---------------------------------------------------------------- helpers
def _px(v):
    return Emu(int(round(v * EMU_PX)))


def _pt(v):
    return Pt(v * 0.5)            # 144 px/in, 72 pt/in


def _hex(c):
    c = (c or "").lstrip("#").upper()
    return c if len(c) == 6 else "000000"


def _alpha(parent, pct):
    """Append <a:alpha> to the srgbClr inside parent's a:solidFill."""
    sf = parent.find(qn("a:solidFill"))
    if sf is None:
        return
    clr = sf.find(qn("a:srgbClr"))
    if clr is None:
        return
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))}))


def _em(ch):
    """Glyph advance of `ch` in em units — fallback only when no font file.

    思源黑体的实测值（size 1000）：CJK 满格 1.0、数字 0.555、句点 0.278、
    % 0.921、括号 0.338、空格 0.224、大写拉丁约 0.608。Noto Sans 的拉丁字形
    与此近似，故兜底模型对两种字体通用。
    """
    o = ord(ch)
    if o >= 0x2E80 or ch in "·—…“”‘’":
        return 1.0
    if ch.isdigit():
        return 0.555
    if ch in ".,:;'!|":
        return 0.278
    if ch == "%":
        return 0.921
    if ch in "+-":
        return 0.555
    if ch in "()[]{}<>/":
        return 0.338
    if ch == " ":
        return 0.224
    return 0.608


_PIL_FONT_CACHE = {}


def _pil_font(size, bold=False, cjk=False):
    """Load the right bundled font for a script (CJK→思源, 拉丁→Noto Sans)."""
    key = (int(round(size)), bool(bold), cjk)
    if key not in _PIL_FONT_CACHE:
        path = (font_file(bold) if cjk else latin_font_file(bold))
        font = None
        if path:
            try:
                font = ImageFont.truetype(path, max(4, key[0]))
            except (OSError, ValueError):
                # unreadable/missing font file -> em-model fallback only.
                # NOTE: never use a bare `except Exception` here: it once
                # silently swallowed a NameError and disabled real-font
                # measurement entirely (text wrapped in PowerPoint).
                font = None
        _PIL_FONT_CACHE[key] = font
    return _PIL_FONT_CACHE[key]


def measure(s, size, bold=False):
    """Real glyph-advance width (px) — per character with its OWN font.

    CJK chars are measured from Source Han Sans CN, Latin/digits from Noto
    Sans, so the overflow protection matches what PowerPoint actually renders
    (it picks the typeface per character by script). The em-model in _em() is
    only the last-resort fallback when a font file is unavailable.
    """
    s = str(s)
    total = 0.0
    for ch in s:
        cjk = is_cjk(ch)
        f = _pil_font(size, bold, cjk=cjk)
        if f is not None:
            total += f.getlength(ch)
        else:
            total += _em(ch) * size
    return total


def est_w(s, size, bold=False):
    return measure(s, size, bold)


def _wrap_est(s, size, w, bold=False):
    """Greedy width wrap via real measurement (CJK breaks anywhere)."""
    lines = []
    for para in str(s).split("\n"):
        cur = ""
        for ch in para:
            if measure(cur + ch, size, bold) <= w or not cur:
                cur += ch
            else:
                lines.append(cur)
                cur = ch
        lines.append(cur)
    return lines


def fit_size(s, size, w, h=None, lh=1.35, max_lines=None, min_size=12,
             bold=False, safety=0.97):
    """Largest size <= `size` whose (wrapped) text fits box w x h px.

    Width budget = w * safety (3% headroom for renderer rounding), height
    budget = h. Guards against the 'fixed font overflows / wraps' failure
    mode: long KPI numbers shrink to one line; long paragraphs shrink
    until the wrapped block fits the box.
    """
    s = str(s)
    w_budget = w * safety
    while size > min_size:
        lines = _wrap_est(s, size, w_budget, bold)
        if max_lines is not None and len(lines) > max_lines:
            size -= 2
            continue
        if h is not None and len(lines) * size * lh > h:
            size -= 2
            continue
        break
    return max(size, min_size)


def rect(slide, x, y, w, h, fill=None, alpha=None, radius=None,
         line=None, line_alpha=None, line_w=1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(kind, _px(x), _px(y), _px(w), _px(h))
    if radius:
        try:
            sp.adjustments[0] = min(0.5, float(radius))
        except Exception:
            pass
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(_hex(fill))
        if alpha is not None:
            _alpha(sp._element.spPr, alpha)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(_hex(line))
        sp.line.width = Pt(line_w)
        if line_alpha is not None:
            ln = sp._element.spPr.find(qn("a:ln"))
            if ln is not None:
                _alpha(ln, line_alpha)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def oval(slide, cx, cy, r, fill, alpha=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, _px(cx - r), _px(cy - r),
                                _px(2 * r), _px(2 * r))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(_hex(fill))
    if alpha is not None:
        _alpha(sp._element.spPr, alpha)
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def bg(slide, pal):
    """Full-bleed vertical gradient + soft glow circles."""
    g = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, _px(W), _px(H))
    try:
        g.fill.gradient()
        stops = g.fill.gradient_stops
        n = len(stops)
        for i in range(n):
            stops[i].color.rgb = RGBColor.from_string(
                _hex(pal["top"] if i == 0 else
                     pal["bottom"] if i == n - 1 else pal["top"]))
        # pin the angle via XML: python-pptx's gradient_angle setter writes
        # 270 deg for an input of 90; ang=5400000 (90 deg) puts stop 0 on top
        grad = g._element.spPr.find(qn("a:gradFill"))
        lin = grad.find(qn("a:lin")) if grad is not None else None
        if lin is not None:
            lin.set("ang", "5400000")
    except Exception:
        g.fill.solid()
        g.fill.fore_color.rgb = RGBColor.from_string(_hex(pal["bottom"]))
    g.line.fill.background()
    g.shadow.inherit = False
    oval(slide, W - 140, 80, 300, pal["gold"], 16)
    oval(slide, 30, H - 30, 240, CYAN, 11)
    return g


def _set_run(r, text, size, color, bold=False, alpha=None):
    r.text = text
    f = r.font
    f.size = _pt(size)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(_hex(color))
    f.name = FONT_LATIN                 # latin（数字/字母）→ Noto Sans
    rPr = r._r.get_or_add_rPr()
    # ea（CJK）→ 思源黑体；cs（复杂脚本）→ Noto Sans，与 latin 一致
    for tag, fam in (("a:ea", FONT), ("a:cs", FONT_LATIN)):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", fam)
    if alpha is not None:
        sf = rPr.find(qn("a:solidFill"))
        if sf is not None:
            clr = sf.find(qn("a:srgbClr"))
            if clr is not None:
                clr.append(clr.makeelement(qn("a:alpha"),
                                           {"val": str(int(alpha * 1000))}))


def _fill_tf(tf, s, size, color, bold, align, anchor, lh, alpha, wrap=True):
    tf.word_wrap = wrap
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        try:
            p.line_spacing = max(0.5, lh / 1.2)   # PowerPoint single ~1.2em
        except Exception:
            pass
        _set_run(p.add_run(), ln, size, color, bold, alpha)


def text(slide, x, y, w, h, s, size, color, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, lh=1.35, alpha=None,
         fit=False, max_lines=None):
    if fit:
        # single-line elements get a tighter width budget: the viewer's font
        # may be wider than the measured one, and a wrapped KPI number is
        # the worst visual failure
        safety = 0.92 if max_lines == 1 else 0.97
        size = fit_size(s, size, w, h, lh, max_lines=max_lines, bold=bold,
                        safety=safety)
    tb = slide.shapes.add_textbox(_px(x), _px(y), _px(w), _px(h))
    # hard guarantee: max_lines=1 boxes can never wrap, whatever font the
    # viewer substitutes (fit_size keeps them close enough to stay in-card)
    _fill_tf(tb.text_frame, s, size, color, bold, align, anchor, lh, alpha,
             wrap=max_lines != 1)
    return tb


def shape_text(sp, s, size, color, bold=False, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, lh=1.0):
    """Text attached to an autoshape (moves with the shape when edited)."""
    _fill_tf(sp.text_frame, s, size, color, bold, align, anchor, lh, None)
    return sp


def header(slide, pal, title, lead=None):
    rect(slide, MARGIN, 96, 10, 86, fill=pal["gold"])
    text(slide, MARGIN + 34, 90, W - 2 * MARGIN - 40, 80, title, 58, INK,
         bold=True, lh=1.2, fit=True, max_lines=1)
    if lead:
        text(slide, MARGIN + 34, 188, W - 2 * MARGIN - 120, 60, lead, 30,
             MUTED, lh=1.35, fit=True)
        return 245
    return 202


def footer(slide, page, i, n):
    rect(slide, MARGIN, FOOTER_Y - 26, W - 2 * MARGIN, 1, fill=INK, alpha=55)
    text(slide, MARGIN, FOOTER_Y, 900, 40, page.get("org", ""), 24, INK,
         alpha=150, lh=1.2)
    # named so editor.py can refresh page numbers after add/delete/move
    tb = text(slide, W - MARGIN - 220, FOOTER_Y, 220, 40, f"{i:02d} / {n:02d}",
              24, INK, alpha=150, align=PP_ALIGN.RIGHT, lh=1.2)
    tb.name = "pptxgen_pageno"


# ------------------------------------------------------------- renderers
def render_cover(slide, page, i, n, pal, img_path):
    bg(slide, pal)
    oval(slide, W - 130, 100, 330, pal["gold"], 18)
    oval(slide, W - 40, H - 40, 250, CYAN, 12)
    oval(slide, 40, H - 220, 190, CYAN, 8)
    rect(slide, MARGIN, 440, 150, 10, fill=pal["gold"])
    text(slide, MARGIN, 470, W - 2 * MARGIN, 160, page["title"], 92, INK,
         bold=True, lh=1.18, fit=True, max_lines=1)
    text(slide, MARGIN, 640, W - 2 * MARGIN, 64, page.get("subtitle", ""), 40,
         "CDDCEF", lh=1.3, fit=True, max_lines=1)
    # key-number chips (content-driven: page["chips"])
    x, y = MARGIN, 730
    for t in page.get("chips", []):
        tw = est_w(t, 28)
        cw = tw + 56
        chip = rect(slide, x, y, cw, 64, fill=INK, alpha=12, radius=0.5,
                    line=pal["gold"], line_alpha=80, line_w=1.0)
        shape_text(chip, t, 28, INK)
        x += cw + 24
    meta = "   ·   ".join(t for t in (page.get("date"), page.get("org")) if t)
    if meta:
        text(slide, MARGIN, 856, W - 2 * MARGIN, 44, meta, 30, MUTED,
             alpha=235, lh=1.3)


def render_section(slide, page, i, n, pal, img_path):
    bg(slide, pal)
    oval(slide, W - 240, 220, 330, CYAN, 12)
    oval(slide, W + 40, 640, 240, pal["gold"], 10)
    text(slide, MARGIN, 230, 900, 300, page.get("no", f"{i:02d}"), 210,
         pal["gold"], bold=True, lh=1.1)
    rect(slide, MARGIN, 500, 130, 10, fill=pal["gold"])
    text(slide, MARGIN, 546, W - 2 * MARGIN - 200, 120, page["title"], 78,
         INK, bold=True, lh=1.2, fit=True, max_lines=1)
    if page.get("lead"):
        text(slide, MARGIN, 700, W - 2 * MARGIN - 240, 70, page["lead"], 34,
             MUTED, lh=1.35, fit=True)
    footer(slide, page, i, n)


def render_stats(slide, page, i, n, pal, img_path):
    bg(slide, pal)
    top = header(slide, pal, page["title"], page.get("lead"))
    stats = page["stats"]
    cols = min(len(stats), 4)
    gap = 36
    cw = (W - 2 * MARGIN - gap * (cols - 1)) // cols

    def _card_h(s):
        # content: strip@30 + num slot(74..218) + label slot(218..274) + desc
        d = s.get("desc", "")
        d_h = 0
        if d:
            ds = fit_size(d, 25, cw - 52, 130, 1.3)
            d_h = len(_wrap_est(d, ds, (cw - 52) * 0.97)) * ds * 1.3
        return int(274 + d_h + 34)

    ch = max(_card_h(s) for s in stats)
    avail_bot = FOOTER_Y - 60
    y = top + 56 + max(0, (avail_bot - top - 56 - ch) // 2)
    x = MARGIN
    for s in stats:
        rect(slide, x, y, cw, ch, fill=INK, alpha=12, radius=26.0 / ch,
             line=INK, line_alpha=30, line_w=1.0)
        rect(slide, x + 26, y + 30, 60, 6, fill=pal["gold"])
        text(slide, x + 26, y + 74, cw - 52, 140, s["num"], 104, pal["gold"],
             bold=True, lh=1.15, fit=True, max_lines=1)
        text(slide, x + 26, y + 218, cw - 52, 50, s["label"], 32, INK,
             bold=True, lh=1.25, fit=True, max_lines=1)
        if s.get("desc"):
            text(slide, x + 26, y + 274, cw - 52, ch - 274 - 28, s["desc"],
                 25, MUTED, lh=1.3, fit=True)
        x += cw + gap
    footer(slide, page, i, n)


def render_bullets(slide, page, i, n, pal, img_path):
    bg(slide, pal)
    top = header(slide, pal, page["title"], page.get("lead"))
    items = page["items"]
    avail = FOOTER_Y - 40 - top
    rh = min(190, avail // max(1, len(items)))
    y = top + 16
    for idx, it in enumerate(items):
        dot = oval(slide, MARGIN + 36, y + 42, 34, pal["gold"])
        shape_text(dot, str(idx + 1), 30, DARKTXT, bold=True)
        rect(slide, MARGIN + 80, y + 16, 3, rh - 52, fill=CYAN, alpha=55)
        text(slide, MARGIN + 96, y, W - MARGIN - 96 - MARGIN, 60, it["h"], 38,
             INK, bold=True, lh=1.25, fit=True, max_lines=1)
        text(slide, MARGIN + 96, y + 56, W - MARGIN - 96 - MARGIN, 110,
             it.get("d", ""), 28, MUTED, lh=1.35, fit=True)
        y += rh
    footer(slide, page, i, n)


def render_twocol(slide, page, i, n, pal, img_path):
    bg(slide, pal)
    top = header(slide, pal, page["title"], page.get("lead"))
    cols = [page["left"], page["right"]]
    gap = 48
    cw = (W - 2 * MARGIN - gap) // 2
    bottom = FOOTER_Y - 60
    strip = [CYAN, pal["gold"]]
    item_h = 92
    n_max = max(len(c["items"]) for c in cols)
    card_h = 150 + n_max * item_h + 36
    y0 = top + 12 + max(0, (bottom - top - 12 - card_h) // 2)
    for k, col in enumerate(cols):
        x = MARGIN + k * (cw + gap)
        rect(slide, x, y0, cw, card_h, fill=INK, alpha=12,
             radius=26.0 / card_h, line=INK, line_alpha=28, line_w=1.0)
        rect(slide, x, y0, cw, 10, fill=strip[k], radius=0.5)
        text(slide, x + 36, y0 + 52, cw - 72, 64, col["h"], 40, strip[k],
             bold=True, lh=1.25, fit=True, max_lines=1)
        for j, it in enumerate(col["items"]):
            yy = y0 + 150 + j * item_h + (item_h - 44) // 2
            oval(slide, x + 46, yy + 20, 7, strip[k])
            text(slide, x + 70, yy, cw - 106, 44, it, 28, BODY, lh=1.34,
                 fit=True)
    footer(slide, page, i, n)


def render_image(slide, page, i, n, pal, img_path):
    bg(slide, pal)
    top = header(slide, pal, page["title"], page.get("lead"))
    box_top = top + 16
    box_bottom = FOOTER_Y - 90
    avail_h = box_bottom - box_top
    cap_h = 76 if page.get("caption") else 0
    if img_path and os.path.exists(img_path):
        img = Image.open(img_path)
        max_iw, max_ih = W - 2 * MARGIN, avail_h - 56 - cap_h
        r = min(max_iw / img.width, max_ih / img.height)
        iw, ih = int(img.width * r), int(img.height * r)
        # card hugs the image (no giant side voids for wide/short charts)
        card_w = min(W - 2 * MARGIN, iw + 56)
        card_h = ih + 56 + cap_h
        cx = MARGIN + (W - 2 * MARGIN - card_w) // 2
        cy = box_top + max(0, (avail_h - card_h) // 2)
        rect(slide, cx, cy, card_w, card_h, fill="FFFFFF",
             radius=26.0 / card_h, line=pal["gold"], line_alpha=120, line_w=1.0)
        slide.shapes.add_picture(img_path, _px(cx + (card_w - iw) // 2),
                                 _px(cy + 28), _px(iw), _px(ih))
        if page.get("caption"):
            text(slide, cx + 28, cy + card_h - 62, card_w - 56, 44,
                 page["caption"], 26, "5A6B7E", lh=1.3, fit=True, max_lines=1)
    else:
        card_h = avail_h
        rect(slide, MARGIN, box_top, W - 2 * MARGIN, card_h, fill="FFFFFF",
             radius=26.0 / card_h, line=pal["gold"], line_alpha=120, line_w=1.0)
        _img_panel(slide, MARGIN + 28, box_top + 28, W - 2 * MARGIN - 56,
                   card_h - 56, None, pal)
    footer(slide, page, i, n)


def render_end(slide, page, i, n, pal, img_path):
    bg(slide, pal)
    # decorative circles stay clear of the centered text band (y 500~880)
    oval(slide, W - 130, 100, 330, pal["gold"], 18)
    oval(slide, 140, H - 200, 200, CYAN, 10)
    oval(slide, W - 120, H - 280, 140, pal["gold"], 12)
    rect(slide, W // 2 - 75, 468, 150, 10, fill=pal["gold"])
    text(slide, MARGIN, 540, W - 2 * MARGIN, 130, page.get("title", "谢谢观看"),
         84, INK, bold=True, align=PP_ALIGN.CENTER, lh=1.2, fit=True,
         max_lines=1)
    if page.get("subtitle"):
        text(slide, MARGIN, 680, W - 2 * MARGIN, 64, page["subtitle"], 36,
             "CDDCEF", align=PP_ALIGN.CENTER, lh=1.3, fit=True, max_lines=1)
    if page.get("contact"):
        text(slide, MARGIN, 800, W - 2 * MARGIN, 48, page["contact"], 30,
             MUTED, alpha=230, align=PP_ALIGN.CENTER, lh=1.3, fit=True,
             max_lines=1)


# ---------------------------------------------------------- table helpers
def _plain_table(tbl):
    """Strip the default python-pptx table style (we paint every cell)."""
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    sid = tblPr.makeelement(qn("a:tableStyleId"), {})
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid
    tblPr.append(sid)


def _is_num(s):
    t = str(s).strip().replace(",", "").replace("¥", "").replace("$", "")
    if t.endswith("%"):
        t = t[:-1]
    if t.startswith("+"):
        t = t[1:]
    try:
        float(t)
        return True
    except ValueError:
        return False


def _cell_text(cell, s, size, color, bold=False, align=PP_ALIGN.LEFT,
               col_w=None, row_h=None):
    """Write one table cell with margins, vertical centering and font fit."""
    tf = cell.text_frame
    tf.word_wrap = True
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = _px(14)
    cell.margin_top = cell.margin_bottom = _px(6)
    if col_w and row_h:
        while size > 14:
            lines = _wrap_est(s, size, col_w - 28, bold)
            if len(lines) * size * 1.25 <= row_h - 12:
                break
            size -= 2
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        try:
            p.line_spacing = 1.1
        except Exception:
            pass
        _set_run(p.add_run(), ln, size, color, bold, None)


def _img_panel(slide, x, y, w, h, img_path, pal, radius_frac=26.0):
    """White exhibit card with the image fitted inside; graceful placeholder
    when the image file is missing (keeps the layout intact)."""
    rect(slide, x, y, w, h, fill="FFFFFF", radius=radius_frac / max(h, 1),
         line=pal["gold"], line_alpha=110, line_w=1.0)
    pad = 20
    ix, iy = x + pad, y + pad
    iw, ih = w - 2 * pad, h - 2 * pad
    if img_path and os.path.exists(img_path):
        img = Image.open(img_path)
        r = min(iw / img.width, ih / img.height)
        nw, nh = int(img.width * r), int(img.height * r)
        slide.shapes.add_picture(img_path, _px(ix + (iw - nw) // 2),
                                 _px(iy + (ih - nh) // 2), _px(nw), _px(nh))
    else:
        rect(slide, ix, iy, iw, ih, fill=INK, alpha=14, radius=0.04)
        text(slide, ix, iy + ih // 2 - 24, iw, 48, "图片占位", 26, MUTED,
             align=PP_ALIGN.CENTER, lh=1.2, fit=True, max_lines=1)


def render_table(slide, page, i, n, pal, img_path):
    """Data table: VI-brand header row, zebra body, numeric right-align."""
    bg(slide, pal)
    top = header(slide, pal, page["title"], page.get("lead"))
    head, rows = page["head"], page["rows"]
    ncol = len(head)
    area_x, area_w = MARGIN, W - 2 * MARGIN
    area_top = top + 24
    area_bot = FOOTER_Y - 80
    if page.get("caption"):
        area_bot -= 60
    nrow = len(rows) + 1
    rh = max(44, min(92, (area_bot - area_top) // nrow))
    # column widths proportional to the widest content per column
    weights = []
    for c in range(ncol):
        m = measure(str(head[c]), 26, True)
        for r in rows:
            m = max(m, measure(str(r[c]), 24, False))
        weights.append(m)
    tot = sum(weights) or 1
    col_w = [max(120, int(area_w * wgt / tot)) for wgt in weights]
    col_w[-1] += area_w - sum(col_w)
    gf = slide.shapes.add_table(nrow, ncol, _px(area_x), _px(area_top),
                                _px(area_w), _px(rh * nrow))
    tbl = gf.table
    _plain_table(tbl)
    for c in range(ncol):
        tbl.columns[c].width = _px(col_w[c])
    tbl.rows[0].height = _px(rh + 6)
    for r in range(1, nrow):
        tbl.rows[r].height = _px(rh)
    for c, htxt in enumerate(head):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(_hex(pal["top"]))
        _cell_text(cell, htxt, 26, "FFFFFF", bold=True,
                   align=PP_ALIGN.CENTER, col_w=col_w[c], row_h=rh + 6)
    for r, row in enumerate(rows, 1):
        for c in range(ncol):
            val = str(row[c])
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                "FFFFFF" if r % 2 else "EDF2F9")
            if _is_num(val):
                align = PP_ALIGN.RIGHT
            elif c == 0:
                align = PP_ALIGN.LEFT
            else:
                align = PP_ALIGN.CENTER
            _cell_text(cell, val, 24, "1F2937", align=align,
                       col_w=col_w[c], row_h=rh)
    # accent rule under the header row
    rect(slide, area_x, area_top + rh + 6, area_w, 4, fill=pal["gold"])
    if page.get("caption"):
        text(slide, area_x, area_top + rh * nrow + 18, area_w, 44,
             page["caption"], 24, MUTED, lh=1.3, fit=True, max_lines=1)
    footer(slide, page, i, n)


def render_image_text(slide, page, i, n, pal, img_path):
    """Image + text side by side (page["side"] = "left" | "right")."""
    bg(slide, pal)
    top = header(slide, pal, page["title"], page.get("lead"))
    bottom = FOOTER_Y - 60
    gap = 56
    img_w = int((W - 2 * MARGIN - gap) * 0.52)
    txt_w = W - 2 * MARGIN - gap - img_w
    side = page.get("side", "left")
    ix = MARGIN if side == "left" else MARGIN + txt_w + gap
    tx = MARGIN if side == "right" else MARGIN + img_w + gap
    _img_panel(slide, ix, top + 16, img_w, bottom - top - 16, img_path, pal)
    y = top + 44
    if page.get("h"):
        text(slide, tx, y, txt_w, 60, page["h"], 40, INK, bold=True,
             lh=1.25, fit=True, max_lines=1)
        y += 78
    pts = page.get("points", [])
    chips = page.get("chips") or []
    area_bot = bottom - 84 if chips else bottom - 16
    slot = (area_bot - y) // max(1, len(pts))
    for j, pt in enumerate(pts):
        py = y + j * slot
        if isinstance(pt, dict):
            oval(slide, tx + 10, py + 20, 7, pal["gold"])
            text(slide, tx + 34, py, txt_w - 34, 52, pt.get("h", ""), 30,
                 INK, bold=True, lh=1.25, fit=True, max_lines=1)
            if pt.get("d"):
                text(slide, tx + 34, py + 58, txt_w - 34, slot - 66, pt["d"],
                     26, MUTED, lh=1.35, fit=True)
        else:
            oval(slide, tx + 10, py + 18, 7, pal["gold"])
            text(slide, tx + 34, py, txt_w - 34, slot - 16, str(pt), 28,
                 BODY, lh=1.34, fit=True)
    chips = page.get("chips") or []
    if chips:
        cx = tx
        for t in chips:
            cw_ = est_w(t, 26) + 48
            if cx + cw_ > tx + txt_w:
                break
            rect(slide, cx, bottom - 76, cw_, 56, fill=INK, alpha=14,
                 radius=0.5, line=pal["gold"], line_alpha=70, line_w=1.0)
            text(slide, cx + 24, bottom - 76 + 14, cw_ - 48, 30, t, 26, INK,
                 lh=1.2, fit=True, max_lines=1)
            cx += cw_ + 20
    footer(slide, page, i, n)


def render_image_grid(slide, page, i, n, pal, img_path):
    """2~4 images in a grid, each with its own caption."""
    bg(slide, pal)
    top = header(slide, pal, page["title"], page.get("lead"))
    bottom = FOOTER_Y - 80
    if page.get("caption"):
        bottom -= 56
    imgs = page.get("images") or ([page["image"]] if page.get("image") else [])
    paths = page.get("_img_paths") or ([img_path] if img_path else [])
    k = len(imgs)
    if k == 4:
        cols, rows_n = 2, 2
    elif k == 3:
        cols, rows_n = 3, 1
    elif k == 2:
        cols, rows_n = 2, 1
    else:
        cols, rows_n = 1, 1
    gap = 36
    cw = (W - 2 * MARGIN - gap * (cols - 1)) // cols
    ch = (bottom - top - gap * (rows_n - 1)) // rows_n
    for idx in range(cols * rows_n):
        r, c = divmod(idx, cols)
        x = MARGIN + c * (cw + gap)
        y = top + r * (ch + gap)
        p = paths[idx] if idx < len(paths) else None
        _img_panel(slide, x, y, cw, ch - 66, p, pal)
        cap = imgs[idx] if idx < k else None
        cap = cap.get("caption", "") if isinstance(cap, dict) else ""
        if cap:
            text(slide, x, y + ch - 58, cw, 48, cap, 24, MUTED,
                 lh=1.3, fit=True, max_lines=1)
    if page.get("caption"):
        text(slide, MARGIN, bottom + 8, W - 2 * MARGIN, 44, page["caption"],
             24, MUTED, lh=1.3, fit=True, max_lines=1)
    footer(slide, page, i, n)


def render_timeline(slide, page, i, n, pal, img_path):
    """Horizontal timeline: time + title + description per node."""
    bg(slide, pal)
    top = header(slide, pal, page["title"], page.get("lead"))
    items = page["items"]
    k = len(items)
    axis_y = top + 210
    rect(slide, MARGIN + 30, axis_y, W - 2 * MARGIN - 60, 4, fill=INK,
         alpha=60)
    span = (W - 2 * MARGIN - 80) // k
    for idx, it in enumerate(items):
        cx = MARGIN + 40 + span * idx + span // 2
        oval(slide, cx, axis_y + 2, 18, pal["gold"])
        oval(slide, cx, axis_y + 2, 9, pal["top"])
        bw = span - 28
        bx = cx - bw // 2
        text(slide, bx, axis_y - 96, bw, 44, it.get("t", ""), 26,
             pal["gold"], bold=True, align=PP_ALIGN.CENTER, lh=1.2,
             fit=True, max_lines=1)
        text(slide, bx, axis_y + 44, bw, 56, it.get("h", ""), 32, INK,
             bold=True, align=PP_ALIGN.CENTER, lh=1.25, fit=True,
             max_lines=1)
        if it.get("d"):
            text(slide, bx, axis_y + 108, bw, 120, it["d"], 22, MUTED,
                 align=PP_ALIGN.CENTER, lh=1.35, fit=True)
    footer(slide, page, i, n)


def render_process(slide, page, i, n, pal, img_path):
    """Step flow: numbered cards with chevrons between them."""
    bg(slide, pal)
    top = header(slide, pal, page["title"], page.get("lead"))
    avail_bot = FOOTER_Y - 90
    steps = page["steps"]
    k = len(steps)
    gap = 76
    cw = (W - 2 * MARGIN - gap * (k - 1)) // k
    ch = 390
    y = top + 30 + max(0, (avail_bot - top - 30 - ch) // 2)
    for idx, st in enumerate(steps):
        x = MARGIN + idx * (cw + gap)
        rect(slide, x, y, cw, ch, fill=INK, alpha=12, radius=0.06,
             line=INK, line_alpha=30, line_w=1.0)
        rect(slide, x + 30, y + 34, 64, 64, fill=pal["gold"], radius=0.5)
        text(slide, x + 30, y + 48, 64, 40, f"{idx + 1}", 32, DARKTXT,
             bold=True, align=PP_ALIGN.CENTER, lh=1.2)
        text(slide, x + 30, y + 126, cw - 60, 60, st.get("h", ""), 34, INK,
             bold=True, lh=1.25, fit=True, max_lines=1)
        if st.get("d"):
            text(slide, x + 30, y + 196, cw - 60, ch - 226, st["d"], 26,
                 MUTED, lh=1.35, fit=True)
        if idx < k - 1:
            ch_shape = slide.shapes.add_shape(
                MSO_SHAPE.CHEVRON, _px(x + cw + 18), _px(y + ch // 2 - 18),
                _px(40), _px(36))
            ch_shape.fill.solid()
            ch_shape.fill.fore_color.rgb = RGBColor.from_string(
                _hex(pal["gold"]))
            ch_shape.line.fill.background()
            ch_shape.shadow.inherit = False
    footer(slide, page, i, n)


def render_quote(slide, page, i, n, pal, img_path):
    """Full-page statement / quote."""
    bg(slide, pal)
    oval(slide, W - 150, 140, 300, pal["gold"], 16)
    oval(slide, 90, H - 160, 240, CYAN, 10)
    text(slide, MARGIN, 150, 400, 300, "\u201c", 200, pal["gold"],
         alpha=90, lh=1.0)
    text(slide, MARGIN + 90, 330, W - 2 * MARGIN - 180, 420,
         page["quote"], 48, INK, lh=1.5, fit=True)
    rect(slide, MARGIN + 90, 800, 120, 8, fill=pal["gold"])
    if page.get("author"):
        text(slide, MARGIN + 90, 840, W - 2 * MARGIN - 180, 56,
             page["author"], 32, "CDDCEF", bold=True, lh=1.3, fit=True,
             max_lines=1)
    if page.get("title"):
        text(slide, MARGIN + 90, 900, W - 2 * MARGIN - 180, 48,
             page["title"], 26, MUTED, lh=1.3, fit=True, max_lines=1)


def render_toc(slide, page, i, n, pal, img_path):
    """Agenda / table of contents with numbered entries."""
    bg(slide, pal)
    top = header(slide, pal, page.get("title", "目录"), page.get("lead"))
    items = page["items"]
    bottom = FOOTER_Y - 70
    ncol = 2 if len(items) > 4 else 1
    col_w = (W - 2 * MARGIN - 80) // ncol
    rows_n = (len(items) + ncol - 1) // ncol
    rh = min(180, (bottom - top - 20) // max(1, rows_n))
    for idx, it in enumerate(items):
        c, r = divmod(idx, rows_n)
        x = MARGIN + c * (col_w + 80)
        y = top + 16 + r * rh
        h = it.get("h", "") if isinstance(it, dict) else str(it)
        d = it.get("d", "") if isinstance(it, dict) else ""
        text(slide, x, y + 8, 130, 80, f"{idx + 1:02d}", 54, pal["gold"],
             bold=True, lh=1.1, fit=True, max_lines=1)
        text(slide, x + 150, y + 16, col_w - 150, 56, h, 36, INK, bold=True,
             lh=1.25, fit=True, max_lines=1)
        if d:
            text(slide, x + 150, y + 78, col_w - 150, rh - 96, d, 26, MUTED,
                 lh=1.35, fit=True)
        rect(slide, x, y + rh - 14, col_w, 1, fill=INK, alpha=40)
    footer(slide, page, i, n)


NATIVE_RENDERERS = {
    "cover": render_cover,
    "section": render_section,
    "bullets": render_bullets,
    "stats": render_stats,
    "twocol": render_twocol,
    "image": render_image,
    "table": render_table,
    "image_text": render_image_text,
    "image_grid": render_image_grid,
    "timeline": render_timeline,
    "process": render_process,
    "quote": render_quote,
    "toc": render_toc,
    "end": render_end,
}


def _resolve_img(ref, base_dir):
    if not ref:
        return None
    p = ref if os.path.isabs(ref) else os.path.join(base_dir, ref)
    return p if os.path.exists(p) else None


def add_page_slide(prs, page, pal, img_path=None):
    """Append one native slide built from a role spec (used by editor ops).

    Page numbers get refreshed by the editor after structural ops; pass
    "no" explicitly for section pages so the big number stays right.
    """
    role = page["role"]
    if role not in NATIVE_RENDERERS:
        raise ValueError("unknown role %r; available: %s"
                         % (role, ", ".join(sorted(NATIVE_RENDERERS))))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    n = len(prs.slides._sldIdLst)
    NATIVE_RENDERERS[role](slide, page, n, n, pal, img_path)
    return slide


# ---------------------------------------------------------------- pipeline
def _palette(vi):
    return {
        "gold": (vi.get("accent") or "#FFB545").lstrip("#"),
        "top": (vi.get("brand") or "#14395F").lstrip("#"),
        "bottom": (vi.get("dark") or "#0A1E33").lstrip("#"),
    }


def generate_native(template_path, content, out_path, vi_override=None,
                    base_dir=None, remap_fonts=True, embed_fonts=False):
    """Route 3 entry: build a fully editable .pptx from structured content.

    remap_fonts: 把 deck 内任何残留字体名统一为商用免费字体
                （latin/cs→Noto Sans，ea→思源黑体；本路线输出本就只用这套
                 ﻿双轨字体，属幂等保险）。
    embed_fonts: 把思源黑体 + Noto Sans 嵌入 pptx（+18MB），未装这些字体的
                 机器也能原样渲染（opt-in）。
    """
    base_dir = base_dir or (os.getcwd() if template_path is None
                            else os.path.dirname(os.path.abspath(template_path)))
    # 无模板时走内置默认 VI（本项目允许不带 --template 生成）
    vi = default_vi() if template_path is None else extract_vi(template_path)
    if vi_override:
        vi.update(vi_override)
    pal = _palette(vi)

    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 16:9
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]

    slides = content["slides"]
    total = len(slides)
    missing = []
    for i, page in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        role = page["role"]
        if role not in NATIVE_RENDERERS:
            raise ValueError("unknown role %r on slide %d; available: %s"
                             % (role, i, ", ".join(sorted(NATIVE_RENDERERS))))
        img_path = None
        ref = page.get("image")
        if ref:
            img_path = _resolve_img(ref, base_dir)
            if img_path is None:
                missing.append("slide %d: image %r" % (i, ref))
        # multi-image roles (image_grid): resolve page["images"] into a list;
        # entries may be plain path strings or {"src": ..., "caption": ...}
        refs = page.get("images") or ([ref] if ref else [])
        page["_img_paths"] = [
            _resolve_img(r if isinstance(r, str) else r.get("src"), base_dir)
            for r in refs]
        for k, p in enumerate(page["_img_paths"]):
            if p is None and k < len(refs):
                missing.append("slide %d: images[%d] %r"
                               % (i, k, refs[k] if isinstance(refs[k], str)
                                  else refs[k].get("src")))
        NATIVE_RENDERERS[role](slide, page, i, total, pal, img_path)

    if missing:
        sys.stderr.write("WARNING: image file(s) not found (layout kept with "
                         "placeholder panels): %s\n" % "; ".join(missing))

    font_report = remap_presentation(prs) if remap_fonts else None
    try:
        prs.core_properties.title = slides[0].get("title", "PPT")
    except Exception:
        pass
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    embed_report = _embed_fonts(out_path) if embed_fonts else None
    return {"out": out_path, "slides": total, "fonts": font_report,
            "embedded": embed_report}
