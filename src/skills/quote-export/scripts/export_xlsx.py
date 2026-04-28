#!/usr/bin/env python3
"""研学旅游报价Excel导出脚本"""

import argparse
import json
import os
import sys
import tempfile
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print(json.dumps({
        'success': False,
        'error': 'openpyxl未安装，请执行: pip install openpyxl'
    }, ensure_ascii=False))
    sys.exit(1)


# 样式定义
HEADER_FONT = Font(name='微软雅黑', bold=True, size=11)
TITLE_FONT = Font(name='微软雅黑', bold=True, size=14)
DATA_FONT = Font(name='微软雅黑', size=10)
TOTAL_FONT = Font(name='微软雅黑', bold=True, size=11)
CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)
RIGHT_ALIGN = Alignment(horizontal='right', vertical='center')
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
HEADER_FILL = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
TOTAL_FILL = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')


def export_quote_xlsx(quote_data, output_path):
    """导出报价单为Excel文件"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = quote_data.get('course_name', '报价')[:31]  # sheet名最长31字符

    # 列宽
    col_widths = [12, 22, 10, 8, 6, 8, 6, 12, 12, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = 1

    # 标题行
    ws.merge_cells(f'A{row}:J{row}')
    cell = ws.cell(row=row, column=1, value=quote_data.get('title', '研学报价表'))
    cell.font = TITLE_FONT
    cell.alignment = CENTER_ALIGN
    ws.row_dimensions[row].height = 36
    row += 1

    # 信息行
    ws.merge_cells(f'A{row}:C{row}')
    ws.cell(row=row, column=1, value=f"课程名称：{quote_data.get('course_name', '')}")
    ws.merge_cells(f'D{row}:G{row}')
    ws.cell(row=row, column=4, value=f"日期：{quote_data.get('date', '')}")
    ws.merge_cells(f'H{row}:J{row}')
    ws.cell(row=row, column=8, value=f"人数：{quote_data.get('people_count', '')}人")
    for col in range(1, 11):
        ws.cell(row=row, column=col).font = DATA_FONT
    ws.row_dimensions[row].height = 24
    row += 1

    # 空行
    row += 1

    # 表头行
    headers = ['成本类别', '项目', '单价', '数量', '单位', '次数', '单位', '费用小计', '随队老师', '备注']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
        cell.fill = HEADER_FILL
    ws.row_dimensions[row].height = 28
    header_row = row
    row += 1

    # 数据行
    items = quote_data.get('items', [])
    data_start_row = row

    for item in items:
        ws.cell(row=row, column=1, value=item.get('category', '')).font = DATA_FONT
        ws.cell(row=row, column=2, value=item.get('name', '')).font = DATA_FONT
        c = ws.cell(row=row, column=3, value=item.get('unit_price', 0))
        c.font = DATA_FONT
        c.number_format = '#,##0.00'
        ws.cell(row=row, column=4, value=item.get('quantity', 1)).font = DATA_FONT
        ws.cell(row=row, column=5, value=item.get('unit', '')).font = DATA_FONT
        ws.cell(row=row, column=6, value=item.get('frequency', 1)).font = DATA_FONT
        ws.cell(row=row, column=7, value=item.get('freq_unit', '')).font = DATA_FONT
        c = ws.cell(row=row, column=8, value=item.get('subtotal', 0))
        c.font = DATA_FONT
        c.number_format = '#,##0.00'
        c = ws.cell(row=row, column=9, value=item.get('teacher_fee', 0))
        c.font = DATA_FONT
        c.number_format = '#,##0.00'
        ws.cell(row=row, column=10, value=item.get('remark', '')).font = DATA_FONT

        # 应用边框和对齐
        for col in range(1, 11):
            ws.cell(row=row, column=col).border = THIN_BORDER
            if col in (3, 8, 9):
                ws.cell(row=row, column=col).alignment = RIGHT_ALIGN
            elif col in (4, 5, 6, 7):
                ws.cell(row=row, column=col).alignment = CENTER_ALIGN
            else:
                ws.cell(row=row, column=col).alignment = LEFT_ALIGN

        row += 1

    # 合并同类别单元格
    if items:
        current_cat = None
        merge_start = data_start_row
        for i, item in enumerate(items):
            cat = item.get('category', '')
            if cat != current_cat:
                if i > 0 and merge_start < data_start_row + i:
                    ws.merge_cells(f'A{merge_start}:A{data_start_row + i - 1}')
                    cell = ws.cell(row=merge_start, column=1)
                    cell.alignment = CENTER_ALIGN
                current_cat = cat
                merge_start = data_start_row + i
        # 最后一组
        if merge_start < row:
            ws.merge_cells(f'A{merge_start}:A{row - 1}')
            ws.cell(row=merge_start, column=1).alignment = CENTER_ALIGN

    # 合计行
    ws.merge_cells(f'A{row}:G{row}')
    cell = ws.cell(row=row, column=1, value='合计')
    cell.font = TOTAL_FONT
    cell.alignment = CENTER_ALIGN
    c = ws.cell(row=row, column=8, value=quote_data.get('total', 0))
    c.font = TOTAL_FONT
    c.number_format = '#,##0.00'
    c = ws.cell(row=row, column=9, value=quote_data.get('teacher_total', 0))
    c.font = TOTAL_FONT
    c.number_format = '#,##0.00'

    for col in range(1, 11):
        ws.cell(row=row, column=col).border = THIN_BORDER
        ws.cell(row=row, column=col).fill = TOTAL_FILL

    row += 2

    # 审核签字区
    sign_row = row
    ws.cell(row=sign_row, column=1, value='审核').font = Font(name='微软雅黑', bold=True, size=11)
    ws.cell(row=sign_row, column=2, value='成本审核员').font = DATA_FONT
    ws.cell(row=sign_row, column=5, value='研学负责人').font = DATA_FONT
    ws.cell(row=sign_row, column=8, value='财务部审核').font = DATA_FONT

    # 签字线下划线
    row += 1
    for col, width in [(2, 3), (5, 3), (8, 3)]:
        for c in range(col, col + width):
            cell = ws.cell(row=row, column=c, value='')
            cell.border = Border(bottom=Side(style='thin'))

    wb.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description='研学旅游报价Excel导出')
    parser.add_argument('--title', default='贵州天悦旅行社有限公司研学报价表', help='报价表标题')
    parser.add_argument('--course-name', required=True, help='课程名称')
    parser.add_argument('--date', required=True, help='日期')
    parser.add_argument('--people', type=int, required=True, help='人数')
    parser.add_argument('--items', required=True, help='报价项JSON字符串')
    parser.add_argument('--total', type=float, required=True, help='总费用')
    parser.add_argument('--teacher-total', type=float, default=0, help='老师费用合计')
    parser.add_argument('--output', help='输出文件路径（可选）')

    try:
        args = parser.parse_args()
        items = json.loads(args.items)

        if not isinstance(items, list) or len(items) == 0:
            print(json.dumps({
                'success': False,
                'error': 'items必须是非空数组'
            }, ensure_ascii=False))
            sys.exit(1)

        # 构造报价数据
        quote_data = {
            'title': args.title,
            'course_name': args.course_name,
            'date': args.date,
            'people_count': args.people,
            'items': items,
            'total': args.total,
            'teacher_total': args.teacher_total
        }

        # 输出路径
        if args.output:
            output_path = args.output
        else:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"报价单_{args.course_name}_{args.people}人_{timestamp}.xlsx"
            output_path = os.path.join(tempfile.gettempdir(), filename)

        export_quote_xlsx(quote_data, output_path)

        print(json.dumps({
            'success': True,
            'file_path': os.path.abspath(output_path),
            'filename': os.path.basename(output_path),
            'people_count': args.people,
            'total': args.total
        }, ensure_ascii=False))

    except json.JSONDecodeError as e:
        print(json.dumps({
            'success': False,
            'error': f'JSON解析错误: {str(e)}'
        }, ensure_ascii=False))
        sys.exit(1)
    except Exception as e:
        print(json.dumps({
            'success': False,
            'error': str(e)
        }, ensure_ascii=False))
        sys.exit(1)


if __name__ == '__main__':
    main()
