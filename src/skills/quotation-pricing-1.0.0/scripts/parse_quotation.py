#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
分析 spec-to-quotation-list 生成的「室内材料报价清单」Excel，
抽取条目为 items_extracted.json，供后续知识库匹配报价使用。

主表列序（来自 output_format.md）：
A 编号 / B 描述 / C:D 区域(合并) / E 数量 / F 单位 / G 品牌 /
H 图片 / I 单价 / J 合价 / K 数量依据 / L 依据

用法：
  python parse_quotation.py <清单.xlsx 或 目录> [<更多路径>] \
        --json items_extracted.json [--project "项目名"]
"""
import argparse
import json
import os
import re

from openpyxl import load_workbook

GROUP_RE = re.compile(r"^\d+\s")          # "1 Furniture 家具" 分组行
TOTAL_RE = re.compile(r"(total|合计)", re.I)
SECTION_RE = re.compile(r"^■")            # "■ 灯具 Lighting" 分区行


def find_header_row(ws):
    """定位表头行（含 '编号' 与 '依据'）。"""
    for r in range(1, min(ws.max_row, 30) + 1):
        a = ws.cell(r, 1).value
        if a and ("编号" in str(a) or "No." in str(a) or "NO." in str(a)):
            l = ws.cell(r, 12).value
            if l and ("依据" in str(l) or "Ref" in str(l)):
                return r
    return None


def split_desc(b):
    """拆分 B 列描述 -> (标题中英文, 规格要点列表, 中文说明, 供货责任)。"""
    if b is None:
        return "", [], "", ""
    lines = [ln.rstrip() for ln in str(b).split("\n")]
    title = lines[0] if lines else ""
    bullets, cn_note, scope = [], "", ""
    for ln in lines[1:]:
        s = ln.strip()
        if not s:
            continue
        if s[0] in ("•", "·", "-"):
            bullets.append(s)
        elif "中文说明" in s:
            cn_note = s
        elif "Scope" in s or "供货责任" in s:
            scope = s
        elif cn_note == "":
            cn_note = s
    return title, bullets, cn_note, scope


def is_item_row(a, b):
    """判定是否为产品数据行（排除分组/分区/合计行）。"""
    if not a or not b:
        return False
    a = str(a).strip()
    if TOTAL_RE.search(a):
        return False
    if SECTION_RE.match(a):
        return False
    if GROUP_RE.match(a):
        return False
    return True


def parse_file(path, max_rows=3000):
    wb = load_workbook(path, data_only=False)
    ws = wb.worksheets[0]  # 主表
    hrow = find_header_row(ws) or 6
    items = []
    for r in range(hrow + 1, min(ws.max_row, max_rows) + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        if not is_item_row(a, b):
            continue
        code = str(a).strip()
        title, bullets, cn_note, scope = split_desc(b)
        qty = ws.cell(r, 5).value
        uom = ws.cell(r, 6).value
        brand = ws.cell(r, 7).value
        area = ws.cell(r, 3).value
        items.append({
            "file": os.path.basename(path),
            "sheet": ws.title,
            "row": r,
            "code": code,
            "area": area if area is not None else "",
            "title_cn_en": title,
            "spec_bullets": bullets,
            "cn_note": cn_note,
            "scope": scope,
            "uom": uom if uom is not None else "",
            "qty_expr": qty if qty is not None else "",
            "brand": brand if brand is not None else "",
            # 以下字段由 agent 在后续阶段填写
            "in_scope": None,          # 阶段1：用户确认范围后设 true/false
            "unit_price": None,        # 阶段2/3：填写单价
            "source": None,            # 知识库-KB / 知识库-KB(替换) / AI建议参考价 / 未报价
            "kb_match": None,          # 知识库匹配产品名
            "spec_consistent": None,   # 规格是否一致 true/false
            "replacement_desc": None,  # 替换说明（规格差异）
            "replacement_image": None, # 替换产品图片路径（如有）
            "note": None,
        })
    return items


def main():
    ap = argparse.ArgumentParser(description="解析报价清单 -> items_extracted.json")
    ap.add_argument("paths", nargs="+", help="报价清单 xlsx 文件或目录（可多个）")
    ap.add_argument("--json", default="items_extracted.json")
    ap.add_argument("--project", default="")
    args = ap.parse_args()

    items, files = [], []
    for p in args.paths:
        if os.path.isdir(p):
            for fn in sorted(os.listdir(p)):
                if fn.lower().endswith(".xlsx") and not fn.startswith("~"):
                    items += parse_file(os.path.join(p, fn))
                    files.append(fn)
        elif p.lower().endswith(".xlsx"):
            items += parse_file(p)
            files.append(os.path.basename(p))

    out = {
        "project": args.project,
        "files": sorted(set(files)),
        "item_count": len(items),
        "items": items,
    }
    with open(args.json, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

    # 控制台结构摘要
    by_file = {}
    for it in items:
        by_file.setdefault(it["file"], 0)
        by_file[it["file"]] += 1
    print("已抽取条目：%d 条，来自 %d 个文件" % (len(items), len(by_file)))
    for f, c in by_file.items():
        print("  - %s : %d 条" % (f, c))
    print("写出 -> %s" % args.json)


if __name__ == "__main__":
    main()
