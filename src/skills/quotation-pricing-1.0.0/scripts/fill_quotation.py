#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把 items_extracted.json 中 agent 填写的报价写回报价清单 Excel。

写入内容：
  - I 列（单价）= unit_price
  - M 列（报价来源 Source）= source
  - N 列（报价说明/替换 Note）= note 或「替换：<replacement_desc>」
  - 新建「替换与差异 Replacements」sheet：规格不一致/替换项（含图片嵌入，如有）

默认输出 <原名>_priced.xlsx，不破坏原文件。

用法：
  python fill_quotation.py --json items_extracted.json \
        [--indir <清单所在目录>] [--outdir <输出目录>]
"""
import argparse
import json
import os

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment

PRICE_COL = 9   # I
SRC_COL = 13    # M
NOTE_COL = 14   # N
REPL_SHEET = "替换与差异 Replacements"


def find_header_row(ws):
    for r in range(1, min(ws.max_row, 30) + 1):
        a = ws.cell(r, 1).value
        if a and ("编号" in str(a) or "No." in str(a)):
            return r
    return 6


def create_replacements_sheet(wb, items):
    if REPL_SHEET in wb.sheetnames:
        del wb[REPL_SHEET]
    ws = wb.create_sheet(REPL_SHEET)
    headers = ["编号 Code", "原产品 Product", "KB匹配 Match",
               "规格差异/替换说明 Deviation", "单价 UnitPrice",
               "来源 Source", "图片 Image"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c)
        cell.value = h
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D3D3D3")
    r = 2
    for it in items:
        ws.cell(r, 1).value = it.get("code", "")
        ws.cell(r, 2).value = it.get("title_cn_en", "")
        ws.cell(r, 3).value = it.get("kb_match") or ""
        ws.cell(r, 4).value = it.get("replacement_desc") or ""
        up = it.get("unit_price")
        ws.cell(r, 5).value = up if up not in (None, "") else ""
        ws.cell(r, 6).value = it.get("source") or ""
        img = it.get("replacement_image")
        if img and os.path.exists(img):
            try:
                img_obj = XLImage(img)
                maxw = 120
                if img_obj.width > maxw:
                    ratio = maxw / img_obj.width
                    img_obj.width = int(img_obj.width * ratio)
                    img_obj.height = int(img_obj.height * ratio)
                ws.add_image(img_obj, "G%d" % r)
                ws.row_dimensions[r].height = max(
                    ws.row_dimensions[r].height or 20, img_obj.height + 10)
            except Exception as e:  # noqa: BLE001
                ws.cell(r, 7).value = "[图片读取失败:%s]" % e
        r += 1
    widths = [14, 30, 24, 50, 12, 18, 22]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + c)].width = w
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser(description="回填报价到 Excel")
    ap.add_argument("--json", required=True)
    ap.add_argument("--indir", default="", help="清单所在目录（文件名不含路径时）")
    ap.add_argument("--outdir", default="", help="输出目录；默认与原文件同目录")
    args = ap.parse_args()

    data = json.load(open(args.json, encoding="utf-8"))
    by_file = {}
    for it in data["items"]:
        by_file.setdefault(it["file"], []).append(it)

    written = []
    for fname, items in by_file.items():
        path = os.path.join(args.indir, fname) if args.indir else fname
        if not os.path.exists(path):
            print("⚠ 找不到文件：%s，跳过" % fname)
            continue
        wb = load_workbook(path)
        ws = wb.worksheets[0]
        hrow = find_header_row(ws)
        ws.cell(hrow, SRC_COL).value = "报价来源 Source"
        ws.cell(hrow, NOTE_COL).value = "报价说明/替换 Note"

        replaced = []
        for it in items:
            r = it["row"]
            up = it.get("unit_price")
            if it.get("in_scope") and up not in (None, ""):
                try:
                    ws.cell(r, PRICE_COL).value = float(up)
                except (TypeError, ValueError):
                    ws.cell(r, PRICE_COL).value = up
            if it.get("in_scope"):
                ws.cell(r, SRC_COL).value = it.get("source") or ""
                note = it.get("note") or ""
                if it.get("replacement_desc"):
                    note = (note + " | 替换：" + it["replacement_desc"]).strip(" |")
                ws.cell(r, NOTE_COL).value = note
                if it.get("spec_consistent") is False or it.get("replacement_desc"):
                    replaced.append(it)

        if replaced:
            create_replacements_sheet(wb, replaced)

        stem, ext = os.path.splitext(path)
        if args.outdir:
            os.makedirs(args.outdir, exist_ok=True)
            outpath = os.path.join(args.outdir, os.path.basename(stem) + "_priced" + ext)
        else:
            outpath = stem + "_priced" + ext
        wb.save(outpath)
        written.append(outpath)
        print("写出 -> %s （替换差异 %d 项）" % (outpath, len(replaced)))

    print("完成：%d 个文件" % len(written))


if __name__ == "__main__":
    main()
