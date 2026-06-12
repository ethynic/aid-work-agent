#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel 模板导出"""

import json
import os
import re
import shutil
import tempfile
import uuid
from pathlib import Path

from loguru import logger

# 技能目录
SKILL_DIR = Path(__file__).resolve().parent.parent


def _tmp_path(suffix: str) -> str:
    """跨平台临时文件路径：优先项目级 storage/tmp（环境变量注入），回退系统 tmp。
    避免依赖系统 /tmp（部分生产环境无写入权限）。
    """
    base = os.environ.get('SKILL_TMP_DIR')
    if base:
        d = Path(base)
        d.mkdir(parents=True, exist_ok=True)
        return str(d / f"{uuid.uuid4().hex}{suffix}")
    return tempfile.mktemp(suffix=suffix)


def export_with_template(quote_data: dict, template_path: str) -> str:
    """使用模板导出报价单 Excel"""
    try:
        import openpyxl
        from openpyxl.cell.cell import MergedCell
        from openpyxl.utils import get_column_letter
        from copy import copy
    except ImportError:
        return _export_simple(quote_data)

    if not template_path or not Path(template_path).exists():
        template_path = str(SKILL_DIR / "templates" / "default.xlsx")

    if not Path(template_path).exists():
        return _export_simple(quote_data)

    dst = _tmp_path('.xlsx')
    shutil.copy2(template_path, dst)
    wb = openpyxl.load_workbook(dst)
    ws = wb.active

    items_start_row, items_end_row = None, None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and '{{#items}}' in str(cell.value):
                items_start_row = cell.row
            if cell.value and '{{/items}}' in str(cell.value):
                items_end_row = cell.row

    if items_start_row and items_end_row:
        template_data_row = items_start_row + 1

        template_cell_map = {}
        has_teacher_placeholder = False
        for col in range(1, ws.max_column + 1):
            c = ws.cell(row=template_data_row, column=col)
            if isinstance(c, MergedCell):
                continue
            template_cell_map[col] = {
                "value": c.value,
                "font": copy(c.font),
                "alignment": copy(c.alignment),
                "border": copy(c.border),
                "fill": copy(c.fill),
                "number_format": c.number_format,
            }
            if c.value and isinstance(c.value, str) and '{{teacher_subtotal}}' in c.value:
                has_teacher_placeholder = True

        template_row_height = ws.row_dimensions[template_data_row].height

        row_merge_ranges = []
        for merge in list(ws.merged_cells.ranges):
            if merge.min_row == template_data_row and merge.max_row == template_data_row:
                row_merge_ranges.append((merge.min_col, merge.max_col))

        teacher_col = None
        if not has_teacher_placeholder:
            subtotal_col = None
            for col, info in template_cell_map.items():
                if info["value"] and isinstance(info["value"], str) and '{{subtotal}}' in str(info["value"]):
                    subtotal_col = col
                    break
            if subtotal_col:
                candidate = subtotal_col + 1
                if candidate in template_cell_map:
                    val = template_cell_map[candidate]["value"]
                    if val == 0 or val == 0.0:
                        teacher_col = candidate

        category_col = None
        for col, info in template_cell_map.items():
            if info["value"] and '{{category}}' in str(info["value"]):
                category_col = col
                break

        saved_merges = []
        for merge in list(ws.merged_cells.ranges):
            if merge.min_row == template_data_row and merge.max_row == template_data_row:
                continue
            saved_merges.append({
                "min_row": merge.min_row,
                "max_row": merge.max_row,
                "min_col": merge.min_col,
                "max_col": merge.max_col,
            })

        for merge in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(merge))

        ws.delete_rows(items_end_row)
        ws.delete_rows(template_data_row)
        ws.delete_rows(items_start_row)
        total_row_new = items_start_row

        items_list = [item for item in quote_data.get('items', [])
                       if (item.get('quantity') or 0) > 0]
        num_items = len(items_list)

        if num_items > 0:
            ws.insert_rows(total_row_new, amount=num_items)
            if template_row_height:
                for r in range(total_row_new, total_row_new + num_items):
                    ws.row_dimensions[r].height = template_row_height

            for i, item in enumerate(items_list):
                row_num = total_row_new + i
                for col_num, tmpl in template_cell_map.items():
                    cell = ws.cell(row=row_num, column=col_num)
                    if isinstance(cell, MergedCell):
                        continue
                    cell.font = copy(tmpl["font"])
                    cell.alignment = copy(tmpl["alignment"])
                    cell.border = copy(tmpl["border"])
                    cell.fill = copy(tmpl["fill"])
                    cell.number_format = tmpl["number_format"]
                    val = tmpl["value"]
                    if isinstance(val, str) and '{{' in val:
                        val = _replace_placeholders(val, item)
                    if col_num == teacher_col:
                        val = item.get('teacher_subtotal', 0)
                    cell.value = val

                for min_col, max_col in row_merge_ranges:
                    cl_start = get_column_letter(min_col)
                    cl_end = get_column_letter(max_col)
                    ws.merge_cells(f"{cl_start}{row_num}:{cl_end}{row_num}")

            updated_total_row = total_row_new + num_items

            if category_col and num_items > 1:
                cat_letter = get_column_letter(category_col)
                categories = [item.get('category', '') for item in items_list]
                i = 0
                while i < len(categories):
                    cat = categories[i]
                    j = i + 1
                    while j < len(categories) and categories[j] == cat:
                        j += 1
                    if j - i > 1:
                        ws.merge_cells(
                            f"{cat_letter}{total_row_new + i}:"
                            f"{cat_letter}{total_row_new + j - 1}"
                        )
                    i = j

            row_offset = num_items - 3
            for sm in saved_merges:
                orig_min_row = sm["min_row"]
                orig_max_row = sm["max_row"]

                if orig_min_row < items_start_row:
                    new_min_row = orig_min_row
                    new_max_row = orig_max_row
                elif orig_min_row > items_end_row:
                    new_min_row = orig_min_row + row_offset
                    new_max_row = orig_max_row + row_offset
                else:
                    continue

                cl_start = get_column_letter(sm["min_col"])
                cl_end = get_column_letter(sm["max_col"])
                ws.merge_cells(f"{cl_start}{new_min_row}:{cl_end}{new_max_row}")

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if cell.value and isinstance(cell.value, str) and '{{' in str(cell.value):
                cell.value = _replace_placeholders(str(cell.value), quote_data)

    wb.save(dst)
    return dst


def _replace_placeholders(text: str, data: dict) -> str:
    """替换 {{变量名}} 占位符"""
    def replacer(match):
        key = match.group(1).strip()
        val = data.get(key, '')
        if isinstance(val, float):
            return f"{val:,.2f}"
        return str(val) if val is not None else ''

    return re.sub(r'\{\{(\w+)\}\}', replacer, text)


def _export_simple(quote_data: dict) -> str:
    """无模板时的简单导出"""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        dst = _tmp_path('.json')
        with open(dst, 'w', encoding='utf-8') as f:
            json.dump(quote_data, f, ensure_ascii=False, indent=2)
        return dst

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = quote_data.get('course_name', '报价')[:31]

    HEADER_FONT = Font(name='微软雅黑', bold=True, size=11)
    TITLE_FONT = Font(name='微软雅黑', bold=True, size=14)
    DATA_FONT = Font(name='微软雅黑', size=10)
    CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
    BORDER = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))

    col_widths = [12, 22, 10, 8, 6, 8, 6, 12, 12, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = 1
    ws.merge_cells(f'A{row}:J{row}')
    cell = ws.cell(row=row, column=1, value=f"{quote_data.get('company_name', '')}报价表")
    cell.font = TITLE_FONT
    cell.alignment = CENTER
    row += 1

    info = f"课程：{quote_data.get('course_name', '')}    日期：{quote_data.get('start_date', '')}    人数：{quote_data.get('total_people', '')}人    天数：{quote_data.get('trip_days', '')}天"
    if quote_data.get('teacher_count', 0) > 0:
        info += f"    随队老师：{quote_data.get('teacher_count', 0)}人"
    ws.merge_cells(f'A{row}:J{row}')
    ws.cell(row=row, column=1, value=info).font = DATA_FONT
    row += 2

    headers = ['成本类别', '项目', '单价', '数量', '单位', '次数', '单位', '费用小计', '随队老师', '备注']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER
    row += 1

    for item in quote_data.get('items', []):
        teacher_val = item.get('teacher_subtotal', 0)
        vals = [
            item.get('category', ''),
            item.get('name', ''),
            item.get('unit_price', 0),
            item.get('quantity', 1),
            item.get('unit', ''),
            item.get('frequency', 1),
            item.get('freq_unit', ''),
            item.get('subtotal', 0),
            teacher_val if teacher_val != 0 else '-',
            item.get('remark', ''),
        ]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = DATA_FONT
            c.border = BORDER
            if isinstance(v, float):
                c.number_format = '#,##0.00'
        row += 1

    ws.merge_cells(f'A{row}:G{row}')
    c = ws.cell(row=row, column=1, value='合计')
    c.font = HEADER_FONT
    c.alignment = CENTER
    c = ws.cell(row=row, column=8, value=quote_data.get('quote_per_person', 0))
    c.font = HEADER_FONT
    c.number_format = '#,##0.00'
    c = ws.cell(row=row, column=9, value=quote_data.get('teacher_total', 0))
    c.font = HEADER_FONT
    c.number_format = '#,##0.00'

    dst = _tmp_path('.xlsx')
    wb.save(dst)
    return dst
