#!/usr/bin/env python3
"""生成默认报价单模板"""

import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("需要 openpyxl: pip install openpyxl")
    sys.exit(1)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "报价单"

# 样式
TITLE_FONT = Font(name='微软雅黑', bold=True, size=14)
DATA_FONT = Font(name='微软雅黑', size=10)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
BORDER = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
HEADER_FILL = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

col_widths = [12, 22, 10, 8, 6, 8, 6, 12, 30]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

row = 1

# 标题
ws.merge_cells(f'A{row}:I{row}')
cell = ws.cell(row=row, column=1, value='{{company_name}}报价表')
cell.font = TITLE_FONT
cell.alignment = CENTER
row += 1

# 信息行
ws.merge_cells(f'A{row}:I{row}')
ws.cell(row=row, column=1,
        value='课程：{{course_name}}    日期：{{start_date}}    人数：{{total_people}}人    天数：{{trip_days}}天').font = DATA_FONT
row += 2

# 表头
headers = ['成本类别', '项目', '单价', '数量', '单位', '次数', '单位', '费用小计', '备注']
for col, h in enumerate(headers, 1):
    c = ws.cell(row=row, column=col, value=h)
    c.font = Font(name='微软雅黑', bold=True, size=11)
    c.alignment = CENTER
    c.border = BORDER
    c.fill = HEADER_FILL
row += 1

# items 起始标记
ws.cell(row=row, column=1, value='{{#items}}')
row += 1

# 数据行模板
data_vars = ['{{category}}', '{{name}}', '{{unit_price}}', '{{quantity}}', '{{unit}}',
             '{{frequency}}', '{{freq_unit}}', '{{subtotal}}', '{{remark}}']
for col, v in enumerate(data_vars, 1):
    c = ws.cell(row=row, column=col, value=v)
    c.font = DATA_FONT
    c.border = BORDER
row += 1

# items 结束标记
ws.cell(row=row, column=1, value='{{/items}}')
row += 2

# 合计行
ws.merge_cells(f'A{row}:G{row}')
ws.cell(row=row, column=1, value='合计').font = Font(name='微软雅黑', bold=True, size=11)
ws.cell(row=row, column=8, value='{{cost_per_person}}').font = Font(name='微软雅黑', bold=True, size=11)
row += 2

# 汇总信息
ws.cell(row=row, column=1, value='人均成本：{{cost_per_person}} 元').font = DATA_FONT
row += 1
ws.cell(row=row, column=1, value='人均报价：{{quote_per_person}} 元').font = DATA_FONT
row += 1
ws.cell(row=row, column=1, value='整团报价：{{quote_total}} 元').font = DATA_FONT
row += 2

# 签字区
ws.cell(row=row, column=1, value='审核签字：________').font = DATA_FONT
ws.cell(row=row, column=4, value='成本审核员：________').font = DATA_FONT
ws.cell(row=row, column=7, value='负责人：________').font = DATA_FONT

# 保存
output_path = Path(__file__).parent / "default.xlsx"
wb.save(str(output_path))
print(f"默认模板已生成: {output_path}")
