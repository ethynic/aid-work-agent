#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
aggregate_chart.py - 聚合统计 + 图表生成，输出多 Sheet Excel

功能：
- 按指定列分组聚合（sum/mean/count/min/max）
- 排序
- 生成 openpyxl 原生图表（bar/line/pie/scatter）
- 输出多 Sheet Excel：原始数据 + 汇总表 + 图表

调用方式：
  python scripts/aggregate_chart.py --file "data.xlsx" \
    --aggregations '{"group_by": ["部门"], "metrics": {"金额": ["sum"]}, "charts": [{"type": "bar", "title": "汇总", "x": "部门", "y": "金额_sum"}]}' \
    --output-dir "uploads/user123" --filename "报表.xlsx"
"""

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).parent))
from excel_utils import (
    CHART_TYPES,
    HEADER_FILL,
    HEADER_FONT,
    THIN_BORDER,
    CENTER_ALIGN,
    LEFT_ALIGN,
    DATA_FONT,
    create_chart,
    generate_output_path,
    output_error,
    output_json,
    read_data_file,
    write_formatted_excel,
)
from openpyxl.utils import get_column_letter


def build_aggregated_df(df: pd.DataFrame, config: Dict[str, Any]) -> pd.DataFrame:
    """
    根据聚合配置构建汇总 DataFrame。

    config 结构:
      group_by: ["列1", "列2"]
      metrics: {"金额": ["sum", "mean"], "人数": ["count"]}
      sort_by: {"金额_sum": "desc"}  (可选)
    """
    group_by = config.get("group_by", [])
    metrics = config.get("metrics", {})

    if not group_by:
        raise ValueError("group_by 不能为空")

    # 校验列存在
    missing_group = [c for c in group_by if c not in df.columns]
    if missing_group:
        raise ValueError(f"group_by 列不存在: {missing_group}")

    missing_metric = [c for c in metrics if c not in df.columns]
    if missing_metric:
        raise ValueError(f"metrics 列不存在: {missing_metric}")

    # 构建 agg 字典
    agg_dict: Dict[str, List[str]] = {}
    for col, funcs in metrics.items():
        valid_funcs = [f for f in funcs if f in ("sum", "mean", "count", "min", "max", "median", "std")]
        if valid_funcs:
            agg_dict[col] = valid_funcs

    if not agg_dict:
        raise ValueError("没有有效的聚合指标")

    # 转换数值列（聚合前需要数值类型）
    for col in agg_dict:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    # 执行聚合
    grouped = df.groupby(group_by, dropna=False).agg(agg_dict)

    # 展平多级列名: (金额, sum) → 金额_sum
    grouped.columns = ["_".join(col).strip("_") for col in grouped.columns.values]
    result = grouped.reset_index()

    # 四舍五入
    numeric_cols = result.select_dtypes(include="number").columns
    result[numeric_cols] = result[numeric_cols].round(2)

    # 排序
    sort_by = config.get("sort_by", {})
    if sort_by:
        for col, direction in sort_by.items():
            if col in result.columns:
                ascending = direction.lower() in ("asc", "ascending")
                result = result.sort_values(by=col, ascending=ascending)

    result = result.reset_index(drop=True)
    return result


def build_chart_sheet(
    wb: openpyxl.Workbook,
    chart_config: Dict[str, Any],
    agg_df: pd.DataFrame,
) -> str:
    """
    创建图表 Sheet：写入聚合数据 + 嵌入图表。

    Returns:
        Sheet 名称
    """
    chart_type = chart_config.get("type", "bar")
    title = chart_config.get("title", "图表")
    x_col = chart_config.get("x", "")
    y_col = chart_config.get("y", "")

    if not x_col or not y_col:
        return ""

    # 在 agg_df 中查找列
    if x_col not in agg_df.columns:
        # 尝试模糊匹配
        matches = [c for c in agg_df.columns if x_col.lower() in c.lower()]
        if not matches:
            return ""
        x_col = matches[0]

    if y_col not in agg_df.columns:
        matches = [c for c in agg_df.columns if y_col.lower() in c.lower()]
        if not matches:
            return ""
        y_col = matches[0]

    # Sheet 名称
    sheet_name = f"图表-{title}"[:31]

    ws = wb.create_sheet(title=sheet_name)

    # 写入数据（图表需要数据源在同一 Sheet）
    x_col_idx = list(agg_df.columns).index(x_col) + 1
    y_col_idx = list(agg_df.columns).index(y_col) + 1

    # 写表头
    ws.cell(row=1, column=x_col_idx, value=x_col).font = HEADER_FONT
    ws.cell(row=1, column=x_col_idx).fill = HEADER_FILL
    ws.cell(row=1, column=x_col_idx).border = THIN_BORDER
    ws.cell(row=1, column=y_col_idx, value=y_col).font = HEADER_FONT
    ws.cell(row=1, column=y_col_idx).fill = HEADER_FILL
    ws.cell(row=1, column=y_col_idx).border = THIN_BORDER

    # 写数据
    for row_idx, (_, row) in enumerate(agg_df.iterrows(), 2):
        x_val = row.get(x_col, "")
        y_val = row.get(y_col, 0)
        if pd.isna(y_val):
            y_val = 0
        ws.cell(row=row_idx, column=x_col_idx, value=x_val).border = THIN_BORDER
        ws.cell(row=row_idx, column=y_col_idx, value=float(y_val) if y_val else 0).border = THIN_BORDER

    # 创建图表
    end_row = len(agg_df) + 1
    try:
        chart = create_chart(
            chart_type=chart_type,
            ws_data=ws,
            title=title,
            x_col=x_col_idx,
            y_col=y_col_idx,
            start_row=1,
            end_row=end_row,
        )
        chart.width = 20
        chart.height = 12

        # 图表放在数据右侧
        chart_anchor = f"{get_column_letter(max(x_col_idx, y_col_idx) + 2)}2"
        ws.add_chart(chart, chart_anchor)
    except Exception as e:
        # 图表创建失败不影响数据输出
        pass

    return sheet_name


def main():
    parser = argparse.ArgumentParser(description="聚合统计 + 图表生成")
    parser.add_argument("--file", required=True, help="输入文件路径")
    parser.add_argument("--aggregations", required=True, help="聚合配置 JSON 字符串")
    parser.add_argument("--output-dir", required=True, help="输出目录")
    parser.add_argument("--filename", default="汇总报表.xlsx", help="输出文件名")
    parser.add_argument("--sheet", default=None, help="Excel Sheet 名称")
    args = parser.parse_args()

    # 解析配置
    try:
        config = json.loads(args.aggregations)
    except json.JSONDecodeError as e:
        output_error(f"aggregations JSON 格式错误: {e}")

    # 读取数据
    try:
        df, meta = read_data_file(args.file, sheet_name=args.sheet)
    except Exception as e:
        output_error(f"读取文件失败: {e}")

    # 聚合
    try:
        agg_df = build_aggregated_df(df, config)
    except Exception as e:
        output_error(f"聚合失败: {e}")

    # 生成 Excel
    output_path = generate_output_path(args.output_dir, args.filename)

    # 用 write_formatted_excel 写原始数据 + 汇总数据
    data_frames = [("原始数据", df), ("汇总统计", agg_df)]
    try:
        result = write_formatted_excel(data_frames, output_path)
    except Exception as e:
        output_error(f"生成 Excel 失败: {e}")

    # 添加图表 Sheet（需要打开已有 workbook）
    chart_sheets = []
    chart_configs = config.get("charts", [])
    if chart_configs:
        wb = openpyxl.load_workbook(str(output_path))
        for chart_cfg in chart_configs:
            sheet_name = build_chart_sheet(wb, chart_cfg, agg_df)
            if sheet_name:
                chart_sheets.append(sheet_name)
        wb.save(str(output_path))
        wb.close()

    # 汇总信息
    all_sheets = result["sheets"] + chart_sheets
    group_count = len(agg_df)

    # 提取一些摘要统计
    summary = {"groups": group_count, "charts": len(chart_sheets)}
    # 尝试添加数值列总和
    numeric_cols = agg_df.select_dtypes(include="number").columns
    for col in numeric_cols:
        if "sum" in col.lower():
            summary[col] = float(agg_df[col].sum())

    output_json({
        "success": True,
        "file_path": str(output_path),
        "filename": output_path.name,
        "sheets": all_sheets,
        "summary": summary,
        "message": f"已生成汇总报表: {group_count} 个分组, {len(chart_sheets)} 个图表",
    })


if __name__ == "__main__":
    main()
