"""通用 Excel 导出工具

提供 export_table_to_excel() 函数，所有业务表导出共用。
"""
import io
from typing import List, Optional
from loguru import logger
import openpyxl
from fastapi.responses import StreamingResponse


def export_table_to_excel(
    table_name: str,
    tenant_id: str,
    columns: List[str],
    filename: str,
    where_clause: str = "",
    where_params: Optional[tuple] = None,
    db_label: Optional[str] = None,
) -> StreamingResponse:
    """通用单表 Excel 导出

    Args:
        table_name: 数据库表名（如 bs_travel_quote_vehicles）
        tenant_id: 租户 ID
        columns: 导出列列表（不含数字 id，含 uuid）
        filename: 下载文件名（不含 .xlsx 扩展名）
        where_clause: 额外 WHERE 条件（如 AND source_type = %s）
        where_params: 额外 WHERE 条件的参数
        db_label: 数据库标签，默认使用默认连接

    Returns:
        StreamingResponse with Excel file
    """
    from src.db.database import get_db_connection

    col_str = ", ".join(f'"{c}"' for c in columns)

    params = [tenant_id]
    if where_params:
        params.extend(where_params)

    sql = f'SELECT {col_str} FROM {table_name} WHERE tenant_id = %s {where_clause} ORDER BY id'
    logger.info(f'导出表 {table_name}, tenant={tenant_id}, columns={len(columns)}')

    conn_fn = get_db_connection(db_label) if db_label else get_db_connection()
    with conn_fn as conn:
        cursor = conn.cursor()
        cursor.execute(sql, params)
        rows = cursor.fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = filename[:31]  # Excel sheet name max 31 chars

    # 写列头
    for col_idx, col_name in enumerate(columns, 1):
        ws.cell(row=1, column=col_idx, value=col_name)

    # 写数据行
    for row_idx, row in enumerate(rows, 2):
        for col_idx, col_name in enumerate(columns, 1):
            value = row[col_name] if isinstance(row, dict) else row[col_idx - 1]
            ws.cell(row=row_idx, column=col_idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"}
    )
